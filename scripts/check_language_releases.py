#!/usr/bin/env python3
"""Check national college-enrollment releases without substituting regional data."""

import argparse
import io
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin

from bs4 import BeautifulSoup
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from us_china_exchange.language_sources import MLA_DATA, ReviewClient
from us_china_exchange.model import json_text
from us_china_exchange.pipeline import write_atomic


MLA_REPORTS = ("https://www.mla.org/Resources/Guidelines-and-Data/Reports-and-Professional-Guidelines/"
               "Enrollments-in-Languages-Other-Than-English-in-United-States-Institutions-of-Higher-Education")
MLA_DATABASE = "https://apps.mla.org/flsurvey_search"


def check_mla(client, data):
    body, report_asset = client.get(MLA_REPORTS)
    report = BeautifulSoup(body, "html.parser")
    report_years = []
    for link in report.select("a[href]"):
        label = link.get_text(" ", strip=True)
        match = re.fullmatch(r"(20\d{2})\s+(?:(?:Final|Preliminary)\s+)?Report", label)
        if match:
            report_years.append(int(match[1]))
    if not report_years:
        raise ValueError("No national MLA enrollment report links found")
    body, database_asset = client.get(MLA_DATABASE)
    database = BeautifulSoup(body, "html.parser")
    workbook_url = next((urljoin(MLA_DATABASE, link["href"]) for link in database.select("a[href]")
                         if re.search(r"Historical-language-enrollments.*\.xlsx(?:\?|$)", link["href"], re.I)), MLA_DATA)
    body, workbook_asset = client.get(workbook_url)
    book = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    try:
        rows = book["FL history"].iter_rows(values_only=True)
        columns = list(next(rows))
        year_index, term_index = columns.index("SRVY_YEAR"), columns.index("TERM")
        years = set()
        for row in rows:
            if row[term_index] != "Fall":
                continue
            year = row[year_index]
            if isinstance(year, bool) or not isinstance(year, int):
                raise ValueError("The national workbook's fall survey-year coding changed")
            years.add(year)
        years = sorted(years)
    finally:
        book.close()
    if not years:
        raise ValueError("No actual fall enrollment observations found in the national workbook")
    published = {}
    for language in ("chinese", "japanese", "korean"):
        observed = [row["year"] for row in data["observations"]
                    if row["series_id"] == "mla:" + language and row["value"] is not None]
        if not observed:
            raise ValueError("No published national " + language + " course enrollments to compare")
        published[language] = max(observed)
    displayed = min(published.values())
    available = max(max(report_years), years[-1])
    if available < max(published.values()):
        raise ValueError("The national source now ends before the published series; review the source change")
    return {"source": "Modern Language Association", "scope": "Nationwide U.S. college course enrollments",
            "published_latest_year": displayed, "published_years_by_language": published,
            "report_latest_year": max(report_years),
            "workbook_fall_years": years, "available_latest_year": available,
            "status": "update_available" if available > displayed else "latest_released_data",
            "assets": [report_asset, database_asset, workbook_asset]}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT / "data/local/language-release-check.json")
    parser.add_argument("--offline", action="store_true", help="Use checksummed cached sources")
    args = parser.parse_args()
    result = {"checked_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
              "check_mode": "cached" if args.offline else "live"}
    try:
        data = json.loads((ROOT / "data/site/language_learning.json").read_text())
        result.update(check_mla(ReviewClient(ROOT / "data/local/language-current/raw", refresh=not args.offline), data))
    except Exception as exc:
        result.update(status="check_failed", error=str(exc))
        write_atomic(args.output, json_text(result))
        raise
    write_atomic(args.output, json_text(result))
    if result["status"] == "update_available":
        print(f'::warning title=New national language-enrollment release::MLA data through '
              f'{result["available_latest_year"]} are available; the dashboard ends in '
              f'{result["published_latest_year"]}. Review the new national release before updating the charts.')
    print(json_text({k: v for k, v in result.items() if k != "assets"}), end="")


if __name__ == "__main__":
    main()
