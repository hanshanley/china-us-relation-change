import copy
import csv
import io
import importlib.util
import hashlib
import json
import re
from collections import defaultdict
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import parse_qs, urlsplit
from xml.etree import ElementTree

import pytest
from bs4 import BeautifulSoup
from PIL import Image, ImageColor
from openpyxl import Workbook

from us_china_exchange.business_analysis import BENCHMARK_CHART, CHARTS, COLORS, RETIRED_CHARTS, build_business, chart_groups, svg_business_chart, validate_business_data
from us_china_exchange.business_sources import CBS_NETHERLANDS_METHODS, CBS_NETHERLANDS_REPORT, DENMARK_METADATA, DENMARK_DATA, DENMARK_QUERY, FINLAND_EUROSTAT_METHODS, POLAND_EUROSTAT_METHODS, SPAIN_EUROSTAT_METHODS, NORDIC_QUERIES, SWEDEN_COUNTRIES, SWEDEN_EUROSTAT_METHODS, SWEDEN_QUALITY, count_cell, country_counts, parse_cbs_netherlands, parse_denmark, parse_eurostat, parse_germany, parse_norway, parse_nordic, parse_ons
from us_china_exchange.render import build
from us_china_exchange.theme import CHART_THEMES, PARTNER_COLORS


ROOT = Path(__file__).resolve().parents[1]


def published():
    return json.loads((ROOT / "data/site/business_presence.json").read_text())


def eurostat():
    codes = {"freq": "A", "nace_r2": "B-S_X_O", "indic_sbs": "ENT_NR", "partner": "CN_X_HK", "geo": "EU27_2020"}
    return {
        "id": list(codes) + ["time"], "size": [1] * 5 + [3],
        "dimension": dict({k: {"category": {"index": {v: 0}}} for k, v in codes.items()},
                          time={"category": {"index": {"2021": 0, "2022": 1, "2023": 2}}}),
        "extension": {"id": "FATS_OUT_PART"}, "value": {"0": 6900, "1": 7191, "2": 0},
        "status": {"0": "e", "2": "|C"},
    }


def irish_fixture(dataset):
    if dataset == "fats_out2":
        activity, indicator, metric, code = "nace_r1", "indic_bp", "ENT", "C-O_X_L"
        years, counts = list(range(2007, 2010)), [None] * 3
    elif dataset == "fats_out2_r2":
        activity, indicator, metric, code = "nace_r2", "indic_bp", "ENT", "B-S_X_O"
        years = list(range(2010, 2021))
        counts = [36, 27, 26, 27, 30, 36, 79, 93, 111, 109, 98]
    else:
        activity, indicator, metric, code = "nace_r2", "indic_sbs", "ENT_NR", "B-S_X_O"
        years, counts = [2021, 2022, 2023], [101, 98, 122]
    codes = {"freq": "A", activity: code, indicator: metric, "partner": "CN_X_HK", "geo": "IE"}
    labels = {"geo": "Ireland", "partner": "China except Hong Kong",
              indicator: "Enterprises - number" if dataset == "fats_out_part" else "Number of enterprises"}
    data = {
        "class": "dataset", "source": "ESTAT", "updated": "2026-10-25T00:00:00Z",
        "id": [*codes, "time"], "size": [1] * len(codes) + [len(years)],
        "dimension": {key: {"category": {"index": {value: 0}, "label": {value: labels.get(key, value)}}}
                      for key, value in codes.items()},
        "extension": {"id": dataset.upper()},
        "value": {str(index): value for index, value in enumerate(counts) if value is not None},
        "status": {str(index): "|C" for index, value in enumerate(counts) if value is None},
    }
    data["dimension"]["time"] = {"category": {"index": {str(year): index for index, year in enumerate(years)}}}
    native = (",".join(codes) + "\\TIME_PERIOD\t" + "\t".join(map(str, years)) + "\n"
              + ",".join(codes.values()) + "\t"
              + "\t".join(": @C" if value is None else str(value) for value in counts) + "\n")
    return data, native.encode()


def test_distinct_parent_and_affiliate_checkpoints_remain_different_units():
    data = validate_business_data(published())
    rows = {(r["series_id"], r["year"]): r for r in data["observations"]}
    assert [rows[("bea_parents:china", y)]["value"] for y in (2004, 2009, 2014, 2019)] == [581, 841, 1244, 1134]
    assert rows[("bea_affiliates:china", 2000)]["value"] == 458
    assert rows[("bea_affiliates:china", 2023)]["value"] == 1950
    assert rows[("bea_affiliates:china", 2023)]["status"] == "provisional"
    assert rows[("eurostat_affiliates:china", 2023)]["value"] is None
    assert rows[("eurostat_affiliates:china", 2023)]["status"] == "suppressed"
    assert rows[("eurostat_affiliates:china", 2021)]["status"] == "estimated"
    assert ("bea_parents:china", 2023) not in rows
    assert all(sid.endswith(":china") for sid in data["series"])
    assert [rows[("eurostat_historical_eu:china", y)]["value"] for y in (2009, 2011, 2012)] == [5688, 6607, 6860]
    assert [rows[("eurostat_uk:china", y)]["value"] for y in (2009, 2010, 2018, 2019)] == [691, 681, 709, 531]
    assert rows[("eurostat_uk:china", 2019)]["status"] == "estimated"
    assert rows[("eurostat_norway:china", 2023)]["value"] == 88
    assert len([r for r in rows.values() if not r["series_id"].startswith(
        ("tillvaxtanalys_", "statfin_", "eurostat_france:", "eurostat_netherlands:", "eurostat_finland:",
         "eurostat_italy", "eurostat_sweden:", "eurostat_spain:", "dst_denmark:", "eurostat_austria", "eurostat_poland:",
         "eurostat_luxembourg:", "eurostat_portugal", "eurostat_czechia:", "eurostat_belgium",
         "snb_swiss_liechtenstein:", "ine_spain", "eurostat_hungary", "eurostat_ireland",
         "bsi_slovenia:", "eurostat_slovenia:"))]) == 119
    assert rows[("ssb_norway:china", 2023)]["value"] == 89
    assert rows[("ssb_norway:china", 2024)]["value"] == 80


def test_irish_native_counts_reject_wrong_cubes_flags_and_confidential_reconstruction():
    from us_china_exchange.business_sources import parse_irish_counts

    for dataset, count in (("fats_out2", 3), ("fats_out2_r2", 11), ("fats_out_part", 3)):
        data, native = irish_fixture(dataset)
        assert len(parse_irish_counts(data, dataset, native)) == count
    original, native = irish_fixture("fats_out2_r2")
    mutations = (
        lambda data: data["dimension"]["geo"]["category"].update(index={"LU": 0}),
        lambda data: data["dimension"]["partner"]["category"].update(index={"HK": 0}),
        lambda data: data["dimension"]["indic_bp"]["category"].update(index={"EMP": 0}),
        lambda data: data["dimension"]["indic_bp"]["category"]["label"].update(ENT="Persons employed"),
        lambda data: data["dimension"]["nace_r2"]["category"].update(index={"TOTAL": 0}),
        lambda data: data["dimension"]["time"]["category"]["index"].update({"2011": 0}),
        lambda data: data["dimension"]["time"]["category"]["index"].update({"2011": True}),
        lambda data: data["extension"].update(id="fats_in2"),
        lambda data: data["value"].update({"0": 36.5}),
        lambda data: data["value"].update({"0": True}),
        lambda data: data["status"].update({"0": "Z"}),
        lambda data: data["status"].update({"0": "b"}),
        lambda data: data["value"].update({"99": 1}),
        lambda data: data["id"].append("time"),
    )
    for mutate in mutations:
        changed = copy.deepcopy(original)
        mutate(changed)
        with pytest.raises(ValueError):
            parse_irish_counts(changed, "fats_out2_r2", native)
    for changed in (native.replace(b"CN_X_HK,IE", b"CN_X_HK,LU"),
                    native.replace(b"\t36\t", b"\t37\t", 1)):
        with pytest.raises(ValueError):
            parse_irish_counts(original, "fats_out2_r2", changed)
    confidential, native = irish_fixture("fats_out2")
    confidential["value"]["0"] = 0
    with pytest.raises(ValueError):
        parse_irish_counts(confidential, "fats_out2", native)


def test_irish_collector_preserves_foreign_data_and_keeps_2023_outside_the_comparable_window(monkeypatch):
    from us_china_exchange import business_sources

    background = ("FATS data are compiled according to the UCI concept; 51% or greater; "
                  "Irish resident institutional unit.")
    quality = (
        "The unit of observation is the foreign affiliate. 2022 is directly comparable "
        "with 2021, 2020, 2019, 2018, 2017, 2016, 2015, 2014, 2013, 2012, 2011 and 2010. "
        "Exclude zero employment and zero turnover and ownership less than 50 per cent. "
        "The response rate for OFATS was 52% in 2022; data was received from UCIs in a previous year. "
        "The number of foreign affiliates controlled by Irish multinationals is underestimated.")
    monkeypatch.setattr(business_sources, "pdf_text", lambda _: [quality])

    class Client:
        def get(self, url):
            if url == business_sources.IRELAND_BACKGROUND:
                content = ("<p>" + background + "</p>").encode()
            elif url == business_sources.IRELAND_QUALITY:
                content = b"original quality report"
            else:
                dataset = next(dataset for dataset in ("fats_out2_r2", "fats_out_part", "fats_out2")
                               if dataset + "/" in url or dataset + "?" in url)
                data, native = irish_fixture(dataset)
                content = native if "/sdmx/" in url else json.dumps(data).encode()
            return content, dict(url=url, request_url=url, sha256=hashlib.sha256(content).hexdigest(),
                                 retrieved_at="2026-10-25T00:00:00+00:00")

    spec = importlib.util.spec_from_file_location("irish_collector", ROOT / "scripts/build_business_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    old = published()
    refreshed = module.collect_ireland(Client(), old)
    validate_business_data(refreshed)
    assert module.collect_ireland(Client(), refreshed) == refreshed
    foreign = lambda data: [row for row in data["observations"]
                            if not row["series_id"].startswith("eurostat_ireland")]
    assert foreign(refreshed) == foreign(old)
    for identity, definition in old["series"].items():
        if not identity.startswith("eurostat_ireland"):
            assert refreshed["series"][identity] == definition
    for identity, source in old["sources"].items():
        if identity not in {"fats_out2_IE", "fats_out2_r2_IE", "fats_out_part_IE"}:
            assert refreshed["sources"][identity] == source
    rows = [row for row in refreshed["observations"] if row["series_id"].startswith("eurostat_ireland")]
    assert len(rows) == 17 and sum(row["value"] is not None for row in rows) == 14
    assert [row["year"] for row in rows if row["status"] == "suppressed"] == [2007, 2008, 2009]
    comparable = [row for row in rows if row["series_id"] == "eurostat_ireland:china"]
    assert [row["year"] for row in comparable] == list(range(2010, 2023))
    assert len({row["comparability"] for row in comparable}) == 1
    assert next(row for row in rows if row["year"] == 2023)["series_id"] == "eurostat_ireland_unreviewed:china"
    assert refreshed["context"]["geography"].count(business_sources.IRELAND_CONTEXT) == 1
    prior = copy.deepcopy(refreshed)
    prior["series"] = {key: value for key, value in prior["series"].items()
                       if not key.startswith("eurostat_ireland")}
    prior["sources"] = {key: value for key, value in prior["sources"].items()
                        if key not in {"fats_out2_IE", "fats_out2_r2_IE", "fats_out_part_IE"}}
    prior["observations"] = foreign(prior)
    prior["coverage"] = [row for row in prior["coverage"] if not row["id"].startswith("eurostat_ireland")]
    prior["context"]["geography"] = prior["context"]["geography"].replace(business_sources.IRELAND_CONTEXT, "")
    monkeypatch.setattr(module, "collect_business", lambda _: refreshed)
    added = module.collect_all(Client(), prior)
    assert {key: value for key, value in added.items() if key not in {"coverage", "context"}} == {
        key: value for key, value in refreshed.items() if key not in {"coverage", "context"}}
    assert added["coverage"] == prior["coverage"] + [
        row for row in refreshed["coverage"] if row["id"].startswith("eurostat_ireland")]
    assert added["context"] == dict(
        prior["context"], geography=prior["context"]["geography"] + business_sources.IRELAND_CONTEXT)
    assert module.collect_all(Client(), added) == added
    assert module.collect_all(Client(), refreshed) == refreshed
    for changed in (background.replace("UCI concept", "resident-parent concept"),
                    background.replace("51% or greater", "any shareholding")):
        with pytest.raises(ValueError):
            business_sources.verify_irish_methods(changed, quality)
    with pytest.raises(ValueError):
        business_sources.verify_irish_methods(background, quality.replace("2022 is directly comparable", "unreviewed"))


def test_irish_archive_cannot_be_relabelled_or_spliced_into_current_comparisons():
    data = validate_business_data(published())
    assert all(not sid.startswith("eurostat_ireland") for chart in CHARTS.values() for sid in chart["series"])
    for change in ("join2023", "generic2021break", "zero_confidential", "wrong_source", "all_parents", "wrong_host"):
        changed = copy.deepcopy(data)
        if change == "join2023":
            row = next(row for row in changed["observations"] if row["series_id"] == "eurostat_ireland_unreviewed:china")
            row["series_id"] = "eurostat_ireland:china"
        elif change == "generic2021break":
            row = next(row for row in changed["observations"]
                       if row["series_id"] == "eurostat_ireland:china" and row["year"] == 2021)
            row["comparability"] = "fats_out_part:IE:0"
        elif change == "zero_confidential":
            row = next(row for row in changed["observations"] if row["series_id"] == "eurostat_ireland_legacy:china")
            row.update(value=0, status="observed")
        elif change == "wrong_source":
            row = next(row for row in changed["observations"] if row["series_id"] == "eurostat_ireland:china")
            row["source_id"] = "fats_out2_r2_FR"
        else:
            scope = changed["series"]["eurostat_ireland:china"]["scope"]
            changed["series"]["eurostat_ireland:china"]["scope"] = scope.replace(
                "Irish-resident UCI", "all Irish parents") if change == "all_parents" else scope.replace(
                    "China except Hong Kong", "China including Hong Kong")
        with pytest.raises(ValueError):
            validate_business_data(changed)


@pytest.mark.parametrize("new_populations", [False, True])
def test_default_business_refresh_preserves_reviewed_context_coverage_and_unreturned_history(tmp_path, monkeypatch, new_populations):
    import runpy
    import sys
    from us_china_exchange import business_sources
    from us_china_exchange.model import json_text

    complete = published()
    prior = copy.deepcopy(complete)
    if new_populations:
        removed = {"snb_swiss_liechtenstein:china", "ine_spain:china", "ine_spain_units:china",
                   "eurostat_luxembourg:china", "eurostat_hungary:china", "eurostat_hungary_legal:china",
                   "eurostat_ireland_legacy:china", "eurostat_ireland:china", "eurostat_ireland_unreviewed:china",
                   "bsi_slovenia:china", "eurostat_slovenia:china"}
        prior["series"] = {key: value for key, value in prior["series"].items() if key not in removed}
        prior["observations"] = [row for row in prior["observations"] if row["series_id"] not in removed]
        prior["coverage"] = [row for row in prior["coverage"] if row["id"] not in removed]
        prior["context"]["geography"] = prior["context"]["geography"].split(business_sources.SNB_CONTEXT_NOTE)[0]
    fresh = copy.deepcopy(complete)
    fresh["series"].pop("bea_parents:china")
    fresh["observations"] = [row for row in fresh["observations"] if row["series_id"] != "bea_parents:china"]
    fresh["coverage"] = [dict(row, description="Automatic replacement coverage") for row in fresh["coverage"]
                         if row["id"] != "bea_parents:china"]
    fresh["context"]["geography"] = "Automatic replacement geography"
    revised = next(row for row in fresh["observations"]
                   if row["series_id"] == "eurostat_germany:china" and row["year"] == 2007)
    revised["value"] += 1
    monkeypatch.setattr(business_sources, "collect_business", lambda client: fresh)
    folder = tmp_path / "data/site"
    folder.mkdir(parents=True)
    (folder / "business_presence.json").write_text(json_text(prior))
    script = tmp_path / "scripts/build_business_data.py"
    script.parent.mkdir()
    script.write_text((ROOT / "scripts/build_business_data.py").read_text())
    monkeypatch.setattr(sys, "argv", [str(script), "--data-only"])
    runpy.run_path(str(script), run_name="__main__")
    actual = validate_business_data(json.loads((folder / "business_presence.json").read_text()))
    expected = copy.deepcopy(complete)
    next(row for row in expected["observations"]
         if row["series_id"] == "eurostat_germany:china" and row["year"] == 2007)["value"] += 1
    assert actual == expected


def test_italian_collection_preserves_other_countries_and_separates_statistical_units(monkeypatch):
    from us_china_exchange import business_sources
    from us_china_exchange.business_sources import ITALY_EUROSTAT_METHODS, ISTAT_ITALY_REPORT

    spec = importlib.util.spec_from_file_location("italian_collector", ROOT / "scripts/build_business_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    def workbook(file_contents):
        data = json.loads(file_contents)
        counts = {"Totale": data["count"], "Industria": 500,
                  "Servizi": data["count"] - 500 + int(data["wrong_sector_sum"])}
        sheets = []
        for name, count in counts.items():
            rows = [[f"Imprese a controllo nazionale residenti all'estero - Anno {data['year']}"],
                    ["PAESI", "Addetti" if data["wrong_national_unit"] else "Imprese"],
                    ["Hong Kong" if data["wrong_national_country"] else "Cina", count],
                    ["Totale primi 40 paesi", count], ["Cina", 4.4]]
            sheets.append(SimpleNamespace(name=name, nrows=len(rows), row_values=lambda i, rows=rows: rows[i],
                                          cell_value=lambda i, j, rows=rows: rows[i][j]))
        return SimpleNamespace(sheets=lambda: sheets)

    monkeypatch.setattr(business_sources.xlrd, "open_workbook", workbook)
    monkeypatch.setattr(business_sources, "pdf_text", lambda content: [content.decode()])

    class Client:
        wrong_country = False
        wrong_methods = False
        wrong_national_methods = False
        wrong_workbook_year = False
        wrong_national_unit = False
        wrong_national_country = False
        wrong_sector_sum = False
        wrong_overlap_2019 = False
        wrong_overlap_2020 = False

        def get(self, url):
            if url == ITALY_EUROSTAT_METHODS:
                body = (b"Unrelated methods" if self.wrong_methods else
                        b"ISTAT-Italian National Statistical Institute. "
                        b"In previous years the legal unit was used. 2021 onwards")
            elif url == ISTAT_ITALY_REPORT:
                text = ("Nel caso del presente Report \u00e8 sempre l\u2019unit\u00e0 giuridica. "
                        "Outward, i primi dati sono relativi al 2021. "
                        "sottoposte a controllo ultimo nazionale; settori da B a N e P -Q-R-S")
                body = ("Enterprises only" if self.wrong_national_methods else text).encode()
            elif "annuarioistatice.istat.it" in url:
                edition = int(re.search(r"/(?:epo|tavole)/(\d{4})/", url).group(1))
                year = edition - 3
                if url.endswith(".html"):
                    body = (f'<table><tr><td>Imprese a controllo nazionale residenti all\'estero per '
                            f'principali paesi e macro-settore - Anno {year}</td>'
                            f'<td><a href="/tavole/{edition}/italy.xls">XLS</a></td></tr></table>').encode()
                else:
                    count = {2019: 1141, 2020: 1133, 2021: 1113, 2022: 1169, 2023: 1122}[year]
                    count += int(getattr(self, f"wrong_overlap_{year}", False))
                    body = json.dumps({"year": year + int(self.wrong_workbook_year), "count": count,
                                       "wrong_national_unit": self.wrong_national_unit,
                                       "wrong_national_country": self.wrong_national_country,
                                       "wrong_sector_sum": self.wrong_sector_sum}).encode()
            else:
                dataset = url.split("/data/")[1].split("?")[0]
                indicator = "indic_sbs" if dataset == "fats_out_part" else "indic_bp"
                activity = "nace_r1" if dataset == "fats_out2" else "nace_r2"
                codes = {"freq": "A", activity: "C-O_X_L" if activity == "nace_r1" else "B-S_X_O",
                         indicator: "ENT_NR" if indicator == "indic_sbs" else "ENT",
                         "partner": "CN_X_HK", "geo": "FR" if self.wrong_country else "IT"}
                values = ({2007: 792} if dataset == "fats_out2" else
                          {2010: 923, 2019: 1141, 2020: 1133} if dataset == "fats_out2_r2" else {2023: 669})
                data = {"id": list(codes) + ["time"], "size": [1] * len(codes) + [len(values)],
                        "dimension": {k: {"category": {"index": {v: 0}}} for k, v in codes.items()},
                        "extension": {"id": dataset}, "value": {str(i): v for i, v in enumerate(values.values())},
                        "updated": "2026-10-25"}
                data["dimension"]["time"] = {"category": {"index": {str(y): i for i, y in enumerate(values)}}}
                if dataset != "fats_out_part":
                    data["status"] = {str(i): "e" for i in range(len(values))}
                body = json.dumps(data).encode()
            return body, {"url": url, "request_url": url, "sha256": hashlib.sha256(body).hexdigest(),
                          "retrieved_at": "2026-10-25T00:00:00+00:00"}

    prior = published()
    prior["series"] = {k: v for k, v in prior["series"].items() if not k.startswith("eurostat_italy")}
    prior["observations"] = [r for r in prior["observations"] if not r["series_id"].startswith("eurostat_italy")]
    prior["sources"] = {k: v for k, v in prior["sources"].items()
                        if not k.endswith("_IT") and not k.startswith("istat_italy_legal_")}
    prior["coverage"] = [r for r in prior["coverage"] if not r["id"].startswith("eurostat_italy")]
    before = copy.deepcopy(prior)
    updated = validate_business_data(module.collect_italy(Client(), prior))
    assert prior == before
    actual = {(r["series_id"], r["year"]): r for r in updated["observations"]}
    assert all(actual[r["series_id"], r["year"]] == r for r in prior["observations"])
    assert all(updated["series"][k] == v for k, v in prior["series"].items())
    assert all(updated["sources"][k] == v for k, v in prior["sources"].items())
    assert actual["eurostat_italy_legal:china", 2007]["status"] == "estimated"
    assert actual["eurostat_italy:china", 2023]["status"] == "observed"
    assert updated["series"]["eurostat_italy_legal:china"]["unit"] == "legal units"
    assert updated["series"]["eurostat_italy:china"]["unit"] == "enterprises"
    assert [actual["eurostat_italy_legal:china", y]["value"] for y in (2021, 2022, 2023)] == [1113, 1169, 1122]
    assert all(actual["eurostat_italy_legal:china", y]["source_id"] == f"istat_italy_legal_{y}"
               for y in (2021, 2022, 2023))
    assert module.collect_italy(Client(), updated) == updated
    for attribute in ("wrong_country", "wrong_methods", "wrong_national_methods", "wrong_workbook_year",
                      "wrong_national_unit", "wrong_national_country", "wrong_sector_sum",
                      "wrong_overlap_2019", "wrong_overlap_2020"):
        client = Client()
        setattr(client, attribute, True)
        with pytest.raises(ValueError):
            module.collect_italy(client, prior)


def test_italian_legal_unit_history_uses_national_continuation_not_enterprise_counts():
    data = published()
    points = chart_groups(data, "italy-affiliates-china")[0]["points"]
    assert [point["year"] for point in points] == list(range(2007, 2024))
    assert [point["value"] for point in points[-3:]] == [1113, 1169, 1122]
    legacy = [r for r in data["observations"] if r["series_id"] == "eurostat_italy_legal:china"]
    assert [r["year"] for r in legacy] == list(range(2007, 2024))
    assert all(r["status"] == "estimated" for r in legacy)
    assert not any("eurostat_italy:china" in chart["series"] for chart in CHARTS.values())
    for sid, year in (("eurostat_italy_legal:china", 2024), ("eurostat_italy:china", 2020)):
        invalid = copy.deepcopy(data)
        next(r for r in invalid["observations"] if r["series_id"] == sid)["year"] = year
        with pytest.raises(ValueError, match="must not be spliced"):
            validate_business_data(invalid)
    invalid = copy.deepcopy(data)
    next(r for r in invalid["observations"]
         if r["series_id"] == "eurostat_italy_legal:china" and r["year"] == 2023)["source_id"] = "fats_out_part_IT"
    with pytest.raises(ValueError, match="method/source boundary"):
        validate_business_data(invalid)
    invalid = copy.deepcopy(data)
    invalid["series"]["eurostat_italy_legal:china"]["unit"] = "enterprises"
    with pytest.raises(ValueError, match="counting units"):
        validate_business_data(invalid)
    for theme in CHART_THEMES:
        svg = ElementTree.fromstring(svg_business_chart(data, "italy-affiliates-china", theme))
        bars = [node for node in svg.iter() if node.get("class") == "annual-column"]
        assert [int(node.get("data-year")) for node in bars] == list(range(2007, 2024))
        assert [int(node.get("data-value")) for node in bars[-3:]] == [1113, 1169, 1122]


def test_uk_early_refresh_preserves_later_observations_and_classification_boundaries():
    spec = importlib.util.spec_from_file_location("business_uk_history_cli", ROOT / "scripts/build_business_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    class Client:
        wrong_country = False
        wrong_unit = False
        wrong_year = False

        def get(self, url):
            assert "/fats_out2?" in url and "geo=UK" in url and "indic_bp=ENT" in url
            codes = {"freq": "A", "nace_r1": "C-O_X_L", "indic_bp": "EMP" if self.wrong_unit else "ENT",
                     "partner": "CN_X_HK", "geo": "FR" if self.wrong_country else "UK"}
            years = [2007, 2008, 2010 if self.wrong_year else 2009]
            data = {
                "id": list(codes) + ["time"], "size": [1] * len(codes) + [3],
                "dimension": {key: {"category": {"index": {value: 0}}} for key, value in codes.items()},
                "extension": {"id": "FATS_OUT2"}, "value": {"2": 691},
                "label": "Outward FATS - main variables (2007-2009)", "updated": "2024-10-14"}
            data["dimension"]["time"] = {"category": {"index": {str(year): i for i, year in enumerate(years)}}}
            body = json.dumps(data).encode()
            return body, {"url": url, "request_url": url, "sha256": hashlib.sha256(body).hexdigest(),
                          "retrieved_at": "2026-10-25T00:00:00+00:00"}

    sid = "eurostat_uk:china"
    prior = published()
    prior["observations"] = [row for row in prior["observations"]
                             if row["series_id"] != sid or row["year"] >= 2010]
    prior["sources"].pop("fats_out2_UK", None)
    before = copy.deepcopy(prior)
    updated = validate_business_data(module.collect_uk_history(Client(), prior))
    assert prior == before
    old_rows = {(row["series_id"], row["year"]): row for row in prior["observations"]}
    new_rows = {(row["series_id"], row["year"]): row for row in updated["observations"]}
    assert all(new_rows[key] == value for key, value in old_rows.items())
    assert len(new_rows) == len(old_rows) + 3
    assert new_rows[(sid, 2009)]["value"] == 691
    assert all(new_rows[(sid, year)]["value"] is None
               and new_rows[(sid, year)]["status"] == "missing" for year in (2007, 2008))
    assert new_rows[(sid, 2009)]["comparability"] != new_rows[(sid, 2010)]["comparability"]
    assert all(updated["series"][key] == value for key, value in prior["series"].items() if key != sid)
    assert all(updated["sources"][key] == value for key, value in prior["sources"].items())
    assert updated["context"] == prior["context"]
    assert module.collect_uk_history(Client(), updated) == updated
    for attribute in ("wrong_country", "wrong_unit", "wrong_year"):
        client = Client()
        setattr(client, attribute, True)
        with pytest.raises(ValueError):
            module.collect_uk_history(client, prior)
    for year, wrong_source in ((2009, "fats_out2_r2_UK"), (2010, "fats_out2_UK")):
        invalid = copy.deepcopy(updated)
        next(row for row in invalid["observations"]
             if row["series_id"] == sid and row["year"] == year)["source_id"] = wrong_source
        with pytest.raises(ValueError, match="method/source boundary"):
            validate_business_data(invalid)


def cbs_netherlands():
    return b"""<p>* voorlopige cijfers</p>
<figure>5.3.7 Buitenlandse bedrijven in Nederland
<table><tr><th>Land</th><th>2024*</th></tr><tr><td>China</td><td>590</td></tr></table></figure>
<figure>5.4.1 Top 10 landen met Nederlandse dochterondernemingen (dochterondernemingen)
<table><tr><th>Land</th><th>2024*</th><th>2023</th></tr>
<tr><td>China</td><td>620</td><td>574</td></tr></table></figure>"""


def test_cbs_dutch_counts_are_outward_enterprises_not_inward_companies_or_employees():
    body = cbs_netherlands()
    assert parse_cbs_netherlands(body) == {2024: (620, "provisional"), 2023: (574, "observed")}
    for old, new in ((b"(dochterondernemingen)", b"(FTE)"),
                     (b"<th>2024*</th><th>2023", b"<th>2025*</th><th>2024"),
                     (b"<td>620</td>", b"<td>.</td>"),
                     (b"<td>China</td><td>620", b"<td>Japan</td><td>620"),
                     (b"* voorlopige cijfers", b"** nader voorlopige cijfers")):
        with pytest.raises(ValueError):
            parse_cbs_netherlands(body.replace(old, new))


def test_netherlands_selective_collection_preserves_other_countries_and_rejects_conflicts():
    spec = importlib.util.spec_from_file_location("business_netherlands_cli", ROOT / "scripts/build_business_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    class Client:
        conflict = False
        wrong_country = False

        def get(self, url):
            if url == CBS_NETHERLANDS_REPORT:
                body = cbs_netherlands()
                if self.conflict:
                    body = body.replace(b"<td>574</td>", b"<td>575</td>")
            elif url == CBS_NETHERLANDS_METHODS:
                body = ("UCI in Nederland; SBI 2008 sectie B t/m S exclusief O; "
                        "financi\u00eble en niet financi\u00eble; vanaf verslagperiode 2007; "
                        "Voor de cijfers van 2021 is een nieuwe methodiek gekozen").encode()
            else:
                dataset = url.split("/data/")[1].split("?")[0]
                indicator = "indic_sbs" if dataset == "fats_out_part" else "indic_bp"
                activity = "nace_r1" if dataset == "fats_out2" else "nace_r2"
                codes = {"freq": "A", activity: "C-O_X_L" if activity == "nace_r1" else "B-S_X_O",
                         indicator: "ENT_NR" if indicator == "indic_sbs" else "ENT",
                         "partner": "CN_X_HK", "geo": "UK" if self.wrong_country else "NL"}
                values = ([(2007, 465), (2008, 592), (2009, 761)] if dataset == "fats_out2" else
                          [(2019, 695), (2020, 689)] if dataset == "fats_out2_r2" else
                          [(2021, 681), (2022, 714), (2023, 574)])
                d = {"id": list(codes) + ["time"], "size": [1] * len(codes) + [len(values)],
                     "dimension": {k: {"category": {"index": {v: 0}}} for k, v in codes.items()},
                     "extension": {"id": dataset}, "value": {str(i): v for i, (_, v) in enumerate(values)},
                     "label": "Reviewed enterprise population", "updated": "2026-03-31"}
                d["dimension"]["time"] = {"category": {"index": {str(y): i for i, (y, _) in enumerate(values)}}}
                body = json.dumps(d).encode()
            return body, {"url": url, "request_url": url, "sha256": hashlib.sha256(body).hexdigest(),
                          "retrieved_at": "2026-10-25T00:00:00+00:00"}

    old = published()
    before = copy.deepcopy(old)
    updated = validate_business_data(module.collect_netherlands(Client(), old))
    assert old == before
    sid = "eurostat_netherlands:china"
    assert [r for r in updated["observations"] if r["series_id"] != sid] == [
        r for r in old["observations"] if r["series_id"] != sid]
    assert all(updated["series"][k] == v for k, v in old["series"].items() if k != sid)
    points = {r["year"]: r for r in updated["observations"] if r["series_id"] == sid}
    assert (points[2024]["value"], points[2024]["status"]) == (620, "provisional")
    assert points[2019]["comparability"] != points[2024]["comparability"]
    assert points[2023]["comparability"] == points[2024]["comparability"]
    assert module.collect_netherlands(Client(), updated) == updated
    for attribute in ("conflict", "wrong_country"):
        client = Client()
        setattr(client, attribute, True)
        with pytest.raises(ValueError):
            module.collect_netherlands(client, old)


def test_dutch_history_has_every_reported_year_and_explicit_method_boundaries():
    data = published()
    points = chart_groups(data, "netherlands-affiliates-china")[0]["points"]
    assert [p["year"] for p in points] == list(range(2007, 2025))
    assert [points[i]["value"] for i in (0, 12, 16, 17)] == [465, 695, 574, 620]
    assert points[2]["comparability"] != points[3]["comparability"]
    assert points[13]["comparability"] != points[14]["comparability"]
    for theme in CHART_THEMES:
        svg = ElementTree.fromstring(svg_business_chart(data, "netherlands-affiliates-china", theme))
        assert len([n for n in svg.iter() if n.get("class") == "annual-column"]) == 18
        assert not [n for n in svg.iter() if n.get("class") in {"data-line", "covid-band"}]
        assert "2024 provisional" in "".join(svg.itertext())
    wrong = copy.deepcopy(data)
    next(r for r in wrong["observations"] if r["series_id"] == "eurostat_netherlands:china" and
         r["year"] == 2024)["source_id"] = "fats_out2_r2_NL"
    with pytest.raises(ValueError, match="boundary"):
        validate_business_data(wrong)


def test_german_workbook_counts_entities_not_employment_and_does_not_splice_other_vintages():
    book = Workbook()
    book.active.title = "DATA"
    for r in [[], [], [None, 45291, 45657],
              ["BBFOFATS.A.DE.ENT.BTSXO.CN", 1929, 1911],
              ["BBFOFATS.A.DE.EMP.BTSXO.CN", 543341, 550000],
              ["BBFOFATS.A.DE.ENT.BTSXO.HK", 345, 318]]:
        book.active.append(r)
    book.create_sheet("Methodology").append([
        "Foreign affiliates ultimately controlled by German investors; more than 50%; total assets exceed €3 million"])
    book.create_sheet("Economic activity").append([
        "BTSXO", "All economic activities (without agriculture, public administration, private households)"])

    def parse():
        stream = io.BytesIO()
        book.save(stream)
        return parse_germany(stream.getvalue())

    assert parse() == [(2023, 1929, "observed"), (2024, 1911, "provisional")]
    book["DATA"].cell(4, 3, "X")
    assert parse()[1] == (2024, None, "missing")
    book["DATA"].cell(4, 1, "BBFOFATS.A.DE.EMP.BTSXO.CN")
    with pytest.raises(ValueError, match="enterprise-count"):
        parse()
    groups = (chart_groups(published(), "germany-affiliates-china-eurostat")
              + chart_groups(published(), "germany-affiliates-china"))
    assert [g["id"] for g in groups] == ["eurostat_germany:china", "bundesbank_germany:china"]
    assert [[p["year"] for p in g["points"]] for g in groups] == [
        list(range(2007, 2024)), list(range(2015, 2019)) + list(range(2021, 2025))]
    assert [next(p["value"] for p in g["points"] if p["year"] == 2023) for g in groups] == [1921, 1929]
    eurostat_years = {p["year"]: p for p in groups[0]["points"]}
    assert eurostat_years[2009]["comparability"] != eurostat_years[2010]["comparability"]
    assert eurostat_years[2020]["comparability"] != eurostat_years[2021]["comparability"]


def test_german_archived_report_rejects_other_years_units_and_control_populations(monkeypatch):
    from us_china_exchange import business_sources

    method = ("November 2020; ultimately controlled by German investors; more than 50%; "
              "total assets exceed \u20ac3 million; branches and permanent establishments; NACE Rev. 2; "
              "FATS do not include Section A: Agriculture and Section O: Public administration; "
              "latest figures must in all cases be regarded as provisional")
    table = ("1.1. Number of the affiliates abroad\nCountries of investment 2015 2016 2017 2018\n"
             "China, People's Republic of 1)\n1 646 1 688 1 735 1 735\nHong Kong\n1 Excluding Hong Kong.")
    monkeypatch.setattr(business_sources, "pdf_text", lambda _: [method, table])
    assert business_sources.parse_germany_history(b"%PDF-test") == [
        (2015, 1646, "observed"), (2016, 1688, "observed"),
        (2017, 1735, "observed"), (2018, 1735, "provisional")]
    with pytest.raises(ValueError, match="reviewed PDF"):
        business_sources.parse_germany_history(b"PK-current-workbook")
    for wrong_method, wrong_table in (
            (method.replace("more than 50%", "more than 10%"), table),
            (method.replace("3 million", "30 million"), table),
            (method.replace("Section O:", "Section P:"), table),
            (method, table.replace("Number of the affiliates abroad", "Employees of the affiliates abroad")),
            (method, table.replace("2015 2016 2017 2018", "2014 2015 2016 2017")),
            (method, table.replace("China, People's Republic", "Japan")),
            (method, table.replace("1 Excluding Hong Kong.", "Including Hong Kong."))):
        monkeypatch.setattr(business_sources, "pdf_text", lambda _, m=wrong_method, t=wrong_table: [m, t])
        with pytest.raises(ValueError):
            business_sources.parse_germany_history(b"%PDF-test")


def test_german_legacy_eurostat_collection_keeps_country_method_and_native_populations(monkeypatch):
    from us_china_exchange import business_sources

    prior = published()
    native = [r for r in prior["observations"] if r["series_id"] == "bundesbank_germany:china"]
    monkeypatch.setattr(business_sources, "parse_germany", lambda _: [
        (r["year"], r["value"], r["status"]) for r in native if r["year"] >= 2021])
    monkeypatch.setattr(business_sources, "parse_germany_history", lambda _: [
        (r["year"], r["value"], r["status"]) for r in native if r["year"] <= 2018])

    class Client:
        wrong = None

        def get(self, url):
            if url == business_sources.BUNDESBANK_FATS:
                body = b'<a href="/statistik-outward-fats-data.xlsx">Workbook</a>'
            elif "fats_out_simsofat_de.htm" in url:
                body = (
                    b"Compiling agency: Deutsche Bundesbank. "
                    b"an enterprise or branch; ultimate (direct or indirect) control. "
                    b"Comparable since 2007 for country breakdown. "
                    b"Collection system covers all companies with a balance sheet over 3 million Euro.")
                if self.wrong == "method":
                    body = body.replace(b"2007", b"2010")
            elif "eurostat/api/" in url:
                dataset = url.split("/data/")[1].split("?")[0]
                legacy = dataset == "fats_out2"
                codes = {
                    "freq": "A", "nace_r1" if legacy else "nace_r2": "C-O_X_L" if legacy else "B-S_X_O",
                    "indic_sbs" if dataset == "fats_out_part" else "indic_bp":
                        "ENT_NR" if dataset == "fats_out_part" else "ENT",
                    "partner": "CN_X_HK", "geo": "DE",
                }
                values = {2007: 847, 2008: 987, 2009: 1043} if legacy else {
                    r["year"]: r["value"] for r in prior["observations"]
                    if r["series_id"] == "eurostat_germany:china"
                    and r["year"] >= 2010
                    and (r["year"] <= 2020) == (dataset == "fats_out2_r2")}
                if legacy:
                    if self.wrong == "year":
                        values[2004] = values.pop(2007)
                    elif self.wrong == "truncated":
                        values.pop(2007)
                    elif self.wrong in {"partner", "geo", "indic_bp"}:
                        codes[self.wrong] = {
                            "partner": "CN", "geo": "FR", "indic_bp": "EMP"}[self.wrong]
                data = {
                    "extension": {"id": dataset}, "id": list(codes) + ["time"],
                    "size": [1] * len(codes) + [len(values)],
                    "dimension": {k: {"category": {"index": {v: 0}}} for k, v in codes.items()},
                    "value": {str(i): v for i, v in enumerate(values.values())},
                }
                data["dimension"]["time"] = {"category": {"index": {
                    str(year): i for i, year in enumerate(values)}}}
                body = json.dumps(data).encode()
            else:
                body = b"Native source is exercised by the separate workbook/PDF tests."
            return body, {
                "url": url, "request_url": url, "sha256": hashlib.sha256(body).hexdigest(),
                "retrieved_at": "2026-10-25T00:00:00+00:00"}

    client = Client()
    definitions, observations, sources = business_sources.german_affiliates(client)
    history = sorted((r for r in observations if r["series_id"] == "eurostat_germany:china"),
                     key=lambda r: r["year"])
    assert [(r["year"], r["value"]) for r in history[:3]] == [(2007, 847), (2008, 987), (2009, 1043)]
    assert [r["source_id"] for r in history[:3]] == ["fats_out2_DE"] * 3
    by_year = {r["year"]: r for r in history}
    assert by_year[2009]["comparability"] != by_year[2010]["comparability"]
    assert by_year[2020]["comparability"] != by_year[2021]["comparability"]
    assert "NACE C-O except L before 2010" in definitions["eurostat_germany:china"]["scope"]
    assert any("fats_out_simsofat_de.htm" in a["url"] for a in sources["fats_out2_DE"]["assets"])
    actual_native = {r["year"]: r for r in observations if r["series_id"] == "bundesbank_germany:china"}
    assert set(actual_native) == {r["year"] for r in native}
    assert all(actual_native[r["year"]]["value"] == r["value"]
               and actual_native[r["year"]]["status"] == r["status"] for r in native)
    for wrong in ("partner", "geo", "indic_bp", "year", "truncated", "method"):
        client.wrong = wrong
        with pytest.raises(ValueError):
            business_sources.german_affiliates(client)


def test_german_refresh_retains_all_other_country_context_and_release_vintages(monkeypatch):
    spec = importlib.util.spec_from_file_location("german_history_cli", ROOT / "scripts/build_business_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    prior = published()
    before = copy.deepcopy(prior)
    german = {"bundesbank_germany:china", "eurostat_germany:china"}
    points = [r for r in prior["observations"] if r["series_id"] in german]
    source_ids = {r["source_id"] for r in points}
    monkeypatch.setattr(module, "german_affiliates", lambda _: (
        {k: prior["series"][k] for k in german}, copy.deepcopy(points),
        {k: prior["sources"][k] for k in source_ids}))
    updated = validate_business_data(module.collect_germany(None, prior))
    assert prior == before
    assert updated["observations"] == prior["observations"]
    assert updated["series"] == prior["series"] and updated["sources"] == prior["sources"]
    assert updated["context"] == prior["context"]
    assert module.collect_germany(None, updated) == updated
    native = {r["year"]: r for r in points if r["series_id"] == "bundesbank_germany:china"}
    assert native[2018]["status"] == native[2024]["status"] == "provisional"
    assert set(native) == {*range(2015, 2019), *range(2021, 2025)}
    invalid = copy.deepcopy(updated)
    next(r for r in invalid["observations"] if r["series_id"] == "bundesbank_germany:china")["year"] = 2019
    with pytest.raises(ValueError, match="year/source"):
        validate_business_data(invalid)
    invalid = copy.deepcopy(updated)
    next(r for r in invalid["observations"] if r["series_id"] == "eurostat_germany:china"
         and r["year"] == 2007)["source_id"] = "fats_out2_r2_DE"
    with pytest.raises(ValueError, match="boundary"):
        validate_business_data(invalid)
    invalid = copy.deepcopy(updated)
    next(r for r in invalid["observations"] if r["series_id"] == "eurostat_germany:china"
         and r["year"] == 2007)["year"] = 2004
    with pytest.raises(ValueError, match="begins in 2007"):
        validate_business_data(invalid)


def test_eurostat_confidential_zero_is_not_published_and_wrong_populations_fail():
    data = eurostat()
    points = parse_eurostat(data, "fats_out_part")
    assert [r[1:3] for r in points] == [(6900, "estimated"), (7191, "observed"), (None, "suppressed")]
    data["dimension"]["partner"]["category"]["index"] = {"CN": 0}
    with pytest.raises(ValueError, match="different geography"):
        parse_eurostat(data, "fats_out_part")
    data = eurostat()
    data["status"]["0"] = "x"
    with pytest.raises(ValueError, match="Unknown"):
        parse_eurostat(data, "fats_out_part")
    data = eurostat()
    data["status"]["1"] = "b"
    points = parse_eurostat(data, "fats_out_part")
    assert points[0][-1] != points[1][-1]
    with pytest.raises(ValueError, match="count"):
        count_cell(1.5)
    data = eurostat()
    data["dimension"]["geo"]["category"]["index"] = {"UK": 0}
    assert parse_eurostat(data, "fats_out_part", "UK")[0][1] == 6900
    with pytest.raises(ValueError, match="different geography"):
        parse_eurostat(data, "fats_out_part")


def test_norwegian_history_uses_own_full_population_not_a_eurostat_endpoint_splice():
    data = {
        "source": "Statistics Norway", "class": "dataset",
        "extension": {"px": {"tableid": "09827"}},
        "id": ["Land", "ContentsCode", "Tid"], "size": [1, 1, 2],
        "dimension": {
            "Land": {"category": {"index": {"CN": 0}}},
            "ContentsCode": {"category": {"index": {"Foretak": 0},
                "unit": {"Foretak": {"base": "enterprises", "decimals": 0}}}},
            "Tid": {"category": {"index": {"2023": 0, "2024": 1}}}},
        "value": [89, 80],
    }
    assert parse_norway(data) == [(2023, 89, "observed"), (2024, 80, "observed")]
    data["value"][1] = None
    assert parse_norway(data)[1] == (2024, None, "missing")
    data["dimension"]["Land"]["category"]["index"] = {"HK": 0}
    with pytest.raises(ValueError, match="country"):
        parse_norway(data)
    groups = chart_groups(published(), "norway-affiliates-china")
    assert len(groups) == 1 and groups[0]["id"] == "ssb_norway:china"
    assert [p["year"] for p in groups[0]["points"]] == list(range(2008, 2025))


def nordic_fixture(country):
    if country == "sweden":
        dimensions = ["Tid", "Etableringsland", "variabel"]
        years, geo, metric, label = ["2009", "2010", "2024"], "cn", "2", "Antal dotterbolag"
        publisher = "Myndigheten f\u00f6r tillv\u00e4xtpolitiska utv\u00e4rderingar och analyser"
        values, flags = [None, 433, 885], {"0": ".."}
    else:
        dimensions = ["timeperiod_y", "valtio_27_20180101", "contentscode"]
        years, geo, metric, label = ["2020", "2021", "2024"], "156", "ofats_count", "Number of affiliates"
        publisher = "Statistics Finland, finnish affiliates abroad"
        values, flags = [250, 249, 242], {}
    return {
        "class": "dataset", "source": publisher, "id": dimensions, "size": [3, 1, 1],
        "dimension": {
            dimensions[0]: {"category": {"index": {year: i for i, year in enumerate(years)}}},
            dimensions[1]: {"category": {"index": {geo: 0}}},
            dimensions[2]: {"category": {"index": {metric: 0}, "label": {metric: label}}},
        }, "value": values, "status": flags,
    }


@pytest.mark.parametrize("country,values", [("sweden", [None, 433, 885]), ("finland", [250, 249, 242])])
def test_nordic_parsers_keep_reported_counts_and_reject_wrong_country_or_unit(country, values):
    data = nordic_fixture(country)
    assert [point[1] for point in parse_nordic(data, country)] == values
    if country == "sweden":
        assert parse_nordic(data, country)[0][2] == "missing"
    else:
        assert parse_nordic(data, country)[0][3] != parse_nordic(data, country)[1][3]
    for dimension, replacement in [(1, "hk"), (2, "parent_groups")]:
        changed = copy.deepcopy(data)
        changed["dimension"][changed["id"][dimension]]["category"]["index"] = {replacement: 0}
        with pytest.raises(ValueError, match="country or affiliate counting unit"):
            parse_nordic(changed, country)
    data["status"]["1"] = "unreviewed"
    with pytest.raises(ValueError, match="Unreviewed"):
        parse_nordic(data, country)


def test_nordic_selective_collection_preserves_every_other_country_and_is_idempotent():
    spec = importlib.util.spec_from_file_location("business_collector", ROOT / "scripts/build_business_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    prior = published()
    prior["observations"] = [r for r in prior["observations"]
                             if not r["series_id"].startswith(("tillvaxtanalys_", "statfin_", "eurostat_france:", "eurostat_finland:"))]
    before = copy.deepcopy(prior)

    def asset(url, checksum="b" * 64):
        return dict(url=url, request_url=url, sha256=checksum, retrieved_at="2026-10-25T00:00:00+00:00")

    def post(url, query):
        country = "sweden" if url == SWEDEN_COUNTRIES else "finland"
        assert query == NORDIC_QUERIES[country]
        return json.dumps(nordic_fixture(country)).encode(), asset(url)

    client = SimpleNamespace(post_json=post, get=lambda url: (
        b"Reviewed methodology", asset(url, "5d7b77577caf3a9fb159c106bd7fe4e9196ef50f8e09b903b7cbb678d5301289")
        if url == SWEDEN_QUALITY else asset(url)))
    result = module.collect_nordic(client, prior)
    retained = {(r["series_id"], r["year"]): r for r in result["observations"]}
    assert len(retained) == len(prior["observations"]) + 6
    assert all(retained[r["series_id"], r["year"]] == r for r in prior["observations"])
    assert prior == before
    validate_business_data(result)
    assert module.collect_nordic(client, result) == result


def test_new_nordic_histories_retain_missing_years_and_source_definition_boundaries():
    data = published()
    sweden = {r["year"]: r for r in data["observations"] if r["series_id"] == "tillvaxtanalys_sweden:china"}
    assert set(sweden) == set(range(2000, 2025))
    assert all(sweden[y]["value"] is None and sweden[y]["status"] == "missing" for y in range(2000, 2010))
    assert sum(row["value"] is not None for row in sweden.values()) == 15
    for previous, current in [(2017, 2018), (2018, 2019), (2020, 2021)]:
        assert sweden[previous]["comparability"] != sweden[current]["comparability"]
    for chart, expected_years in [("sweden-affiliates-china", range(2010, 2025)),
                                   ("finland-affiliates-china", range(2020, 2025))]:
        doc = ElementTree.fromstring(svg_business_chart(data, chart))
        bars = [node for node in doc.iter() if node.get("class") == "annual-column"]
        assert [int(node.get("data-year")) for node in bars] == list(expected_years)
        assert not any(node.get("class") == "data-line" for node in doc.iter())


def test_selective_country_refreshes_preserve_coverage_order_across_modes(monkeypatch):
    spec = importlib.util.spec_from_file_location("business_refresh_collector", ROOT / "scripts/build_business_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    prior = published()

    def selected(prefixes):
        definitions = {sid: value for sid, value in prior["series"].items() if sid.startswith(prefixes)}
        observations = [row for row in prior["observations"] if row["series_id"] in definitions]
        source_ids = {row["source_id"] for row in observations}
        return definitions, observations, {sid: prior["sources"][sid] for sid in source_ids}

    monkeypatch.setattr(module, "nordic_affiliates", lambda client: selected(("tillvaxtanalys_", "statfin_")))
    monkeypatch.setattr(module, "french_affiliates", lambda client: selected(("eurostat_france:",)))
    monkeypatch.setattr(module, "netherlands_affiliates", lambda client: selected(("eurostat_netherlands:",)))
    def early_uk(client):
        definitions, observations, sources = selected(("eurostat_uk:",))
        return definitions, [row for row in observations if row["year"] <= 2009], sources

    monkeypatch.setattr(module, "uk_early_affiliates", early_uk)
    monkeypatch.setattr(module, "finnish_eurostat_affiliates", lambda client: selected(("eurostat_finland:",)))
    monkeypatch.setattr(module, "swedish_eurostat_affiliates", lambda client: selected(("eurostat_sweden:",)))
    monkeypatch.setattr(module, "spanish_affiliates", lambda client: selected(("eurostat_spain:",)))
    monkeypatch.setattr(module, "danish_affiliates", lambda client: selected(("dst_denmark:",)))
    monkeypatch.setattr(module, "italian_affiliates", lambda client: selected(("eurostat_italy",)))
    monkeypatch.setattr(module, "austrian_affiliates", lambda client: selected(("eurostat_austria",)))
    monkeypatch.setattr(module, "polish_affiliates", lambda client: selected(("eurostat_poland:",)))
    monkeypatch.setattr(module, "luxembourg_affiliates", lambda client: selected(("eurostat_luxembourg:",)))
    monkeypatch.setattr(module, "portuguese_affiliates", lambda client: selected(("eurostat_portugal",)))
    monkeypatch.setattr(module, "czech_affiliates", lambda client: selected(("eurostat_czechia:",)))
    monkeypatch.setattr(module, "belgian_affiliates", lambda client: selected(("eurostat_belgium",)))
    monkeypatch.setattr(module, "snb_affiliates", lambda client: selected(("snb_swiss_liechtenstein:",)))
    monkeypatch.setattr(module, "hungarian_affiliates", lambda client: selected(("eurostat_hungary",)))
    monkeypatch.setattr(module, "irish_affiliates", lambda client: selected(("eurostat_ireland",)))
    monkeypatch.setattr(module, "slovenian_affiliates", lambda client: selected(("bsi_slovenia:", "eurostat_slovenia:")))
    assert module.collect_nordic(None, prior) == prior
    assert module.collect_france(None, prior) == prior
    assert module.collect_netherlands(None, prior) == prior
    assert module.collect_finland_history(None, prior) == prior
    assert module.collect_uk_history(None, prior) == prior
    assert module.collect_sweden_history(None, prior) == prior
    assert module.collect_spain(None, prior) == prior
    assert module.collect_denmark(None, prior) == prior
    assert module.collect_italy(None, prior) == prior
    assert module.collect_austria(None, prior) == prior
    assert module.collect_poland(None, prior) == prior
    assert module.collect_luxembourg(None, prior) == prior
    assert module.collect_portugal(None, prior) == prior
    assert module.collect_czechia(None, prior) == prior
    assert module.collect_belgium(None, prior) == prior
    assert module.collect_swiss_reporting_area(None, prior) == prior
    assert module.collect_hungary(None, prior) == prior
    assert module.collect_ireland(None, prior) == prior
    assert module.collect_slovenia(None, prior) == prior
    for collectors in (
        (module.collect_uk_history, module.collect_netherlands, module.collect_nordic, module.collect_france,
         module.collect_finland_history, module.collect_sweden_history, module.collect_spain, module.collect_denmark,
         module.collect_italy, module.collect_austria, module.collect_poland,
         module.collect_portugal, module.collect_czechia, module.collect_belgium, module.collect_swiss_reporting_area,
         module.collect_luxembourg, module.collect_hungary, module.collect_ireland, module.collect_slovenia),
        (module.collect_hungary, module.collect_luxembourg, module.collect_swiss_reporting_area, module.collect_belgium,
         module.collect_czechia, module.collect_portugal, module.collect_poland,
         module.collect_slovenia),
        (module.collect_swiss_reporting_area, module.collect_belgium, module.collect_czechia, module.collect_portugal, module.collect_poland,
         module.collect_austria, module.collect_italy, module.collect_denmark, module.collect_spain,
         module.collect_sweden_history, module.collect_finland_history,
         module.collect_france, module.collect_nordic, module.collect_netherlands, module.collect_uk_history,
         module.collect_slovenia),
    ):
        refreshed = prior
        for collector in collectors:
            refreshed = collector(None, refreshed)
            assert refreshed == prior


def test_spanish_selective_refresh_preserves_missingness_estimates_and_other_populations():
    spec = importlib.util.spec_from_file_location("business_spain_collector", ROOT / "scripts/build_business_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    class Client:
        wrong_country = False
        truncated_history = False

        def get(self, url):
            if url == SPAIN_EUROSTAT_METHODS:
                body = b"Spanish UCI; 2021 statistical enterprise; financial activities from Bank of Spain"
            else:
                dataset = url.split("/data/")[1].split("?")[0]
                indicator = "indic_sbs" if dataset == "fats_out_part" else "indic_bp"
                activity = "nace_r1" if dataset == "fats_out2" else "nace_r2"
                codes = {"freq": "A", activity: "C-O_X_L" if activity == "nace_r1" else "B-S_X_O",
                         indicator: "ENT_NR" if indicator == "indic_sbs" else "ENT",
                         "partner": "CN_X_HK", "geo": "FR" if self.wrong_country else "ES"}
                values = ([(2007, None), (2008, None), (2009, 58)] if dataset == "fats_out2" else
                          list(zip(range(2010, 2021), (59, 84, 87, 92, 112, 105, 111, 126, 132, 127, 126)))
                          if dataset == "fats_out2_r2" else
                          [(2021, 105), (2022, 102), (2023, 100)])
                if self.truncated_history and dataset == "fats_out2_r2":
                    values = values[1:]
                data = {
                    "id": list(codes) + ["time"], "size": [1] * len(codes) + [len(values)],
                    "dimension": {key: {"category": {"index": {value: 0}}} for key, value in codes.items()},
                    "extension": {"id": dataset},
                    "value": {str(index): value for index, (_, value) in enumerate(values) if value is not None},
                    "updated": "2026-03-31"}
                data["dimension"]["time"] = {"category": {"index": {
                    str(year): index for index, (year, _) in enumerate(values)}}}
                if dataset == "fats_out_part":
                    data["status"] = {"1": "e", "2": "e"}
                body = json.dumps(data).encode()
            return body, {"url": url, "request_url": url, "sha256": hashlib.sha256(body).hexdigest(),
                          "retrieved_at": "2026-10-25T00:00:00+00:00"}

    prior = published()
    before = copy.deepcopy(prior)
    result = validate_business_data(module.collect_spain(Client(), prior))
    sid = "eurostat_spain:china"
    assert prior == before
    assert [row for row in result["observations"] if row["series_id"] != sid] == [
        row for row in prior["observations"] if row["series_id"] != sid]
    points = {row["year"]: row for row in result["observations"] if row["series_id"] == sid}
    assert (points[2007]["value"], points[2008]["status"]) == (None, "missing")
    assert (points[2023]["value"], points[2023]["status"]) == (100, "estimated")
    assert points[2009]["comparability"] != points[2019]["comparability"]
    assert points[2019]["comparability"] == points[2020]["comparability"]
    assert points[2020]["comparability"] != points[2021]["comparability"]
    assert any(asset["url"] == SPAIN_EUROSTAT_METHODS
               for asset in result["sources"]["fats_out_part_ES"]["assets"])
    assert module.collect_spain(Client(), result) == result
    wrong = Client()
    wrong.wrong_country = True
    with pytest.raises(ValueError, match="different geography"):
        module.collect_spain(wrong, prior)
    incomplete = Client()
    incomplete.truncated_history = True
    with pytest.raises(ValueError, match="history is incomplete"):
        module.collect_spain(incomplete, prior)


def test_spanish_chart_retains_actual_years_and_rejects_a_cross_regime_source():
    data = published()
    points = chart_groups(data, "spain-affiliates-china")[0]["points"]
    assert [point["year"] for point in points] == list(range(2009, 2024))
    assert (points[0]["value"], points[-1]["value"], points[-1]["status"]) == (58, 100, "estimated")
    for theme in CHART_THEMES:
        chart = ElementTree.fromstring(svg_business_chart(data, "spain-affiliates-china", theme))
        columns = [node for node in chart.iter() if node.get("class") == "annual-column"]
        assert [int(node.get("data-year")) for node in columns] == list(range(2009, 2024))
        assert not any(node.get("class") in {"data-line", "covid-band"} for node in chart.iter())
    wrong = copy.deepcopy(data)
    row = next(row for row in wrong["observations"] if row["series_id"] == "eurostat_spain:china"
               and row["year"] == 2023)
    row["source_id"] = "fats_out2_r2_ES"
    with pytest.raises(ValueError, match="method/source boundary"):
        validate_business_data(wrong)


def denmark_fixture():
    years = [str(year) for year in range(2007, 2025)]
    metadata = {
        "id": "OFATS4", "suppressedDataValue": "0",
        "variables": [
            {"id": "LAND", "values": [{"id": "CN", "text": "China"}, {"id": "HK", "text": "Hong Kong"},
                                      {"id": "MO", "text": "Macao"}]},
            {"id": "ENHED", "values": [{"id": "DSANTAL", "text": "Affiliates"},
                                       {"id": "DSANSAT", "text": "Employees"}]},
            {"id": "Tid", "values": [{"id": year, "text": year} for year in years]}]}
    dimensions = {name: {"category": {"index": {code: 0}}}
                  for name, code in (("LAND", "CN"), ("ENHED", "DSANTAL"), ("ContentsCode", "OFATS4"))}
    dimensions["ContentsCode"]["category"]["unit"] = {"OFATS4": {"base": "Number", "decimals": 0}}
    dimensions["Tid"] = {"category": {"index": {year: index for index, year in enumerate(years)}}}
    dimensions.update(id=["LAND", "ENHED", "ContentsCode", "Tid"], size=[1, 1, 1, 18])
    data = {"dataset": {
        "dimension": dimensions, "extension": {"px": {"tableid": "OFATS4"}},
        "updated": "2026-01-20T07:00:00Z",
        "value": [217, 228, 272, 320, 353, 395, 417, 468, 487, 482, 468, 464,
                  485, 488, 481, 466, 480, 460]}}
    return data, metadata


def test_denmark_parser_preserves_country_counts_and_masks_ambiguous_zero():
    data, metadata = denmark_fixture()
    points = parse_denmark(data, metadata)
    assert [point[0] for point in points] == list(range(2007, 2025))
    assert (points[0][1], points[12][1], points[-1][1]) == (217, 485, 460)
    assert points[2][3] != points[3][3] and points[12][3] == points[-1][3]
    ambiguous = copy.deepcopy(data)
    ambiguous["dataset"]["value"][-1] = 0
    assert parse_denmark(ambiguous, metadata)[-1][1:3] == (None, "missing")
    for kind in ("country", "employees", "scaled_unit", "truncated", "quality_flag"):
        wrong = copy.deepcopy(data)
        if kind == "country":
            wrong["dataset"]["dimension"]["LAND"]["category"]["index"] = {"HK": 0}
        elif kind == "employees":
            wrong["dataset"]["dimension"]["ENHED"]["category"]["index"] = {"DSANSAT": 0}
        elif kind == "scaled_unit":
            wrong["dataset"]["dimension"]["ContentsCode"]["category"]["unit"]["OFATS4"]["base"] = "Thousands"
        elif kind == "truncated":
            wrong["dataset"]["value"].pop()
        else:
            wrong["dataset"]["status"] = {"0": "p"}
        with pytest.raises(ValueError):
            parse_denmark(wrong, metadata)


def test_danish_selective_refresh_preserves_other_countries_and_is_idempotent():
    spec = importlib.util.spec_from_file_location("business_denmark_collector", ROOT / "scripts/build_business_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    data, metadata = denmark_fixture()

    def response(url, payload):
        body = json.dumps(payload).encode()
        return body, {"url": url, "request_url": url, "sha256": hashlib.sha256(body).hexdigest(),
                      "retrieved_at": "2026-10-25T00:00:00+00:00"}

    def post(url, query):
        if url == DENMARK_METADATA:
            assert query == {"table": "OFATS4", "format": "JSON", "lang": "en"}
            return response(url, metadata)
        assert url == DENMARK_DATA and query == DENMARK_QUERY
        return response(url, data)

    client = SimpleNamespace(post_json=post, get=lambda url: response(url, "Reviewed national methodology"))
    prior = published()
    before = copy.deepcopy(prior)
    result = validate_business_data(module.collect_denmark(client, prior))
    assert prior == before
    assert [row for row in result["observations"] if row["series_id"] != "dst_denmark:china"] == [
        row for row in prior["observations"] if row["series_id"] != "dst_denmark:china"]
    assert module.collect_denmark(client, result) == result


def test_danish_chart_keeps_national_resident_parent_scope_and_all_years():
    data = published()
    points = chart_groups(data, "denmark-affiliates-china")[0]["points"]
    assert [point["year"] for point in points] == list(range(2007, 2025))
    assert (points[0]["value"], points[-1]["value"]) == (217, 460)
    for theme in CHART_THEMES:
        svg = ElementTree.fromstring(svg_business_chart(data, "denmark-affiliates-china", theme))
        columns = [node for node in svg.iter() if node.get("class") == "annual-column"]
        assert [int(node.get("data-year")) for node in columns] == list(range(2007, 2025))
        assert not any(node.get("class") in {"data-line", "covid-band"} for node in svg.iter())
    wrong = copy.deepcopy(data)
    wrong["series"]["dst_denmark:china"]["scope"] = "Danish-UCI enterprises in China"
    with pytest.raises(ValueError, match="resident-parent"):
        validate_business_data(wrong)


def test_austrian_collection_preserves_other_countries_and_rejects_wrong_source_populations():
    spec = importlib.util.spec_from_file_location("austrian_business_collector", ROOT / "scripts/build_business_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    class Client:
        wrong = None

        def get(self, url):
            asset = dict(url=url, request_url=url, sha256="b" * 64,
                         retrieved_at="2026-10-25T00:00:00+00:00")
            if "fats_out_simsofat_at.htm" in url:
                country = "France" if self.wrong == "method" else "Austria"
                return (f"Statistics {country}; from 2020 to 2021 data; "
                        "legal unit to enterprise").encode(), asset
            dataset = urlsplit(url).path.rsplit("/", 1)[-1]
            first, last = {"fats_out2": (2007, 2009), "fats_out2_r2": (2010, 2020),
                           "fats_out_part": (2021, 2023)}[dataset]
            years = list(range(first, last + 1))
            codes = {"freq": "A", "nace_r1" if dataset == "fats_out2" else "nace_r2":
                     "C-O_X_L" if dataset == "fats_out2" else "B-S_X_O",
                     "indic_sbs" if dataset == "fats_out_part" else "indic_bp":
                     "EMP" if self.wrong == "unit" else "ENT_NR" if dataset == "fats_out_part" else "ENT",
                     "partner": "CN_X_HK", "geo": "IT" if self.wrong == "country" else "AT"}
            if self.wrong == "year":
                years = years[1:]
            data = {
                "id": list(codes) + ["time"], "size": [1] * len(codes) + [len(years)],
                "dimension": {key: {"category": {"index": {value: 0}}} for key, value in codes.items()},
                "extension": {"id": dataset.upper()}, "label": dataset, "updated": "2024-10-14",
                "value": {str(i): year - 1900 for i, year in enumerate(years)}}
            data["dimension"]["time"] = {"category": {"index": {str(year): i for i, year in enumerate(years)}}}
            return json.dumps(data).encode(), asset

    prior = published()
    prior["series"] = {sid: definition for sid, definition in prior["series"].items()
                       if not sid.startswith("eurostat_austria")}
    prior["observations"] = [row for row in prior["observations"]
                             if not row["series_id"].startswith("eurostat_austria")]
    prior["sources"] = {sid: source for sid, source in prior["sources"].items() if not sid.endswith("_AT")}
    prior["coverage"] = [entry for entry in prior["coverage"] if not entry["id"].startswith("eurostat_austria")]
    before = copy.deepcopy(prior)
    updated = validate_business_data(module.collect_austria(Client(), prior))
    assert prior == before
    retained = {(row["series_id"], row["year"]): row for row in updated["observations"]}
    assert all(retained[row["series_id"], row["year"]] == row for row in prior["observations"])
    assert all(updated["series"][sid] == definition for sid, definition in prior["series"].items())
    assert all(updated["sources"][sid] == source for sid, source in prior["sources"].items())
    assert len(updated["observations"]) == len(prior["observations"]) + 17
    assert [row["year"] for row in updated["observations"]
            if row["series_id"] == "eurostat_austria_legal:china"] == list(range(2007, 2021))
    assert [row["year"] for row in updated["observations"]
            if row["series_id"] == "eurostat_austria:china"] == [2021, 2022, 2023]
    assert module.collect_austria(Client(), updated) == updated
    for wrong in ("country", "unit", "year", "method"):
        client = Client()
        client.wrong = wrong
        with pytest.raises(ValueError):
            module.collect_austria(client, prior)


def test_austrian_legal_unit_archive_cannot_refresh_the_enterprise_chart():
    data = published()
    for theme in CHART_THEMES:
        chart = ElementTree.fromstring(svg_business_chart(data, "austria-affiliates-china", theme))
        bars = [node for node in chart.iter() if node.get("class") == "annual-column"]
        assert [(int(node.get("data-year")), int(node.get("data-value"))) for node in bars] == [
            (2021, 155), (2022, 166), (2023, 175)]
    for sid, year, wrong_sid in (
        ("eurostat_austria_legal:china", 2020, "eurostat_austria:china"),
        ("eurostat_austria:china", 2021, "eurostat_austria_legal:china"),
    ):
        invalid = copy.deepcopy(data)
        next(row for row in invalid["observations"]
             if row["series_id"] == sid and row["year"] == year)["series_id"] = wrong_sid
        with pytest.raises(ValueError, match="must not be spliced"):
            validate_business_data(invalid)
    for field, wrong in (("unit", "enterprises"), ("measure", "controlled_enterprises")):
        invalid = copy.deepcopy(data)
        invalid["series"]["eurostat_austria_legal:china"][field] = wrong
        with pytest.raises(ValueError):
            validate_business_data(invalid)


@pytest.mark.parametrize("chart,axis_labels", [
    ("spain-affiliates-china", ["0", "75", "150"]),
    ("italy-affiliates-china", ["0", "750", "1.5k"]),
])
def test_static_annual_count_axis_uses_house_headroom_without_changing_points(chart, axis_labels):
    data = published()
    expected = chart_groups(data, chart)[0]["points"]
    for theme in CHART_THEMES:
        svg = ElementTree.fromstring(svg_business_chart(data, chart, theme))
        labels = [node.text for node in svg.iter()
                  if node.get("x") == "100" and node.get("text-anchor") == "end"]
        assert labels == axis_labels
        columns = [node for node in svg.iter() if node.get("class") == "annual-column"]
        points = [node for node in svg.iter() if node.get("class") == "business-point"]
        assert [(int(node.get("data-year")), float(node.get("data-reported-value")),
                 node.get("data-status")) for node in points] == [
                     (point["year"], point["value"], point["status"]) for point in expected]
        assert len(columns) == len(expected)
        assert all(float(node.get("y")) + float(node.get("height")) == pytest.approx(420)
                   for node in columns)


def test_finnish_eurostat_refresh_preserves_national_population_and_explicit_yty_break():
    spec = importlib.util.spec_from_file_location("finnish_business_collector", ROOT / "scripts/build_business_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    prior = published()
    prior["observations"] = [row for row in prior["observations"] if row["series_id"] != "eurostat_finland:china"]
    before = copy.deepcopy(prior)

    def get(url):
        asset = dict(url=url, request_url=url, sha256="b" * 64, retrieved_at="2026-10-25T00:00:00+00:00")
        if url == FINLAND_EUROSTAT_METHODS:
            return b"Reviewed Finland country methodology", asset
        dataset = urlsplit(url).path.split("/")[-1]
        codes = {name: values[0] for name, values in parse_qs(urlsplit(url).query).items() if name != "lang"}
        assert codes["geo"] == "FI" and codes["partner"] == "CN_X_HK"
        points = {
            "fats_out2": [(2007, 149), (2008, 133), (2009, 163)],
            "fats_out2_r2": [(2010, 164), (2012, 205), (2013, 194), (2020, 208)],
            "fats_out_part": [(2021, 205), (2023, 185)],
        }[dataset]
        data = {"id": list(codes) + ["time"], "size": [1] * len(codes) + [len(points)],
                "extension": {"id": dataset.upper()}, "updated": "2026-03-31T23:00:00+0200",
                "dimension": {name: {"category": {"index": {code: 0}}} for name, code in codes.items()},
                "value": {str(index): value for index, (_, value) in enumerate(points)}}
        data["dimension"]["time"] = {"category": {"index": {str(year): index for index, (year, _) in enumerate(points)}}}
        return json.dumps(data).encode(), asset

    client = SimpleNamespace(get=get)
    result = module.collect_finland_history(client, prior)
    assert prior == before
    retained = {(row["series_id"], row["year"]): row for row in result["observations"]}
    assert all(retained[row["series_id"], row["year"]] == row for row in prior["observations"])
    rows = {row["year"]: row for row in result["observations"] if row["series_id"] == "eurostat_finland:china"}
    for previous, current in [(2009, 2010), (2012, 2013), (2020, 2021)]:
        assert rows[previous]["comparability"] != rows[current]["comparability"]
    assert rows[2008]["value"] == 133 and ":pilot" not in rows[2007]["comparability"]
    validate_business_data(result)
    assert module.collect_finland_history(client, result) == result


def test_finnish_eurostat_bars_do_not_fill_national_history_or_cross_source_boundaries():
    data = published()
    for chart, expected_years in [("finland-affiliates-china-eurostat", range(2007, 2024)),
                                   ("finland-affiliates-china", range(2020, 2025))]:
        svg = ElementTree.fromstring(svg_business_chart(data, chart))
        bars = [node for node in svg.iter() if node.get("class") == "annual-column"]
        assert [int(node.get("data-year")) for node in bars] == list(expected_years)
        assert not any(node.get("class") == "data-line" for node in svg.iter())
    eurostat_group = chart_groups(data, "finland-affiliates-china-eurostat")[0]
    national_group = chart_groups(data, "finland-affiliates-china")[0]
    assert eurostat_group["id"] != national_group["id"]
    assert eurostat_group["points"][-1]["value"] != national_group["points"][-2]["value"]
    row = next(row for row in data["observations"] if row["series_id"] == "eurostat_finland:china")
    row["source_id"] = "finland_13k8"
    with pytest.raises(ValueError, match="method/source"):
        validate_business_data(data)


def test_swedish_eurostat_refresh_preserves_confidentiality_and_country_specific_breaks():
    spec = importlib.util.spec_from_file_location("swedish_business_collector", ROOT / "scripts/build_business_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    prior = published()
    prior["observations"] = [row for row in prior["observations"] if row["series_id"] != "eurostat_sweden:china"]
    before = copy.deepcopy(prior)

    def get(url):
        asset = dict(url=url, request_url=url, sha256="c" * 64, retrieved_at="2026-10-25T00:00:00+00:00")
        if url == SWEDEN_EUROSTAT_METHODS:
            return b"Reviewed Sweden country methodology", asset
        dataset = urlsplit(url).path.split("/")[-1]
        codes = {name: values[0] for name, values in parse_qs(urlsplit(url).query).items() if name != "lang"}
        assert codes["geo"] == "SE" and codes["partner"] == "CN_X_HK"
        points, flags = {
            "fats_out2": ([(2007, 186), (2008, 320), (2009, 0)], {"2": "|C"}),
            "fats_out2_r2": ([(2010, 407), (2017, 702), (2018, 665), (2020, 679)], {"2": "b"}),
            "fats_out_part": ([(2021, 523), (2023, 602)], {}),
        }[dataset]
        data = {"id": list(codes) + ["time"], "size": [1] * len(codes) + [len(points)],
                "extension": {"id": dataset.upper()}, "updated": "2026-03-31T23:00:00+0200",
                "dimension": {name: {"category": {"index": {code: 0}}} for name, code in codes.items()},
                "value": {str(index): value for index, (_, value) in enumerate(points)}, "status": flags}
        data["dimension"]["time"] = {"category": {"index": {str(year): index for index, (year, _) in enumerate(points)}}}
        return json.dumps(data).encode(), asset

    client = SimpleNamespace(get=get)
    result = validate_business_data(module.collect_sweden_history(client, prior))
    assert prior == before
    retained = {(row["series_id"], row["year"]): row for row in result["observations"]}
    assert all(retained[row["series_id"], row["year"]] == row for row in prior["observations"])
    rows = {row["year"]: row for row in result["observations"] if row["series_id"] == "eurostat_sweden:china"}
    assert rows[2009]["value"] is None and rows[2009]["status"] == "suppressed"
    for previous, current in [(2007, 2008), (2008, 2010), (2017, 2018), (2020, 2021)]:
        assert rows[previous]["comparability"] != rows[current]["comparability"]
    assert module.collect_sweden_history(client, result) == result


def test_swedish_eurostat_bars_leave_confidential_year_empty_and_national_spike_unchanged():
    data = published()
    for theme in CHART_THEMES:
        svg = ElementTree.fromstring(svg_business_chart(data, "sweden-affiliates-china-eurostat", theme))
        bars = [node for node in svg.iter() if node.get("class") == "annual-column"]
        assert [int(node.get("data-year")) for node in bars] == [year for year in range(2007, 2024) if year != 2009]
        assert not any(node.get("class") == "data-line" for node in svg.iter())
    assert next(row["value"] for row in data["observations"]
                if row["series_id"] == "tillvaxtanalys_sweden:china" and row["year"] == 2022) == 1133
    row = next(row for row in data["observations"] if row["series_id"] == "eurostat_sweden:china")
    row["source_id"] = "sweden_country_table"
    with pytest.raises(ValueError, match="method/source"):
        validate_business_data(data)


def test_belgian_refresh_preserves_other_countries_and_rejects_wrong_populations():
    spec = importlib.util.spec_from_file_location("belgian_collector", ROOT / "scripts/build_business_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    from us_china_exchange.business_sources import BELGIUM_EUROSTAT_METHODS

    class Client:
        wrong = None

        def get(self, url):
            asset = dict(url=url, request_url=url, sha256="a" * 64, retrieved_at="2026-10-25T00:00:00+00:00")
            if url == BELGIUM_EUROSTAT_METHODS:
                text = ("As of reference year 2009; From 2022 onwards; Before reference year 2022; "
                        "30 000 000; 75 000 000; 15 000 000")
                if self.wrong == "method":
                    text = text.replace("From 2022 onwards", "From 2021 onwards")
                return text.encode(), asset
            dataset = urlsplit(url).path.split("/")[-1]
            codes = {name: values[0] for name, values in parse_qs(urlsplit(url).query).items() if name != "lang"}
            assert codes["geo"] == "BE" and codes["partner"] == "CN_X_HK"
            first, counts = {
                "fats_out2": (2007, [63, 68, 76]),
                "fats_out2_r2": (2010, [82, 78, 99, 97, 97, 81, 85, 86, 83, 90, 89]),
                "fats_out_part": (2021, [65, 72, 80]),
            }[dataset]
            points = list(zip(range(first, first + len(counts)), counts))
            if self.wrong == "geo":
                codes["geo"] = "NL"
            elif self.wrong == "partner":
                codes["partner"] = "HK"
            elif self.wrong == "metric":
                codes["indic_sbs" if dataset == "fats_out_part" else "indic_bp"] = "EMP"
            elif self.wrong == "year-gap" and dataset == "fats_out2_r2":
                points = [point for point in points if point[0] != 2016]
            data = {"id": list(codes) + ["time"], "size": [1] * len(codes) + [len(points)],
                    "extension": {"id": dataset.upper()}, "updated": "2026-03-31T23:00:00+0200",
                    "dimension": {name: {"category": {"index": {code: 0}}} for name, code in codes.items()},
                    "value": {str(index): value for index, (_, value) in enumerate(points)}}
            data["dimension"]["time"] = {"category": {
                "index": {str(year): index for index, (year, _) in enumerate(points)}}}
            if self.wrong == "new-break" and dataset == "fats_out_part":
                data["status"] = {"2": "b"}
            return json.dumps(data).encode(), asset

    prior = published()
    before = copy.deepcopy(prior)
    result = validate_business_data(module.collect_belgium(Client(), prior))
    assert prior == before
    assert [row for row in result["observations"] if not row["series_id"].startswith("eurostat_belgium")] == [
        row for row in prior["observations"] if not row["series_id"].startswith("eurostat_belgium")]
    for key in ("series", "sources"):
        for name, value in prior[key].items():
            if not name.startswith("eurostat_belgium") and not name.endswith("_BE"):
                assert result[key][name] == value
    rows = {row["year"]: row for row in result["observations"] if row["series_id"].startswith("eurostat_belgium")}
    assert rows[2008]["series_id"] == "eurostat_belgium_preuci:china"
    assert rows[2009]["series_id"] == rows[2021]["series_id"] == "eurostat_belgium_legal:china"
    assert rows[2022]["series_id"] == rows[2023]["series_id"] == "eurostat_belgium:china"
    assert (rows[2022]["value"], rows[2023]["value"]) == (72, 80)
    for previous, current in ((2009, 2010), (2013, 2014), (2015, 2016), (2020, 2021)):
        assert rows[previous]["comparability"] != rows[current]["comparability"]
    assert module.collect_belgium(Client(), result) == result
    for failure in ("geo", "partner", "metric", "year-gap", "method", "new-break"):
        client = Client()
        client.wrong = failure
        with pytest.raises(ValueError):
            module.collect_belgium(client, prior)


def test_hungarian_refresh_preserves_history_confidentiality_and_producer_unit_rules(monkeypatch):
    spec = importlib.util.spec_from_file_location("hungarian_collector", ROOT / "scripts/build_business_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    from us_china_exchange.business_sources import HUNGARY_CONTEXT_NOTE, HUNGARY_EUROSTAT_METHODS, HUNGARY_OUTWARD_METHODS, HUNGARY_OUTWARD_TABLE

    class Client:
        wrong = None

        def get(self, url):
            asset = dict(url=url, request_url=url, sha256="a" * 64, retrieved_at="2026-10-25T00:00:00+00:00")
            if url == HUNGARY_EUROSTAT_METHODS:
                text = ("ultimate controlling institutional unit; 2007-2009, 2010-2023; "
                        "40.000 EUR; 10 percent; HUF 10 million; "
                        "minority share provides enough control; no equally shared controlled enterprises; "
                        "non-responding enterprises were not included")
                if self.wrong == "window":
                    text = text.replace("2010-2023", "2010-2024")
                if self.wrong == "frame":
                    text = text.replace("40.000 EUR", "40.001 EUR")
                return text.encode(), asset
            if url == HUNGARY_OUTWARD_METHODS:
                text = ("legal unit level until reference year 2021; "
                        "from reference year 2022 enterprise level; "
                        "From reference year 2021; 2019/2152; "
                        "reference population covers the market producers as defined by ESA 2010")
                if self.wrong == "unit":
                    text = text.replace("reference year 2022", "reference year 2018")
                if self.wrong == "ebs":
                    text = text.replace("reference year 2021", "reference year 2020")
                return text.encode(), asset
            if url == HUNGARY_OUTWARD_TABLE:
                text = ("Business statistics on domestically controlled foreign affiliates. "
                        "From reference year 2022 the statistical unit is enterprise")
                if self.wrong == "table":
                    text = text.replace("2022", "2021")
                return text.encode(), asset
            dataset = urlsplit(url).path.split("/")[-1]
            codes = {name: values[0] for name, values in parse_qs(urlsplit(url).query).items() if name != "lang"}
            assert codes["geo"] == "HU" and codes["partner"] == "CN_X_HK"
            first, last = {"fats_out2": (2007, 2009), "fats_out2_r2": (2010, 2020),
                           "fats_out_part": (2021, 2023)}[dataset]
            counts = {2010: 4, 2018: 3, 2019: 4, 2020: 8, 2021: 9, 2022: 6, 2023: 6}
            years = [year for year in range(first, last + 1)
                     if not (self.wrong == "year-gap" and year == 2016)]
            if self.wrong == "geo":
                codes["geo"] = "AT"
            if self.wrong == "partner":
                codes["partner"] = "HK"
            if self.wrong == "metric":
                codes["indic_sbs" if dataset == "fats_out_part" else "indic_bp"] = "EMP"
            data = {"id": list(codes) + ["time"], "size": [1] * len(codes) + [len(years)],
                    "extension": {"id": dataset.upper()}, "updated": "2026-03-31T23:00:00+0200",
                    "dimension": {name: {"category": {"index": {code: 0}}} for name, code in codes.items()},
                    "value": {str(index): counts.get(year, 0) for index, year in enumerate(years)},
                    "status": {str(index): "|C" for index, year in enumerate(years) if year not in counts}}
            data["dimension"]["time"] = {"category": {"index": {
                str(year): index for index, year in enumerate(years)}}}
            if self.wrong == "new-break" and dataset == "fats_out2_r2":
                data["status"][str(years.index(2018))] = "b"
            return json.dumps(data).encode(), asset

    prior = published()
    prior["series"] = {sid: definition for sid, definition in prior["series"].items()
                       if not sid.startswith("eurostat_hungary")}
    prior["observations"] = [r for r in prior["observations"] if not r["series_id"].startswith("eurostat_hungary")]
    prior["sources"] = {sid: source for sid, source in prior["sources"].items() if not sid.endswith("_HU")}
    prior["coverage"] = [r for r in prior["coverage"] if not r["id"].startswith("eurostat_hungary")]
    prior["context"]["geography"] = prior["context"]["geography"].replace(HUNGARY_CONTEXT_NOTE, "")
    before = copy.deepcopy(prior)
    result = validate_business_data(module.collect_hungary(Client(), prior))
    assert prior == before
    assert [r for r in result["observations"] if not r["series_id"].startswith("eurostat_hungary")] == [
        r for r in prior["observations"] if not r["series_id"].startswith("eurostat_hungary")]
    for key in ("series", "sources"):
        for name, value in prior[key].items():
            if not name.startswith("eurostat_hungary") and not name.endswith("_HU"):
                assert result[key][name] == value
    points = {r["year"]: r for r in result["observations"] if r["series_id"].startswith("eurostat_hungary")}
    assert sorted(points) == list(range(2007, 2024))
    assert [(year, r["value"]) for year, r in sorted(points.items()) if r["value"] is not None] == [
        (2010, 4), (2018, 3), (2019, 4), (2020, 8), (2021, 9), (2022, 6), (2023, 6)]
    assert sum(r["status"] == "suppressed" and r["value"] is None for r in points.values()) == 10
    assert all(r["series_id"] == ("eurostat_hungary_legal:china" if year < 2022 else "eurostat_hungary:china")
               for year, r in points.items())
    assert len({points[year]["comparability"] for year in (2009, 2010, 2021, 2022)}) == 4
    assert module.collect_hungary(Client(), result) == result
    monkeypatch.setattr(module, "collect_business", lambda client: copy.deepcopy(result))
    assert module.collect_all(None, prior) == result
    for failure in ("unit", "ebs", "table", "window", "frame", "geo", "partner", "metric", "year-gap", "new-break"):
        client = Client()
        client.wrong = failure
        with pytest.raises(ValueError):
            module.collect_hungary(client, prior)


def test_hungarian_current_bars_do_not_splice_legal_units_or_fill_confidential_counts():
    data = published()
    for theme in CHART_THEMES:
        svg = ElementTree.fromstring(svg_business_chart(data, "hungary-affiliates-china", theme))
        bars = [node for node in svg.iter() if node.get("class") == "annual-column"]
        assert [(int(node.get("data-year")), int(node.get("data-value"))) for node in bars] == [(2022, 6), (2023, 6)]
        assert not any(node.get("class") in {"data-line", "covid-band"} for node in svg.iter())
        assert max(float(node.get("y")) for node in svg.iter()
                   if node.tag.endswith("}text") and node.get("y") is not None) <= 590
    assert not any("eurostat_hungary_legal:china" in chart["series"] for chart in CHARTS.values())
    for year, wrong_sid in ((2004, "eurostat_hungary_legal:china"), (2021, "eurostat_hungary:china"),
                            (2022, "eurostat_hungary_legal:china")):
        wrong = copy.deepcopy(data)
        point = next(r for r in wrong["observations"]
                     if r["series_id"].startswith("eurostat_hungary") and r["year"] == (2007 if year == 2004 else year))
        point.update(year=year, series_id=wrong_sid)
        with pytest.raises(ValueError, match="must not be spliced"):
            validate_business_data(wrong)
    wrong = copy.deepcopy(data)
    wrong["series"]["eurostat_hungary_legal:china"].update(unit="enterprises", measure="controlled_enterprises")
    with pytest.raises(ValueError, match="counting units"):
        validate_business_data(wrong)
    wrong = copy.deepcopy(data)
    wrong["series"]["eurostat_hungary:china"]["scope"] = "Hungary-resident UCI; China except Hong Kong; all firms"
    with pytest.raises(ValueError, match="reporting-frame"):
        validate_business_data(wrong)
    wrong = copy.deepcopy(data)
    point = next(r for r in wrong["observations"]
                 if r["series_id"].startswith("eurostat_hungary") and r["status"] == "suppressed")
    point.update(value=0, status="observed")
    with pytest.raises(ValueError, match="confidential"):
        validate_business_data(wrong)


def test_belgian_current_bars_exclude_archives_and_reject_cross_unit_relabelling():
    data = published()
    for theme in CHART_THEMES:
        svg = ElementTree.fromstring(svg_business_chart(data, "belgium-affiliates-china", theme))
        bars = [node for node in svg.iter() if node.get("class") == "annual-column"]
        assert [(int(node.get("data-year")), int(node.get("data-value"))) for node in bars] == [(2022, 72), (2023, 80)]
        assert not any(node.get("class") in {"data-line", "covid-band"} for node in svg.iter())
    assert not any(sid in {"eurostat_belgium_preuci:china", "eurostat_belgium_legal:china"}
                   for chart in CHARTS.values() for sid in chart["series"])
    for year, wrong_sid in ((2008, "eurostat_belgium_legal:china"),
                            (2021, "eurostat_belgium:china"), (2022, "eurostat_belgium_legal:china")):
        wrong = copy.deepcopy(data)
        next(row for row in wrong["observations"]
             if row["series_id"].startswith("eurostat_belgium") and row["year"] == year)["series_id"] = wrong_sid
        with pytest.raises(ValueError, match="must not be spliced"):
            validate_business_data(wrong)
    wrong = copy.deepcopy(data)
    wrong["series"]["eurostat_belgium_preuci:china"]["scope"] = (
        "Belgium-resident UCI; China except Hong Kong; thresholded FDI reporting-unit frame")
    with pytest.raises(ValueError, match="pre-UCI"):
        validate_business_data(wrong)
    wrong = copy.deepcopy(data)
    wrong["series"]["eurostat_belgium_legal:china"]["unit"] = "enterprises"
    wrong["series"]["eurostat_belgium_legal:china"]["measure"] = "controlled_enterprises"
    with pytest.raises(ValueError, match="counting units"):
        validate_business_data(wrong)
    wrong = copy.deepcopy(data)
    next(row for row in wrong["observations"] if row["series_id"] == "eurostat_belgium:china")["source_id"] = "fats_out_part_AT"
    with pytest.raises(ValueError, match="method/source"):
        validate_business_data(wrong)


def test_snb_public_document_transport_uses_cache_or_a_successful_javascript_shell_only(tmp_path, monkeypatch):
    from contextlib import nullcontext
    from subprocess import CalledProcessError
    from playwright import sync_api
    from us_china_exchange.language_sources import ReviewClient

    url, marker = "https://data.snb.ch/de/topics/aube/doc/explanations_aube_amne", "3.3. Anzahl Unternehmen"
    rendered = f"<html><body><h2>{marker}</h2>Reviewed method</body></html>".encode()
    shell = b"<html><body><app-root></app-root></body></html>"
    client = ReviewClient(tmp_path)
    asset = dict(url=url, request_url=url, sha256=hashlib.sha256(shell).hexdigest(),
                 retrieved_at="2026-10-25T00:00:00+00:00")
    state = {"launches": 0, "closed": 0, "status": 200}

    def launch():
        state["launches"] += 1
        return browser

    page = SimpleNamespace(
        url=url, goto=lambda *args, **kwargs: SimpleNamespace(status=state["status"]),
        content=lambda: rendered.decode(),
        get_by_text=lambda *args, **kwargs: SimpleNamespace(first=SimpleNamespace(wait_for=lambda **kwargs: None)))
    browser = SimpleNamespace(new_page=lambda: page, close=lambda: state.update(closed=state["closed"] + 1))
    monkeypatch.setattr(sync_api, "sync_playwright", lambda: nullcontext(
        SimpleNamespace(chromium=SimpleNamespace(launch=launch))))
    monkeypatch.setattr(client, "get", lambda _: (rendered, asset))
    assert client.get_public_document(url, marker) == (rendered, asset)
    assert state["launches"] == 0
    monkeypatch.setattr(client, "get", lambda _: (shell, asset))
    body, provenance = client.get_public_document(url, marker)
    assert body == rendered and provenance["sha256"] == hashlib.sha256(rendered).hexdigest()
    assert provenance["request_url"] == url and state["closed"] == 1
    monkeypatch.setattr(client, "get", lambda _: (b"<html>Changed method</html>", asset))
    with pytest.raises(ValueError, match="content changed"):
        client.get_public_document(url, marker)
    assert state["launches"] == 1
    monkeypatch.setattr(client, "get", lambda _: (shell, asset))
    state["status"] = 403
    with pytest.raises(ValueError, match="access failed"):
        client.get_public_document(url, marker)
    assert state["closed"] == 2

    def denied(_):
        raise CalledProcessError(22, ["curl", url])

    monkeypatch.setattr(client, "get", denied)
    with pytest.raises(CalledProcessError):
        client.get_public_document(url, marker)
    assert state["launches"] == 2


def snb_native_fixture(data):
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, delimiter=";", lineterminator="\n")
    writer.writerows([["CubeId", "opanmuausuland"], ["PublishingDate", "2025-12-12 09:00"],
                      [], ["Date", "d0", "Value"]])
    for series in data["timeseries"]:
        code = series["metadata"]["key"].split("{")[1][:-1]
        writer.writerows([[point["date"], code, point["value"]] for point in series["values"]])
    return stream.getvalue().encode()


def slovenia_fixture(overrides=None):
    from us_china_exchange import business_sources as module

    overrides = overrides or {}
    counts = dict(zip(range(2009, 2026), (7, 5, 6, 7, 7, 7, 7, 7, 8, 20, 12, 12, 11, 13, 13, 9, 11)))
    countries = ["Total world (W1)", "China (CN)"] + [f"Other country {i}" for i in range(55)]
    quoted = lambda values: ",".join(json.dumps(value) for value in values)
    header = (
        'CODEPAGE="utf-8";\nMATRIX="6";\nSUBJECT-AREA="EOT_OFATS";\nUNITS="number";\n'
        'TITLE="Foreign companies ultimately controled by Slovenian companies";\n'
        'STUB="Date","Breakdown","Code of breakdown";\nHEADING="Items";\n'
        'LAST-UPDATED="20260907 09:42";\n'
        'NOTEX="reporters interest exceeds 50%; In 2017 reporting threshold EUR 2 million; '
        'in 2023 reporting threshold was set at EUR 5 million";\n')
    for name, values in (("Date", [str(year) for year in counts]), ("Breakdown", ["Countries"]),
                         ("Code of breakdown", countries),
                         ("Items", ["Number of enterprises*", "Number of persons employed"])):
        header += f'VALUES("{name}")={quoted(values)};\n'
    cells = []
    for year in counts:
        for country in countries:
            cells.append(str(overrides.get(year, counts[year])) + " 123" if country == "China (CN)" else "1 2.25")
    native = (header + "DATA=\n" + "\n".join(cells) + ";\n").encode()
    bodies = {module.SLOVENIA_COUNTS: native, module.SLOVENIA_METHODS: b"Reviewed national methods",
              module.SLOVENIA_REPORT: b"%PDF Reviewed current report"}
    hashes = {"national": hashlib.sha256(native).hexdigest(),
              "methods": hashlib.sha256(bodies[module.SLOVENIA_METHODS]).hexdigest(),
              "report": hashlib.sha256(bodies[module.SLOVENIA_REPORT]).hexdigest()}
    for dataset, years, activity, indicator in (
            ("fats_out2", range(2007, 2010), ("nace_r1", "C-O_X_L"), ("indic_bp", "ENT")),
            ("fats_out2_r2", range(2010, 2021), ("nace_r2", "B-S_X_O"), ("indic_bp", "ENT")),
            ("fats_out_part", range(2021, 2024), ("nace_r2", "B-S_X_O"), ("indic_sbs", "ENT_NR"))):
        codes = {"freq": "A", activity[0]: activity[1], indicator[0]: indicator[1], "partner": "CN_X_HK", "geo": "SI"}
        values = [4 if year == 2007 else 6 if year == 2008 else None if year == 2009 else counts[year] for year in years]
        flags = {str(i): "|C" if year == 2009 else "b" for i, year in enumerate(years) if year in {2009, 2017}}
        data = {"id": list(codes) + ["time"], "size": [1] * 5 + [len(years)],
                "dimension": {name: {"category": {"index": {code: 0}}} for name, code in codes.items()},
                "extension": {"id": dataset}, "value": {str(i): value for i, value in enumerate(values) if value is not None},
                "status": flags, "updated": "2026-03-31"}
        data["dimension"]["time"] = {"category": {"index": {str(year): i for i, year in enumerate(years)}}}
        url = module.EUROSTAT + dataset + "?lang=EN&freq=A&geo=SI&partner=CN_X_HK&" + "=".join(indicator) + "&" + "=".join(activity)
        body = json.dumps(data).encode()
        bodies[url], hashes[dataset] = body, hashlib.sha256(body).hexdigest()
        key = ".".join(codes.values())
        native_url = "https://ec.europa.eu/eurostat/api/dissemination/sdmx/2.1/data/" + dataset + "/" + key + "?format=TSV&compressed=false"
        cells = [": @C" if value is None else str(value) + (" b" if year == 2017 else "") for year, value in zip(years, values)]
        tsv = (",".join(codes) + "\\TIME_PERIOD\t" + "\t".join(map(str, years)) + "\n"
               + ",".join(codes.values()) + "\t" + "\t".join(cells) + "\n").encode()
        bodies[native_url], hashes[dataset + "-native"] = tsv, hashlib.sha256(tsv).hexdigest()
    return bodies, hashes


def test_slovenia_national_parser_preserves_full_query_precision_and_masks_without_unit_substitution():
    from us_china_exchange import business_sources as module

    bodies, _ = slovenia_fixture()
    body = bodies[module.SLOVENIA_COUNTS]
    points = module.parse_slovenia_national(body)
    assert [point[0] for point in points] == list(range(2009, 2026))
    assert [point[1] for point in points][-2:] == [9, 11]
    assert points[2019 - 2009][4] != points[2025 - 2009][4]
    masked, _ = slovenia_fixture({2009: '"..."'})
    assert module.parse_slovenia_national(masked[module.SLOVENIA_COUNTS])[0][1:4] == (None, "missing", '"..."')
    for old, new in ((b'UNITS="number"', b'UNITS="EUR"'),
                     (b'"China (CN)"', b'"Hong Kong (HK)"'),
                     (b'"Number of enterprises*"', b'"Employees"'),
                     (b'"2025"', b'"2026"'), (b"EUR 5 million", b"EUR 2 million")):
        with pytest.raises(ValueError):
            module.parse_slovenia_national(body.replace(old, new))
    for count in (-1, 1.5):
        changed, _ = slovenia_fixture({2025: count})
        with pytest.raises(ValueError, match="Slovenian"):
            module.parse_slovenia_national(changed[module.SLOVENIA_COUNTS])


def test_slovenia_refresh_preserves_other_populations_and_rejects_changed_originals(monkeypatch):
    from us_china_exchange import business_sources as sources

    spec = importlib.util.spec_from_file_location("slovenia_collector", ROOT / "scripts/build_business_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    bodies, hashes = slovenia_fixture()
    monkeypatch.setattr(sources, "SLOVENIA_REVIEWED_ASSETS", hashes)

    class Client:
        def get(self, url):
            body = bodies[url]
            return body, dict(url=url, request_url=url, sha256=hashlib.sha256(body).hexdigest(),
                              retrieved_at="2026-10-25T00:00:00+00:00")

    prior, before = published(), copy.deepcopy(published())
    result = validate_business_data(module.collect_slovenia(Client(), prior))
    assert prior == before
    selected = {"bsi_slovenia:china", "eurostat_slovenia:china"}
    assert [row for row in result["observations"] if row["series_id"] not in selected] == [
        row for row in prior["observations"] if row["series_id"] not in selected]
    assert len([row for row in result["observations"] if row["series_id"] in selected]) == 34
    assert next(row for row in result["observations"] if row["series_id"] == "eurostat_slovenia:china" and row["year"] == 2009)["value"] is None
    assert module.collect_slovenia(Client(), result) == result
    bodies[sources.SLOVENIA_COUNTS] += b"\nChanged original"
    with pytest.raises(ValueError, match="source asset"):
        module.collect_slovenia(Client(), prior)


def test_slovenia_chart_keeps_national_thresholds_and_archive_confidentiality_separate():
    data = validate_business_data(published())
    groups = chart_groups(data, "slovenia-affiliates-china")
    assert len(groups) == 1 and groups[0]["id"] == "bsi_slovenia:china"
    assert [point["year"] for point in groups[0]["points"]] == list(range(2009, 2026))
    assert [point["value"] for point in groups[0]["points"]][-2:] == [9, 11]
    assert all(point["denominator"] is None for point in groups[0]["points"])
    for theme in CHART_THEMES:
        svg = ElementTree.fromstring(svg_business_chart(data, "slovenia-affiliates-china", theme))
        bars = [node for node in svg.iter() if node.tag.endswith("rect") and node.get("data-year")]
        assert [int(node.get("data-year")) for node in bars] == list(range(2009, 2026))
    wrong = copy.deepcopy(data)
    point = next(row for row in wrong["observations"] if row["series_id"] == "bsi_slovenia:china" and row["year"] == 2025)
    point["comparability"] = "bsi:SN11:CN:EUR2m2017:levels-only"
    with pytest.raises(ValueError, match="threshold"):
        validate_business_data(wrong)
    wrong = copy.deepcopy(data)
    point = next(row for row in wrong["observations"] if row["series_id"] == "eurostat_slovenia:china" and row["year"] == 2009)
    point.update(value=0, status="observed")
    with pytest.raises(ValueError, match="confidentiality"):
        validate_business_data(wrong)


def snb_fixture():
    from us_china_exchange import business_sources as module

    nodes = [{"id": "CN", "name": "China"}, {"id": "HK", "name": "Hongkong"}]
    extras = [{"id": f"ZZ{i:02}", "name": f"Other {i}"} for i in range(90)]
    metadata = {"cubeId": module.SNB_CUBE, "dimensions": [{"id": "d0", "dimensionItems": [
        {"id": "root", "name": "Alle Länder", "dimensionItems": [
            {"id": "asia", "name": "Asien", "dimensionItems": nodes}, *extras]}]}]}
    data = {"timeseries": []}
    counts = [543, 598, 622, 639, 664, 716, 724, 749, 799, 816, 792]
    for node in nodes + extras:
        code = node["id"]
        label = "Alle Länder - " + ("Asien - " if code in {"CN", "HK"} else "") + node["name"]
        data["timeseries"].append({
            "header": [{"dim": "Länder/Ländergruppe", "dimItem": label}],
            "metadata": {"key": f"EPB@SNB.{module.SNB_CUBE}{{{code}}}", "frequency": "P1Y", "scale": ""},
            "values": [{"date": str(year), "value": counts[index] if code == "CN" else
                        7.5 if code == "ZZ00" else 131.00001 if code == "ZZ01" else 10}
                       for index, year in enumerate(range(2014, 2025))]})
    bodies = {
        module.SNB_API + "dimensions/de": json.dumps(metadata).encode(),
        module.SNB_API + "data/json/de": json.dumps(data).encode(),
        module.SNB_API + "data/csv/de": snb_native_fixture(data),
        module.SNB_API + "lastUpdate": json.dumps({
            "editionDate": "20251212_0900", "publicSinceDate": "20251212_0900"}).encode(),
        module.SNB_METHODS: (
            "Die AMNE/FATS erfassen nur Mehrheitsbeteiligungen; Beteiligungsverhältnisse grösser als 50%; "
            "nur Daten schweizerisch beherrschter Unternehmen; in der Schweiz und dem Fürstentum Liechtenstein; "
            "gruppenweit 10 Mio. Franken übersteigt; Die Anzahl Unternehmen umfasst alle direkten und "
            "indirekten Mehrheitsbeteiligungen inkl. rechtlich unselbständiger Filialen und Zweigniederlassungen").encode(),
        module.SNB_CONTROL_METHODS: (
            "Unternehmen gelten dann als schweizerisch beherrscht; letztlich Berechtigte im Inland").encode(),
        module.SNB_COUNTRY_METHODS: (
            "Liechtenstein, das statistisch zum Inland gehört; Geonomenklatur von Eurostat; "
            "China CN Hongkong HK Macau MO").encode(),
    }
    return metadata, data, bodies


def test_snb_refresh_preserves_old_populations_and_rejects_changed_source_contracts():
    from us_china_exchange import business_sources as sources

    spec = importlib.util.spec_from_file_location("snb_collector", ROOT / "scripts/build_business_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    _, _, bodies = snb_fixture()

    class Client:
        wrong_request = False

        def __init__(self, payloads):
            self.bodies = payloads

        def get(self, url):
            body = self.bodies[url]
            return body, dict(url=url, request_url="https://example.com/other" if self.wrong_request else url,
                              sha256=hashlib.sha256(body).hexdigest(), retrieved_at="2026-10-25T00:00:00+00:00")

    prior, before = published(), copy.deepcopy(published())
    result = validate_business_data(module.collect_swiss_reporting_area(Client(bodies), prior))
    assert prior == before
    assert [row for row in result["observations"] if row["series_id"] != sources.SNB_SERIES] == [
        row for row in prior["observations"] if row["series_id"] != sources.SNB_SERIES]
    for field in ("series", "sources"):
        for key, value in prior[field].items():
            if key not in {sources.SNB_SERIES, "snb_opanmuausuland_20251212"}:
                assert result[field][key] == value
    selected = [row for row in result["observations"] if row["series_id"] == sources.SNB_SERIES]
    assert [(row["year"], row["value"]) for row in selected] == list(zip(
        range(2014, 2025), [543, 598, 622, 639, 664, 716, 724, 749, 799, 816, 792]))
    assert {row["status"] for row in selected} == {"reported"}
    assert module.collect_swiss_reporting_area(Client(bodies), result) == result
    wrong = Client(bodies)
    wrong.wrong_request = True
    with pytest.raises(ValueError, match="source asset"):
        module.collect_swiss_reporting_area(wrong, prior)
    for url, old, new in (
            (sources.SNB_METHODS, "Fürstentum Liechtenstein", "Ausland"),
            (sources.SNB_METHODS, "10 Mio.", "1 Mio."),
            (sources.SNB_METHODS, "grösser als 50%", "ab 10%"),
            (sources.SNB_COUNTRY_METHODS, "China CN", "China HK"),
            (sources.SNB_API + "lastUpdate", "20251212_0900", "20261211_0900")):
        changed = dict(bodies)
        changed[url] = changed[url].replace(old.encode(), new.encode())
        with pytest.raises(ValueError):
            module.collect_swiss_reporting_area(Client(changed), prior)
    assert prior == before


def test_snb_native_parser_preserves_other_country_precision_and_never_zero_fills_china():
    from us_china_exchange.business_sources import parse_snb_counts

    metadata, data, _ = snb_fixture()
    native = snb_native_fixture(data)
    expected = parse_snb_counts(metadata, data, native)
    assert expected[0] == (2014, 543) and expected[-1] == (2024, 792)
    changed = copy.deepcopy(data)
    changed["timeseries"][2]["values"][0]["value"] = None
    assert parse_snb_counts(metadata, changed, snb_native_fixture(changed)) == expected
    for field, value in (("key", "EPB@SNB.opanmuauspland{CN}"), ("frequency", "P1M"), ("scale", "1000")):
        changed = copy.deepcopy(data)
        changed["timeseries"][0]["metadata"][field] = value
        with pytest.raises(ValueError):
            parse_snb_counts(metadata, changed, native)
    for value in (None, 543.5, -1, True):
        changed = copy.deepcopy(data)
        changed["timeseries"][0]["values"][0]["value"] = value
        with pytest.raises(ValueError):
            parse_snb_counts(metadata, changed, snb_native_fixture(changed))
    changed = copy.deepcopy(data)
    changed["timeseries"][2]["values"][0]["value"] = 8
    with pytest.raises(ValueError, match="without rounding"):
        parse_snb_counts(metadata, changed, native)
    changed = copy.deepcopy(data)
    changed["timeseries"][0]["values"][0]["date"] = "2025"
    with pytest.raises(ValueError, match="measurement years"):
        parse_snb_counts(metadata, changed, native)
    changed = copy.deepcopy(data)
    changed["timeseries"][0]["header"][0]["dimItem"] = "Alle Länder - Asien - Hongkong"
    with pytest.raises(ValueError, match="country labels"):
        parse_snb_counts(metadata, changed, native)


def test_snb_combined_frame_count_units_and_reporting_status_survive_both_themes():
    sid, chart = "snb_swiss_liechtenstein:china", "swiss-liechtenstein-affiliates-china"
    data = validate_business_data(published())
    for theme in CHART_THEMES:
        svg = ElementTree.fromstring(svg_business_chart(data, chart, theme))
        bars = [node for node in svg.iter() if node.get("class") == "annual-column"]
        assert [int(node.get("data-year")) for node in bars] == list(range(2014, 2025))
        assert (int(bars[0].get("data-value")), int(bars[-1].get("data-value"))) == (543, 792)
        points = [node for node in svg.iter() if node.get("class") == "business-point"]
        assert len(points) == 11 and {node.get("data-status") for node in points} == {"reported"}
        text = " ".join(svg.itertext())
        assert "Switzerland + Liechtenstein" in text and "Affiliates and branches" in text
        assert not any(node.get("class") in {"data-line", "covid-band"} for node in svg.iter())
    for field, value in (("scope", "Switzerland; all enterprises"), ("label", "Switzerland"),
                          ("unit", "enterprises"), ("measure", "controlled_enterprises"),
                          ("period_basis", "reference year")):
        wrong = copy.deepcopy(data)
        wrong["series"][sid][field] = value
        with pytest.raises(ValueError):
            validate_business_data(wrong)
    for field, value in (("year", 2013), ("year", 2025), ("status", "observed"),
                          ("status", "provisional"), ("source_id", "fats_out_part")):
        wrong = copy.deepcopy(data)
        next(row for row in wrong["observations"] if row["series_id"] == sid)[field] = value
        with pytest.raises(ValueError):
            validate_business_data(wrong)
    wrong = copy.deepcopy(data)
    wrong["sources"]["snb_opanmuausuland_20251212"]["assets"][0]["url"] = (
        "https://data.snb.ch/api/cube/opanmuauspland/dimensions/de")
    with pytest.raises(ValueError, match="source provenance"):
        validate_business_data(wrong)


def test_polish_selective_refresh_preserves_older_countries_and_the_2018_unit_break():
    spec = importlib.util.spec_from_file_location("polish_business_collector", ROOT / "scripts/build_business_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    prior = published()
    before = copy.deepcopy(prior)

    class Client:
        wrong = None

        def get(self, url):
            asset = dict(url=url, request_url=url, sha256="d" * 64, retrieved_at="2026-10-25T00:00:00+00:00")
            if url == POLAND_EUROSTAT_METHODS:
                return b"Reviewed Polish country methodology", asset
            dataset = urlsplit(url).path.split("/")[-1]
            codes = {name: values[0] for name, values in parse_qs(urlsplit(url).query).items() if name != "lang"}
            assert codes["geo"] == "PL" and codes["partner"] == "CN_X_HK"
            first, values, flags = {
                "fats_out2": (2007, [None, 13, 20], {}),
                "fats_out2_r2": (2010, [22, 25, 33, 24, 28, 34, 35, 32, 31, 29, 35], {"8": "b"}),
                "fats_out_part": (2021, [42, 42, 32], {}),
            }[dataset]
            points = list(zip(range(first, first + len(values)), values))
            if self.wrong == "geo":
                codes["geo"] = "DE"
            elif self.wrong == "metric":
                codes["indic_sbs" if dataset == "fats_out_part" else "indic_bp"] = "EMP"
            elif self.wrong == "missing-year" and dataset == "fats_out2_r2":
                points = [point for point in points if point[0] != 2015]
            elif self.wrong == "new-break" and dataset == "fats_out_part":
                flags = {"2": "b"}
            data = {"id": list(codes) + ["time"], "size": [1] * len(codes) + [len(points)],
                    "extension": {"id": dataset.upper()}, "updated": "2026-03-31T23:00:00+0200",
                    "dimension": {name: {"category": {"index": {code: 0}}} for name, code in codes.items()},
                    "value": {str(index): value for index, (_, value) in enumerate(points) if value is not None},
                    "status": flags}
            data["dimension"]["time"] = {"category": {
                "index": {str(year): index for index, (year, _) in enumerate(points)}}}
            return json.dumps(data).encode(), asset

    result = validate_business_data(module.collect_poland(Client(), prior))
    assert prior == before
    assert [row for row in result["observations"] if row["series_id"] != "eurostat_poland:china"] == [
        row for row in prior["observations"] if row["series_id"] != "eurostat_poland:china"]
    for sid, definition in prior["series"].items():
        if sid != "eurostat_poland:china":
            assert result["series"][sid] == definition
    rows = {row["year"]: row for row in result["observations"] if row["series_id"] == "eurostat_poland:china"}
    assert rows[2007]["value"] is None and rows[2007]["status"] == "missing"
    assert rows[2017]["comparability"] != rows[2018]["comparability"]
    assert rows[2020]["comparability"] == rows[2021]["comparability"]
    assert rows[2020]["source_id"] != rows[2021]["source_id"]
    assert "Source flags: b" in rows[2018]["notes"]
    assert module.collect_poland(Client(), result) == result
    for failure in ("geo", "metric", "missing-year", "new-break"):
        invalid = Client()
        invalid.wrong = failure
        with pytest.raises(ValueError):
            module.collect_poland(invalid, prior)


def test_polish_bars_preserve_missingness_and_reject_the_broader_equity_population():
    data = published()
    for theme in CHART_THEMES:
        svg = ElementTree.fromstring(svg_business_chart(data, "poland-affiliates-china", theme))
        bars = [node for node in svg.iter() if node.get("class") == "annual-column"]
        assert [int(node.get("data-year")) for node in bars] == list(range(2008, 2024))
        assert not any(node.get("class") == "data-line" for node in svg.iter())
    wrong = copy.deepcopy(data)
    wrong["series"]["eurostat_poland:china"]["scope"] = "All Polish foreign equity interests including Hong Kong"
    with pytest.raises(ValueError, match="Polish ultimate-control"):
        validate_business_data(wrong)
    wrong = copy.deepcopy(data)
    next(row for row in wrong["observations"] if row["series_id"] == "eurostat_poland:china")["source_id"] = "fats_out2_ES"
    with pytest.raises(ValueError, match="method/source"):
        validate_business_data(wrong)


@pytest.mark.parametrize("geo,collector,prefix", [
    ("PT", "collect_portugal", "eurostat_portugal"),
    ("CZ", "collect_czechia", "eurostat_czechia:"),
])
def test_portuguese_czech_refresh_preserves_countries_confidentiality_and_control_breaks(geo, collector, prefix):
    from us_china_exchange.business_sources import CZECHIA_EUROSTAT_METHODS, PORTUGAL_EUROSTAT_METHODS

    spec = importlib.util.spec_from_file_location("pt_cz_business_collector", ROOT / "scripts/build_business_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    prior, before = published(), published()

    class Client:
        wrong = None

        def get(self, url):
            asset = dict(url=url, request_url=url, sha256="d" * 64, retrieved_at="2026-10-25T00:00:00+00:00")
            if url in {PORTUGAL_EUROSTAT_METHODS, CZECHIA_EUROSTAT_METHODS}:
                text = ("Banco de Portugal; UCI approach first implemented for 2008 data; "
                        "In 2014, the source data for OFATS compilation was changed; "
                        "In 2020, some companies were added" if geo == "PT" else
                        "Czech National Bank; in 2007 due to the inclusion of the UCI principle; "
                        "Starting from 2008 NACE Rev. 2 was implemented; In reference year 2021 a new Regulation")
                text += "; the statistical unit is the enterprise as defined in line with the Regulation"
                if self.wrong == "unit":
                    text = text.replace("is the enterprise as defined", "is the legal unit as defined")
                return (b"Unreviewed methodology" if self.wrong == "methodology" else text.encode()), asset
            dataset = urlsplit(url).path.split("/")[-1]
            codes = {name: values[0] for name, values in parse_qs(urlsplit(url).query).items() if name != "lang"}
            values = {
                "PT": {"fats_out2": [3, 3, 4],
                       "fats_out2_r2": [4, 3, None, 6, 7, 8, 11, 10, 11, 16, 11],
                       "fats_out_part": [12, 20, 41]},
                "CZ": {"fats_out2": [4, 5, 5],
                       "fats_out2_r2": [7, 7, 8, 10, 10, 5, 5, 8, 6, 10, 7],
                       "fats_out_part": [None, 9, 7]},
            }[geo][dataset]
            first = {"fats_out2": 2007, "fats_out2_r2": 2010, "fats_out_part": 2021}[dataset]
            points = list(zip(range(first, first + len(values)), values))
            if self.wrong == "geo":
                codes["geo"] = "DE"
            elif self.wrong == "metric":
                codes["indic_sbs" if dataset == "fats_out_part" else "indic_bp"] = "EMP"
            elif self.wrong == "missing-year" and dataset == "fats_out2_r2":
                points = [point for point in points if point[0] != 2015]
            flags = {str(i): "|C" for i, (_, value) in enumerate(points) if value is None}
            if self.wrong == "new-break" and dataset == "fats_out_part":
                flags["2"] = "b"
            response = {
                "id": list(codes) + ["time"], "size": [1] * len(codes) + [len(points)],
                "extension": {"id": dataset.upper()}, "updated": "2026-03-31T23:00:00+0200",
                "dimension": {name: {"category": {"index": {code: 0}}} for name, code in codes.items()},
                "value": {str(i): value for i, (_, value) in enumerate(points) if value is not None},
                "status": flags}
            response["dimension"]["time"] = {"category": {
                "index": {str(year): i for i, (year, _) in enumerate(points)}}}
            return json.dumps(response).encode(), asset

    refresh = getattr(module, collector)
    updated = validate_business_data(refresh(Client(), prior))
    assert prior == before
    assert [r for r in updated["observations"] if not r["series_id"].startswith(prefix)] == [
        r for r in prior["observations"] if not r["series_id"].startswith(prefix)]
    assert {sid: value for sid, value in updated["sources"].items() if not sid.endswith("_" + geo)} == {
        sid: value for sid, value in prior["sources"].items() if not sid.endswith("_" + geo)}
    rows = {r["year"]: r for r in updated["observations"] if r["series_id"].startswith(prefix)}
    assert sorted(rows) == list(range(2007, 2024))
    assert rows[2012 if geo == "PT" else 2021]["status"] == "suppressed"
    assert rows[2012 if geo == "PT" else 2021]["value"] is None
    assert rows[2019]["comparability"] != rows[2023]["comparability"]
    if geo == "PT":
        assert rows[2007]["series_id"] == "eurostat_portugal_pre_uci:china"
        assert rows[2008]["series_id"] == "eurostat_portugal:china"
        for previous, current in ((2009, 2010), (2013, 2014), (2019, 2020)):
            assert rows[previous]["comparability"] != rows[current]["comparability"]
        assert rows[2020]["comparability"] == rows[2021]["comparability"]
    else:
        assert rows[2007]["comparability"] != rows[2008]["comparability"]
        assert rows[2020]["comparability"] != rows[2021]["comparability"]
    assert refresh(Client(), updated) == updated
    for wrong in ("geo", "metric", "missing-year", "new-break", "methodology", "unit"):
        invalid = Client()
        invalid.wrong = wrong
        with pytest.raises(ValueError):
            refresh(invalid, prior)


@pytest.mark.parametrize("chart,sid,first,confidential", [
    ("portugal-affiliates-china", "eurostat_portugal:china", 2008, 2012),
    ("czechia-affiliates-china", "eurostat_czechia:china", 2007, 2021),
])
def test_portuguese_czech_exports_keep_suppression_gaps_and_verified_control_populations(chart, sid, first, confidential):
    data = published()
    for theme in CHART_THEMES:
        svg = ElementTree.fromstring(svg_business_chart(data, chart, theme))
        bars = [node for node in svg.iter() if node.get("class") == "annual-column"]
        assert [int(node.get("data-year")) for node in bars] == [
            year for year in range(first, 2024) if year != confidential]
        assert not any(node.get("class") == "data-line" for node in svg.iter())
    invalid = copy.deepcopy(data)
    invalid["series"][sid]["scope"] = "All minority foreign investment interests including Hong Kong"
    with pytest.raises(ValueError, match="ultimate-control"):
        validate_business_data(invalid)
    invalid = copy.deepcopy(data)
    next(row for row in invalid["observations"] if row["series_id"] == sid)["year"] = first - 1
    with pytest.raises(ValueError, match="pre-UCI|verified 2007"):
        validate_business_data(invalid)


def test_ons_parser_uses_enterprises_not_employment_and_verifies_year():
    book = Workbook()
    book.remove(book.active)
    for year, name, count in [(2018, "Table 1", 709), (2019, "Table 2", 531)]:
        sheet = book.create_sheet(name)
        sheet.cell(1, 1, f"UK Outward Foreign Affiliate Statistics: {year}")
        sheet.cell(9, 2, "TOTAL ACTIVITY")
        sheet.cell(10, 2, "ENT")
        sheet.append(["China", count, 124879, 15247499])
    stream = io.BytesIO()
    book.save(stream)
    assert parse_ons(stream.getvalue()) == [(2018, 709, "estimated"), (2019, 531, "estimated")]
    book["Table 2"].cell(1, 1, "2020")
    stream = io.BytesIO()
    book.save(stream)
    with pytest.raises(ValueError, match="year"):
        parse_ons(stream.getvalue())


def test_france_keeps_its_missing_cell_and_each_reported_population_boundary():
    data = published()
    rows = {r["year"]: r for r in data["observations"] if r["series_id"] == "eurostat_france:china"}
    assert set(rows) == set(range(2007, 2024))
    assert rows[2008]["value"] is None and rows[2008]["status"] == "missing"
    assert rows[2023]["value"] == 2153
    assert "China except Hong Kong" in data["series"]["eurostat_france:china"]["scope"]
    for previous, current in [(2007, 2009), (2009, 2010), (2015, 2016), (2020, 2021)]:
        assert rows[previous]["comparability"] != rows[current]["comparability"]
    doc = ElementTree.fromstring(svg_business_chart(data, "france-affiliates-china"))
    bars = [n for n in doc.iter() if n.get("class") == "annual-column"]
    assert [int(n.get("data-year")) for n in bars] == [y for y in range(2007, 2024) if y != 2008]
    assert not any(n.get("class") == "data-line" for n in doc.iter())


def test_earlier_eu_membership_and_classification_are_not_relabelled_as_current():
    data = eurostat()
    data["id"][1] = "nace_r1"
    data["id"][2] = "indic_bp"
    del data["dimension"]["nace_r2"], data["dimension"]["indic_sbs"]
    data["dimension"]["nace_r1"] = {"category": {"index": {"C-O_X_L": 0}}}
    data["dimension"]["indic_bp"] = {"category": {"index": {"ENT": 0}}}
    data["dimension"]["geo"]["category"]["index"] = {"EU27_2007": 0}
    data["dimension"]["time"]["category"]["index"] = {"2007": 0, "2008": 1, "2009": 2}
    data["extension"]["id"] = "FATS_OUT2"
    data["value"], data["status"] = {"2": 5688}, {}
    points = parse_eurostat(data, "fats_out2", "EU27_2007")
    assert [(r[0], r[1]) for r in points] == [(2007, None), (2008, None), (2009, 5688)]
    with pytest.raises(ValueError, match="different geography"):
        parse_eurostat(data, "fats_out2")


def test_count_parser_skips_contents_but_rejects_duplicate_country_rows():
    def sheet(name, rows):
        return SimpleNamespace(name=name, nrows=len(rows), row_values=lambda i: rows[i],
                               cell_value=lambda i, j: rows[i][j])
    contents = sheet("Contents", [["Number of U.S. Parents", ""]])
    table = sheet("Table I.K2", [["Number of U.S. Parents", "All industries"], ["China", 1134], ["Europe", 2701]])
    assert country_counts([contents, table], "Number of U.S. Parents") == {"China": (1134, "observed"), "Europe": (2701, "observed")}
    with pytest.raises(ValueError, match="Ambiguous"):
        country_counts([table, table], "Number of U.S. Parents")


def test_annual_histories_and_method_boundaries_are_not_endpoint_snapshots():
    data = published()
    doc = ElementTree.fromstring(svg_business_chart(data, BENCHMARK_CHART))
    assert not [n for n in doc.iter() if n.get("class") == "data-line"]
    assert len([n for n in doc.iter() if n.get("class") == "annual-column"]) == 4
    doc = ElementTree.fromstring(svg_business_chart(data, "us-affiliates-china"))
    assert not [n for n in doc.iter() if n.get("class") == "data-line"]
    assert len([n for n in doc.iter() if n.get("class") == "annual-column"]) == 24
    assert len(chart_groups(data, "us-affiliates-china")[0]["points"]) == 24
    assert sum(len(g["points"]) for g in chart_groups(data, "eu-affiliates-china")) == 12
    assert len(chart_groups(data, "uk-affiliates-china")[0]["points"]) == 11
    assert len(chart_groups(data, "norway-affiliates-china")[0]["points"]) == 17
    doc = ElementTree.fromstring(svg_business_chart(data, "eu-affiliates-china"))
    path = next(n for n in doc.iter() if n.get("class") == "data-line" and n.get("data-series-id") == "eurostat_affiliates:china")
    assert path.get("d").count("M") == 2
    assert "2023 confidential" in "".join(doc.itertext())
    assert "2023" in [n.text for n in doc.iter() if n.get("class") == "year-tick"]


def test_complete_business_downloads_exact_points_backgrounds_and_stale_cleanup(tmp_path):
    data = published()
    for base in (tmp_path / "outputs/figures/business_presence", tmp_path / "site/figures/business_presence"):
        (base / "substack").mkdir(parents=True)
        for chart_id in RETIRED_CHARTS:
            (base / "substack" / (chart_id + ".svg")).write_text("Obsolete wrong-destination chart")
    assert build_business(tmp_path, data, defaultdict(list)) == 25
    assert not any((tmp_path / "outputs/figures/business_presence/substack" / (cid + ".svg")).exists()
                   for cid in RETIRED_CHARTS)
    assert json.loads((tmp_path / "site/downloads/business-presence.json").read_text()) == data
    with (tmp_path / "site/downloads/business-observations.csv").open() as f:
        raw = list(csv.DictReader(f))
    assert len(raw) == len(data["observations"])
    legal = [r for r in raw if r["series_id"] == "eurostat_italy_legal:china"]
    assert len(legal) == 17 and all(r["unit"] == "legal units" and r["status"] == "estimated" for r in legal)
    for actual, expected in zip(raw, data["observations"]):
        assert all(actual[k] == ("" if v is None else str(v)) for k, v in expected.items())
        assert actual["unit"] == data["series"][expected["series_id"]]["unit"]
        assert json.loads(actual["source_assets"]) == data["sources"][expected["source_id"]]["assets"]
    with (tmp_path / "site/downloads/business-chart-values.csv").open() as f:
        plotted = list(csv.DictReader(f))
    wanted = {(cid, g["id"], p["year"]): p["plotted_value"]
              for cid in CHARTS for g in chart_groups(data, cid) for p in g["points"]}
    assert len(plotted) == len(wanted)
    for r in plotted:
        assert float(r["plotted_value"]) == wanted[(r["chart_id"], r["series_id"], int(r["year"]))]
    base = tmp_path / "outputs/figures/business_presence"
    for path in base.glob("*/*.svg"):
        doc = ElementTree.fromstring(path.read_text())
        points = [n for n in doc.iter() if n.get("class") == "business-point"]
        expected = {k[1:]: v for k, v in wanted.items() if k[0] == path.stem}
        assert len(points) == len(expected)
        for p in points:
            assert float(p.get("data-value")) == expected[(p.get("data-series-id"), int(p.get("data-year")))]
            assert p.get("fill") == COLORS[p.get("data-series-id")]
        for suffix in (".png", ".svg"):
            p = path.with_suffix(suffix)
            assert p.read_bytes() == (tmp_path / "site/figures/business_presence" / p.relative_to(base)).read_bytes()
        with Image.open(path.with_suffix(".png")) as image:
            assert image.size == (3000, 1800)
            assert image.getpixel((0, 0)) == (*ImageColor.getrgb(CHART_THEMES[path.parent.name]["background"]), 255)
    page = BeautifulSoup((tmp_path / "site/business-presence.html").read_text(), "html.parser")
    assert "different counting units" in page.header.text
    assert "EU27" in page.header.text and "not matched" in page.header.text
    table = page.select("#annual-values tbody tr")
    assert len(table) == 26 and table[0].th.text == "2000" and table[-1].th.text == "2025"
    assert page.select_one("#business-benchmark-background").get("open") is None
    assert page.select_one("#" + BENCHMARK_CHART).find_parent("details")["id"] == "business-benchmark-background"
    assert not page.select('#us-affiliates-europe,#us-affiliates-index,#us-parent-companies')
    custom = base / "substack/custom.svg"
    custom.write_text("User-owned")
    data["observations"] = [r for r in data["observations"] if not r["series_id"].startswith("eurostat")]
    build_business(tmp_path, data, defaultdict(list))
    assert not list(base.glob("*/eu-affiliates-china.*"))
    assert custom.read_text() == "User-owned"


def test_invalid_business_input_is_rejected_before_writes_and_deleted_input_withdraws_exports(tmp_path):
    path = tmp_path / "data/site"
    path.mkdir(parents=True)
    (path / "observations.json").write_text(json.dumps(json.loads((ROOT / "data/site/observations.json").read_text())[:1]))
    data = published()
    data["series"]["bea_parents:china"]["unit"] = "affiliates"
    frozen = path / "business_presence.json"
    frozen.write_text(json.dumps(data))
    before = {p.relative_to(tmp_path): p.read_bytes() for p in tmp_path.rglob("*") if p.is_file()}
    with pytest.raises(ValueError, match="counting units"):
        build(tmp_path)
    assert before == {p.relative_to(tmp_path): p.read_bytes() for p in tmp_path.rglob("*") if p.is_file()}
    frozen.write_text(json.dumps(published()))
    with pytest.raises(ValueError, match="current business evidence"):
        build(tmp_path)
    (path / "business_current.json").write_text((ROOT / "data/site/business_current.json").read_text())
    assert build(tmp_path)["business_charts"] == 26
    page = BeautifulSoup((tmp_path / "site/business-presence.html").read_text(), "html.parser")
    assert not page.select('a[href="language-learning.html"]')
    frozen.unlink()
    assert build(tmp_path)["business_charts"] == 0
    assert not (tmp_path / "site/business-presence.html").exists()
    assert not list((tmp_path / "site/downloads").glob("business-*"))
    assert not list((tmp_path / "outputs/figures/business_presence").glob("*/*.*"))


def test_publication_is_business_only_with_complete_data_original_downloads_and_provenance(tmp_path):
    spec = importlib.util.spec_from_file_location("business_publication", ROOT / "scripts/build_business_publication.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    output = tmp_path / "published"
    manifest = module.build_publication(ROOT, output, "verified-source-revision")
    page = BeautifulSoup((output / "index.html").read_text(), "html.parser")
    assert page.select_one('link[rel="canonical"]')["href"] == manifest["canonical_url"]
    assert "Interactive analysis" in page.select_one(".publication-header nav").text
    assert not page.select('a[href="language-learning.html"],iframe')
    assert len(page.select("svg[role=img]")) == 28
    assert not page.select("#business-controls,#annual-values,#module-chart,#uk-affiliates-china,#eu-affiliates-china")
    assert len(page.select(".business-card.native-chart,.business-card .native-chart")) == 28
    assert len(page.select(".business-count-card")) == 22
    assert len(page.select("#luxembourg-affiliates-china .annual-column")) == 15
    assert len(page.select("#hungary-affiliates-china .annual-column")) == 2
    assert len(page.select("#slovenia-affiliates-china .annual-column")) == 17
    assert manifest["business_charts"] == 28
    assert len(page.select("#business-inward-data table tbody tr")) == 25
    assert page.select_one("#business-data").get("open") is None
    assert set(p.name for p in (output / "downloads").iterdir()) == {
        "business-presence.json", "business-observations.csv", "business-chart-values.csv", "business-current.json",
        "business-current-observations.csv", "business-current-chart-values.csv",
        "business-inward.json", "business-inward-observations.csv", "business-inward-chart-values.csv"}
    data = json.loads((output / "downloads/business-presence.json").read_text())
    assert data == published() and manifest["historical_observations"] == len(data["observations"])
    assert manifest["observations"] == 59
    assert manifest["inward_observations"] == 50
    assert manifest["source_revision"] == "verified-source-revision"
    for name, checksum in manifest["files"].items():
        assert hashlib.sha256((output / name).read_bytes()).hexdigest() == checksum
    with (output / "downloads/business-observations.csv").open() as f:
        assert len(list(csv.DictReader(f))) == len(data["observations"])
    for link in page.select('a[href^="figures/"]'):
        assert link.has_attr("data-original-chart") and (output / link["href"].split("?")[0]).is_file()
    assert not (output / "business-presence.html").exists()
    assert not any("weibo" in str(p.relative_to(output)) or "language" in str(p.relative_to(output))
                   for p in output.rglob("*"))
    custom = output / "figures/business_presence/substack/custom.svg"
    custom.write_text("User-owned")
    source = tmp_path / "source"
    (source / "data/site").mkdir(parents=True)
    (source / "src/us_china_exchange").mkdir(parents=True)
    (source / "src/us_china_exchange/static").symlink_to(ROOT / "src/us_china_exchange/static", target_is_directory=True)
    data["observations"] = [r for r in data["observations"] if not r["series_id"].startswith("eurostat")]
    (source / "data/site/business_presence.json").write_text(json.dumps(data))
    (source / "data/site/business_current.json").write_text((ROOT / "data/site/business_current.json").read_text())
    module.build_publication(source, output, "next-revision")
    assert not list((output / "figures/business_presence").glob("*/eu-affiliates-china.*"))
    assert custom.read_text() == "User-owned"
    assert not list((output / "figures/business_presence").glob("*/*ubo-us-affiliates.*"))
    assert not list((output / "downloads").glob("business-inward*"))


def spanish_national_px(year, count=134):
    variables = ["N\u00famero de filiales", "Personas ocupadas", "Volumen de negocio"] if year < 2021 else [
        "N\u00famero de filiales", "Personas ocupadas", "Cifra de negocios",
        "Costes de personal", "Inversi\u00f3n material bruta"]
    text = (
        'CODEPAGE="iso-8859-15";\nSOURCE="Instituto Nacional de Estad\u00edstica";\n'
        'MATRIX="01001";\nCREATION-DATE="20261025";\nSTUB="Pa\u00eds";\n'
        'HEADING="Principales variables";\nVALUES("Pa\u00eds")="TOTAL","China";\n'
        'VALUES("Principales variables")=' + ",".join(f'"{v}"' for v in variables) + ";\n"
        'UNITS="cifras econ\u00f3micas en miles de euros";\nDATA=' +
        " ".join(map(str, [5000, 90000, 800000, 70000, 60000][:len(variables)] +
                     [count, 26841, 5057324, 493464, 670596][:len(variables)])) + ";\n")
    return text.encode("iso-8859-15")


def test_spanish_national_refresh_preserves_other_counts_and_separates_units(monkeypatch):
    from us_china_exchange import business_sources

    spec = importlib.util.spec_from_file_location("spanish_national_collector", ROOT / "scripts/build_business_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    tables = {year: (url, spanish_national_px(year, 55 if year == 2009 else year - 2000))
              for year, (url, _) in business_sources.SPAIN_NATIONAL_TABLES.items()}
    reviewed = {year: (url, hashlib.sha256(body).hexdigest()) for year, (url, body) in tables.items()}
    method_url = business_sources.SPAIN_NATIONAL_METHODS[0]
    methods = b"Reviewed national methodology fixture"
    monkeypatch.setattr(business_sources, "SPAIN_NATIONAL_TABLES", reviewed)
    monkeypatch.setattr(business_sources, "SPAIN_NATIONAL_METHODS", (method_url, hashlib.sha256(methods).hexdigest()))

    class Client:
        overrides = {}

        def get(self, url):
            bodies = {url: body for url, body in tables.values()}
            bodies[method_url] = methods
            body = self.overrides.get(url, bodies[url])
            return body, {"url": url, "request_url": url, "sha256": hashlib.sha256(body).hexdigest(),
                          "retrieved_at": "2026-10-25T00:00:00+00:00"}

    prior = published()
    prior["series"] = {k: v for k, v in prior["series"].items() if not k.startswith("ine_spain")}
    prior["observations"] = [r for r in prior["observations"] if not r["series_id"].startswith("ine_spain")]
    prior["sources"] = {k: v for k, v in prior["sources"].items() if not k.startswith("ine_spain_")}
    prior["coverage"] = [r for r in prior["coverage"] if not r["id"].startswith("ine_spain")]
    before = copy.deepcopy(prior)
    updated = validate_business_data(module.collect_spain_national(Client(), prior))
    assert prior == before
    actual = {(r["series_id"], r["year"]): r for r in updated["observations"]}
    assert all(actual[r["series_id"], r["year"]] == r for r in prior["observations"])
    assert all(updated["series"][k] == v for k, v in prior["series"].items())
    assert all(updated["sources"][k] == v for k, v in prior["sources"].items())
    assert all(row in updated["coverage"] for row in prior["coverage"])
    assert len(updated["observations"]) == len(prior["observations"]) + 16
    assert [r["year"] for r in updated["observations"] if r["series_id"] == "ine_spain_units:china"] == list(range(2009, 2021))
    assert [r["year"] for r in updated["observations"] if r["series_id"] == "ine_spain:china"] == list(range(2021, 2025))
    assert actual["ine_spain_units:china", 2009]["value"] == 55
    assert actual["ine_spain:china", 2024]["value"] == 24
    assert updated["series"]["ine_spain_units:china"]["unit"] == "affiliate units"
    assert updated["series"]["ine_spain:china"]["unit"] == "enterprises"
    assert module.collect_spain_national(Client(), updated) == updated
    assert all(chart_groups(updated, chart_id) == chart_groups(prior, chart_id) for chart_id in CHARTS)
    for year in (2009, 2020, 2021, 2024):
        sid = "ine_spain_units:china" if year < 2021 else "ine_spain:china"
        point = actual[sid, year]
        assert point["source_id"] == f"ine_spain_{year}"
        assert point["comparability"] == "ine:spain:china:" + ("affiliate_units" if year < 2021 else "affiliate_enterprises")
        assert "15.2 also claims temporal continuity" in point["notes"]
        assert "without assuming retrospective harmonization" in point["notes"]
        assert "CN_X_HK equivalence not established" in updated["series"][sid]["scope"]
    for sid, year in (("ine_spain:china", 2019), ("ine_spain_units:china", 2021)):
        invalid = copy.deepcopy(updated)
        next(r for r in invalid["observations"] if r["series_id"] == sid)["year"] = year
        with pytest.raises(ValueError, match="must not be spliced"):
            validate_business_data(invalid)
    for field, value in (("unit", "enterprises"), ("measure", "controlled_enterprises"),
                         ("scope", "Spanish-resident UCI; China except Hong Kong")):
        invalid = copy.deepcopy(updated)
        invalid["series"]["ine_spain_units:china"][field] = value
        with pytest.raises(ValueError):
            validate_business_data(invalid)
    for field, value in (("source_id", "fats_out_part_ES"), ("source_id", "ine_spain_2023"),
                         ("comparability", "ine:spain:china:affiliate_units")):
        invalid = copy.deepcopy(updated)
        next(r for r in invalid["observations"] if r["series_id"] == "ine_spain:china" and r["year"] == 2024)[field] = value
        with pytest.raises(ValueError):
            validate_business_data(invalid)
    for url, replacement in ((method_url, methods + b"changed"),
                             (tables[2019][0], tables[2020][1]),
                             (tables[2024][0], spanish_national_px(2024, 25))):
        client = Client()
        client.overrides = {url: replacement}
        with pytest.raises(ValueError, match="changed"):
            module.collect_spain_national(client, prior)


def test_spanish_national_parser_rejects_wrong_identity_units_and_missingness():
    from us_china_exchange.business_sources import parse_spanish_national

    current = spanish_national_px(2024)
    assert parse_spanish_national(current, 2024) == 134
    assert parse_spanish_national(spanish_national_px(2009, 55), 2009) == 55
    assert parse_spanish_national(spanish_national_px(2024, 0), 2024) == 0
    for old, new in ((b'Instituto Nacional de Estad', b'Other publisher'),
                     (b'MATRIX="01001"', b'MATRIX="02001"'),
                     (b'STUB="Pa', b'STUB="Other'),
                     (b'"China"', b'"Hong Kong"'),
                     (b'"China"', b'"TOTAL"'),
                     ("N\u00famero de filiales".encode("iso-8859-15"), b"Persons employed"),
                     (b'134 26841', b': 26841'),
                     (b'134 26841', b'NaN 26841'),
                     (b'134 26841', b'1.5 26841'),
                     (b'134 26841', b'-1 26841'),
                     (b'134 26841', b'134'),
                     (b'CODEPAGE="iso-8859-15"', b'CODEPAGE="utf-8"')):
        with pytest.raises(ValueError):
            parse_spanish_national(current.replace(old, new), 2024)
    for year in (2008, 2025, 2009.0, True):
        with pytest.raises(ValueError, match="reference year"):
            parse_spanish_national(current, year)
    with pytest.raises(ValueError, match="unit regime"):
        parse_spanish_national(current, 2020)


def test_luxembourg_refresh_preserves_declared_continuity_and_rejects_wrong_scope():
    from us_china_exchange.business_sources import LUXEMBOURG_EUROSTAT_METHODS

    spec = importlib.util.spec_from_file_location("luxembourg_collector", ROOT / "scripts/build_business_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    prior = published()
    snapshot = copy.deepcopy(prior)

    class Client:
        wrong = None

        def get(self, url):
            asset = dict(url=url, request_url=url, sha256="e" * 64, retrieved_at="2026-10-25T00:00:00+00:00")
            if url == LUXEMBOURG_EUROSTAT_METHODS:
                text = (
                    "National reference metadata Luxembourg; LU - Grand Duchy of Luxembourg; "
                    "statistical unit of FATS is the enterprise as defined; "
                    "SPEs are not covered by our FDI survey; "
                    "All large and medium sized enterprises and many small entities are covered; "
                    "There is no break in series since there were no changes in the methodology; "
                    "Length of comparable time series: from 2009 up to now; "
                    "Number of enterprises and employment variables are recorded in absolute figures; "
                    "We apply the UCI approach")
                if self.wrong == "method-spe":
                    text = text.replace("SPEs are not covered", "SPEs are covered")
                elif self.wrong == "method-unit":
                    text = text.replace("enterprise as defined", "legal unit as defined")
                elif self.wrong == "method-window":
                    text = text.replace("from 2009 up to now", "from 2007 up to now")
                return text.encode(), asset
            family = urlsplit(url).path.split("/")[-1]
            codes = {k: v[0] for k, v in parse_qs(urlsplit(url).query).items() if k != "lang"}
            first, values = {
                "fats_out2": (2007, [None, None, 8]),
                "fats_out2_r2": (2010, [10, 7, 13, 13, 7, 7, 7, 7, 8, 10, 10]),
                "fats_out_part": (2021, [9, 11, 11]),
            }[family]
            points, flags = list(zip(range(first, first + len(values)), values)), {}
            if self.wrong == "geo":
                codes["geo"] = "DE"
            elif self.wrong == "metric":
                codes["indic_sbs" if family == "fats_out_part" else "indic_bp"] = "EMP"
            elif self.wrong == "missing-year" and family == "fats_out2_r2":
                points = [p for p in points if p[0] != 2015]
            elif self.wrong == "pre2009-zero" and family == "fats_out2":
                points[0] = (2007, 0)
            elif self.wrong == "break" and family == "fats_out_part":
                flags = {"0": "b"}
            data = {
                "id": list(codes) + ["time"], "size": [1] * len(codes) + [len(points)],
                "extension": {"id": family}, "updated": "2026-10-25T00:00:00+0000",
                "dimension": {k: {"category": {"index": {v: 0}}} for k, v in codes.items()},
                "value": {str(i): value for i, (_, value) in enumerate(points) if value is not None},
                "status": flags,
            }
            data["dimension"]["time"] = {"category": {"index": {
                str(year): i for i, (year, _) in enumerate(points)}}}
            return json.dumps(data).encode(), asset

    result = validate_business_data(module.collect_luxembourg(Client(), prior))
    assert prior == snapshot
    assert [r for r in result["observations"] if r["series_id"] != "eurostat_luxembourg:china"] == [
        r for r in prior["observations"] if r["series_id"] != "eurostat_luxembourg:china"]
    rows = [r for r in result["observations"] if r["series_id"] == "eurostat_luxembourg:china"]
    assert [r["year"] for r in rows] == list(range(2009, 2024))
    assert len({r["comparability"] for r in rows}) == 1
    assert len({r["source_id"] for r in rows}) == 3
    assert module.collect_luxembourg(Client(), result) == result
    for failure in ("geo", "metric", "missing-year", "pre2009-zero", "break",
                    "method-spe", "method-unit", "method-window"):
        invalid = Client()
        invalid.wrong = failure
        with pytest.raises(ValueError):
            module.collect_luxembourg(invalid, prior)

def test_luxembourg_model_preserves_all_bars_and_source_population_guards():
    data = published()
    for theme in CHART_THEMES:
        doc = ElementTree.fromstring(svg_business_chart(data, "luxembourg-affiliates-china", theme))
        bars = [node for node in doc.iter() if node.get("class") == "annual-column"]
        assert [int(node.get("data-year")) for node in bars] == list(range(2009, 2024))
        assert [int(node.get("data-value")) for node in bars] == [8, 10, 7, 13, 13, 7, 7, 7, 7, 8, 10, 10, 9, 11, 11]
    wrong = copy.deepcopy(data)
    wrong["series"]["eurostat_luxembourg:china"]["scope"] = "All resident parents and Hong Kong holdings"
    with pytest.raises(ValueError, match="Luxembourg control"):
        validate_business_data(wrong)
    wrong = copy.deepcopy(data)
    next(r for r in wrong["observations"] if r["series_id"] == "eurostat_luxembourg:china")["comparability"] = "invented2021break"
    with pytest.raises(ValueError, match="Luxembourg historical"):
        validate_business_data(wrong)
    wrong = copy.deepcopy(data)
    next(r for r in wrong["observations"] if r["series_id"] == "eurostat_luxembourg:china")["source_id"] = "fats_out2_FR"
    with pytest.raises(ValueError, match="method/source"):
        validate_business_data(wrong)
