import copy
import csv
import io
import importlib.util
import json
import zipfile
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from xml.etree import ElementTree

import pytest
from us_china_exchange.theme import png_dimensions
from bs4 import BeautifulSoup
from openpyxl import Workbook
from PIL import Image

from us_china_exchange import language_sources
from us_china_exchange.language_analysis import CHARTS, RETIRED_CHARTS, chart_points, svg_language_chart, validate_language_data
from us_china_exchange.model import digest
from us_china_exchange.render import build_languages
from us_china_exchange.render import build


ROOT = Path(__file__).resolve().parents[1]


@pytest.mark.parametrize("report_year,fall_year,expected,chinese_year", [
    (2021, 2021, "latest_released_data", 2021),
    (2024, 2021, "update_available", 2021),
    (2021, 2024, "update_available", 2021),
    (2024, 2024, "update_available", 2024),
])
def test_national_release_check_detects_new_data_without_using_summer_or_regional_years(
        report_year, fall_year, expected, chinese_year):
    spec = importlib.util.spec_from_file_location("language_releases", ROOT / "scripts/check_language_releases.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    book = Workbook()
    book.active.title = "FL history"
    book.active.append(["SRVY_YEAR", "TERM"])
    book.active.append([fall_year, "Fall"])
    book.active.append([2026, "Summer"])
    stream = io.BytesIO()
    book.save(stream)
    url = "https://www.mla.org/Historical-language-enrollments-1958-2024.xlsx"
    assets = {
        module.MLA_REPORTS: f'<p>Next data gathering: 2026</p><a href="report.pdf">{report_year} Report</a>',
        module.MLA_DATABASE: f'<a href="{url}">Complete data set</a>',
        url: stream.getvalue(),
    }
    client = SimpleNamespace(get=lambda address: (assets[address], {"request_url": address}))
    data = {"observations": [
        {"series_id": "mla:chinese", "year": chinese_year, "value": 46492},
        {"series_id": "mla:japanese", "year": 2021, "value": 65661},
        {"series_id": "mla:korean", "year": 2021, "value": 19270},
        {"series_id": "ccc:chinese", "year": 2026, "value": 2936},
        {"series_id": "mla_matched:chinese", "year": 2024, "value": 100},
    ]}
    report = module.check_mla(client, data)
    assert report["status"] == expected
    assert report["published_latest_year"] == 2021
    assert report["published_years_by_language"]["chinese"] == chinese_year
    assert report["available_latest_year"] == max(report_year, fall_year)
    assert report["workbook_fall_years"] == [fall_year]


def published():
    return json.loads((ROOT / "data/site/language_learning.json").read_text())


def test_us_annual_ap_parser_selects_national_counts_and_the_actual_reporting_year(monkeypatch):
    caption = ("Advanced Placement (AP) Examinations and Scores, by Subject, "
               "Texas and the United States, ")
    page = (caption + "2021-22\nExam Texas exams (N) U.S. exams (N)\n"
            "Chinese Language and Culture 741 13,186 0.1 0.3 85.2 85.2 3.9 3.9\n"
            "Japanese Language and Culture 149 2,450 <0.1 0.1 53.7 72.8 2.9 3.5\n"
            "Note. Data are based on public and nonpublic school examinees.")
    pages = ["Contents\n" + caption + "2021-22 ........ 9",
             page.replace("2021-22", "2020-21").replace("13,186", "11,256"), page]
    monkeypatch.setattr(language_sources, "pdf_text", lambda body: pages)
    assert language_sources.parse_ap_us_annual(b"fixture", 2022) == [
        ("chinese", 13186, 3), ("japanese", 2450, 3)]
    pages[-1] = page.replace("public and nonpublic", "public")
    with pytest.raises(ValueError, match="population"):
        language_sources.parse_ap_us_annual(b"fixture", 2022)
    pages[-1] = page.replace("Texas exams (N) U.S. exams (N)", "U.S. exams (N) Texas exams (N)")
    with pytest.raises(ValueError, match="columns"):
        language_sources.parse_ap_us_annual(b"fixture", 2022)
    pages[-1] = page.replace("741 13,186", "13,186 741")
    with pytest.raises(ValueError, match="smaller than Texas"):
        language_sources.parse_ap_us_annual(b"fixture", 2022)
    pages[-1] = page + "\nChinese Language and Culture 741 13,186 0.1 0.3 85.2 85.2 3.9 3.9"
    with pytest.raises(ValueError, match="ambiguous"):
        language_sources.parse_ap_us_annual(b"fixture", 2022)
    pages[-1] = page.replace("Japanese Language and Culture", "Other languages")
    with pytest.raises(ValueError, match="missing"):
        language_sources.parse_ap_us_annual(b"fixture", 2022)


def test_combined_us_annual_ap_report_keeps_year_sections_across_page_breaks(monkeypatch):
    header = ("Advanced Placement (AP) Examinations and Scores, by Subject, "
              "Texas and the United States, 2012-13 and 2013-14\n"
              "Examination Texas U.S. Texas U.S. Texas U.S. Texas U.S.\n")
    note = "\nNote. Data are based on public and nonpublic school examinees."
    pages = [
        header + "2012-13\nChinese Language and Culture 498 8,078 0.1 0.2 94.6 93.9 4.5 4.4" + note,
        header + "Japanese Language and Culture 94 1,855 <0.1 <0.1 62.8 72.9 3.1 3.4\n"
        "2013-14\nChinese Language and Culture 549 8,669 0.1 0.2 95.1 93.4 4.4 4.4" + note,
        header + "Japanese Language and Culture 134 1,942 <0.1 <0.1 56.7 72.7 2.9 3.4" + note]
    monkeypatch.setattr(language_sources, "pdf_text", lambda body: pages)
    assert language_sources.parse_ap_us_annual(b"fixture", 2013) == [
        ("chinese", 8078, 1), ("japanese", 1855, 2)]
    assert language_sources.parse_ap_us_annual(b"fixture", 2014) == [
        ("chinese", 8669, 2), ("japanese", 1942, 3)]


def test_us_annual_ap_counts_preserve_separate_populations_and_cannot_exceed_global_totals():
    data = validate_language_data(published())
    values = {(r["series_id"], r["year"]): r["value"] for r in data["observations"]}
    assert [values["ap_us_annual:chinese", year] for year in (2019, 2022)] == [12153, 13186]
    assert [values["ap_us_annual:japanese", year] for year in (2019, 2022)] == [2113, 2450]
    assert not any(sid.startswith("ap_us_annual:") for chart in CHARTS.values() for sid in chart["series"])
    data["series"]["ap_us_annual:chinese"]["scope"] = "Global annual participants"
    with pytest.raises(ValueError, match="all-school"):
        validate_language_data(data)
    data = published()
    next(r for r in data["observations"] if r["series_id"] == "ap_us_annual:chinese" and r["year"] == 2022)["value"] = values["ap:global_total", 2022] + 1
    with pytest.raises(ValueError, match="global total"):
        validate_language_data(data)


def test_selective_us_annual_ap_refresh_preserves_global_cohort_and_other_records(monkeypatch):
    spec = importlib.util.spec_from_file_location("language_builder", ROOT / "scripts/build_language_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    prior = published()
    definitions = {sid: definition for sid, definition in prior["series"].items() if sid.startswith("ap_us_annual:")}
    observations = [r for r in prior["observations"] if r["series_id"] in definitions]
    sources = {r["source_id"]: prior["sources"][r["source_id"]] for r in observations}
    monkeypatch.setattr(module, "ap_us_annual", lambda client: (definitions, observations, sources))
    assert module.collect_ap_us_annual(None, prior) == prior


def test_article_snapshot_preserves_the_matched_sample_and_requires_rechecking_changed_pdf(monkeypatch):
    selected = json.loads((ROOT / "data/reviewed/language-snapshot-2020.json").read_text())
    content = b"snapshot fixture"
    selected["source_sha256"] = digest(content)
    monkeypatch.setattr(language_sources, "pdf_text", lambda body: [
        "1,308 matched institutions; 2016 and 2020 course enrollments, not students."])
    rows = language_sources.parse_mla_snapshot(content, selected)
    assert rows == [
        ("chinese", 2016, 33561), ("chinese", 2020, 26528),
        ("japanese", 2016, 46936), ("japanese", 2020, 43268),
        ("korean", 2016, 9033), ("korean", 2020, 11323)]
    with pytest.raises(ValueError, match="review the table image"):
        language_sources.parse_mla_snapshot(b"changed publisher PDF", selected)
    selected["cohort_institutions"] = 2547
    with pytest.raises(ValueError, match="scope"):
        language_sources.parse_mla_snapshot(content, selected)


def test_hesa_article_table_keeps_academic_years_and_rounded_total_not_component_sum(monkeypatch):
    pages = ["Chinese subjects 101164 101165 101166 101167; apportioned; nearest 5",
             "Table 7: HESA total\n"
             "2014/15 1440 850 190 395\n2015/16 1385 895 175 315\n"
             "2016/17 1420 925 180 315\n2017/18 1325 885 170 270\n"
             "2018/19 1225 805 160 260\n2019/20 1150 785 160 205\n"
             "2020/21 1115 760 145 205\n2021/22 1055 735 115 215\n"
             "According to these figures, 2012/13 and 2013/14 were recorded before; "
             "fallen from 1535 in 2012 to 1055 in 2021/22."]
    monkeypatch.setattr(language_sources, "pdf_text", lambda body: pages)
    values = language_sources.parse_hesa_bacs(b"fixture")
    assert values == [(2013, 1535), (2015, 1440), (2016, 1385), (2017, 1420),
                      (2018, 1325), (2019, 1225), (2020, 1150), (2021, 1115), (2022, 1055)]
    assert 2014 not in dict(values)
    pages[1] = pages[1].replace("2021/22 1055", "2021/23 1055")
    with pytest.raises(ValueError, match="academic-year"):
        language_sources.parse_hesa_bacs(b"fixture")


def test_article_observations_remain_separate_from_the_national_and_us_populations():
    data = validate_language_data(published())
    values = {(r["series_id"], r["year"]): r["value"] for r in data["observations"]}
    assert values["mla_snapshot:chinese", 2016] == 33561
    assert values["mla_snapshot:chinese", 2020] == 26528
    assert values["mla:chinese", 2016] == 54248
    assert values["mla:chinese", 2021] == 46492
    assert values["hesa_chinese_studies:total", 2022] == 1055
    hesa = [r for r in data["observations"] if r["series_id"] == "hesa_chinese_studies:total"]
    assert len(hesa) == 8
    assert ("hesa_chinese_studies:total", 2013) not in values
    assert values["hesa_chinese_studies:quoted_2012", 2013] == 1535
    assert "unverified" in data["series"]["hesa_chinese_studies:quoted_2012"]["notes"]
    assert {r["comparability"] for r in hesa} == {
        "hesa_chinese_studies:jacs", "hesa_chinese_studies:hecos"}
    data["series"]["mla_snapshot:chinese"]["scope"] = "U.S. national census"
    with pytest.raises(ValueError, match="matched population"):
        validate_language_data(data)
    data = published()
    data["series"]["hesa_chinese_studies:total"]["unit"] = "unique Mandarin learners"
    with pytest.raises(ValueError, match="subject allocations"):
        validate_language_data(data)
    data = published()
    next(r for r in data["observations"] if r["series_id"] == "hesa_chinese_studies:quoted_2012")["comparability"] = "hesa_chinese_studies:jacs"
    with pytest.raises(ValueError, match="subject allocations"):
        validate_language_data(data)


def test_school_pulse_requires_national_denominator_and_keeps_unreportable_korean_null():
    book = Workbook()
    book.active.title = "About"
    book.active.append(["A sample of 4,000 public schools; experimental data"])
    sheet = book.create_sheet("All Schools")
    sheet.append(["Most Recent Month", "Subgroup", "VarName", "Respondents", "Description",
                  "Category", "Jan24 Pct", "Jan24 SE", "Jan24 Flag"])
    for variable, label, value, error in [
            ("Lang1_3", "Chinese (Mandarin or Cantonese)", 3, .4),
            ("Lang1_9", "Japanese", 1, .3), ("Lang1_10", "Korean", "‡", "†")]:
        sheet.append([datetime(2024, 1, 1), "All public schools", variable, "All Schools",
                      "Which of the following languages are taught as independent classes at your school?",
                      label, value, error, " "])

    def parse():
        stream = io.BytesIO()
        book.save(stream)
        return language_sources.parse_school_pulse(stream.getvalue())

    assert [(r[0], r[1], r[2]) for r in parse()] == [
        ("chinese", 3, "estimated"), ("japanese", 1, "estimated"), ("korean", None, "suppressed")]
    sheet.cell(2, 4, "Schools offering world languages")
    with pytest.raises(ValueError, match="denominator"):
        parse()
    sheet.cell(2, 4, "All Schools")
    book["About"].title = "OA_About"
    with pytest.raises(ValueError, match="outlying-area"):
        parse()
    groups = chart_points(published(), "school-language-availability")
    assert [[p["value"] for p in g["points"]] for g in groups] == [[3], [1], []]
    assert groups[2]["unavailable_label"] == "Not reportable"
    assert all(p["reported_unit"] == "percent" and p["year"] == 2024 for g in groups for p in g["points"])
    assert any(r["series_id"] == "school_availability:korean" and r["value"] == 43
               for r in published()["observations"])


def test_school_parser_skips_contents_and_keeps_enrollment_and_availability_populations_separate(monkeypatch):
    pages = ["Contents: Enrollment for Major Languages by State",
             "Enrollment for Major Languages by State\nTotal 26,045 130,411 227,086 1,289,004 330,898 67,909",
             "Distribution of Foreign Language Programs\nNumber of HS programs\nChinese 1144 6.43\nJapanese 433 2.44\nKorean 43 0.24",
             "In 43 schools across 22 states, 936 students are enrolled in Korean language classes."]
    monkeypatch.setattr(language_sources, "pdf_text", lambda body: pages)
    asset = published()["sources"]["k12_2017"]["assets"][0]
    client = SimpleNamespace(get=lambda url: (b"fixture", asset))
    definitions, observations, _ = language_sources.school_comparisons(client)
    assert {r["series_id"]: r["value"] for r in observations} == {
        "schools:k12_japanese_estimate": 67909, "schools:korean_respondent_enrollment": 936,
        "school_availability:chinese": 1144, "school_availability:japanese": 433, "school_availability:korean": 43}
    assert "respondent subset" in definitions["schools:korean_respondent_enrollment"]["scope"]
    pages[2] = pages[2].replace("Korean 43", "Korean 44")
    with pytest.raises(ValueError, match="availability"):
        language_sources.school_comparisons(client)


def test_enrollment_charts_use_all_college_languages_but_do_not_treat_korean_respondents_as_national():
    data = published()
    groups = chart_points(data, "college-enrollments")
    assert [next(p["value"] for p in g["points"] if p["year"] == 2021) for g in groups] == [46492, 65661, 19270]
    schools = chart_points(data, "school-language-enrollments")
    assert [g["points"][0]["value"] for g in schools] == [227086, 67909]
    assert all(g["points"][0]["year"] == 2015 and g["points"][0]["status"] == "estimated" for g in schools)
    assert not any(g["id"] == "schools:korean_respondent_enrollment" for cid in CHARTS for g in chart_points(data, cid))
    next(r for r in data["observations"] if r["series_id"] == "schools:k12_estimate").update(value=None, status="missing")
    doc = ElementTree.fromstring(svg_language_chart(data, "school-language-enrollments"))
    assert "Unavailable" in "".join(doc.itertext())
    assert [n.get("data-value") for n in doc.iter() if n.get("class") == "language-point"] == ["67909"]


def test_reviewed_checkpoints_and_derived_values_are_not_new_learner_counts():
    data = validate_language_data(published())
    values = {(r["series_id"], r["year"]): r["value"] for r in data["observations"]}
    assert [values[("mla:chinese", y)] for y in (2002, 2013, 2021)] == [34153, 61997, 46492]
    assert [values[("ipeds:" + award, 2022)] for award in ("bachelors", "masters", "doctorates")] == [364, 26, 14]
    assert values[("ap:global_total", 2025)] == 18518
    assert values[("ap:standard_group", 2025)] == 4777
    assert [values[("ap_cohort:us_public", y)] for y in (2007, 2008, 2010, 2013)] == [1056, 1992, 3428, 5684]
    assert next(r for r in data["observations"] if (r["series_id"], r["year"]) == (
        "ap_cohort:us_public", 2008))["source_id"] == "ap_cohort_2014"
    assert {r["year"] for r in data["observations"] if r["series_id"] == "ap:global_total"} == set(range(2007, 2026))
    assert {r["year"] for r in data["observations"] if r["series_id"] == "ipeds:bachelors"} == set(range(2000, 2026))
    shares = {r["year"]: r for r in chart_points(data, "college-share")[0]["points"]}
    assert shares[2013]["plotted_value"] == pytest.approx(61997 / 1561176 * 100)
    assert shares[2021]["plotted_value"] == pytest.approx(46492 / 1182562 * 100)
    indexed = chart_points(data, "college-benchmarks")[0]["points"]
    assert next(r for r in indexed if r["year"] == 2021)["plotted_value"] == pytest.approx(46492 / 61997 * 100)
    assert 2020 not in shares
    assert not any(r["value"] == 34341 for r in data["observations"])


def test_review_client_distinguishes_get_and_different_public_query_bodies(tmp_path, monkeypatch):
    from types import SimpleNamespace
    from us_china_exchange import language_sources

    requests = []

    def run(command, **kwargs):
        payload = (json.loads(command[command.index("--data-binary") + 1])
                   if "--data-binary" in command else None)
        requests.append(payload)
        return SimpleNamespace(stdout=json.dumps({"query": payload}).encode() + b"\nhttps://example.test/table")

    monkeypatch.setattr(language_sources.subprocess, "run", run)
    client = language_sources.ReviewClient(tmp_path)
    url = "https://example.test/table"
    assert json.loads(client.get(url)[0]) == {"query": None}
    first, asset = client.post_json(url, {"country": "CN"})
    assert json.loads(first) == {"query": {"country": "CN"}}
    assert json.loads(client.post_json(url, {"country": "HK"})[0]) == {"query": {"country": "HK"}}
    assert client.post_json(url, {"country": "CN"}) == (first, asset)
    assert requests == [None, {"country": "CN"}, {"country": "HK"}]
    assert asset["request_url"] == asset["url"] == url


def test_nationwide_comparisons_keep_chinese_japanese_and_korean_separate():
    data = published()
    charts = ["college-enrollments", "college-share", "college-benchmarks", "matched-college-benchmarks",
              "chinese-bachelors", "chinese-graduate-awards", "language-doctorates"]
    for chart in charts:
        groups = chart_points(data, chart)
        assert len(groups) == 3
        assert [group["color"] for group in groups] == ["#7B2D26", "#3D6F8C", "#4A7C59"]
        assert not any("all_languages" in group["id"] for group in groups)
    for chart, expected in [("chinese-bachelors", [292, 464, 44]),
                            ("chinese-graduate-awards", [28, 21, 1]),
                            ("language-doctorates", [11, 6, 1])]:
        groups = chart_points(data, chart)
        assert [g["label"] for g in groups] == ["Chinese", "Japanese", "Korean"]
        assert [next(p["value"] for p in g["points"] if p["year"] == 2025) for g in groups] == expected
        assert all(g["period_basis"] == "academic ending year" for g in groups)
    assert not any(sid.startswith("nsc_program:") for sid in data["series"])


@pytest.mark.parametrize("chart", [
    "chinese-bachelors", "chinese-graduate-awards", "language-doctorates", "ap-chinese-global"])
@pytest.mark.parametrize("theme", ["themarginoferror", "substack"])
def test_degree_charts_remove_dots_but_retain_every_interactive_observation(chart, theme):
    data = published()
    svg = BeautifulSoup(svg_language_chart(data, chart, theme), "html.parser").svg
    markers = chart == "ap-chinese-global"
    circles = svg.select("circle.language-point")
    assert circles and all(float(p["r"]) == (4 if markers else 0) for p in circles)
    spec = importlib.util.spec_from_file_location("degree_marker_publication", ROOT / "scripts/build_exchange_publication.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    model = module.plot_model(svg, {}, {})
    assert model.get("point_markers", True) == markers
    expected_groups = chart_points(data, chart)
    assert len(model["groups"]) == len(expected_groups)
    for actual, expected in zip(model["groups"], expected_groups):
        assert actual["id"] == expected["id"]
        assert [(p["year"], p["value"]) for p in actual["points"]] == [
            (p["year"], p["plotted_value"]) for p in expected["points"]]
        if not markers:
            assert not any(p["connect"] for p in actual["points"] if p["year"] == 2001)
            if chart == "language-doctorates":
                assert not any(p["connect"] for p in actual["points"] if p["year"] == 2010)


def test_language_population_validation_rejects_global_to_us_relabelling_and_bad_subsets():
    data = published()
    data["series"]["ap:global_total"]["scope"] = "U.S.-only learners"
    with pytest.raises(ValueError, match="relabelled"):
        validate_language_data(data)
    data = published()
    next(r for r in data["observations"] if r["series_id"] == "ap:standard_group")["value"] = 10 ** 8
    with pytest.raises(ValueError, match="subset"):
        validate_language_data(data)
    data = published()
    next(r for r in data["observations"] if (r["series_id"], r["year"]) == (
        "ap:japanese_standard_group", 2025))["value"] = 10000
    with pytest.raises(ValueError, match="subset"):
        validate_language_data(data)
    data = published()
    data["observations"].append(copy.deepcopy(data["observations"][0]))
    with pytest.raises(ValueError, match="Duplicate"):
        validate_language_data(data)
    data = published()
    data["context"]["nexus_response_scope"]["us_program_responses"] += 1
    with pytest.raises(ValueError, match="geographies"):
        validate_language_data(data)


def test_zero_baselines_missing_rounds_and_doctoral_breaks_are_not_silently_connected():
    data = published()
    next(r for r in data["observations"] if (r["series_id"], r["year"]) == ("mla:chinese", 2013))["value"] = 0
    assert chart_points(data, "college-benchmarks")[0]["points"] == []
    data = published()
    next(r for r in data["observations"] if (r["series_id"], r["year"]) == ("mla:chinese", 2016)).update(
        value=None, status="missing")
    doc = ElementTree.fromstring(svg_language_chart(data, "college-enrollments"))
    path = next(n for n in doc.iter() if n.get("class") == "data-line")
    assert path.get("d").count("M") == 2
    doc = ElementTree.fromstring(svg_language_chart(published(), "language-doctorates"))
    path = next(n for n in doc.iter() if n.get("data-series-id") == "ipeds:doctorates" and n.get("class") == "data-line")
    assert path.get("d").count("M") == 3


def test_all_language_exports_and_complete_csv_round_trip(tmp_path):
    data = published()
    for base in ["site/figures/language_learning", "outputs/figures/language_learning"]:
        for theme in ["themarginoferror", "substack"]:
            path = tmp_path / base / theme
            path.mkdir(parents=True)
            for name in RETIRED_CHARTS:
                (path / f"{name}.svg").write_text("Retired chart")
    gallery = defaultdict(list)
    count = build_languages(tmp_path, data, gallery)
    assert count == len(CHARTS)
    assert json.loads((tmp_path / "site/downloads/language-learning.json").read_text()) == data
    downloaded = list(csv.DictReader((tmp_path / "site/downloads/language-observations.csv").open()))
    assert len(downloaded) == len(data["observations"])
    for actual, expected in zip(downloaded, data["observations"]):
        assert all(actual[k] == ("" if v is None else str(v)) for k, v in expected.items())
        assert actual["scope"] == data["series"][expected["series_id"]]["scope"]
        assert json.loads(actual["source_assets"]) == data["sources"][expected["source_id"]]["assets"]
    plotted = list(csv.DictReader((tmp_path / "site/downloads/language-chart-values.csv").open()))
    expected = {(cid, group["id"], p["year"]): p for cid in CHARTS for group in chart_points(data, cid) for p in group["points"]}
    assert len(plotted) == len(expected)
    for r in plotted:
        p = expected[(r["chart_id"], r["series_id"], int(r["year"]))]
        assert float(r["plotted_value"]) == p["plotted_value"]
    figures = tmp_path / "outputs/figures/language_learning"
    for path in figures.glob("*/*.svg"):
        assert path.read_bytes() == (tmp_path / "site/figures/language_learning" / path.relative_to(figures)).read_bytes()
        with Image.open(path.with_suffix(".png")) as image:
            assert image.size == png_dimensions(path.read_text())
        doc = ElementTree.fromstring(path.read_text())
        points = [n for n in doc.iter() if n.get("class") == "language-point"]
        wanted = {(group["id"], p["year"]): p for group in chart_points(data, path.stem) for p in group["points"]}
        assert len(points) == len(wanted)
        for point in points:
            assert float(point.get("data-value")) == wanted[(point.get("data-series-id"), int(point.get("data-year")))]["plotted_value"]
        if path.parent.name == "themarginoferror":
            other = ElementTree.fromstring((path.parent.parent / "substack" / path.name).read_text())
            next(n for n in other.iter() if n.get("class") == "chart-background").set("fill", "#F7F5F0")
            assert ElementTree.tostring(doc) == ElementTree.tostring(other)
    page = BeautifulSoup((tmp_path / "site/language-learning.html").read_text(), "html.parser")
    assert "global, not U.S.-only" in page.header.text
    assert "not the full 2021 census" in page.text
    assert not page.select("#ap-language-standard-group")
    assert "standard-group" not in page.select_one("#ap-chinese-global").text
    assert not list(tmp_path.rglob("ap-language-standard-group.*"))
    before = {p.relative_to(tmp_path): p.read_bytes() for p in tmp_path.rglob("*") if p.is_file()}
    build_languages(tmp_path, data, defaultdict(list))
    assert before == {p.relative_to(tmp_path): p.read_bytes() for p in tmp_path.rglob("*") if p.is_file()}
    data["observations"] = [r for r in data["observations"] if not r["series_id"].startswith("ap:")]
    custom = figures / "substack/custom.svg"
    custom.write_text("User figure")
    build_languages(tmp_path, data, defaultdict(list))
    assert not list(figures.glob("*/ap-chinese-global.*"))
    assert custom.read_text() == "User figure"


def test_invalid_language_input_cannot_partially_rewrite_publication_and_removed_input_withdraws_exports(tmp_path):
    source = tmp_path / "data/site"
    source.mkdir(parents=True)
    canonical = json.loads((ROOT / "data/site/observations.json").read_text())
    (source / "observations.json").write_text(json.dumps(canonical[:1]))
    language_path = source / "language_learning.json"
    data = published()
    data["series"]["mla:korean"]["unit"] = "awards"
    language_path.write_text(json.dumps(data))
    (tmp_path / "site").mkdir()
    (tmp_path / "site/index.html").write_text("Last valid publication")
    before = {p.relative_to(tmp_path): p.read_bytes() for p in tmp_path.rglob("*") if p.is_file()}
    with pytest.raises(ValueError, match="units"):
        build(tmp_path)
    assert before == {p.relative_to(tmp_path): p.read_bytes() for p in tmp_path.rglob("*") if p.is_file()}
    data = published()
    next(r for r in data["observations"] if r["series_id"] == "schools:k12_estimate").update(value=None, status="missing")
    language_path.write_text(json.dumps(data))
    assert build(tmp_path)["language_charts"] == len(CHARTS)
    custom = tmp_path / "outputs/figures/language_learning/substack/custom.svg"
    custom.write_text("User-owned")
    language_path.unlink()
    assert build(tmp_path)["language_charts"] == 0
    assert not (tmp_path / "site/language-learning.html").exists()
    assert not list((tmp_path / "site/downloads").glob("language-*"))
    assert not list((tmp_path / "site/figures/language_learning").glob("*/*.png"))
    assert list((tmp_path / "outputs/figures/language_learning").glob("*/*.svg")) == [custom]


def zipped_csv(name, rows, bom=False, extra=None):
    stream = io.StringIO()
    writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
    writer.writeheader()
    writer.writerows(rows)
    content = stream.getvalue().encode("utf-8-sig" if bom else "cp1252")
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as z:
        z.writestr(name, content)
        if extra:
            z.writestr(extra, content.replace(b"5", b"9"))
    return archive.getvalue()


@pytest.mark.parametrize("year", [2000, 2022, 2025])
def test_ipeds_filters_geography_second_majors_revisions_and_legacy_totals(year):
    legacy = year == 2000
    hd = [
        dict(unitid="1", stabbr="CA", deggrant="1", pset4flg="1"),
        dict(unitid="2", stabbr="PR", deggrant="1", pset4flg="1"),
        dict(unitid="3", stabbr="CA", deggrant="1", pset4flg="2"),
    ]
    records = []
    for unit, major, value in [("1", "1", "5"), ("1", "2", "99"), ("2", "1", "99"), ("3", "1", "99")]:
        r = dict(unitid=unit, cipcode="16.0301", awlevel="05")
        if legacy:
            if major == "2":
                continue
            r.update(crace15="2" if unit == "1" else value, crace16="3" if unit == "1" else value)
        else:
            r.update(majornum=major, ctotalt=value, xctotalt="R")
        records.append(r)
    for cip, value in [("16.0302", 7), ("16.0303", 11), ("16.0300", 999)]:
        r = dict(unitid="1", cipcode=cip, awlevel="05")
        if legacy:
            r.update(crace15=str(value), crace16="0")
        else:
            r.update(majornum="1", ctotalt=str(value), xctotalt="R")
        records.append(r)
    directory = zipped_csv("hd.csv", hd, bom=True)
    completions = zipped_csv("c_a.csv" if legacy else "c_a_rv.csv", records,
                             extra=None if legacy else "c_a.csv")

    def get(url):
        if year == 2025:
            assert url.startswith("https://nces.ed.gov/ipeds/complete-data-files/")
        body = directory if ("HD" in url or "hd.zip" in url) else completions
        return body, dict(url=url, request_url=url, sha256=digest(body), retrieved_at="2026-10-22T14:00:00+00:00")

    _, observations, sources = language_sources.ipeds(SimpleNamespace(get=get), [year])
    bachelor = next(r for r in observations if r["series_id"] == "ipeds:bachelors")
    assert bachelor["value"] == 5
    assert next(r for r in observations if r["series_id"] == "ipeds_japanese:bachelors")["value"] == 7
    assert next(r for r in observations if r["series_id"] == "ipeds_korean:bachelors")["value"] == 11
    assert all(r["value"] != 999 for r in observations)
    assert "no_major_split" in bachelor["comparability"] if legacy else True
    assert next(r for r in observations if r["series_id"] == "ipeds:masters")["value"] is None
    assert "CRACE15+CRACE16" in sources[f"ipeds_{year}"]["locator"] if legacy else "_rv.csv" in sources[f"ipeds_{year}"]["locator"]
    if not legacy:
        completions = zipped_csv("c_a.csv", [{k: v for k, v in r.items() if k != "majornum"} for r in records])
        with pytest.raises(ValueError, match="first-major"):
            language_sources.ipeds(SimpleNamespace(get=get), [year])


def test_ap_parser_does_not_confuse_total_with_passing_or_standard_counts(monkeypatch):
    text = ("Global AP Chinese Language and Culture May 2025\nTotal Group\nNumber of Students 18,518 -\n"
            "3 or Higher / % 16,518 89.2\nStandard Group**\nNumber of Students 4,777 -")
    monkeypatch.setattr(language_sources, "pdf_text", lambda body: [text])
    assert language_sources.parse_ap_chinese(b"fixture", 2025) == (18518, 4777)
    with pytest.raises(ValueError, match="geography/year"):
        language_sources.parse_ap_chinese(b"fixture", 2024)
    monkeypatch.setattr(language_sources, "pdf_text", lambda body: [text.replace("Chinese", "Japanese")])
    with pytest.raises(ValueError, match="subject"):
        language_sources.parse_ap_chinese(b"fixture", 2025)
    assert language_sources.parse_ap_language(b"fixture", 2025, "Japanese") == (18518, 4777)


def test_ap_history_ignores_empty_participation_counts_in_percentage_only_years():
    header = "<tr>" + "".join(f"<th>{s}</th>" for s in ["Year", "5", "4", "3", "2", "1", "3+", "Test Takers", "Mean Score"]) + "</tr>"
    def table(group, count):
        return (f"<table><caption>Past AP Japanese Language and Culture Score Distributions ({group} Group)</caption>" +
                header + "<tr>" + "".join(f"<td>{s}</td>" for s in ["2025", "43%", "11%", "20%", "8%", "18%", "74%", count, "3.55"]) +
                "</tr><tr>" + "".join(f"<td>{s}</td>" for s in ["2026", "47%", "10%", "15%", "7%", "21%", "72%", "", ""]) + "</tr></table>")
    body = (table("Total", "3,269") + table("Standard", "1,446")).encode()
    assert language_sources.parse_ap_history(body, "Japanese") == {"total": {2025: 3269}, "standard": {2025: 1446}}
    with pytest.raises(ValueError, match="tables missing"):
        language_sources.parse_ap_history(body, "Chinese")


def test_ap_language_benchmarks_have_matching_years_populations_and_colours():
    data = published()
    for chart_id in ["ap-chinese-global"]:
        groups = chart_points(data, chart_id)
        assert len(groups) == 2
        assert [g["color"] for g in groups] == ["#7B2D26", "#3D6F8C"]
        assert all({p["year"] for p in g["points"]} == set(range(2007, 2026)) for g in groups)
        doc = ElementTree.fromstring(svg_language_chart(data, chart_id))
        years = [int(n.text) for n in doc.iter() if n.get("class") == "year-tick"]
        assert years[0] == 2007 and years[-1] == 2025 and 2026 not in years
    values = {(r["series_id"], r["year"]): r["value"] for r in data["observations"]}
    assert values[("ap:japanese_total", 2025)] == 3269
    assert values[("ap:japanese_standard_group", 2025)] == 1446
    assert not any(s.startswith("ap:") and "korean" in s for s in data["series"])


def test_us_ap_cohorts_exclude_standard_group_counts_and_use_later_reviewed_vintages(monkeypatch):
    heading = "Chinese Language\nand Culture\nScore Distributions\nU.S. public schools only; graduating classes\n"
    tables = {
        2011: heading + "".join(f"{score} N/A N/A {10 - i} / 20.0% {20 - i} / 20.0%\n"
                                for i, score in enumerate("54321")) +
              "".join(f"{score} N/A N/A 1 / 20.0% 2 / 20.0%\n" for score in "54321"),
        2013: heading + "".join(f"{year}\n" + "".join(f"{start+i}\n" for i in range(5))
                                for year, start in [(2007, 1), (2011, 11), (2012, 21)]),
        2014: heading + "".join(f"{score} — — 5 20.0% 30 20.0% 40 20.0%\n" for score in "54321"),
    }
    monkeypatch.setattr(language_sources, "pdf_text", lambda body: [
        "Contents: Chinese Language and Culture", tables[int(body)]])

    def get(url):
        year = 2011 if "2011" in url else 2013 if "9th-annual" in url else 2014
        body = str(year).encode()
        return body, dict(url=url, request_url=url, sha256=digest(body), retrieved_at="2026-10-22T14:00:00+00:00")

    definitions, observations, _ = language_sources.ap_us_cohorts(SimpleNamespace(get=get))
    assert definitions["ap_cohort:us_public"]["period_basis"] == "graduating class"
    assert {r["year"]: r["value"] for r in observations} == {
        2007: 15, 2008: 25, 2009: 40, 2010: 90, 2011: 65, 2012: 150, 2013: 200}


def test_mla_matched_panel_keeps_explicit_zeros_but_excludes_nonresponse_and_new_entrants(monkeypatch):
    headers = ["INST_ID", "SRVY_YEAR", "TERM", "LANGUAGE", "ALL LEVEL TOTAL", "N_RESP", "ZERO_ERL"]
    rows = []
    for year, institution, counts in [
        (2013, 1, [10, 2, 1, 20]), (2013, 2, [5, 0, 0, 5]),
        (2021, 1, [8, 3, 2, 16]), (2021, 3, [4, 0, 0, 0]), (2021, 4, [100, 0, 0, 0]),
        (2013, 6, [50, None, 0, 0]), (2021, 6, [30, 2, 0, 0]),
    ]:
        rows += [[institution, year, "Fall", language, count, None, None]
                 for language, count in zip(["CHINESE", "JAPANESE", "KOREAN", "SPANISH"], counts)]
    rows += [[3, 2013, "Fall", None, 0, None, "Z"], [2, 2021, "Fall", None, 0, "NR", None]]
    report = "Table 5 Fall Enrollments in the Fifteen\n"
    for label, lang in [("Mandarin", "CHINESE"), ("Japanese", "JAPANESE"), ("Korean", "KOREAN"),
                        ("Spanish", "SPANISH"), ("Total", None)]:
        values = [0] * 10
        for pos, year in [(7, 2013), (9, 2021)]:
            values[pos] = sum(r[4] for r in rows if r[1] == year and r[5] != "NR" and r[4] is not None
                              and (lang is None or r[3] == lang))
        report += label + " " + " ".join(map(str, values + [0, 0])) + "\n"
    monkeypatch.setattr(language_sources, "pdf_text", lambda body: [report])
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "FL history"
    sheet.append(headers)
    for r in rows:
        sheet.append(r)
    stream = io.BytesIO()
    workbook.save(stream)

    def get(url):
        body = stream.getvalue() if url.endswith(".xlsx") else b"PDF"
        return body, dict(url=url, request_url=url, sha256=digest(body), retrieved_at="2026-10-22T14:00:00+00:00")

    definitions, observations, _ = language_sources.mla(SimpleNamespace(get=get))
    assert "Matched set of 2" in definitions["mla_matched:chinese"]["scope"]
    assert [(r["year"], r["value"]) for r in observations if r["series_id"] == "mla_matched:chinese"] == [(2013, 10), (2021, 12)]
