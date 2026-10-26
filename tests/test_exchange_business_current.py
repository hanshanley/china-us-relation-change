"""Current Business evidence keeps survey bases, periods and money distinct."""

import copy
import csv
import io
import json
from collections import defaultdict
from pathlib import Path

import pytest
from bs4 import BeautifulSoup
from openpyxl import Workbook

from us_china_exchange import business_current


ROOT = Path(__file__).resolve().parents[1]


def current():
    return json.loads((ROOT / "data/site/business_current.json").read_text())


@pytest.mark.parametrize("case", ["unchanged", "no_uk", "current_uk", "missing_uk"])
def test_uk_coverage_keeps_historical_missingness_without_a_stale_or_duplicate_chart(case):
    archive = json.loads((ROOT / "data/site/business_presence.json").read_text())
    if case == "no_uk":
        archive["observations"] = [r for r in archive["observations"] if r["series_id"] != "eurostat_uk:china"]
    elif case == "current_uk":
        prior = next(r for r in archive["observations"] if r["series_id"] == "eurostat_uk:china" and r["year"] == 2019)
        archive["observations"].append(dict(prior, year=2023))
    elif case == "missing_uk":
        for row in archive["observations"]:
            if row["series_id"] == "eurostat_uk:china":
                row.update(value=None, status="missing")
    before = copy.deepcopy(archive)
    markup = business_current.uk_historical_coverage(archive)
    assert archive == before
    if case in {"no_uk", "current_uk"}:
        assert markup == ""
    else:
        page = BeautifulSoup(markup, "html.parser")
        assert len(page.select("tbody tr")) == 14
        assert not page.select("svg,img,.plot-data,.comparison-chart")
        if case == "missing_uk":
            assert "No UK enterprise counts are reported" in page.text
            assert all(cell["data-uk-value"] == "" for cell in page.select("[data-uk-value]"))
        else:
            assert page.select_one('[data-uk-year="2019"] [data-uk-value]').text == "531"
            assert page.select_one('[data-uk-year="2007"] [data-uk-value]').text == "Not reported"
            assert [int(row["data-uk-year"]) for row in page.select('[data-uk-break="true"]')] == [
                2010, 2013, 2014, 2015, 2018]


def test_position_parser_preserves_direction_units_and_missingness():
    book = Workbook()
    sheet = book.active
    sheet.title = "2025"
    for row in [
        ["U.S. Direct Investment Position Abroad on a Historical-Cost Basis by Account for Selected Countries, 2025"],
        ["[Millions of dollars]"], [None, "Total"], ["China", 120547], ["Hong Kong", 103737],
    ]:
        sheet.append(row)

    def parse():
        stream = io.BytesIO()
        book.save(stream)
        return business_current.parse_positions(stream.getvalue())

    assert [(r["series_id"], r["value"]) for r in parse()] == [
        ("bea_position:CHN", 120547000000), ("bea_position:HKG", 103737000000)]
    sheet.cell(4, 2, "(D)")
    with pytest.raises(ValueError, match="suppressed"):
        parse()
    sheet.cell(4, 2, -42)
    assert parse()[0]["value"] == -42000000
    sheet.cell(1, 1, "Foreign Direct Investment Position in the United States, 2025")
    with pytest.raises(ValueError, match="Unexpected BEA"):
        parse()


def test_survey_bases_performance_years_and_investment_units_are_not_interchangeable():
    data = business_current.validate(current())
    assert data["surveys"]["amcham_2026"]["overall_n"] == 262
    assert data["surveys"]["amcham_2026"]["item_n"] is None
    assert data["surveys"]["amcham_2026"]["field_end"] is None
    german = data["surveys"]["ahk_2026"]
    assert (german["item_n"], german["overall_n"], german["field_end"], german["preliminary"]) == (
        381, 488, "2026-09-29", True)
    profit = next(r for r in data["observations"] if r["series_id"] == "amcham:profitable")
    assert profit["year"] == 2025
    assert data["sources"][profit["source_id"]]["release_date"] == "2026-09-10"
    for mutate, message in [
        (lambda d: d["surveys"]["ahk_2026"].update(item_n=489), "denominator"),
        (lambda d: d["series"]["bea_position:CHN"].update(unit="enterprises"), "counts or flows"),
        (lambda d: d["observations"][-1].update(value=101), "percentage"),
        (lambda d: d["observations"][0].update(year=2026), "reference year"),
        (lambda d: d["observations"][0].pop("raw_sha256"), "workbook provenance"),
    ]:
        invalid = copy.deepcopy(data)
        mutate(invalid)
        with pytest.raises(ValueError, match=message):
            business_current.validate(invalid)


def test_full_position_history_preserves_reported_years_units_and_each_workbook():
    data = current()
    for geo, first, last in [("CHN", 11140, 120547), ("HKG", 27447, 103737)]:
        rows = sorted((r for r in data["observations"] if r["series_id"] == "bea_position:" + geo),
                      key=lambda r: r["year"])
        assert [r["year"] for r in rows] == list(range(2000, 2026))
        assert [rows[0]["value"], rows[-1]["value"]] == [first * 1_000_000, last * 1_000_000]
        assert len({r["raw_sha256"] for r in rows}) == 3
        assert all(r["source_url"].startswith("https://apps.bea.gov/")
                   and r["retrieved_at"] and "cell B" in r["notes"] for r in rows)
    book = Workbook()
    sheet = book.active
    sheet.title = "2000"
    for row in [
        ["U.S. Direct Investment Position Abroad on a Historical-Cost Basis, 2000"],
        ["[Millions of dollars]"], [None, "All industries"], ["China", "11,140"], ["Hong Kong", "-27,447"],
    ]:
        sheet.append(row)
    book.create_sheet("Notes")
    stream = io.BytesIO()
    book.save(stream)
    assert [r["value"] for r in business_current.parse_positions(stream.getvalue(), all_industries=True)] == [
        11140000000, -27447000000]
    sheet.cell(4, 2, "10-20")
    stream = io.BytesIO()
    book.save(stream)
    with pytest.raises(ValueError, match="suppressed or ranged"):
        business_current.parse_positions(stream.getvalue(), all_industries=True)


def test_component_exports_complete_counts_and_current_evidence_without_stale_charts(tmp_path):
    data = current()
    folder = tmp_path / "data/site"
    folder.mkdir(parents=True)
    (folder / "business_current.json").write_text(json.dumps(data))
    archive = json.loads((ROOT / "data/site/business_presence.json").read_text())
    stale = tmp_path / "site/figures/business_presence/themarginoferror/eu-affiliates-china.svg"
    stale.parent.mkdir(parents=True)
    stale.write_text("Rejected stale chart")
    gallery = defaultdict(list)
    assert business_current.build_current_business(tmp_path, archive, gallery) == 26
    assert not stale.exists()
    page = BeautifulSoup((tmp_path / "site/business-presence.html").read_text(), "html.parser")
    assert not page.select("#business-controls,#annual-values,#eu-affiliates-china")
    assert len(page.select(".business-card")) == 26
    count_cards = page.select(".business-count-card")
    assert len(count_cards) == 22
    assert not page.select("#business-count-references")
    assert page.select("#affiliate-presence .business-count-card") == count_cards
    assert all(card.find_parent("details") is None for card in count_cards)
    assert "Earlier reference" not in page.text
    for cid, period, source in (
        ("us-affiliates-china", "2000-2023", "BEA"),
        ("germany-affiliates-china", "2015-2024", "Bundesbank"),
        ("germany-affiliates-china-eurostat", "2007-2023", "Eurostat"),
        ("norway-affiliates-china", "2008", "Statistics Norway"),
    ):
        card = page.select_one("#" + cid)
        assert card.find_parent(id="affiliate-presence")
        assert period in card.select_one(".chart-context").text
        assert source in card.select_one(".chart-context").text
    for card in count_cards:
        figure = tmp_path / "site" / card.img["src"].replace(".png", ".svg")
        svg = BeautifulSoup(figure.read_text(), "html.parser")
        assert svg.select_one(".annual-column") and not svg.select(".data-line")
        if card["id"] == "luxembourg-affiliates-china":
            assert len(svg.select(".annual-column")) == 15
        if card["id"] == "hungary-affiliates-china":
            assert len(svg.select(".annual-column")) == 2
    assert not page.select("#uk-affiliates-china,#eu-affiliates-china,#us-parent-companies-china")
    uk = page.select_one("#uk-business-coverage")
    assert uk is not None and uk.find_parent(id="affiliate-presence")
    assert "United Kingdom" in uk.h3.text and "historical" in uk.h3.text.lower()
    assert "2019" in uk.select_one(".chart-context").text
    assert not uk.select("svg,img,.comparison-chart,.plot-data")
    uk_rows = sorted((r for r in archive["observations"] if r["series_id"] == "eurostat_uk:china"),
                     key=lambda r: r["year"])
    assert [(int(row["data-uk-year"]), row.select_one("[data-uk-value]")["data-uk-value"])
            for row in uk.select("tbody tr")] == [
                (row["year"], "" if row["value"] is None else str(row["value"])) for row in uk_rows]
    assert [int(row["data-uk-year"]) for row in uk.select("tbody tr")
            if row.select_one("[data-uk-value]").text == "Not reported"] == [2007, 2008, 2020]
    assert {a["href"] for a in uk.select("a[download]")} == {
        "downloads/business-presence.json", "downloads/business-observations.csv"}
    assert "agriculture" in uk.text and "like-for-like" in uk.text
    assert json.loads((tmp_path / "site/downloads/business-current.json").read_text()) == data
    assert json.loads((tmp_path / "site/downloads/business-presence.json").read_text()) == archive
    records = list(csv.DictReader((tmp_path / "site/downloads/business-current-observations.csv").open()))
    assert len(records) == len(data["observations"]) == 59
    assert all(json.loads(row["source_assets"]) and row["source_locator"] for row in records)
    historical = list(csv.DictReader((tmp_path / "site/downloads/business-observations.csv").open()))
    assert len(historical) == len(archive["observations"])
    assert all(row["definition_notes"] and json.loads(row["source_provenance"]) for row in historical)
    for cid in business_current.CHARTS:
        svg = BeautifulSoup(business_current.svg_current_chart(data, cid), "html.parser")
        if cid.startswith("us-investment"):
            assert svg.svg["data-plot-style"] == "columns"
            assert len(svg.select(".annual-column")) == 26 and not svg.select(".data-line")
        else:
            assert svg.svg["data-axis-max"] == "100"
            for bar in svg.select(".business-point"):
                assert float(bar["width"]) == pytest.approx(float(bar["data-value"]) * 8.4)
