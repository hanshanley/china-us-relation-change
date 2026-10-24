import io
import copy
import json
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

import pytest
from bs4 import BeautifulSoup
from openpyxl import Workbook

from us_china_exchange.model import change
from us_china_exchange.pipeline import refresh
from us_china_exchange.render import svg_chart
from us_china_exchange import tourism_backfill
from us_china_exchange.tourism_backfill import (
    parse_archived_macau_services, parse_dsec_visitors, parse_early_hong_kong, parse_early_macau, parse_hong_kong_profile,
    parse_macau_months, parse_top_ten,
)


ROOT = Path(__file__).resolve().parents[1]


def test_reviewed_hktb_months_reconcile_stay_types_and_successive_cumulative_totals():
    content = (ROOT / "data/reviewed/hktb-monthly-2026-08.json").read_bytes()
    prior = json.loads((ROOT / "data/site/observations.json").read_text())
    rows, _, notes = tourism_backfill.reviewed_hktb_months(content, prior, "2026-10-24T16:00:00+00:00")
    assert len(rows) == 8
    latest = next(r for r in rows if r["frequency"] == "ytd" and r["period"] == "2026-08")
    assert latest["value"] == 737530 and latest["comparability"] == "hktb:us_nationality:all_mode"
    assert latest["raw_sha256"] == tourism_backfill.digest(content)
    assert latest["source_url"] == tourism_backfill.HK_REVIEWED_INPUT
    assert "transcription, not publisher PDFs" in notes
    data = json.loads(content)
    data["reports"][0]["overnight"][0] += 1
    with pytest.raises(ValueError, match="overnight"):
        tourism_backfill.reviewed_hktb_months(json.dumps(data).encode(), prior, latest["retrieved_at"])
    data = json.loads(content)
    data["reports"][0]["total"][0] += 1
    data["reports"][0]["overnight"][0] += 1
    with pytest.raises(ValueError, match="preceding"):
        tourism_backfill.reviewed_hktb_months(json.dumps(data).encode(), prior, latest["retrieved_at"])
    changed = copy.deepcopy(prior)
    next(r for r in changed if r["series_id"] == latest["series_id"]
         and r["frequency"] == "ytd" and r["period"] == "2026-06")["value"] += 1
    with pytest.raises(ValueError, match="preceding"):
        tourism_backfill.reviewed_hktb_months(content, changed, latest["retrieved_at"])


def test_archived_services_require_reviewed_source_and_two_independent_controls(monkeypatch):
    profile = (
        "Macau Trade Summary "
        "U.S. goods and services trade with Macau totaled an estimated $864.0 million in 2024. "
        "U.S. total goods trade (exports plus imports) with Macau was an estimated $390 million in 2024. "
        "U.S. total services trade (exports plus imports) with Macau totaled an estimated $474 billion in 2024. "
        "U.S. services exports to Macau in 2024 were $440 million. "
        "U.S. services imports from Macau in 2024 were $34 million. "
        "U.S. services trade surplus with Macau was $406 million in 2024. Stay in the Know")
    with pytest.raises(ValueError, match="statistical passage"):
        parse_archived_macau_services(profile)
    monkeypatch.setattr(tourism_backfill, "digest", lambda content: (
        "ada54d78f1505dec3823d6f1e8fb35b7ee9b7ce7ea8a4942899ab10e7e2ba642"))
    rows = parse_archived_macau_services(profile)
    assert {(r["origin"], r["destination"], r["year"]): r["value"] for r in rows} == {
        ("USA", "MAC", 2024): 440000000, ("MAC", "USA", 2024): 34000000}
    assert all(r["status"] == "estimated" and "$474 billion" in r["notes"] for r in rows)
    for invalid in (profile.replace("$390 million", "$391 million"),
                    profile.replace("$406 million", "$500 million"),
                    profile.replace("$474 billion", "$475 billion")):
        with pytest.raises(ValueError):
            parse_archived_macau_services(invalid)


def test_archival_services_are_not_compared_with_a_newer_revision_vintage():
    rows = json.loads((ROOT / "data/site/observations.json").read_text())
    for origin in ("USA", "MAC"):
        selected = {r["year"]: r for r in rows if r["source_id"] == "ustr" and r["origin"] == origin}
        assert set(selected) == {2024, 2025}
        assert change(selected[2025], selected[2024])["percent"] is None
        assert "/20251218202809/" in selected[2024]["source_url"]
        chart = BeautifulSoup(svg_chart(list(selected.values())), "html.parser")
        assert chart.select_one(".data-line")["d"].count("M") == 2
        assert "L" not in chart.select_one(".data-line")["d"]
        bars = chart.select(".snapshot-bar")
        assert len(bars) == 2
        assert all(115 < float(bar["x"]) < 950 - float(bar["width"]) for bar in bars)
        footer = " ".join(n.text for n in chart.select(".figure-source"))
        assert "Release-vintage" in footer and "Population-definition" not in footer


def test_manual_services_review_is_stable_under_normal_daily_refresh(tmp_path, monkeypatch):
    rows = [r for r in json.loads((ROOT / "data/site/observations.json").read_text()) if r["source_id"] == "ustr"]
    current = [dict(r) for r in rows if r["year"] == 2025]
    archive = [dict(r) for r in rows if r["year"] == 2024]
    directory = tmp_path / "data/site"
    directory.mkdir(parents=True)
    (directory / "observations.json").write_text(json.dumps(rows))
    (directory / "sources.json").write_text("{}")
    normal = lambda client: (current, "Current USTR vintage", "Current source controls")
    monkeypatch.setattr(tourism_backfill, "ustr_macau_services", normal)
    monkeypatch.setattr(tourism_backfill, "parse_archived_macau_services", lambda content: [dict(r) for r in archive])
    first = archive[0]
    client = SimpleNamespace(now="2026-10-24T12:00:00+00:00", get=lambda url: (
        b"reviewed", {"url": first["source_url"], "sha256": first["raw_sha256"],
                      "retrieved_at": first["retrieved_at"]}))
    refresh(tmp_path, adapters={"ustr": tourism_backfill.macau_services_history}, client=client)
    before = (directory / "observations.json").read_bytes()
    report = refresh(tmp_path, adapters={"ustr": normal}, client=client)
    assert report["publication_changed"] is False
    assert (directory / "observations.json").read_bytes() == before
    assert json.loads(before) == rows


def test_early_macau_uses_residence_totals_not_mode_counts_shares_or_nationality():
    header = ("TABLE 1 : VISITOR ARRIVALS BY PLACE OF RESIDENCE\n"
              "2001 2002 2001 2002 2001 2002 2001 2002")
    row = ("US  51 379  50 170  16 998  20 850  5 230  5 884  73 607  0.72  76 904  0.67  4.48\n"
           "Estados Unidos da America\nUnited States of America")
    assert parse_early_macau(header, row) == {2001: 73607, 2002: 76904}
    with pytest.raises(ValueError, match="residence table"):
        parse_early_macau(header.replace("PLACE OF RESIDENCE", "NATIONALITY"), row)
    with pytest.raises(ValueError, match="reconcile"):
        parse_early_macau(header, row.replace("76 904", "86 648"))
    with pytest.raises(ValueError, match="counts"):
        parse_early_macau(header, row.replace("50 170", "50 17"))
    with pytest.raises(ValueError, match="U.S. residence"):
        parse_early_macau(header, row.replace("United States of America", "Canada"))


def test_early_hong_kong_preserves_seven_digit_count_and_residence_columns():
    plain = (
        "Visitor Arrivals by Country/Territory of Residence\n"
        "1992 1997 1998 1999 2000 2001 2002\n"
        "United States of 694 290 861 217 828 326 858 925 966 008 935 717 1 000 844\nAmerica")
    assert parse_early_hong_kong(plain) == {2000: 966008, 2001: 935717, 2002: 1000844}
    with pytest.raises(ValueError, match="digit grouping"):
        parse_early_hong_kong(plain.replace("1 000 844", "1 844"))
    with pytest.raises(ValueError, match="year columns"):
        parse_early_hong_kong(plain.replace("2002", "2003"))


def test_hong_kong_profile_reads_hong_kong_not_macau_or_package_tour_counts():
    plain = (
        "MARKET PROFILE : U.S.A.\nTotal i) 171 1,777 93,814 147,941 162,460 82,121 7.4%\n"
        "Arrivals to Hong Kong 1,365 25,987 596,958 884,261 1,026,030 552,528 12.9%\n"
        "Total package tours 0 0 517 2,688 2,765 1,103")
    assert parse_hong_kong_profile(plain)[-1] == 552528
    with pytest.raises(ValueError, match="another country"):
        parse_hong_kong_profile(plain.replace("U.S.A.", "CANADA"))
    with pytest.raises(ValueError, match="six-column"):
        parse_hong_kong_profile(plain.replace(" 552,528", " unavailable"))


def test_macau_monthly_and_ytd_columns_allow_different_source_market_ranks():
    halves = [
        "2026 Aug\nOther market 30,738\nRepublic of Korea\n"
        "Country label 11,713\nUnited States of America\n(-0.4)\n0.3",
        "2026 Jan - Aug\nOther market 143,174\nThailand\n"
        "Other market 110,543\nIndonesia\nCountry label 109,693\nUnited States of America\n(+5.5)\n0.4"]
    assert parse_top_ten(halves, 2026, 8) == {"monthly": 11713, "ytd": 109693}
    missing = [halves[0].replace("United States of America", "Another country"), halves[1]]
    assert parse_top_ten(missing, 2026, 8) == {"ytd": 109693}
    with pytest.raises(ValueError, match="period caption"):
        parse_top_ten(halves[::-1], 2026, 8)
    with pytest.raises(ValueError, match="exceed"):
        parse_top_ten([halves[0].replace("11,713", "200,000"), halves[1]], 2026, 8)
    with pytest.raises(ValueError, match="ambiguous"):
        parse_top_ten([halves[0] + "\n1\nUnited States of America", halves[1]], 2026, 8)


def test_dsec_country_workbook_extends_past_stale_top_ten_report_with_two_controls():
    book = Workbook()
    book.remove(book.active)
    totals = {
        1: ("VISITOR ARRIVALS", [8322, 10378, 112330, 120071]),
        2: ("VISITOR ARRIVALS BY SEA AND", [2344, 3135, 31690, 33974]),
        3: ("VISITOR ARRIVALS BY LAND AND", [4850, 5931, 67141, 71980]),
        4: ("VISITOR ARRIVALS BY AIR AND", [1128, 1312, 13499, 14117]),
        5: ("SAME-DAY VISITOR ARRIVALS", [4006, 5193, 56096, 61657]),
        9: ("OVERNIGHT VISITOR ARRIVALS", [4316, 5185, 56234, 58414]),
    }
    for number, (title, values) in totals.items():
        sheet = book.create_sheet(str(number))
        sheet.cell(3, 1, title + " BY ISSUING PLACE OF TRAVEL DOCUMENT")
        sheet.cell(5, 7, "Jan. - Sept.")
        sheet.cell(10, 3, "USA")
        for column, year, value in zip((4, 5, 7, 9), (2025, 2026, 2025, 2026), values):
            sheet.cell(6, column, datetime(year, 9, 1))
            sheet.cell(9, column, "No.")
            sheet.cell(10, column, value)
        sheet.cell(10, 6, 24.7)
        sheet.cell(10, 8, .4)

    def content():
        stream = io.BytesIO()
        book.save(stream)
        return stream.getvalue()

    rows = parse_dsec_visitors(content(), 2026, 9)
    assert {(r["year"], r["frequency"]): r["value"] for r in rows} == {
        (2025, "monthly"): 8322, (2025, "ytd"): 112330,
        (2026, "monthly"): 10378, (2026, "ytd"): 120071}
    assert all(r["period"].endswith("-09") and "travel document" in r["definition"] for r in rows)
    with pytest.raises(ValueError, match="periods"):
        parse_dsec_visitors(content(), 2026, 8)
    book["2"].cell(10, 4, 2345)
    with pytest.raises(ValueError, match="reconcile"):
        parse_dsec_visitors(content(), 2026, 9)
    book["2"].cell(10, 4, 2344)
    book["9"].cell(10, 9, 58415)
    with pytest.raises(ValueError, match="reconcile"):
        parse_dsec_visitors(content(), 2026, 9)
    published = json.loads((ROOT / "data/site/observations.json").read_text())
    current = next(r for r in published if r["source_id"] == "macau_tourism"
                   and r["period"] == "2026-09" and r["frequency"] == "ytd")
    assert current["value"] == 120071 and current["source_url"].startswith("https://www.dsec.gov.mo/getAttachment/")


def test_hong_kong_recovery_keeps_population_break_and_matched_halves():
    rows = json.loads((ROOT / "data/site/observations.json").read_text())
    hk = [r for r in rows if r["source_id"] == "hktb"]
    annual = {r["year"]: r for r in hk if r["frequency"] == "annual"}
    assert set(annual) == set(range(2000, 2026))
    assert annual[2002]["value"] == 1000844
    assert annual[2023]["value"] == 596958
    assert change(annual[2025], annual[2020])["reason"] == "Definitions or periods differ"
    chart = BeautifulSoup(svg_chart(hk), "html.parser")
    assert "Population-definition breaks: 2023" in chart.get_text(" ", strip=True)
    assert chart.select_one(".data-line")["d"].count("M") == 2
    assert "Definitions differ" in svg_chart(hk, indexed=True)
    partial = {r["year"]: r for r in hk if r["frequency"] == "ytd" and r["period"].endswith("-06")}
    assert {y: r["value"] for y, r in partial.items()} == {2025: 489417, 2026: 552528}
    assert change(partial[2026], partial[2025])["absolute"] == 63111
    latest = {r["year"]: r for r in hk if r["frequency"] == "ytd" and r["period"].endswith("-08")}
    assert {y: r["value"] for y, r in latest.items()} == {2025: 659003, 2026: 737530}
    assert change(latest[2026], latest[2025])["absolute"] == 78527


def test_complete_macau_profile_fills_a_top_ten_omission_without_subtraction():
    plain = "VISITOR ARRIVALS FROM U.S.A.\n" + "\n".join(
        month + " 3,409 6,427 1,405 11,241 12.1% 1,214"
        for month in ("Jan", "Feb", "Mar", "Apr", "May", "Jun"))
    assert parse_macau_months(plain)[5] == 11241
    with pytest.raises(ValueError, match="reconcile"):
        parse_macau_months(plain.replace("May 3,409", "May 3,410"))
    rows = json.loads((ROOT / "data/site/observations.json").read_text())
    may = next(r for r in rows if r["source_id"] == "macau_tourism"
               and r["period"] == "2025-05" and r["frequency"] == "monthly")
    assert may["value"] == 11241 and "monthly profile" in may["source_vintage"]


def test_macau_latest_month_and_matching_cumulative_counts_are_not_confused():
    rows = json.loads((ROOT / "data/site/observations.json").read_text())
    selected = {(r["year"], r["frequency"]): r for r in rows
                if r["source_id"] == "macau_tourism" and r["period"].endswith("-08")}
    assert {k: r["value"] for k, r in selected.items()} == {
        (2025, "monthly"): 11755, (2025, "ytd"): 104008,
        (2026, "monthly"): 11713, (2026, "ytd"): 109693}
    assert change(selected[2026, "ytd"], selected[2025, "ytd"])["absolute"] == 5685
    monthly = {r["period"]: r["value"] for r in rows if r["source_id"] == "macau_tourism"
               and r["year"] == 2025 and r["frequency"] == "monthly"}
    assert set(monthly) == {f"2025-{month:02d}" for month in range(1, 13)}
    december = next(r["value"] for r in rows if r["source_id"] == "macau_tourism"
                    and r["period"] == "2025-12" and r["frequency"] == "ytd")
    annual = next(r["value"] for r in rows if r["source_id"] == "macau_tourism"
                  and r["year"] == 2025 and r["frequency"] == "annual")
    assert sum(monthly.values()) == december == annual
