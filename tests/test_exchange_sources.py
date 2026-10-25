import io
import json
from datetime import datetime
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from us_china_exchange.sources import COUNTRIES, parse_bea, parse_ntto, parse_taiwan
from us_china_exchange import sources
from us_china_exchange import travel_sources
from us_china_exchange.model import change, validate
from us_china_exchange.travel_sources import parse_monitor_card


def test_ustr_services_are_direct_country_flows_not_goods_or_derived_prior_years():
    content = b"""<h1>Macau Trade Summary</h1>
    U.S. goods exports to Macau in 2025 were $208.6 million.
    U.S. services trade (exports plus imports) with Macau totaled an estimated $642.0 million in 2025.
    U.S. services exports to Macau in 2025 were $434.0 million, down 1.6 percent ($7.0 million) from 2024.
    U.S. services imports from Macau in 2025 were $208.0 million, up 511.8 percent ($174.0 million) from 2024.
    The U.S. services trade surplus with Macau was $226 million in 2025."""
    rows, year = sources.parse_ustr_macau_services(content)
    assert year == 2025 and len(rows) == 2
    assert {(r["origin"], r["destination"]): r["value"] for r in rows} == {
        ("USA", "MAC"): 434_000_000, ("MAC", "USA"): 208_000_000}
    assert all(r["status"] == "estimated" and r["unit"] == "USD" for r in rows)
    with pytest.raises(ValueError, match="different annual periods"):
        sources.parse_ustr_macau_services(content.replace(b"Macau in 2025 were $208.0", b"Macau in 2024 were $208.0"))
    with pytest.raises(ValueError, match="rounding precision"):
        sources.parse_ustr_macau_services(content.replace(b"$642.0", b"$652.0"))
    with pytest.raises(ValueError, match="identity"):
        sources.parse_ustr_macau_services(content.replace(b"Macau Trade Summary", b"Hong Kong Trade Summary"))


def test_siat_cards_guard_destination_annual_year_and_imported_partial_boundary():
    text = ("Total Traveler Estimate\nto Destination\n842,000\n"
            "2025 Q1 - 2025 Q4 United States Travelers to China, People's Republic of\n"
            "SIAT Respondent Count\nto Destination\n1,113\nLast Update:06-Oct-2026 11:40\n"
            "Last Quarter Imported:2026 Q2")
    assert parse_monitor_card(text, "China, People's Republic of", 2025, "annual")[:2] == (842000, 1113)
    historical = text.replace("2025 Q1 - 2025 Q4", "2012 Q1 - 2012 Q4").replace(
        "842,000", "1,140,000")
    assert parse_monitor_card(historical, "China, People's Republic of", 2012, "annual")[:2] == (1140000, 1113)
    pandemic = text.replace("2025 Q1 - 2025 Q4", "2020 Q1 - 2020 Q4").replace(
        "842,000", "89,000").replace("1,113", "357")
    assert parse_monitor_card(pandemic, "China, People's Republic of", 2020, "annual")[:2] == (89000, 357)
    with pytest.raises(ValueError, match="source clarification"):
        parse_monitor_card(pandemic.replace("89,000", "0"), "China, People's Republic of", 2020, "annual")
    with pytest.raises(ValueError, match="filters"):
        parse_monitor_card(historical, "China, People's Republic of", 2013, "annual")
    with pytest.raises(ValueError, match="filters"):
        parse_monitor_card(text, "Hong Kong", 2025, "annual")
    with pytest.raises(ValueError, match="filters"):
        parse_monitor_card(text, "China, People's Republic of", 2025, "ytd")
    half = text.replace("2025 Q4", "2025 Q2").replace("842,000", "442,000").replace("1,113", "504")
    assert parse_monitor_card(half, "China, People's Republic of", 2025, "ytd")[:2] == (442000, 504)
    partial = text.replace("2025", "2026").replace("842,000", "474,000").replace("1,113", "756")
    assert parse_monitor_card(partial, "China, People's Republic of", 2026, "ytd")[:2] == (474000, 756)
    with pytest.raises(ValueError, match="annual"):
        parse_monitor_card(partial, "China, People's Republic of", 2026, "annual")
    with pytest.raises(ValueError, match="imported-quarter"):
        parse_monitor_card(partial.replace("Imported:2026 Q2", "Imported:2026 Q3"),
                           "China, People's Republic of", 2026, "ytd")
    with pytest.raises(ValueError, match="vintage"):
        parse_monitor_card(partial.replace("06-Oct-2026", "07-Oct-2026"),
                           "China, People's Republic of", 2026, "ytd")


def test_historical_siat_parser_keeps_mainland_hong_kong_units_and_missingness(monkeypatch):
    header = "Region/Country " + " ".join(map(str, range(2000, 2010)))
    china = "China (POC) 644 682 725 562 1,067 1,295 1,327 1,374 1,201 1,182"
    hong_kong = "Hong Kong 832 732 749 489 738 921 904 968 924 788"
    table = header + "\n2000 - 2009 Historical Outbound Visitation (one or more nights)\n"
    table += china + "\n" + hong_kong + "\nTaiwan 671 656 632 489 629 662 693 687 647 576"
    class Page:
        def __init__(self, text):
            self.text = text
        def extract_text(self):
            return self.text
    class Reader:
        pages = [Page("2009 United States Resident Travel Abroad (000)"), Page(table)]
    monkeypatch.setattr(travel_sources, "PdfReader", lambda _: Reader())
    rows = travel_sources.parse_historical_siat(b"source", 2009)
    assert len(rows) == 20
    assert {(row["destination"], row["year"]): row["value"] for row in rows
            if row["year"] == 2000} == {("CHN", 2000): 644000, ("HKG", 2000): 832000}
    Reader.pages[-1] = Page(table.replace("832 732", "** 732"))
    missing = next(row for row in travel_sources.parse_historical_siat(b"source", 2009)
                   if row["destination"] == "HKG" and row["year"] == 2000)
    assert missing["value"] is None and missing["status"] == "missing"
    for invalid in (table.replace("2000 2001", "1999 2001"),
                    table.replace("644 682", "644"),
                    table.replace(china, china + "\n" + china),
                    table.replace("China (POC)", "China including Hong Kong")):
        Reader.pages[-1] = Page(invalid)
        with pytest.raises(ValueError):
            travel_sources.parse_historical_siat(b"changed source", 2009)


def test_historical_siat_requires_matching_releases_and_current_2012_overlap(monkeypatch):
    from us_china_exchange.model import digest

    modern = [{"destination": geo, "year": 2012, "frequency": "annual", "value": value}
              for geo, value in (("CHN", 1140000), ("HKG", 627000))]
    def parse(body, report_year):
        return [{"destination": geo, "year": year, "value": value if year == 2012 else year * 1000,
                 "status": "estimated"}
                for geo, value in (("CHN", 1140000), ("HKG", 627000))
                for year in range(report_year - 9, report_year + 1)]
    def get(url):
        body = url.encode()
        return body, {"url": url, "sha256": digest(body), "retrieved_at": "2026-10-24T23:00:00+00:00"}
    monkeypatch.setattr(travel_sources, "parse_historical_siat", parse)
    rows = travel_sources.historical_siat(SimpleNamespace(get=get), modern)
    assert len(rows) == 24
    assert {(row["destination"], row["year"]) for row in rows} == {
        (geo, year) for geo in ("CHN", "HKG") for year in range(2000, 2012)}
    assert all(row["source_id"] == "ntto_siat" and row["unit"] == "trips" for row in rows)
    assert all(row["source_url"].startswith("https://web.archive.org/web/") and row["raw_sha256"] for row in rows)
    modern[0]["value"] += 1000
    with pytest.raises(ValueError, match="2012 overlap"):
        travel_sources.historical_siat(SimpleNamespace(get=get), modern)
    modern[0]["value"] -= 1000
    def conflicting(body, report_year):
        rows = parse(body, report_year)
        if report_year == 2011:
            rows[0]["value"] += 1000
        return rows
    monkeypatch.setattr(travel_sources, "parse_historical_siat", conflicting)
    with pytest.raises(ValueError, match="overlaps disagree"):
        travel_sources.historical_siat(SimpleNamespace(get=get), modern)


@pytest.fixture
def china_2021_profile():
    book = Workbook()
    sheet = book.active
    sheet.title = "CHINA"
    for cell, value in {"C5": "CHINA", "C14": "U.S. resident travel to China",
                        "C63": "Volume Trends (in thousands)", "C66": "Total Departures",
                        "C94": "Respondents", "AA94": 267}.items():
        sheet[cell] = value
    for column, year in enumerate(range(2000, 2022), 6):
        sheet.cell(64, column, year)
        sheet.cell(66, column, 5 if year == 2021 else year - 1999).number_format = "#,##0"
    sheet["AB64"] = 2021
    sheet["AB66"] = 1568
    def content():
        stream = io.BytesIO()
        book.save(stream)
        return stream.getvalue()
    return sheet, content


def test_2021_profile_uses_chronological_column_not_stale_duplicate_year_helper(china_2021_profile):
    sheet, content = china_2021_profile
    values, sample = travel_sources.parse_china_2021_profile(content())
    assert values[2021] == 5000 and values[2000] == 1000 and sample == 267
    for cell, invalid in [("C5", "China including Hong Kong"), ("C14", "U.S. citizen direct flights"),
                          ("C63", "Volume Trends (in millions)"), ("AA64", 2020),
                          ("AA66", None), ("AA66", 0), ("AA66", True), ("AA94", 0)]:
        previous = sheet[cell].value
        sheet[cell] = invalid
        with pytest.raises(ValueError):
            travel_sources.parse_china_2021_profile(content())
        sheet[cell] = previous


def test_2021_profile_requires_all_historical_controls_and_original_vintage(
        china_2021_profile, monkeypatch):
    sheet, content = china_2021_profile
    body = content()
    checksum = travel_sources.digest(body)
    monkeypatch.setattr(travel_sources, "CHINA_2021_SHA256", checksum)
    asset = {"url": travel_sources.CHINA_2021_PROFILE, "sha256": checksum,
             "retrieved_at": "2026-10-25T02:30:08+00:00"}
    client = SimpleNamespace(get=lambda url: (body, asset))
    retained = [{"destination": "CHN", "frequency": "annual", "year": year,
                 "value": (year - 1999) * 1000} for year in range(2000, 2021)]
    rows = travel_sources.china_2021_siat(client, retained)
    assert len(rows) == 1 and rows[0]["value"] == 5000 and rows[0]["year"] == 2021
    assert rows[0]["series_id"] == "ntto_siat:USA:CHN:air_visitors"
    assert rows[0]["status"] == "estimated" and rows[0]["raw_sha256"] == checksum
    with pytest.raises(ValueError, match="overlaps"):
        travel_sources.china_2021_siat(client, retained[1:])
    retained[0]["value"] += 1000
    with pytest.raises(ValueError, match="overlaps"):
        travel_sources.china_2021_siat(client, retained)
    with pytest.raises(ValueError, match="vintage"):
        travel_sources.china_2021_siat(SimpleNamespace(get=lambda url: (body + b"changed", asset)), retained)


def test_nbs_arrivals_preserve_country_identity_units_and_excel_precision():
    records = [
        ["Number of Oversea Visitor Arrivals by Country/Region"],
        ["(10 000 person-times)"], ["Region", 2003, 2004],
        ["United States", 82.2511, 130.8627],
        ["Canada", 21, 22],
    ]
    assert sources.parse_nbs_visitor_records(records, 2003, 2004) == {2003: 822511, 2004: 1308627}
    records[3][1] = ""
    with pytest.raises(ValueError, match="missingness"):
        sources.parse_nbs_visitor_records(records, 2003, 2004)
    records[1][0] = "(persons)"
    with pytest.raises(ValueError, match="units"):
        sources.parse_nbs_visitor_records(records, 2003, 2004)


def xlsx(book):
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def expanded_services():
    cells = []
    for index, kind in enumerate(("Exports of services", "Imports of services", "Balance on services")):
        for offset, year in enumerate(range(2018, 2026)):
            column = index * 8 + offset + 3
            cells.extend([{"Row_ID": "1", "Column_ID": str(column), "Cell_Value": kind},
                          {"Row_ID": "2", "Column_ID": str(column), "Cell_Value": str(year)}])
    for line, country in enumerate(("China", "Hong Kong", "Macau", "Taiwan"), 3):
        cells.append({"Row_ID": str(line), "Column_ID": "2", "Cell_Value": country})
        for offset, year in enumerate(range(2018, 2026)):
            values = (441, 34, 406) if country == "Macau" and year == 2024 else (434, 208, 226)
            for index, value in enumerate(values):
                cells.append({"Row_ID": str(line), "Column_ID": str(index * 8 + offset + 3),
                              "Cell_Value": str(value)})
    table = {"Title": "Table 2.4. U.S. Trade in Services, Expanded Geographic Detail",
             "Sub_Title": "[Millions of dollars]", "Description": "Release Date: July 7, 2026",
             "Number_Of_Columns": "26", "TD": cells}
    return {"AppId": 62, "Number": 6, "Prompts": [{"Name": "TheTableFlexibleIipIta",
            "PromtData": json.dumps({"Table": json.dumps(table)})}]}, table


def test_expanded_bea_services_use_direct_macau_history_and_independent_rounding_controls():
    data, table = expanded_services()
    rows, _ = sources.parse_bea_expanded(json.dumps(data))
    assert len(rows) == 16 and {r["year"] for r in rows} == set(range(2018, 2026))
    assert {r["source_id"] for r in rows} == {"bea_expanded"}
    assert next(r["value"] for r in rows if r["origin"] == "USA" and r["year"] == 2024) == 441_000_000
    assert next(r["value"] for r in rows if r["origin"] == "MAC" and r["year"] == 2025) == 208_000_000
    cell = next(c for c in table["TD"] if (c["Row_ID"], c["Column_ID"]) == ("5", "9"))
    cell["Cell_Value"] = "(*)"
    data["Prompts"][0]["PromtData"] = json.dumps({"Table": json.dumps(table)})
    rows, _ = sources.parse_bea_expanded(json.dumps(data))
    missing = next(r for r in rows if r["origin"] == "USA" and r["year"] == 2024)
    assert missing["value"] is None and missing["status"] == "missing"
    cell["Cell_Value"] = "451"
    data["Prompts"][0]["PromtData"] = json.dumps({"Table": json.dumps(table)})
    with pytest.raises(ValueError, match="rounding precision"):
        sources.parse_bea_expanded(json.dumps(data))
    table["Sub_Title"] = "[Billions of dollars]"
    data["Prompts"][0]["PromtData"] = json.dumps({"Table": json.dumps(table)})
    with pytest.raises(ValueError, match="units"):
        sources.parse_bea_expanded(json.dumps(data))


def test_bea_scale_directions_and_matched_quarter_sums():
    book = Workbook()
    book.remove(book.active)
    for n in [4, 5, 7, 8]:
        sheet = book.create_sheet(f"Table {n}")
        for cells in [
            ["Last updated September 3, 2026"], ["Title"], ["Trade"],
            ["[Millions of dollars, quarters seasonally adjusted]"], [],
            ["Period", "China", "Hong Kong", "Taiwan"], ["Annual"],
            ["1999", 1, 2, 3], ["2000", 4, 5, 6], ["2020", 10, 20, 30], ["Quarterly"],
            ["2020 1", 1, 2, 3], ["2020 2", 4, 5, 6],
            ["2026 1 (R)", 7, 8, 9], ["2026 2", 10, 11, 12],
        ]:
            sheet.append(cells)
    rows, vintage = parse_bea(xlsx(book))
    assert len(rows) == 48
    assert min(r["year"] for r in rows) == 2000
    assert next(r for r in rows if r["year"] == 2000 and r["origin"] == "USA"
                and r["destination"] == "CHN" and r["measure"] == "goods")["value"] == 4_000_000
    assert {r["origin"] for r in rows} == {"USA", "CHN", "HKG", "TWN"}
    row = next(r for r in rows if r["origin"] == "USA" and r["destination"] == "CHN"
               and r["measure"] == "services" and r["frequency"] == "ytd" and r["year"] == 2026)
    assert row["value"] == 17_000_000
    assert row["period"] == "2026-06"
    assert "seasonally adjusted" in row["definition"]


def test_bea_ambiguous_early_labels_use_independently_labelled_archive():
    def content(archived=False):
        book = Workbook()
        book.remove(book.active)
        for n in [4, 5, 7, 8]:
            sheet = book.create_sheet(f"Table {n}")
            for cells in [
                ["Last updated December 5, 2024" if archived else "Last updated September 3, 2026"],
                ["Title"], ["Trade"], ["Millions of dollars"], [],
                ["Period", "China", "Hong Kong", "Taiwan"], ["Annual"],
                [2000, 1, 2, 3], [2001 if archived else 2000, 4, 5, 6],
                [2020, 10, 20, 30], ["Quarterly"], ["2020 1", 1, 2, 3],
            ]:
                sheet.append(cells)
        return xlsx(book)

    current, archived = content(), content(True)
    with pytest.raises(sources.BEAEarlyHistoryError):
        parse_bea(current)

    def get(url):
        page = b'<a href="https://example.test/geo-time-series.xlsx">History</a>'
        body = page if url == sources.BEA else archived if url == sources.BEA_HISTORY else current
        return body, {"url": url, "sha256": "a" * 64, "retrieved_at": "2026-10-21T00:00:00+00:00"}

    rows, _, notes = sources.bea(SimpleNamespace(get=get))
    history = [r for r in rows if r["frequency"] == "annual" and r["year"] < 2002]
    assert len(history) == 24
    assert {r["source_url"] for r in history} == {sources.BEA_HISTORY}
    assert {r["source_vintage"] for r in history} == {"Last updated December 5, 2024"}
    assert {r["value"] for r in history if r["year"] == 2001} == {4_000_000, 5_000_000, 6_000_000}
    assert all(r["source_url"] != sources.BEA_HISTORY for r in rows if r["year"] == 2020)
    assert "repeats the 2000 label" in notes


def test_ntto_mixed_headers_single_digit_months_and_provisional_ytd():
    book = Workbook()
    sheet = book.active
    sheet.title = "Monthly"
    sheet.append(["#", "Country", "Region", datetime(2020, 1, 2),
                  datetime(2020, 2, 3), "2022-1", "2022-2",
                  "2026-01", "2026-02\nPreliminary"])
    for i, country in enumerate(COUNTRIES):
        sheet.append([i, country, "Asia", 10, 20, 40, 50, 60, 70])
    rows, _ = parse_ntto(xlsx(book), monthly=True)
    assert len([r for r in rows if r["frequency"] == "monthly"]) == 24
    assert next(r for r in rows if r["period"] == "2022-02" and r["origin"] == "MAC")["value"] == 50
    current = next(r for r in rows if r["year"] == 2026 and r["frequency"] == "ytd")
    assert current["value"] == 130
    assert current["status"] == "provisional"


def test_taiwan_all_purpose_and_leisure_are_separate():
    content = "年別,細分,合計,觀光\n2020,美國 USA,\"82,872\",\"26,634\"\n".encode("utf-8-sig")
    rows = parse_taiwan(content)
    assert {r["measure"]: r["value"] for r in rows} == {
        "visitors": 82872, "leisure_visitors": 26634}
    with pytest.raises(ValueError):
        parse_taiwan(b"year,country\n2020,USA\n")


def test_ntto_retains_annual_history_from_2000():
    book = Workbook()
    sheet = book.active
    sheet.title = "Annual"
    sheet.append(["Country", 1999, 2000, 2001])
    for country in COUNTRIES:
        sheet.append([country, 9, 10, 11])
    rows, _ = parse_ntto(xlsx(book))
    assert {r["year"] for r in rows} == {2000, 2001}
    assert {r["value"] for r in rows if r["year"] == 2000} == {10}


@pytest.mark.parametrize("edition,basis,share", [
    (2008, "PLACE OF RESIDENCE", "100.00%"),
    (2024, "Place of Issue of Travel Document", "100.0%"),
])
def test_macau_profile_reconciles_modes_and_excludes_monthly_rows(monkeypatch, edition, basis, share):
    text = "VISITOR ARRIVALS FROM U.S.A.\n" + basis + "\n"
    years = range(edition - 5, edition + 1)
    text += "\n".join(f"{year} 3 104 64 171 {share}" for year in years)
    text += "\nJanuary 1 2 3 6 4.0%"
    page = SimpleNamespace(extract_text=lambda: text)
    monkeypatch.setattr(sources, "PdfReader", lambda stream: SimpleNamespace(pages=[page]))
    assert sources.parse_macau_market(b"%PDF fixture", edition) == dict.fromkeys(years, 171)
    text = text.replace(f"{edition - 2} 3 104 64 171", f"{edition - 2} 3 104 64 172")
    with pytest.raises(ValueError, match="reconcile"):
        sources.parse_macau_market(b"%PDF fixture", edition)
    text = text.replace(basis, "Nationality")
    with pytest.raises(ValueError, match="issuing-document profile"):
        sources.parse_macau_market(b"%PDF fixture", edition)


def test_macau_classification_break_is_not_compared_with_2020():
    rows = {year: sources.macau_observation(year, 100, "")
            for year in [2007, 2008, 2014, 2015, 2016, 2020, 2025]}
    assert change(rows[2007], rows[2008])["absolute"] is None
    assert change(rows[2014], rows[2020])["absolute"] is None
    assert change(rows[2015], rows[2020])["absolute"] is None
    assert change(rows[2016], rows[2020])["absolute"] == 0
    assert change(rows[2025], rows[2020])["absolute"] == 0


def test_macau_discovers_historical_and_latest_complete_profiles(monkeypatch):
    from us_china_exchange import tourism_backfill

    paths = [
        "/VAHightlights/2015/Dec.pdf", "/VAHightlights/2025/Dec.pdf",
        "/MarketPerformance/2015/Market Performance  2015.pdf",
        "/MarketPerformance/2023/Market Performance 2023.pdf",
        "/MarketPerformance/2024/Market Performance 2024.pdf",
        "/MarketPerformance/2025/Market Performance 2025 Jan-Sep.pdf",
        "/MarketPerformance/2008/M12_MarketingPerformance_2008_1246941849.pdf",
        "/MarketPerformance/2010/M12_MarketPerformance_2010_1298431536.pdf",
    ]
    archive = "".join(f'<option value="{path}"></option>' for path in paths).encode()
    requested = []

    def get(url):
        requested.append(url)
        return (archive if url == sources.MACAU_INDEX else b"%PDF fixture",
                {"url": url, "sha256": "a" * 64, "retrieved_at": "2026-10-21T00:00:00+00:00"})

    monkeypatch.setattr(sources, "parse_macau", lambda content, year: 999)
    monkeypatch.setattr(sources, "parse_macau_market",
                        lambda content, edition: dict.fromkeys(range(edition - 5, edition + 1), edition))
    monkeypatch.setattr(tourism_backfill, "dsec_recent_visitors",
                        lambda client: ([], "DSEC checked separately", "Direct monthly controls."))
    rows, _, _ = sources.macau(SimpleNamespace(get=get))
    rows = {row["year"]: row for row in rows}
    assert rows[2010]["value"] == rows[2015]["value"] == 2015
    assert rows[2024]["value"] == 2024
    assert rows[2025]["value"] == 999
    assert rows[2003]["value"] == 2008
    assert rows[2015]["source_url"].endswith(paths[2])
    assert rows[2024]["source_url"].endswith(paths[4])
    assert not any("2023" in url or "Jan-Sep" in url for url in requested)
    december = sources.macau_observation(2025, 1001, "Direct DSEC January-December count")
    monkeypatch.setattr(tourism_backfill, "dsec_recent_visitors",
                        lambda client: ([december], "DSEC December", "Full year, not annualized"))
    revised, _, _ = sources.macau(SimpleNamespace(get=get))
    assert len({(r["series_id"], r["period"], r["frequency"]) for r in revised}) == len(revised)
    assert next(r["value"] for r in revised if r["year"] == 2025) == 1001


def test_macau_rejects_conflicting_extracted_us_totals(monkeypatch):
    page = SimpleNamespace(extract_text=lambda: "Visitor Arrivals 2025\nUSA 100 \nU.S.A. 200 ")
    monkeypatch.setattr(sources, "PdfReader", lambda stream: SimpleNamespace(pages=[page]))
    with pytest.raises(ValueError, match="Ambiguous"):
        sources.parse_macau(b"%PDF fixture", 2025)


def test_iie_pandemic_population_break_is_not_applied_to_study_abroad(monkeypatch):
    text = "\n".join(f"{year - 1}/{str(year)[-2:]}\n100\n10" for year in range(2018, 2026))
    text += "\nUndergraduate\n10\n20\nGraduate\n30\n40\nNon-Degree\n5\n6\nOPT\n1\n2\n"
    page = SimpleNamespace(extract_text=lambda: text)
    monkeypatch.setattr(sources, "PdfReader", lambda stream: SimpleNamespace(pages=[page]))
    rows = sources.parse_iie(b"%PDF fixture", "CHN")
    students = {r["year"]: r for r in rows if r["measure"] == "international_students"}
    abroad = {r["year"]: r for r in rows if r["measure"] == "study_abroad"}
    assert change(students[2021], students[2020])["absolute"] is None
    assert change(students[2024], students[2023])["absolute"] is None
    assert change(students[2023], students[2021])["absolute"] == 0
    assert change(students[2025], students[2020])["absolute"] == 0
    assert "regardless of visa status" in students[2021]["definition"]
    assert "excludes students enrolled online" in students[2025]["definition"]
    assert change(abroad[2021], abroad[2020])["absolute"] == 0


def uis_client(**overrides):
    def get(url, params=None):
        record = dict(value=12, magnitude=None, qualifier=None, footnotes=[], year=2024)
        record.update(overrides)
        if params is not None:
            record.update(indicatorId=params["indicator"], geoUnit=params["geoUnit"])
        payload = [] if params is None else {"records": [record]}
        return json.dumps(payload).encode(), {
            "url": url, "sha256": "a" * 64, "retrieved_at": "2026-10-21T00:00:00+00:00"}
    return SimpleNamespace(get=get, json=json.loads)


@pytest.mark.parametrize("magnitude,qualifier,value,status,expected", [
    (None, None, 12, "observed", 12),
    ("NIL", None, 0, "observed", 0),
    ("NA", None, 0, "not_applicable", None),
    ("SUPP", None, None, "suppressed", None),
    ("INCLUDED", None, None, "included_elsewhere", None),
    ("LOWREL", "NAT_EST", 12, "low_reliability", 12),
    ("INCLUDES", None, 12, "observed", 12),
    (None, "NAT_EST", 12.5, "estimated", 12.5),
    (None, "UIS_EST", 12.5, "estimated", 12.5),
])
def test_uis_preserves_source_flags_without_inventing_zeros(magnitude, qualifier, value, status, expected):
    rows, _, _ = sources.uis(uis_client(magnitude=magnitude, qualifier=qualifier, value=value))
    assert validate(rows)
    assert all(r["status"] == status and r["value"] == expected for r in rows)
    assert all(str(qualifier) in r["notes"] for r in rows)
    if magnitude:
        assert all(f"Magnitude: {magnitude}; source value: {value}" in r["notes"] for r in rows)


@pytest.mark.parametrize("overrides", [
    {"magnitude": "UNKNOWN"}, {"qualifier": "UNKNOWN"},
    {"magnitude": "NA", "value": 12}, {"magnitude": "NIL", "value": True},
    {"magnitude": "SUPP", "value": 12}, {"magnitude": "LOWREL", "value": None},
])
def test_uis_unknown_or_inconsistent_flags_fail_closed(overrides):
    with pytest.raises(ValueError, match="UIS"):
        sources.uis(uis_client(**overrides))


def test_iie_adapter_is_blocked_before_any_network_call():
    def get(url):
        pytest.fail("Restricted IIE PDFs must not be requested")
    with pytest.raises(sources.SourcePermissionRequired, match="automatic retrieval is paused"):
        sources.iie(SimpleNamespace(get=get))


def comtrade_client(overrides=None):
    def get(url, params):
        records = []
        for flow in ["X", "M"]:
            row = dict(typeCode="C", freqCode="A", period=str(params["period"]), refYear=params["period"],
                       reporterCode=842, partnerCode=446, cmdCode="TOTAL", flowCode=flow,
                       fobvalue=100, cifvalue=125)
            row.update(overrides or {})
            records.append(row)
        return json.dumps({"data": records}).encode(), {
            "url": url, "sha256": "a" * 64, "retrieved_at": "2026-10-21T00:00:00+00:00"}
    return SimpleNamespace(get=get, json=json.loads)


@pytest.mark.parametrize("overrides", [
    {"period": "1999"}, {"refYear": 1999}, {"freqCode": "M"}, {"typeCode": "S"},
])
def test_comtrade_rejects_wrong_returned_period_frequency_or_type(monkeypatch, overrides):
    monkeypatch.setattr(sources.time, "sleep", lambda seconds: None)
    with pytest.raises(ValueError, match="wrong period, frequency or trade type"):
        sources.comtrade(comtrade_client(overrides))


def test_comtrade_retains_valid_annual_flow_valuations(monkeypatch):
    monkeypatch.setattr(sources.time, "sleep", lambda seconds: None)
    monkeypatch.setattr(sources, "datetime", SimpleNamespace(now=lambda: SimpleNamespace(year=2002)))
    rows, _, _ = sources.comtrade(comtrade_client())
    assert {(r["year"], r["origin"], r["destination"], r["value"]) for r in rows} == {
        (2000, "USA", "MAC", 100), (2000, "MAC", "USA", 125),
        (2001, "USA", "MAC", 100), (2001, "MAC", "USA", 125)}
