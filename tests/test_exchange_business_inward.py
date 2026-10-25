"""Reverse counts never become parent-company totals or one uniform population."""

import copy
import csv
import io
import json
from collections import defaultdict
from pathlib import Path

import pytest
from openpyxl import Workbook
from bs4 import BeautifulSoup

from us_china_exchange import business_current, business_inward


ROOT = Path(__file__).resolve().parents[1]


def published():
    return json.loads((ROOT / "data/site/business_inward.json").read_text())


def test_count_column_ownership_threshold_and_suppression_are_not_interchangeable():
    source = copy.deepcopy(published()["sources"]["bea_mousa_2024"])
    source["cells"] = {"CHN": "B6", "HKG": "B7"}
    book = Workbook()
    sheet = book.active
    sheet.title = source["worksheet"]
    for row in [
        ["U.S. Affiliate Activities: Preliminary 2024 Statistics, Majority-Owned Affiliates"],
        [],
        ["Number of Affiliates With Assets, Sales, or Net Income (or Loss) Greater Than $20 Million, Country of UBO by Employment Size Class"],
        [None, "Total", None, None, None, "Number of companies consolidated"],
        [],
        ["China", 360, None, None, None, 1062],
        ["Hong Kong", 137, None, None, None, 522],
    ]:
        sheet.append(row)

    def parse():
        stream = io.BytesIO()
        book.save(stream)
        return business_inward.parse_counts(stream.getvalue(), source)

    assert parse() == {"CHN": (360, "provisional"), "HKG": (137, "provisional")}
    sheet.cell(6, 2, "(D)")
    sheet.cell(7, 2, 0)
    assert parse() == {"CHN": (None, "suppressed"), "HKG": (0, "provisional")}
    sheet.cell(6, 2, "350-400")
    with pytest.raises(ValueError, match="range"):
        parse()
    sheet.cell(6, 2, 360)
    source["size_threshold_usd"] = 15_000_000
    with pytest.raises(ValueError, match="threshold"):
        parse()
    source["size_threshold_usd"] = 20_000_000
    sheet.cell(3, 1, sheet.cell(3, 1).value.replace("Country of UBO", "Country of foreign parent"))
    with pytest.raises(ValueError, match="country of UBO"):
        parse()


def test_history_retains_all_years_and_four_actual_population_blocks():
    data = business_inward.validate(published())
    assert len(data["observations"]) == 50
    for geo in ("CHN", "HKG"):
        rows = [row for row in data["observations"] if row["series_id"] == "bea_mousa:" + geo]
        assert [row["year"] for row in rows] == list(range(2000, 2025))
        assert len({row["comparability"] for row in rows}) == 4
        for row in rows:
            expected = ("nonbank_only", 3_000_000) if row["year"] <= 2001 else (
                ("nonbank_only", 10_000_000) if row["year"] <= 2006 else
                ("including_banks", 15_000_000) if row["year"] <= 2011 else ("including_banks", 20_000_000))
            assert (row["bank_coverage"], row["size_threshold_usd"]) == expected
    for mutate in [
        lambda d: d["series"]["bea_mousa:CHN"].update(destination="CHN"),
        lambda d: d["series"]["bea_mousa:CHN"].update(unit="companies"),
        lambda d: d["observations"][0].update(size_threshold_usd=30_000_000),
        lambda d: d["observations"][0].update(comparability="all years are comparable"),
        lambda d: d["observations"].pop(),
    ]:
        invalid = copy.deepcopy(data)
        mutate(invalid)
        with pytest.raises(ValueError):
            business_inward.validate(invalid)


def test_complete_exports_preserve_every_count_and_its_original_locator(tmp_path):
    data = published()
    business_inward.export(data, tmp_path)
    assert json.loads((tmp_path / "business-inward.json").read_text()) == data
    with (tmp_path / "business-inward-observations.csv").open() as stream:
        rows = list(csv.DictReader(stream))
    assert [(row["series_id"], int(row["year"]), int(row["value"])) for row in rows] == [
        (row["series_id"], row["year"], row["value"]) for row in data["observations"]]
    assert all(row["source_url"] and len(row["raw_sha256"]) == 64 and row["cell"].startswith("B")
               and row["country_basis"] == "UBO" and row["size_threshold_usd"] for row in rows)


def test_current_definition_charts_keep_the_complete_earlier_history_and_withdraw_cleanly(tmp_path):
    data = published()
    current = json.loads((ROOT / "data/site/business_current.json").read_text())
    for geo, chart in [("CHN", "us-investment-mainland"), ("HKG", "us-investment-hong-kong")]:
        assert business_inward.chart_groups(data, geo)[0]["color"] != business_current.chart_groups(current, chart)[0]["color"]
    gallery = defaultdict(list)
    markup, count = business_inward.build(tmp_path, data, gallery)
    assert count == 2
    page = BeautifulSoup(markup, "html.parser")
    assert len(page.select("#business-inward-data table tbody tr")) == 25
    assert "not joined" in page.get_text()
    for cid in business_inward.CHARTS:
        for theme in ("themarginoferror", "substack"):
            path = tmp_path / "site/figures/business_presence" / theme / (cid + ".svg")
            svg = BeautifulSoup(path.read_text(), "html.parser")
            assert [int(bar["data-year"]) for bar in svg.select(".annual-column")] == list(range(2012, 2025))
            assert not svg.select(".data-line,.covid-band,.covid-label")
    rows = list(csv.DictReader((tmp_path / "site/downloads/business-inward-chart-values.csv").open()))
    assert len(rows) == 26 and all(row["comparability"] == "including_banks:USD20000000" for row in rows)
    data["observations"][-1].update(value=None, status="suppressed")
    svg = BeautifulSoup(business_inward.svg_count_chart(data, "hong-kong-ubo-us-affiliates", "substack"), "html.parser")
    assert len(svg.select(".annual-column")) == 12 and "2024" in svg.get_text()
    custom = tmp_path / "site/figures/business_presence/substack/custom.svg"
    custom.write_text("User-owned")
    business_inward.withdraw(tmp_path)
    assert not list((tmp_path / "site/downloads").glob("business-inward*"))
    assert not list((tmp_path / "site/figures/business_presence").glob("*/*ubo-us-affiliates.*"))
    assert custom.read_text() == "User-owned"
