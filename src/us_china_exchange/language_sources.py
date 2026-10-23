"""Explicitly reviewed public language-learning aggregates, never a daily scraper."""

import csv
import fcntl
import io
import json
import math
import re
import subprocess
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin

from bs4 import BeautifulSoup
from openpyxl import load_workbook
from pypdf import PdfReader
import xlrd

from .model import digest, json_text
from .pipeline import write_atomic


MLA_REPORT = "https://www.mla.org/content/download/191324/file/Enrollments-in-Languages-Other-Than-English-in-US-Institutions-of-Higher-Education-Fall-2021.pdf"
MLA_DATA = "https://www.mla.org/content/download/191323/file/Historical-language-enrollments-1958-2021.xlsx"
MLA_SNAPSHOT = "https://www.mla.org/content/download/191129/file/Fall-2022-NL-snapshot.pdf"
AP_ARCHIVE = "https://reports.collegeboard.org/ap-program-results/data-archive"
AP_COURSES = "https://apstudents.collegeboard.org/courses"
AP_GROUP_CONTEXT = {
    "all_participants": "Everyone taking that language's AP exam, including the standard-group subset.",
    "standard_group": "A subset whose answers did not report regularly speaking or hearing the exam language, or living for at least one month where it is spoken. Generally school-trained, but not a U.S.-only or ancestry group.",
    "relationship": "The standard group is already included in all participants. Never add them together.",
    "years": "Years are annual exam administrations, not graduating classes. Verified counts run from 2007 through 2025. The 2026 score-percentage rows do not provide participant counts.",
    "korean": "College Board lists AP Chinese and AP Japanese, but no AP Korean exam. Korean benchmarks use MLA college enrollments separately.",
}
IPEDS = "https://nces.ed.gov/ipeds/datacenter/data/"
IPEDS_CURRENT = "https://nces.ed.gov/ipeds/complete-data-files/"
NEXUS = "https://nexusmandarin.org/publications-1"
STATES = set("AL AK AZ AR CA CO CT DE DC FL GA HI ID IL IN IA KS KY LA ME MD MA MI MN MS MO MT NE NV NH NJ NM NY NC ND OH OK OR PA RI SC SD TN TX UT VT VA WA WV WI WY".split())


class ReviewClient:
    def __init__(self, cache, refresh=False):
        self.cache = Path(cache)
        self.cache.mkdir(parents=True, exist_ok=True)
        self.index = self.cache / "assets.json"
        self.assets = json.loads(self.index.read_text()) if self.index.exists() else {}
        self.refresh = refresh

    def get(self, url):
        prior = self.assets.get(url)
        if prior and not self.refresh:
            path = self.cache / prior["sha256"]
            if path.is_file() and digest(path.read_bytes()) == prior["sha256"]:
                return path.read_bytes(), prior
        result = subprocess.run(
            ["curl", "--fail", "--location", "--silent", "--show-error", "--retry", "2",
             "--max-time", "180", "--write-out", "\n%{url_effective}", url],
            capture_output=True, check=True)
        content, effective = result.stdout.rsplit(b"\n", 1)
        return self.remember(url, content, effective.decode())

    def get_public_download(self, url):
        try:
            return self.get(url)
        except subprocess.CalledProcessError:
            from tempfile import TemporaryDirectory
            from playwright.sync_api import Error, sync_playwright

            with TemporaryDirectory() as directory, sync_playwright() as playwright:
                browser = playwright.chromium.launch()
                try:
                    page = browser.new_page()
                    with page.expect_download(timeout=120000) as pending:
                        try:
                            page.goto(url, timeout=120000)
                        except Error as exc:
                            if "Download" not in str(exc):
                                raise
                    download = pending.value
                    path = Path(directory) / "public-download"
                    download.save_as(path)
                    return self.remember(url, path.read_bytes(), download.url)
                finally:
                    browser.close()

    def remember(self, url, content, effective):
        asset = {"url": effective, "request_url": url, "sha256": digest(content),
                 "retrieved_at": datetime.now(timezone.utc).isoformat(timespec="seconds")}
        write_atomic(self.cache / asset["sha256"], content)
        self.assets[url] = asset
        with (self.cache / "assets.lock").open("a") as lock:
            fcntl.flock(lock, fcntl.LOCK_EX)
            current = json.loads(self.index.read_text()) if self.index.exists() else {}
            current[url] = asset
            write_atomic(self.index, json_text(current))
            fcntl.flock(lock, fcntl.LOCK_UN)
        return content, asset


def source(name, vintage, assets, locator, notes=""):
    return {"name": name, "vintage": vintage, "assets": assets, "locator": locator, "notes": notes}


def series(label, measure, unit, scope, cadence, basis="calendar", notes=""):
    return {"label": label, "measure": measure, "unit": unit, "scope": scope,
            "cadence": cadence, "period_basis": basis, "notes": notes}


def row(sid, year, value, source_id, status="observed", notes="", comparability=""):
    return {"series_id": sid, "year": year, "value": value, "source_id": source_id,
            "status": status, "notes": notes, "comparability": comparability or sid}


def pdf_text(content):
    if not content.startswith(b"%PDF"):
        raise ValueError("Expected a language-source PDF")
    return [p.extract_text() for p in PdfReader(io.BytesIO(content)).pages]


def mla(client):
    content, report_asset = client.get(MLA_REPORT)
    pages = pdf_text(content)
    table = next(t for t in pages if "Table" in t and "Fall Enrollments in the Fifteen" in t)
    years = [1958, 1980, 1990, 1998, 2002, 2006, 2009, 2013, 2016, 2021]
    labels = {"chinese": "Mandarin", "japanese": "Japanese", "korean": "Korean",
              "spanish": "Spanish", "all_languages": "Total"}
    data, definitions = [], {}
    expected = {}
    for language, label in labels.items():
        match = re.search(r"^\s*" + label + r"\s+([\d,].+)$", table, re.MULTILINE)
        if match is None:
            raise ValueError("MLA historical table row missing: " + language)
        values = re.findall(r"\d[\d,]*(?:\.\d+)?", match[1])
        if len(values) != 12:
            raise ValueError("MLA Table 5 reviewed columns changed")
        sid = "mla:" + language
        definitions[sid] = series(
            "Chinese/Mandarin" if language == "chinese" else language.replace("_", " ").capitalize(),
            "course_enrollments", "course registrations",
            "U.S. colleges and universities; undergraduate and graduate fall language course registrations",
            "census", notes="Published 2021-report vintage. Registrations, not unique students; all languages includes ASL and classical languages.")
        for year, value in zip(years, values[:10]):
            if year >= 2000:
                expected[(year, language)] = int(value.replace(",", ""))
                data.append(row(sid, year, expected[(year, language)], "mla_report_2021"))
    sources = {"mla_report_2021": source(
        "Modern Language Association", "Fall 2021 census; report published 2023",
        [report_asset], "Table 5 (PDF page 63)",
        "National figures use one published report vintage, not later workbook revisions. "
        "Only selected numerical aggregates are republished; source report is not redistributed. "
        "MLA network terms limit ordinary use to personal noncommercial use; separate commercial reuse permission is unverified.")}
    content, data_asset = client.get(MLA_DATA)
    sheet = load_workbook(io.BytesIO(content), data_only=True, read_only=True)["FL history"]
    records = sheet.iter_rows(values_only=True)
    headers = next(records)
    required = {"INST_ID", "SRVY_YEAR", "TERM", "LANGUAGE", "ALL LEVEL TOTAL", "N_RESP", "ZERO_ERL"}
    if not required.issubset(headers):
        raise ValueError("MLA institution-level columns changed")
    institutions = defaultdict(set)
    counts = defaultdict(Counter)
    incomplete = defaultdict(set)
    raw_totals = Counter()
    for cells in records:
        r = dict(zip(headers, cells))
        year = r["SRVY_YEAR"]
        if year not in {2013, 2016, 2021} or r["TERM"] != "Fall":
            continue
        if r["N_RESP"] == "NR":
            continue
        institution, value, language = r["INST_ID"], r["ALL LEVEL TOTAL"], r["LANGUAGE"]
        institutions[year].add(institution)
        if value is not None:
            if isinstance(value, bool) or not isinstance(value, (int, float)) or value < 0 or value != int(value):
                raise ValueError("Invalid MLA course count")
            counts[(year, institution)]["all_languages"] += int(value)
            raw_totals[(year, "all_languages")] += int(value)
            normalized = ("chinese" if language in {"CHINESE", "MANDARIN"} else
                          {"JAPANESE": "japanese", "KOREAN": "korean", "SPANISH": "spanish"}.get(language))
            if normalized:
                counts[(year, institution)][normalized] += int(value)
                raw_totals[(year, normalized)] += int(value)
        elif r["ZERO_ERL"] != "Z":
            incomplete[year].add(institution)
    for year in (2013, 2021):
        for language in labels:
            if raw_totals[(year, language)] != expected[(year, language)]:
                raise ValueError(f"MLA {year} workbook/report mismatch requires review: {language}")
    cohort = ((institutions[2013] & institutions[2021])
              - incomplete[2013] - incomplete[2021])
    if not cohort:
        raise ValueError("No matched MLA responding institutions")
    for language in labels:
        sid = "mla_matched:" + language
        definitions[sid] = series(
            definitions["mla:" + language]["label"], "course_enrollments", "course registrations",
            f"Matched set of {len(cohort)} MLA institution IDs responding in both fall 2013 and fall 2021",
            "paired census", notes="Nonrespondents excluded; explicit zero-language respondents retained. "
            "Institutions with any unreported language-count row in either endpoint are excluded, "
            "so missing benchmark or other-language counts are not treated as zero. "
            "An absent language at a responding institution means no reported registrations, not proof of no language activity. "
            "This is not the national census total and excludes institutions outside the matched cohort.")
        for year in (2013, 2021):
            data.append(row(sid, year, sum(counts[(year, i)][language] for i in cohort), "mla_matched_2021"))
    differences = {language: {"report": expected[(2016, language)],
                              "workbook": raw_totals[(2016, language)]}
                   for language in labels if expected[(2016, language)] != raw_totals[(2016, language)]}
    sources["mla_matched_2021"] = source(
        "Modern Language Association", "Historical workbook retrieved for 2013/2021 matched panel",
        [data_asset], "FL history; INST_ID, N_RESP, ZERO_ERL, LANGUAGE, ALL LEVEL TOTAL",
        f"Matched complete-count respondents: {len(cohort)}. 2013/2021 national sums reconcile with the report. "
        "2016 workbook/report revisions are not spliced into national charts: " + json.dumps(differences, sort_keys=True))
    return definitions, data, sources


def zip_csv(content, prefer_revised=True):
    archive = zipfile.ZipFile(io.BytesIO(content))
    members = [n for n in archive.namelist() if n.lower().endswith(".csv")]
    revised = [n for n in members if "_rv." in n.lower()]
    candidates = revised if prefer_revised and revised else [n for n in members if "_rv." not in n.lower()]
    if len(candidates) != 1:
        raise ValueError("Ambiguous IPEDS CSV members")
    member = candidates[0]
    stream = archive.open(member)
    encoding = "utf-8-sig" if stream.read(3) == b"\xef\xbb\xbf" else "cp1252"
    stream.seek(0)
    reader = csv.DictReader(io.TextIOWrapper(stream, encoding=encoding))
    reader.fieldnames = [name.strip().upper() for name in reader.fieldnames]
    return member, reader


def ipeds(client, years):
    definitions = {
        "ipeds:bachelors": series("Bachelor's", "degree_awards", "awards",
                                 "50 states and DC; Title IV degree-granting institutions; CIP 16.0301; first major from 2001",
                                 "annual", "academic ending year",
                                 "Awards, not unique people or all China specialists. NCES reported/imputed counts; "
                                 "certificates excluded. First-major restriction from 2001; 2000 has no separate major-number field."),
        "ipeds:masters": series("Master's", "degree_awards", "awards",
                               "50 states and DC; Title IV degree-granting institutions; CIP 16.0301; first major from 2001",
                               "annual", "academic ending year",
                               "First-major restriction from 2001; 2000 has no separate major-number field."),
        "ipeds:doctorates": series("Doctoral", "degree_awards", "awards",
                                  "50 states and DC; Title IV degree-granting institutions; CIP 16.0301; first major from 2001",
                                  "annual", "academic ending year",
                                  "2000 has no separate major-number field; first-major restriction from 2001. "
                                  "Pre-2010: old doctoral category (9); from 2010: research/scholarship doctorates (17). Definition break is not connected."),
    }
    codes = {"16.0301": "ipeds", "16.0302": "ipeds_japanese", "16.0303": "ipeds_korean"}
    for cip, prefix in codes.items():
        if prefix == "ipeds":
            continue
        language = prefix.removeprefix("ipeds_")
        for name in ("bachelors", "masters", "doctorates"):
            definition = dict(definitions["ipeds:" + name])
            definition["label"] = language.capitalize() + " · " + definition["label"]
            definition["scope"] = definition["scope"].replace("16.0301", cip)
            definition["notes"] = definition["notes"].replace("China specialists", "language specialists")
            definitions[prefix + ":" + name] = definition
    observations, sources = [], {}
    for year in years:
        base = IPEDS_CURRENT if year >= 2025 else IPEDS
        directory = f"fa{year}hd.zip" if year < 2002 else f"HD{year}.zip"
        hd_content, hd_asset = client.get(base + directory)
        hd_member, hd = zip_csv(hd_content)
        required = {"UNITID", "STABBR", "DEGGRANT", "PSET4FLG"}
        if not required.issubset(hd.fieldnames or []):
            raise ValueError("IPEDS institutional filters changed in " + str(year))
        eligible = {r["UNITID"] for r in hd
                    if r["STABBR"] in STATES and r["DEGGRANT"].strip() == "1"
                    and r["PSET4FLG"].strip() == "1"}
        content, asset = client.get(base + f"C{year}_A.zip")
        member, records = zip_csv(content)
        if not {"UNITID", "CIPCODE", "AWLEVEL"}.issubset(records.fieldnames or []):
            raise ValueError("IPEDS completions columns changed")
        total_fields = (["CTOTALT"] if "CTOTALT" in records.fieldnames
                        else ["CRACE15", "CRACE16"])
        if not set(total_fields).issubset(records.fieldnames):
            raise ValueError("IPEDS completion total fields changed")
        no_major_number = "MAJORNUM" not in records.fieldnames
        if no_major_number and year >= 2001:
            raise ValueError("IPEDS first-major identifier missing after 2000")
        totals, cells, flags = Counter(), Counter(), Counter()
        seen = set()
        for r in records:
            prefix = codes.get(r["CIPCODE"].strip())
            if r["UNITID"] not in eligible or prefix is None:
                continue
            if r.get("MAJORNUM", "1").strip() != "1":
                continue
            award = int(r["AWLEVEL"])
            if award not in {5, 7, 9, 17}:
                continue
            key = r["UNITID"], prefix, award
            if key in seen:
                raise ValueError("Duplicate IPEDS first-major award cell")
            seen.add(key)
            value = sum(int(r[field]) for field in total_fields)
            if value < 0:
                raise ValueError("IPEDS blank/suppressed or invalid award count")
            totals[(prefix, award)] += value
            cells[(prefix, award)] += 1
            for field in total_fields:
                flags[prefix + ":" + field + ":" + r.get("X" + field, "not supplied")] += 1
        if not any(prefix == "ipeds" for _, prefix, _ in seen):
            raise ValueError("No verified Chinese degree-award records in " + str(year))
        sid = f"ipeds_{year}"
        provisional = year == 2025 and "_rv" not in member.lower()
        vintage = f"Completions survey {year}" + (
            "; provisional 2024/25 release, July 28, 2026" if provisional else "")
        sources[sid] = source(
            "National Center for Education Statistics, IPEDS", vintage,
            [hd_asset, asset], f"{hd_member}: geographic/Title IV/degree filters; {member}: "
            "CIPCODE=16.0301 (Chinese), 16.0302 (Japanese), 16.0303 (Korean); " +
            ("no MAJORNUM field in source" if no_major_number else "MAJORNUM=1") +
            "; total=" + "+".join(total_fields),
            "Only one data revision used; _RV preferred when present. Source flag cell counts: " +
            json.dumps(dict(flags), sort_keys=True))
        for prefix in codes.values():
            for name, award in (("bachelors", 5), ("masters", 7), ("doctorates", 9 if year < 2010 else 17)):
                value = totals[(prefix, award)] if cells[(prefix, award)] else None
                observations.append(row(
                    prefix + ":" + name, year, value, sid, "reported_or_imputed" if value is not None else "missing",
                    f"Academic year {year - 1}/{str(year)[-2:]}; AWLEVEL={award}; "
                    f"{cells[(prefix, award)]} eligible source cells." +
                    (" Provisional NCES 2024/25 release of July 28, 2026." if provisional else ""),
                    prefix + ":" + name + ":no_major_split" if no_major_number else
                    prefix + ":doctoral_pre2010" if name == "doctorates" and year < 2010 else prefix + ":" + name))
        counts = {language: totals[(prefix, 5)] if cells[(prefix, 5)] else None for language, prefix in
                  [("Chinese", "ipeds"), ("Japanese", "ipeds_japanese"), ("Korean", "ipeds_korean")]}
        print(f"IPEDS {year} bachelor's: {json.dumps(counts)}", flush=True)
    return definitions, observations, sources


def parse_ap_chinese(content, year):
    return parse_ap_language(content, year, "Chinese")


def parse_ap_language(content, year, language):
    text = "\n".join(pdf_text(content))
    if (language.upper() + " LANGUAGE" not in " ".join(text.upper().split())
            or "Global" not in text or str(year) not in text):
        raise ValueError(f"AP {language} report subject/geography/year changed")
    total = re.search(r"Total Group.*?Number of Students\s+([\d,]+)", text, re.DOTALL)
    standard = re.search(r"Standard Group\*+.*?Number of Students\s+([\d,]+)", text, re.DOTALL)
    if not total or not standard:
        raise ValueError("AP reviewed total/standard-group layout changed")
    counts = int(total[1].replace(",", "")), int(standard[1].replace(",", ""))
    if counts[1] > counts[0]:
        raise ValueError("AP standard group cannot exceed total group")
    return counts


def parse_ap_history(content, language):
    soup = BeautifulSoup(content, "html.parser")
    result = {}
    for table in soup.select("table"):
        caption = table.find("caption")
        heading = caption.get_text(" ", strip=True) if caption else ""
        if f"AP {language} Language and Culture" not in heading:
            continue
        group = "standard" if "(Standard Group)" in heading else "total" if "(Total Group)" in heading else None
        if group is None or group in result:
            raise ValueError("AP historical group identity changed")
        rows = table.select("tr")
        columns = [re.sub(r"\s+", "", c.get_text()) for c in rows[0].find_all(["th", "td"])]
        if columns != ["Year", "5", "4", "3", "2", "1", "3+", "TestTakers", "MeanScore"]:
            raise ValueError("AP historical count columns changed")
        values = {}
        for tr in rows[1:]:
            cells = [c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"])]
            if len(cells) != len(columns) or not re.fullmatch(r"20\d{2}", cells[0]):
                raise ValueError("AP historical row layout changed")
            if not cells[7]:
                continue
            if not re.fullmatch(r"\d[\d,]*", cells[7]):
                raise ValueError("AP participation count is not an integer")
            year, value = int(cells[0]), int(cells[7].replace(",", ""))
            if year in values:
                raise ValueError("Duplicate AP historical year")
            values[year] = value
        result[group] = values
    if set(result) != {"total", "standard"}:
        raise ValueError("AP historical total/standard tables missing")
    return result


def ap_benchmarks(client, existing_sources):
    sources, observations, totals = {}, [], {}
    for sid in ["ap_volume_2017", "ap_volume_2023", "ap_volume_2024"]:
        body, _ = client.get(existing_sources[sid]["assets"][0]["request_url"])
        if sid.endswith("2017"):
            book = xlrd.open_workbook(file_contents=body)
            sheet = book.sheet_by_index(0)
            records = [sheet.row_values(i) for i in range(sheet.nrows)]
            header = next(r for r in records if any(str(v).strip() in {"2007", "2007.0"} for v in r))
        else:
            book = load_workbook(io.BytesIO(body), data_only=True, read_only=True)
            records = list(book.active.values)
            start = 2013 if sid.endswith("2023") else 2014
            header = next(r for r in records if any(str(v).strip() == str(start) for v in r))
        subject = next(r for r in records if str(r[0]).strip().lower() in {
            "japanese language", "japanese language and culture", "japanese language & culture"})
        for i, label in enumerate(header):
            label = str(label).strip().removesuffix(".0")
            if re.fullmatch(r"20\d{2}", label):
                value = subject[i]
                if isinstance(value, bool) or not isinstance(value, (float, int)) or not math.isfinite(value) or value < 0 or int(value) != value:
                    raise ValueError("Invalid AP Japanese volume count")
                totals[int(label)] = (int(value), sid)
        existing_sources[sid]["locator"] += "; Japanese Language row"
    if set(totals) != set(range(2007, 2025)):
        raise ValueError("AP Japanese annual history incomplete")
    observations += [row("ap:japanese_total", year, value, sid) for year, (value, sid) in sorted(totals.items())]
    for language in ["Chinese", "Japanese"]:
        url = f"https://apstudents.collegeboard.org/about-ap-scores/score-distributions/ap-{language.lower()}-language-and-culture"
        body, asset = client.get(url)
        history = parse_ap_history(body, language)
        sid = f"ap_{language.lower()}_history"
        sources[sid] = source("College Board", f"Historical AP {language} participation tables", [asset],
                              "Standard Group table; Test Takers column",
                              "Global participation. Empty 2026 count cells are not zero or inferred from score percentages. "
                              "Final subject-specific PDF counts take precedence over retrospective webpage counts.")
        end = 2022 if language == "Chinese" else 2024
        if not set(range(2007, end + 1)) <= set(history["standard"]):
            raise ValueError("AP standard-group annual history incomplete")
        group = "ap:standard_group" if language == "Chinese" else "ap:japanese_standard_group"
        observations += [row(group, year, history["standard"][year], sid) for year in range(2007, end + 1)]
    for year in (2023, 2024, 2025):
        url = f"https://apcentral.collegeboard.org/media/pdf/ap{year % 100}-japanese-language-score-distributions.pdf"
        body, asset = client.get(url)
        total, standard = parse_ap_language(body, year, "Japanese")
        sid = f"ap_japanese_{year}"
        sources[sid] = source("College Board", f"AP Japanese score distributions, May {year}", [asset],
                              "Total Group and Standard Group; Number of Students",
                              "Global final subject report; takes priority over earlier volume or webpage vintages.")
        observations = [r for r in observations if not (r["series_id"].startswith("ap:japanese") and r["year"] == year)]
        observations += [row("ap:japanese_total", year, total, sid),
                         row("ap:japanese_standard_group", year, standard, sid)]
    body, asset = client.get(AP_COURSES)
    catalog = BeautifulSoup(body, "html.parser")
    names = {a.get_text(" ", strip=True) for a in catalog.select("a")}
    if not {"AP Chinese Language and Culture", "AP Japanese Language and Culture"} <= names or any("AP Korean" in n for n in names):
        raise ValueError("Review the current AP language catalogue before publishing the Korean availability note")
    sources["ap_course_catalog"] = source("College Board", "AP course and exam catalogue", [asset],
                                           "World Languages and Cultures subject list",
                                           "Chinese and Japanese are listed; no Korean AP exam is listed. Absence of an exam is not zero Korean learners.")
    definitions = {
        "ap:japanese_total": series("Japanese · all participants", "exam_participation", "exam participants",
                                    "Global AP Japanese Language and Culture exam population, including international participation",
                                    "annual", notes="Not U.S.-only or a count of all Japanese learners."),
        "ap:japanese_standard_group": series("Japanese · standard-group subset", "exam_participation", "exam participants",
                                             "Standard group within global AP Japanese reports; neither regular language use/hearing nor one-month-plus residence reported",
                                             "annual", notes="Subset of all Japanese participants, not a separate geographic, nationality or ancestry group."),
    }
    return definitions, observations, sources


def ap_global(client):
    candidates = {}
    sources = {}
    url = "https://reports.collegeboard.org/media/xls/2017-Exam-Volume-Change_1.xls"
    body, asset = client.get(url)
    book = xlrd.open_workbook(file_contents=body)
    sheet = book.sheet_by_index(0)
    rows = [sheet.row_values(i) for i in range(sheet.nrows)]
    header = next(r for r in rows if any(str(v).strip() in {"2007", "2007.0"} for v in r))
    chinese = next(r for r in rows if str(r[0]).strip().lower() in {
        "chinese language", "chinese language and culture", "chinese language & culture"})
    sources["ap_volume_2017"] = source(
        "College Board", "AP exam-volume report 2007–2017", [asset], sheet.name + "; Chinese Language row",
        "Global administrations, not a U.S.-only learner count.")
    for i, label in enumerate(header):
        label = str(label).strip().removesuffix(".0")
        if re.fullmatch(r"20\d{2}", label):
            value = chinese[i]
            if int(label) >= 2007:
                if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0 or value != int(value):
                    raise ValueError("Invalid AP annual exam count")
                candidates[int(label)] = (int(value), "ap_volume_2017")
    if set(candidates) != set(range(2007, 2018)):
        raise ValueError("Reviewed AP 2007-2017 annual coverage changed")
    for start, end in [(2013, 2023), (2014, 2024)]:
        url = f"https://reports.collegeboard.org/media/xlsx/ap-exam-volume-changes-{start}-{end}.xlsx"
        body, asset = client.get(url)
        book = load_workbook(io.BytesIO(body), data_only=True, read_only=True)
        rows = list(book.active.values)
        header = next((r for r in rows if any(str(c).strip() == str(start) for c in r)), None)
        chinese = next((r for r in rows if str(r[0]).strip().lower() in {
            "chinese language", "chinese language and culture", "chinese language & culture"}), None)
        if header is None or chinese is None:
            raise ValueError("AP reviewed volume-table columns or subject changed")
        source_id = f"ap_volume_{end}"
        sources[source_id] = source("College Board", f"AP exam-volume report through {end}",
                                    [asset], book.active.title + "; Chinese Language row",
                                    "Global AP exam administrations, not a U.S.-only learner count.")
        updated = set()
        for i, label in enumerate(header):
            if re.fullmatch(r"20\d{2}", str(label).strip()) and int(str(label).strip()) >= 2007:
                year, value = int(str(label).strip()), chinese[i]
                if year in updated or isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0 or value != int(value):
                    raise ValueError("Invalid or duplicated AP annual exam count")
                updated.add(year)
                candidates[year] = (int(value), source_id)
        if updated != set(range(start, end + 1)):
            raise ValueError("Reviewed AP annual volume coverage changed")
    observations = [row("ap:global_total", year, value, sid)
                    for year, (value, sid) in sorted(candidates.items())]
    for year in (2023, 2024, 2025):
        url = f"https://apcentral.collegeboard.org/media/pdf/ap{year % 100}-chinese-language-score-distributions.pdf"
        body, asset = client.get(url)
        total, standard = parse_ap_chinese(body, year)
        source_id = f"ap_chinese_{year}"
        sources[source_id] = source(
            "College Board", f"AP Chinese score distributions, May {year}", [asset],
            "Total Group and Standard Group; Number of Students",
            "Global report. Standard group is an exposure/training definition, NOT a U.S.-only or nonheritage demographic count. "
            "Subject-specific count takes priority over earlier program-summary vintages.")
        observations = [r for r in observations if not (r["series_id"] == "ap:global_total" and r["year"] == year)]
        observations += [row("ap:global_total", year, total, source_id),
                         row("ap:standard_group", year, standard, source_id)]
    definitions = {
        "ap:global_total": series("Chinese · all participants", "exam_participation", "exam participants",
                                 "Global AP Chinese Language and Culture exam population, including international participation",
                                 "annual", notes="Not U.S.-only; not all Chinese learners; repeated participants across years are not deduplicated."),
        "ap:standard_group": series("Chinese · standard-group subset", "exam_participation", "exam participants",
                                   "Standard group within global AP Chinese reports; no regular language exposure or one-month-plus residence indicated",
                                   "annual", notes="Generally trained in U.S. schools, but not a geographic or ancestry classification. This is a subset of global total; never add them."),
    }
    extra_definitions, extra_observations, extra_sources = ap_benchmarks(client, sources)
    definitions.update(extra_definitions)
    observations.extend(extra_observations)
    sources.update(extra_sources)
    return definitions, observations, sources


def ap_us_cohorts(client):
    base = "https://secure-media.collegeboard.org/digitalServices/"
    reports = [
        (2011, base + "public/pdf/ap/rtn/AP-Report-2011-Subject-Supplement.pdf"),
        (2013, base + "pdf/ap/rtn/9th-annual/9th-annual-ap-report-subject-supplement-all-subjects.pdf"),
        (2014, base + "pdf/ap/rtn/10th-annual/10th-annual-ap-report-subject-supplement-chinese-language-and-culture.pdf"),
    ]
    candidates, sources = {}, {}
    for vintage, url in reports:
        body, asset = client.get(url)
        page = next((p for p in pdf_text(body) if re.search(r"Chinese Language\s+and\s*Culture", p)
                     and "score distribut" in p.lower()), None)
        text = re.sub(r"\s+", " ", page or "")
        if "U.S. public schools only" not in text or "graduating class" not in text:
            raise ValueError("AP U.S. public-school graduating-cohort definition changed")
        if vintage == 2011:
            matches = re.findall(
                r"(?m)^\s*([1-5])\s+N/A\s+N/A\s+([\d,]+)\s*/\s*[\d.]+%\s+([\d,]+)\s*/", page)
            if len(matches) != 10 or [m[0] for m in matches[:5]] != list("54321"):
                raise ValueError("AP 2011 total-group score columns changed")
            totals = {year: sum(int(m[column].replace(",", "")) for m in matches[:5])
                      for year, column in [(2009, 1), (2010, 2)]}
        elif vintage == 2013:
            totals = {}
            for year in (2007, 2011, 2012):
                match = re.search(r"(?m)^\s*" + str(year) + r"\s*\n((?:\s*\d[\d,]*\s*\n){5})", page)
                if not match:
                    raise ValueError("AP 2013 total-group score columns changed")
                totals[year] = sum(int(v.replace(",", "")) for v in match[1].split())
        else:
            matches = re.findall(
                r"(?m)^\s*([1-5])\s+[—-]\s+[—-]\s+([\d,]+)\s+[\d.]+%\s+([\d,]+)\s+[\d.]+%\s+([\d,]+)\s+[\d.]+%", page)
            if len(matches) != 5 or [m[0] for m in matches] != list("54321"):
                raise ValueError("AP 2014 total-group score columns changed")
            totals = {year: sum(int(m[column].replace(",", "")) for m in matches)
                      for year, column in [(2008, 1), (2012, 2), (2013, 3)]}
        source_id = f"ap_cohort_{vintage}"
        sources[source_id] = source(
            "College Board", f"Report to the Nation {vintage}; U.S. public-school graduating cohorts",
            [asset], "Chinese Language and Culture; sum of five Total Group score counts",
            "Exams taken during high school by each graduating cohort, not an annual exam administration, "
            "all U.S. schools, unique learners or the global series. Later reviewed reports take priority "
            "for overlapping cohorts; the 2008 total was revised across publications.")
        candidates.update({year: (value, source_id) for year, value in totals.items()})
    sid = "ap_cohort:us_public"
    definitions = {sid: series(
        "U.S. public-school cohorts", "cohort_exam_participation", "exams",
        "U.S. public-school graduating classes; Chinese AP exams taken during high school; not private-school or global totals",
        "annual", "graduating class",
        "Year identifies the graduating class, not the year each exam was taken. The exam began in 2007. "
        "Only 2007-2013 cohorts verified here; do not splice into the global annual administration series.")}
    observations = [row(sid, year, value, source_id)
                    for year, (value, source_id) in sorted(candidates.items())]
    return definitions, observations, sources


def schools(client):
    definitions, observations, sources = {}, [], {}
    inventories = [
        ("cal_directory", 2011, "CAL directory",
         "https://webapp.cal.org/Immersion/Doc/Language%20of%20Instruction%20in%20Immersion%20Programs.pdf",
         r"Mandarin\s+(\d+)\s+13\.4", 71,
         "U.S. foreign-language immersion directory; CAL's 2011 inventory definition, not ARC public-school canvass"),
        ("arc_public", 2021, "Public-school DLI",
         "https://www.americancouncils.org/sites/default/files/documents/pages/2021-10/Canvass%20DLI%20-%20October%202021-2_ac.pdf",
         r"Chinese\s+(\d+)\s+", 312,
         "U.S. public-school dual-language/immersion canvass; each school housing a program is counted"),
    ]
    for name, year, label, url, pattern, expected, scope in inventories:
        content, asset = client.get(url)
        match = re.search(pattern, "\n".join(pdf_text(content)))
        if match is None or int(match[1]) != expected:
            raise ValueError("Reviewed school inventory layout/count changed: " + name)
        sid = "schools:" + name
        definitions[sid] = series(label, "program_inventory", "programs", scope, "snapshot",
                                  notes="Different inventory universe and methodology; not a comparable annual growth series.")
        observations.append(row(sid, year, expected, name))
        sources[name] = source("Center for Applied Linguistics" if year == 2011 else "American Councils Research Center",
                               str(year) + " program inventory", [asset], "Language-of-instruction table, page 1",
                               "Program/school counts, not students. Scope differs across inventories.")
    page, page_asset = client.get(NEXUS)
    urls = [urljoin(NEXUS, a["href"]) for a in BeautifulSoup(page, "html.parser").select("a[href]")
            if "Survey%20Report.pdf" in a["href"]]
    if len(set(urls)) != 1:
        raise ValueError("Nexus primary report link missing or ambiguous")
    content, asset = client.get(urls[0])
    pages = pdf_text(content)
    appendix = next(p for p in pages if "State Number of Programs" in p and "Total 335" in p)
    state_counts = re.findall(r"^([A-Z]{2})\s+(\d+)\s*$", appendix, re.MULTILINE)
    if not state_counts or any(s not in STATES for s, _ in state_counts) or sum(int(n) for _, n in state_counts) != 335:
        raise ValueError("Nexus state inventory totals do not reconcile")
    definitions["schools:nexus_all"] = series(
        "Mixed school types", "program_inventory", "programs",
        "Nexus compiled U.S. immersion inventory, including public, charter and private schools", "snapshot",
        notes="335 inventoried programs, NOT the 103 survey responses or a national learner census. "
        "The survey-response enrollment total includes one Hong Kong response and is excluded from U.S. learner counts.")
    observations.append(row("schools:nexus_all", 2025, 335, "nexus_2025"))
    sources["nexus_2025"] = source(
        "Nexus Mandarin", "2025 survey and directory; published 2026", [page_asset, asset],
        "Executive summary page 3; state inventory appendix page 55",
        "103 responses include one in Hong Kong; 34,341 response-associated students are NOT published here as U.S. enrollment. "
        "All source-state inventory counts reconcile to 335.")
    url = "https://www.americancouncils.org/sites/default/files/FLE-report-June17.pdf"
    content, asset = client.get(url)
    table = next(p for p in pdf_text(content) if "227,086" in p and "Enrollment for Major Languages by State" in p)
    if not re.search(r"Total\s+26,045\s+130,411\s+227,086", table):
        raise ValueError("Reviewed K-12 Chinese enrollment estimate changed")
    definitions["schools:k12_estimate"] = series(
        "K-12 Chinese enrollment estimate", "estimated_enrollments", "enrollments",
        "U.S. formal K-12 schools; reported state counts plus model-based estimates for missing state data",
        "snapshot", "2014-15-era reference period",
        "Published June 2017. A mixed-data/model estimate, not an enumerated census; excludes substantial community/heritage and summer networks.")
    observations.append(row("schools:k12_estimate", 2015, 227086, "k12_2017", "estimated",
                            "2014-15-era survey reference period; 2015 is the ending-year plot key, not an exact common measurement date."))
    sources["k12_2017"] = source(
        "American Councils and partner organizations", "June 2017 K-16 survey report",
        [asset], "Table 2, page 8; Chinese column national total",
        "Reported and modeled state enrollment totals. Not an immersion-program count; never summed with later college/AP figures.")
    return definitions, observations, sources


def school_comparisons(client):
    url = "https://www.americancouncils.org/sites/default/files/FLE-report-June17.pdf"
    content, asset = client.get(url)
    pages = pdf_text(content)
    total = next(text for text in pages if "Enrollment for Major Languages by State" in text and "227,086" in text)
    if not re.search(r"Total\s+26,045\s+130,411\s+227,086\s+1,289,004\s+330,898\s+67,909", total):
        raise ValueError("American Councils national enrollment estimate columns changed")
    availability = next(text for text in pages if "Distribution of Foreign Language Programs" in text and "Number of HS programs" in text)
    counts = {}
    for language, expected in [("Chinese", 1144), ("Japanese", 433), ("Korean", 43)]:
        match = re.search(r"^\s*" + language + r"\s+(\d+)\s+[\d.]+", availability, re.MULTILINE)
        if not match or int(match[1]) != expected:
            raise ValueError("Reviewed high-school availability count changed")
        counts[language.lower()] = int(match[1])
    korean = next(text for text in pages if "In 43 schools across 22 states" in text)
    if "936 students" not in korean:
        raise ValueError("Reviewed Korean respondent enrollment changed")
    definitions = {
        "schools:k12_japanese_estimate": series(
            "K-12 Japanese estimate", "estimated_enrollments", "enrollments",
            "U.S. formal K-12 schools; reported state counts plus model-based estimates for missing state data",
            "snapshot", "2014-15-era reference period",
            "2017 report Table 2. Not unique people or U.S. citizens; not comparable with the Korean high-school respondent count."),
        "schools:korean_respondent_enrollment": series(
            "Korean high-school respondents", "reported_enrollments", "enrollments",
            "936 Korean-class enrollments reported by 43 schools across 22 states; respondent subset, not a national K-12 estimate",
            "snapshot", "2014-15-era reference period",
            "2017 report page 22. Deliberately not plotted alongside national Chinese/Japanese K-12 estimates."),
    }
    observations = [row("schools:k12_japanese_estimate", 2015, 67909, "k12_comparisons_2017", "estimated",
                        "School-year 2014-15-era data; 2015 is an ending-year plot key, not a 2017 measurement."),
                    row("schools:korean_respondent_enrollment", 2015, 936, "k12_comparisons_2017",
                        notes="Reported respondent enrollments, not all Korean learners or a national total.")]
    for language, count in counts.items():
        key = "school_availability:" + language
        definitions[key] = series(language.capitalize(), "reported_programs", "reported programs",
                                  "High-school language programs reported in the American Councils survey; schools/districts may contribute multiple language offerings",
                                  "snapshot", "2014-15 survey",
                                  "Table 4, page 11. Availability, not student enrollment, latent demand or a census of unique U.S. schools.")
        observations.append(row(key, 2015, count, "k12_comparisons_2017"))
    sources = {"k12_comparisons_2017": source(
        "American Councils and partner organizations", "National K-12 Foreign Language Enrollment Survey, June 2017",
        [asset], "Table 2 p8 (national estimates); Table 4 p11 (reported high-school programs); p22 (Korean respondent enrollments)",
        "2014-15-era reference period. Chinese/Japanese national totals combine reported data and modeling; "
        "the Korean 936 is a respondent-only count and cannot be substituted for a national estimate. "
        "Availability counts measure reported language programs, not unique people, unique schools or demand.")}
    return definitions, observations, sources
