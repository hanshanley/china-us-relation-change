"""Manually reviewed business-presence aggregates, not bilateral flows."""

import csv
import hashlib
import io
import json
import math
import re
import zipfile
from pathlib import PurePosixPath
from decimal import Decimal, InvalidOperation
from urllib.parse import urljoin, urlsplit

import xlrd
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from .language_sources import pdf_text, row, series, source


BEA_CURRENT = "https://www.bea.gov/sites/default/files/2025-08/selected-mofas-cntry.xls"
BEA_ARCHIVES = {
    2000: "https://apps.bea.gov/international/zip/USDIA2000revised.ZIP",
    2001: "https://apps.bea.gov/international/zip/USDIA2001revised.zip",
    2002: "https://apps.bea.gov/international/zip/USDIA2002revised.zip",
    2003: "https://apps.bea.gov/international/zip/usdia2003revised.zip",
    2004: "https://apps.bea.gov/international/zip/usdia2004revised.zip",
    2005: "https://apps.bea.gov/international/zip/USDIA2005revised.zip",
    2006: "https://apps.bea.gov/international/zip/USDIA2006revised.zip",
    2007: "https://apps.bea.gov/international/zip/USDIA2007revised.zip",
    2008: "https://apps.bea.gov/international/pdf/usdia_2008r/USDIA_08R.zip",
}
PARENT_BOOKS = {
    2009: "https://apps.bea.gov/international/xls/usdia_2009r/Part%20I%20K1-K2.xls",
    2014: "https://apps.bea.gov/international/xls/usdia2014r/PartI-K1-K2.xls",
    2019: "https://apps.bea.gov/international/xls/usdia2019r/PartI-K1-K2.xls",
}
EUROSTAT = "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/"
ITALY_EUROSTAT_METHODS = "https://ec.europa.eu/eurostat/cache/metadata/EN/fats_out_simsofat_it.htm"
ISTAT_ITALY_REPORT = "https://www.istat.it/wp-content/uploads/2025/11/Report_Imprese-multinazionali_Anno-2023.pdf"
CBS_NETHERLANDS_REPORT = "https://longreads.cbs.nl/nederland-handelsland-2026/buitenlandse-investeringen-en-multinationals/"
CBS_NETHERLANDS_METHODS = "https://www.cbs.nl/nl-nl/onze-diensten/methoden/onderzoeksomschrijvingen/korte-onderzoeksbeschrijvingen/outward-fats-statistiek-over-nederlandse-zeggenschappen-over-buitenlandse-bedrijven"
ONS_DATA = "https://www.ons.gov.uk/file?uri=/businessindustryandtrade/internationaltrade/datasets/ukoutwardforeignaffiliatestatisticsofats/2019/ofatspublicationfinal.xlsx"
ONS_COVERAGE = "https://www.ons.gov.uk/aboutus/transparencyandgovernance/freedomofinformationfoi/ukoutwardforeignaffiliatestatisticsofats2015to2025"
BUNDESBANK_FATS = "https://www.bundesbank.de/en/statistics/external-sector/direct-investments/foreign-affiliates-statistics-fats--795240"
GERMANY_FATS_2020 = (
    "https://web.archive.org/web/20250419145839id_/https://www.bundesbank.de/resource/blob/"
    "768898/57e1855ed8172f316354ab4b6d39f6d5/472B63F073F071307366337C94F8C870/statistik-outward-fats-data.pdf")
SSB_COUNTRIES = "https://data.ssb.no/api/v0/en/table/09827"
SWEDEN_COUNTRIES = (
    "https://statistik.tillvaxtanalys.se/PxWeb/api/v1/sv/"
    "Tillv%C3%A4xtanalys%20statistikdatabas/"
    "Svenska%20koncerner%20med%20dotterbolag%20i%20utlandet/"
    "internationella_foretag_koncerner_etableringsland_1996_2024.px")
SWEDEN_QUALITY = (
    "https://www.tillvaxtanalys.se/download/18.2c8ee5971a08641ad00c03e9/"
    "1789409959340/NV1501_Kvalitetsdeklaration_2024.pdf")
FINLAND_COUNTRIES = "https://pxdata.stat.fi/PxWeb/api/v1/en/StatFin/stu/13k8.px"
FINLAND_METHODS = "https://stat.fi/en/documentation/changes-in-statistics/clh08g97cga760bw05x4vrip0"
FINLAND_DOCUMENTATION = "https://stat.fi/en/documentation/documentation-of-statistics/stu"
FINLAND_EUROSTAT_METHODS = "https://ec.europa.eu/eurostat/cache/metadata/EN/fats_out_simsofat_fi.htm"
SWEDEN_EUROSTAT_METHODS = "https://ec.europa.eu/eurostat/cache/metadata/EN/fats_out_simsofat_se.htm"
AUSTRIA_EUROSTAT_METHODS = "https://ec.europa.eu/eurostat/cache/metadata/EN/fats_out_simsofat_at.htm"
POLAND_EUROSTAT_METHODS = "https://ec.europa.eu/eurostat/cache/metadata/EN/fats_out_simsofat_pl.htm"
SNB_CUBE = "opanmuausuland"
SNB_API = "https://data.snb.ch/api/cube/" + SNB_CUBE + "/"
SNB_METHODS = "https://data.snb.ch/de/topics/aube/doc/explanations_aube_amne"
SNB_CONTROL_METHODS = "https://data.snb.ch/de/topics/aube/doc/explanations_aube_fdi"
SNB_COUNTRY_METHODS = "https://data.snb.ch/de/topics/aube/doc/explanations_aube_laend"
SNB_SERIES = "snb_swiss_liechtenstein:china"
SNB_CONTEXT_NOTE = (
    " The SNB Switzerland/Liechtenstein reporting area is a separate qualified supplement, "
    "not a pure-Switzerland population; majority subsidiaries include unincorporated branches "
    "and the current groupwide CHF10m reporting frame is not backdated or treated as all firms.")
SPAIN_EUROSTAT_METHODS = "https://ec.europa.eu/eurostat/cache/metadata/EN/fats_out_simsofat_es.htm"
DENMARK_METADATA = "https://api.statbank.dk/v1/tableinfo"
DENMARK_DATA = "https://api.statbank.dk/v1/data"
DENMARK_METHODS = (
    "https://www.dst.dk/da/Statistik/dokumentation/statistikdokumentation/"
    "danske-datterselskaber-i-udlandet/")
DENMARK_QUERY = {
    "table": "OFATS4", "format": "JSONSTAT", "lang": "en",
    "variables": [{"code": "LAND", "values": ["CN"]},
                  {"code": "ENHED", "values": ["DSANTAL"]},
                  {"code": "Tid", "values": ["*"]}],
}
NORDIC_QUERIES = {
    "sweden": {"query": [
        {"code": "Tid", "selection": {"filter": "item", "values": [str(y) for y in range(2000, 2025)]}},
        {"code": "Etableringsland", "selection": {"filter": "item", "values": ["cn"]}},
        {"code": "variabel", "selection": {"filter": "item", "values": ["2"]}},
    ], "response": {"format": "json-stat2"}},
    "finland": {"query": [
        {"code": "timeperiod_y", "selection": {"filter": "all", "values": ["*"]}},
        {"code": "valtio_27_20180101", "selection": {"filter": "item", "values": ["156"]}},
        {"code": "contentscode", "selection": {"filter": "item", "values": ["ofats_count"]}},
    ], "response": {"format": "json-stat2"}},
}
SSB_CHINA_QUERY = {
    "query": [
        {"code": "Land", "selection": {"filter": "item", "values": ["CN"]}},
        {"code": "ContentsCode", "selection": {"filter": "item", "values": ["Foretak"]}},
        {"code": "Tid", "selection": {"filter": "all", "values": ["*"]}},
    ],
    "response": {"format": "json-stat2"},
}


def count_cell(value):
    if value in {"(D)", "D"}:
        return None, "suppressed"
    if value in {"", ":", None}:
        return None, "missing"
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0 or value != int(value):
        raise ValueError("Invalid business entity count")
    return int(value), "observed"


def parse_nordic(data, country):
    schemas = {
        "sweden": (["Tid", "Etableringsland", "variabel"], "cn", "2",
                   "Myndigheten f\u00f6r tillv\u00e4xtpolitiska utv\u00e4rderingar och analyser",
                   "Antal dotterbolag", {None, ".."}),
        "finland": (["timeperiod_y", "valtio_27_20180101", "contentscode"], "156", "ofats_count",
                    "Statistics Finland, finnish affiliates abroad", "Number of affiliates", {None, ".", "..."}),
    }
    dimensions, geo, metric, publisher, label, flags = schemas[country]
    if data.get("class") != "dataset" or data.get("source") != publisher or data.get("id") != dimensions:
        raise ValueError("Nordic affiliate dataset identity changed")
    time, location, information = [data["dimension"][code]["category"] for code in dimensions]
    if (location["index"] != {geo: 0} or information["index"] != {metric: 0}
            or information["label"].get(metric) != label):
        raise ValueError("Nordic country or affiliate counting unit changed")
    years, values = time["index"], data.get("value")
    statuses = data.get("status", {})
    if (not years or data.get("size") != [len(years), 1, 1] or not isinstance(values, list)
            or len(values) != len(years) or set(years.values()) != set(range(len(years)))
            or any(not isinstance(year, str) or not re.fullmatch(r"\d{4}", year) for year in years)
            or set(statuses) - {str(i) for i in range(len(years))}):
        raise ValueError("Nordic periods, dimensions or source flags changed")
    result = []
    for year, position in sorted(years.items()):
        value, flag = values[position], statuses.get(str(position))
        if flag not in flags or (flag is not None and value is not None):
            raise ValueError("Unreviewed Nordic source flag or flagged numeric value")
        count, status = count_cell(value)
        if flag == "...":
            status = "suppressed"
        y = int(year)
        if country == "sweden":
            boundary = ("pre2010" if y < 2010 else "2010-population" if y < 2018 else
                        "2018-EGR" if y == 2018 else "2019-frame" if y < 2021 else "2021-GGH")
        else:
            boundary = "pre2021" if y < 2021 else "EBS2021"
        result.append((y, count, status, f"{country}:CN:affiliates:{boundary}"))
    return result


def nordic_affiliates(client):
    definitions, observations, sources = {}, [], {}
    for country, api, methodology in [
        ("sweden", SWEDEN_COUNTRIES, SWEDEN_QUALITY),
        ("finland", FINLAND_COUNTRIES, FINLAND_METHODS),
    ]:
        body, asset = client.post_json(api, NORDIC_QUERIES[country])
        points = parse_nordic(json.loads(body), country)
        _, method_asset = client.get(methodology)
        if country == "sweden":
            if method_asset["sha256"] != "5d7b77577caf3a9fb159c106bd7fe4e9196ef50f8e09b903b7cbb678d5301289":
                raise ValueError("Reviewed Swedish affiliate methodology changed")
            sid, source_id, publisher = "tillvaxtanalys_sweden:china", "sweden_country_table", "Swedish Agency for Growth Policy Analysis"
            scope = "Subsidiary companies in China (CN) of Swedish-headed groups; all industries; Hong Kong and Macao separately coded"
            notes = ("Swedish owners hold at least 50% of group voting rights; a subsidiary has >50% ownership "
                     "or dominant influence. Counts at balance-sheet date, not unique parents or offices. "
                     "Subsidiary counts began in 2010; 2000-2009 source '..' cells remain missing. "
                     "2018 EGR coverage/imputation, 2019 frame repair and 2021 GDC-to-GGH ownership boundaries "
                     "remain separate. The API's 9999 update sentinel is not a publication date.")
            vintage = "2024 survey; published 15 September 2026; quality declaration 14 September 2026"
            locator = "Country table; Tid=2000-2024; Etableringsland=cn; variabel=2 (subsidiaries), not 0 (groups); quality declaration pp3-4,9-11"
        else:
            sid, source_id, publisher = "statfin_finland:china", "finland_13k8", "Statistics Finland"
            scope = "Finland-reported foreign affiliates and branches operating in China (156); Hong Kong and Macao separately coded"
            notes = ("Active affiliates and branches controlled by Finland-resident enterprises; an affiliate "
                     "has >50% direct/indirect voting control. All size classes and industries; not unique "
                     "parents or a physical-office census. Financial years ending within the reference calendar year. The 2021 EBS methodology "
                     "boundary remains separate. This supplementary country table starts in 2020; earlier "
                     "national regional aggregates are not allocated to China. Not matched to other countries' thresholds.")
            vintage = f"StatFin 13k8; reference years {points[0][0]}-{points[-1][0]}"
            locator = "13k8; country156 (China); contentscode=ofats_count; all reported years; 2021 methodology-change notice"
        definitions[sid] = series(country.title(), "controlled_enterprises", "enterprises",
                                  scope, "annual", "reference year", notes)
        assets = [asset, method_asset]
        if country == "finland":
            _, definition_asset = client.get(FINLAND_DOCUMENTATION)
            assets.append(definition_asset)
        sources[source_id] = source(publisher, vintage, assets, locator, notes)
        observations.extend(row(sid, year, value, source_id, status, notes, comparability)
                            for year, value, status, comparability in points if year >= 2000)
    return definitions, observations, sources


def uk_early_affiliates(client):
    url = (EUROSTAT + "fats_out2?lang=EN&geo=UK&partner=CN_X_HK&"
           "indic_bp=ENT&nace_r1=C-O_X_L")
    body, asset = client.get(url)
    data = json.loads(body)
    points = parse_eurostat(data, "fats_out2", "UK")
    if [point[0] for point in points] != [2007, 2008, 2009]:
        raise ValueError("Earlier UK outward-affiliate year coverage changed")
    sid, source_id = "eurostat_uk:china", "fats_out2_UK"
    notes = (
        "Earlier Eurostat UK NACE Rev. 1.1 history is separate from NACE Rev. 2 from 2010; "
        "missing source cells are not zero. Later Eurostat history and revised ONS 2018/2019 "
        "estimates remain distinct. The 2018 source/sector boundary is disconnected; ONS "
        "total activity includes agriculture. The 2018 China count agrees across both sources. "
        "No UK outward affiliate statistics from 2020 onward verified; investment flows are not substituted.")
    definitions = {sid: series(
        "United Kingdom", "controlled_enterprises", "enterprises",
        "UK-controlled enterprises in China except Hong Kong; NACE C-O except L in 2009, "
        "NACE B-S except O in 2010-2017, ONS total published activity in 2018-2019",
        "annual", "reference year", notes)}
    sources = {source_id: source(
        "Eurostat", data["label"] + "; updated " + data["updated"], [asset],
        "A; UK; CN_X_HK; nace_r1=C-O_X_L; indic_bp=ENT",
        "Original 2007-2009 outward table; enterprise counts, not parent groups or inward affiliates. "
        "Preserve missing cells and the 2010 classification boundary.")}
    observations = [row(sid, year, value, source_id, status,
                        "Eurostat source flags: " + (flag or "none"), comparability)
                    for year, value, status, flag, comparability in points]
    return definitions, observations, sources


def french_affiliates(client):
    sid, observations, sources = "eurostat_france:china", [], {}
    notes = ("French-controlled enterprises in the source category China except Hong Kong; not unique "
             "parents, offices or INSEE's rounded broader count. Retain the 2007 pilot/2009 frame, "
             "2010 NACE, 2016 source-flag and 2021 EBS boundaries; 2008 is unreported. "
             "No invented 2000-2006 history or growth across those population changes.")
    definitions = {sid: series(
        "France (Eurostat)", "controlled_enterprises", "enterprises",
        "French-controlled enterprises in China except Hong Kong; legacy NACE C-O except L, then NACE B-S except O",
        "annual", "reference year", notes)}
    for dataset, indicator, activity, first, last in [
        ("fats_out2", "indic_bp=ENT", "nace_r1=C-O_X_L", 2007, 2009),
        ("fats_out2_r2", "indic_bp=ENT", "nace_r2=B-S_X_O", 2010, 2020),
        ("fats_out_part", "indic_sbs=ENT_NR", "nace_r2=B-S_X_O", 2021, 9999),
    ]:
        url = EUROSTAT + dataset + "?lang=EN&freq=A&geo=FR&partner=CN_X_HK&" + indicator + "&" + activity
        body, asset = client.get(url)
        data, source_id = json.loads(body), dataset + "_FR"
        sources[source_id] = source(
            "Eurostat", dataset + "; France; source updated " + data.get("updated", "not supplied"),
            [asset], "geo=FR; partner=CN_X_HK; " + indicator + "; " + activity,
            notes + " Original country/category index and observation flags are retained.")
        for year, value, status, flag, comparability in parse_eurostat(data, dataset, "FR"):
            if first <= year <= last:
                if dataset == "fats_out2":
                    comparability += {2007: ":pilot2007", 2008: ":unreported2008", 2009: ":annual2009"}[year]
                observations.append(row(sid, year, value, source_id, status,
                                        notes + " Source flags: " + flag, comparability))
    return definitions, observations, sources


def spanish_affiliates(client):
    sid, observations, sources = "eurostat_spain:china", [], {}
    notes = (
        "Spanish-resident ultimate control; Eurostat's China except Hong Kong host category. "
        "Financial and nonfinancial activities, not unique Spanish parents. "
        "Counts are reported from 2009; 2007/2008 are unreported and no 2000-2006 values are inferred. "
        "Retain the 2010 classification transition and the 2021 change from legal units to "
        "statistical enterprises. Recent estimates retain their original flags. "
        "INE's national nonfinancial population is different and is not appended to this series.")
    definitions = {sid: series(
        "Spain (Eurostat)", "controlled_enterprises", "enterprises",
        "Spanish-resident UCI enterprises in China except Hong Kong; legacy NACE C-O except L, "
        "then NACE B-S except O; financial and nonfinancial activities",
        "annual", "reference year", notes)}
    _, methodology = client.get(SPAIN_EUROSTAT_METHODS)
    for dataset, indicator, activity in [
        ("fats_out2", "indic_bp=ENT", "nace_r1=C-O_X_L"),
        ("fats_out2_r2", "indic_bp=ENT", "nace_r2=B-S_X_O"),
        ("fats_out_part", "indic_sbs=ENT_NR", "nace_r2=B-S_X_O"),
    ]:
        url = EUROSTAT + dataset + "?lang=EN&freq=A&geo=ES&partner=CN_X_HK&" + indicator + "&" + activity
        body, asset = client.get(url)
        data, source_id = json.loads(body), dataset + "_ES"
        sources[source_id] = source(
            "Eurostat", dataset + "; Spain; source updated " + data.get("updated", "not supplied"),
            [asset, methodology], "geo=ES; partner=CN_X_HK; " + indicator + "; " + activity,
            notes + " Original reporter metadata and observation flags are retained.")
        points = parse_eurostat(data, dataset, "ES")
        years = {point[0] for point in points}
        expected = (set(range(2007, 2010)) if dataset == "fats_out2" else
                    set(range(2010, 2021)) if dataset == "fats_out2_r2" else
                    set(range(2021, max(years) + 1)))
        if years != expected or (dataset == "fats_out_part" and max(years) < 2023):
            raise ValueError("Spanish affiliate history is incomplete; preserve the reviewed snapshot")
        observations.extend(
            row(sid, year, value, source_id, status, notes + " Source flags: " + flag, comparability)
            for year, value, status, flag, comparability in points)
    return definitions, observations, sources


def parse_denmark(data, metadata):
    if metadata.get("id") != "OFATS4" or metadata.get("suppressedDataValue") != "0":
        raise ValueError("Danish table identity or suppression convention changed")
    variables = {item["id"]: item for item in metadata.get("variables", [])}
    if set(variables) != {"LAND", "ENHED", "Tid"}:
        raise ValueError("Danish country-count dimensions changed")
    countries = {item["id"]: item["text"] for item in variables["LAND"]["values"]}
    units = {item["id"]: item["text"] for item in variables["ENHED"]["values"]}
    if any(countries.get(code) != name for code, name in (
            ("CN", "China"), ("HK", "Hong Kong"), ("MO", "Macao"))) or units.get("DSANTAL") != "Affiliates":
        raise ValueError("Danish host geography or subsidiary unit changed")
    years = sorted(item["id"] for item in variables["Tid"]["values"])
    if (not years or any(not re.fullmatch(r"\d{4}", year) for year in years)
            or years[0] != "2007" or int(years[-1]) < 2024
            or [int(year) for year in years] != list(range(2007, int(years[-1]) + 1))):
        raise ValueError("Danish reviewed history is incomplete")
    dataset = data.get("dataset", {})
    dimensions = dataset.get("dimension", {})
    if (dataset.get("extension", {}).get("px", {}).get("tableid") != "OFATS4"
            or set(dimensions.get("id", [])) != {"LAND", "ENHED", "ContentsCode", "Tid"}
            or len(dimensions.get("size", [])) != 4 or dataset.get("status")):
        raise ValueError("Danish dataset identity or quality flags changed")
    for name, code in (("LAND", "CN"), ("ENHED", "DSANTAL"), ("ContentsCode", "OFATS4")):
        if (dimensions[name]["category"]["index"] != {code: 0}
                or dimensions["size"][dimensions["id"].index(name)] != 1):
            raise ValueError("Danish response is not the requested China affiliate population")
    unit = dimensions["ContentsCode"]["category"].get("unit", {}).get("OFATS4", {})
    times = dimensions["Tid"]["category"]["index"]
    values = dataset.get("value", [])
    if (unit != {"base": "Number", "decimals": 0} or sorted(times) != years
            or set(times.values()) != set(range(len(years)))
            or dimensions["size"][dimensions["id"].index("Tid")] != len(years)
            or not isinstance(values, list) or len(values) != len(years)):
        raise ValueError("Danish count units or complete time axis changed")
    result = []
    for year in years:
        value, status = count_cell(values[times[year]])
        # The publisher also uses zero for suppressed cells; do not infer zero subsidiaries.
        if value == 0:
            value, status = None, "missing"
        result.append((int(year), value, status, "dst_ofats4:" + ("DB03" if int(year) < 2010 else "DB07")))
    return result


def danish_affiliates(client):
    metadata_body, metadata_asset = client.post_json(
        DENMARK_METADATA, {"table": "OFATS4", "format": "JSON", "lang": "en"})
    body, asset = client.post_json(DENMARK_DATA, DENMARK_QUERY)
    metadata, data = json.loads(metadata_body), json.loads(body)
    points = parse_denmark(data, metadata)
    assets = [asset, metadata_asset]
    for section in ("indhold", "sammenlignelighed", "statistisk-behandling", "praecision-og-paalidelighed"):
        _, method_asset = client.get(DENMARK_METHODS + section)
        assets.append(method_asset)
    sid, source_id = "dst_denmark:china", "dst_ofats4"
    notes = (
        "Subsidiaries in China controlled from Denmark-resident parents, including parents under foreign "
        "ultimate control; not the narrower Danish-UCI Eurostat population. Autonomous subsidiary entities, "
        "not unique parents, plants or branches. Private non-agricultural activities, including finance; "
        "no lower activity threshold. Retain the 2010 DB03/DB07 classification transition. "
        "The source also uses zero for suppressed cells, so no zero is inferred. No pre-2007 values invented.")
    definitions = {sid: series(
        "Denmark (Statistics Denmark)", "controlled_enterprises", "enterprises",
        "Denmark-resident parents; China; Hong Kong and Macao separately coded; foreign ultimate owners "
        "included; private non-agricultural subsidiary entities; branches excluded",
        "annual", "reference year", notes)}
    sources = {source_id: source(
        "Statistics Denmark", "OFATS4; source updated " + data["dataset"].get("updated", "not supplied"),
        assets, "OFATS4; LAND=CN; ENHED=DSANTAL (Affiliates); Tid=all published years; JSONSTAT",
        notes)}
    return definitions, [
        row(sid, year, value, source_id, status, notes, comparability)
        for year, value, status, comparability in points], sources


def finnish_eurostat_affiliates(client):
    sid, observations, sources = "eurostat_finland:china", [], {}
    notes = ("Finnish-UCI controlled enterprises in China except Hong Kong, not unique parents or "
             "the broader all-industry national StatFin population. Keep the 2010 NACE transition, "
             "2013 YTY system break and 2021 EBS boundary. Overlapping national counts differ; "
             "no splice, cross-regime growth rate or inferred 2000-2006/2024/2025 values.")
    definitions = {sid: series(
        "Finland (Eurostat)", "controlled_enterprises", "enterprises",
        "Finnish-UCI controlled enterprises in China except Hong Kong; legacy NACE C-O except L, then market producers NACE B-S except O",
        "annual", "reference year", notes)}
    _, methodology = client.get(FINLAND_EUROSTAT_METHODS)
    for dataset, indicator, activity, first, last in [
        ("fats_out2", "indic_bp=ENT", "nace_r1=C-O_X_L", 2007, 2009),
        ("fats_out2_r2", "indic_bp=ENT", "nace_r2=B-S_X_O", 2010, 2020),
        ("fats_out_part", "indic_sbs=ENT_NR", "nace_r2=B-S_X_O", 2021, 9999),
    ]:
        url = EUROSTAT + dataset + "?lang=EN&freq=A&geo=FI&partner=CN_X_HK&" + indicator + "&" + activity
        body, asset = client.get(url)
        data, source_id = json.loads(body), dataset + "_FI"
        sources[source_id] = source(
            "Eurostat / Statistics Finland", dataset + "; Finland; source updated " + data["updated"],
            [asset, methodology], "geo=FI; partner=CN_X_HK; " + indicator + "; " + activity,
            notes + " Country-specific metadata retains 2007-2012, 2013-2020 and 2021-2023 comparison blocks. "
            "Reference and accounting years may differ. Original observation flags remain explicit.")
        for year, value, status, flag, comparability in parse_eurostat(data, dataset, "FI"):
            if first <= year <= last:
                if dataset == "fats_out2_r2" and year >= 2013:
                    comparability += ":YTY2013"
                observations.append(row(sid, year, value, source_id, status,
                                        notes + " Source flags: " + (flag or "none"), comparability))
    return definitions, observations, sources


def parse_istat_italy(content, year):
    book = xlrd.open_workbook(file_contents=content)
    counts = {}
    for name in ("Totale", "Industria", "Servizi"):
        sheets = [sheet for sheet in book.sheets() if sheet.name.strip().casefold() == name.casefold()]
        if len(sheets) != 1:
            raise ValueError("Italian national total/industry/services sheets changed")
        sheet = sheets[0]
        heading = " ".join(str(cell) for i in range(min(5, sheet.nrows)) for cell in sheet.row_values(i))
        if (not re.search(rf"\bAnno\s+{year}\b", heading)
                or not all(token in heading for token in ("controllo nazionale", "residenti"))):
            raise ValueError("Italian national reference year or control population changed")
        headers = [i for i in range(sheet.nrows)
                   if [str(v).strip() for v in sheet.row_values(i)[:2]] == ["PAESI", "Imprese"]]
        if not headers:
            raise ValueError("Italian national country/count column changed")
        start = headers[0]
        ends = [i for i in range(start + 1, sheet.nrows)
                if str(sheet.cell_value(i, 0)).strip().startswith("Totale primi")]
        if not ends:
            raise ValueError("Italian absolute-count table boundary changed")
        rows = [i for i in range(start + 1, ends[0]) if str(sheet.cell_value(i, 0)).strip() == "Cina"]
        if len(rows) != 1:
            raise ValueError("Italian national China count is missing or ambiguous")
        value, status = count_cell(sheet.cell_value(rows[0], 1))
        if status != "observed":
            raise ValueError("Italian national China count is unavailable")
        counts[name] = value
    if counts["Totale"] != counts["Industria"] + counts["Servizi"]:
        raise ValueError("Italian national industry/services totals do not reconcile")
    return counts["Totale"]


def italian_affiliates(client):
    legacy, current = "eurostat_italy_legal:china", "eurostat_italy:china"
    scope = "Italy-resident UCI; China except Hong Kong; market producers in the selected NACE sectors"
    notes = ("Eurostat changes from legal units to grouped enterprises in 2021; ISTAT national reports "
             "continue counting legal units. The national 2019/2020 controls match both older Eurostat values. "
             "Keep the 2010 NACE and 2021 source boundaries, original estimated flags and separate enterprise "
             "population. No pre-2007 count or cross-source 2019 change is inferred.")
    definitions = {
        legacy: series("Italy, legal units (Eurostat / ISTAT)", "controlled_legal_units", "legal units",
                       scope + "; NACE C-O except L before 2010, B-S except O thereafter; legal units",
                       "annual", "reference year", notes + " National host label: Cina. Legal-unit history, 2007-2023."),
        current: series("Italy (Eurostat enterprises)", "controlled_enterprises", "enterprises",
                        scope + "; NACE B-S except O; enterprise statistical unit from 2021",
                        "annual", "reference year", notes + " Separate 2021-onward enterprise counts remain download-only."),
    }
    body, methodology = client.get(ITALY_EUROSTAT_METHODS)
    text = BeautifulSoup(body, "html.parser").get_text(" ", strip=True)
    if not all(token in text for token in (
            "ISTAT-Italian National Statistical Institute",
            "In previous years the legal unit was used", "2021 onwards")):
        raise ValueError("Italian legal-unit/enterprise methodology changed")
    observations, sources = [], {}
    for dataset, indicator, activity, first, last in [
        ("fats_out2", "indic_bp=ENT", "nace_r1=C-O_X_L", 2007, 2009),
        ("fats_out2_r2", "indic_bp=ENT", "nace_r2=B-S_X_O", 2010, 2020),
        ("fats_out_part", "indic_sbs=ENT_NR", "nace_r2=B-S_X_O", 2021, 9999),
    ]:
        url = EUROSTAT + dataset + "?lang=EN&freq=A&geo=IT&partner=CN_X_HK&" + indicator + "&" + activity
        body, asset = client.get(url)
        data, source_id = json.loads(body), dataset + "_IT"
        sources[source_id] = source(
            "Eurostat / ISTAT", dataset + "; Italy; source updated " + data["updated"],
            [asset, methodology], "geo=IT; partner=CN_X_HK; " + indicator + "; " + activity,
            notes + " Original Italy-specific metadata explicitly rejects comparison across 2021.")
        for year, value, status, flag, comparability in parse_eurostat(data, dataset, "IT"):
            if first <= year <= last:
                sid = legacy if year < 2021 else current
                observations.append(row(sid, year, value, source_id, status,
                                        notes + " Source flags: " + (flag or "none"), comparability))
    body, national_methods = client.get(ISTAT_ITALY_REPORT)
    text = " ".join(" ".join(pdf_text(body)).split())
    if (not all(token in text for token in (
            "Nel caso del presente Report \u00e8 sempre l\u2019unit\u00e0 giuridica",
            "sottoposte a controllo ultimo nazionale"))
            or not re.search(r"Outward.{0,120}primi dati.{0,120}2021", text)
            or not re.search(r"settori da B a N e P\s*-\s*Q\s*-\s*R\s*-\s*S", text)):
        raise ValueError("Italian national legal-unit/control/sector methodology changed")
    national, evidence = {}, {}
    for edition, year in ((2022, 2019), (2023, 2020), (2024, 2021), (2025, 2022), (2026, 2023)):
        page_url = f"https://annuarioistatice.istat.it/epo/{edition}/attiv_intern/attivita.html"
        body, catalogue = client.get(page_url)
        soup = BeautifulSoup(body, "html.parser")
        tables = [tr for tr in soup.select("tr") if all(
            token in " ".join(tr.get_text(" ", strip=True).split())
            for token in ("controllo nazionale residenti", "principali paesi", "macro-settore"))]
        if (len(tables) != 1 or f"Anno {year}" not in tables[0].get_text(" ", strip=True)
                or tables[0].find("a", href=True) is None):
            raise ValueError("Italian yearbook reference year or outward table changed")
        url = urljoin(page_url, tables[0].find("a", href=True)["href"])
        parsed = urlsplit(url)
        if parsed.scheme != "https" or parsed.netloc != "annuarioistatice.istat.it" or not parsed.path.endswith(".xls"):
            raise ValueError("Italian yearbook workbook identity changed")
        body, workbook = client.get(url)
        national[year] = parse_istat_italy(body, year)
        evidence[year] = [catalogue, workbook, national_methods]
    older = {r["year"]: r["value"] for r in observations if r["series_id"] == legacy}
    if any(national[year] != older.get(year) for year in (2019, 2020)):
        raise ValueError("Italian national/Eurostat legal-unit overlaps no longer reconcile")
    for year in (2021, 2022, 2023):
        source_id = f"istat_italy_legal_{year}"
        sources[source_id] = source(
            "ISTAT / ICE", f"ISTAT-ICE {year + 3} yearbook; reference {year}", evidence[year],
            f"Outward control; Totale; PAESI=Cina; Imprese; reference {year}; industry+services checked",
            notes + " National report pages10/12 document legal units, Italian ultimate control and B-N/P-S. "
            "Published census-survey estimates; no Eurostat observation flag is assigned.")
        observations.append(row(
            legacy, year, national[year], source_id, "estimated",
            notes + " ISTAT national legal-unit estimate, not the Eurostat enterprise count.",
            "NACE2:ISTAT_LEGAL"))
    return definitions, observations, sources


def swedish_eurostat_affiliates(client):
    sid, observations, sources = "eurostat_sweden:china", [], {}
    notes = ("Swedish-UCI controlled enterprises in China except Hong Kong; a separate population "
             "from the national subsidiary table. Preserve the metadata's 2008 classification break "
             "and original legacy NACE label/2010 table transition, 2018 legal-unit-to-enterprise/EGR "
             "change and 2021 EBS boundary. Official compilation can include imputation. "
             "2009 is confidential, not zero; no national-series splice or cross-regime growth.")
    definitions = {sid: series(
        "Sweden (Eurostat)", "controlled_enterprises", "enterprises",
        "Swedish-UCI enterprises in China except Hong Kong; selected legacy NACE C-O except L, then market producers NACE B-S except O",
        "annual", "reference year", notes)}
    _, methodology = client.get(SWEDEN_EUROSTAT_METHODS)
    for dataset, indicator, activity, first, last in [
        ("fats_out2", "indic_bp=ENT", "nace_r1=C-O_X_L", 2007, 2009),
        ("fats_out2_r2", "indic_bp=ENT", "nace_r2=B-S_X_O", 2010, 2020),
        ("fats_out_part", "indic_sbs=ENT_NR", "nace_r2=B-S_X_O", 2021, 9999),
    ]:
        url = EUROSTAT + dataset + "?lang=EN&freq=A&geo=SE&partner=CN_X_HK&" + indicator + "&" + activity
        body, asset = client.get(url)
        data, source_id = json.loads(body), dataset + "_SE"
        sources[source_id] = source(
            "Eurostat / Statistics Sweden", dataset + "; Sweden; source updated " + data["updated"],
            [asset, methodology], "geo=SE; partner=CN_X_HK; " + indicator + "; " + activity,
            notes + " GGH is the source's UCI proxy. The original dataset dimensions and flags are retained.")
        for year, value, status, flag, comparability in parse_eurostat(data, dataset, "SE"):
            if first <= year <= last:
                if dataset == "fats_out2":
                    comparability += ":classification2007" if year == 2007 else ":classification2008"
                observations.append(row(sid, year, value, source_id, status,
                                        notes + " Source flags: " + (flag or "none"), comparability))
    return definitions, observations, sources


def austrian_affiliates(client):
    legacy, current = "eurostat_austria_legal:china", "eurostat_austria:china"
    scope = "Austria-resident UCI; China except Hong Kong; market producers in the selected NACE sectors"
    notes = (
        "Austria changes the statistical unit in 2021: legal units through 2020, enterprises from 2021. "
        "Foreign legal-unit data remain an imperfect enterprise proxy; the official metadata retains "
        "limited comparability. Preserve the 2010 NACE boundary. Older legal units are not joined "
        "to the short current enterprise series. Broader national resident-parent counts and the "
        "separate national 2024 control snapshot are not substituted; no pre-2007 or 2019-baseline inference.")
    definitions = {
        legacy: series(
            "Austria, historical legal units", "controlled_legal_units", "legal units",
            scope + "; NACE C-O except L before 2010, B-S except O thereafter; legal units",
            "annual", "reference year", notes + " Historical legal-unit observations are download-only."),
        current: series(
            "Austria (Eurostat enterprises)", "controlled_enterprises", "enterprises",
            scope + "; NACE B-S except O; enterprise statistical unit from 2021",
            "annual", "reference year", notes + " The current supplementary chart starts in 2021."),
    }
    body, methodology = client.get(AUSTRIA_EUROSTAT_METHODS)
    text = " ".join(BeautifulSoup(body, "html.parser").get_text(" ", strip=True).split())
    if not all(token in text for token in (
            "Statistics Austria", "from 2020 to 2021 data", "legal unit", "enterprise")):
        raise ValueError("Austrian legal-unit/enterprise methodology changed")
    observations, sources = [], {}
    for dataset, indicator, activity, first, last in [
        ("fats_out2", "indic_bp=ENT", "nace_r1=C-O_X_L", 2007, 2009),
        ("fats_out2_r2", "indic_bp=ENT", "nace_r2=B-S_X_O", 2010, 2020),
        ("fats_out_part", "indic_sbs=ENT_NR", "nace_r2=B-S_X_O", 2021, None),
    ]:
        url = EUROSTAT + dataset + "?lang=EN&freq=A&geo=AT&partner=CN_X_HK&" + indicator + "&" + activity
        body, asset = client.get(url)
        data, source_id = json.loads(body), dataset + "_AT"
        points = parse_eurostat(data, dataset, "AT")
        years = [point[0] for point in points]
        end = last if last is not None else max(years, default=0)
        if years != list(range(first, end + 1)) or (last is None and end < 2023):
            raise ValueError("Austrian outward-affiliate year coverage changed")
        sources[source_id] = source(
            "Eurostat / Statistics Austria", dataset + "; Austria; source updated " + data["updated"],
            [asset, methodology], "geo=AT; partner=CN_X_HK; " + indicator + "; " + activity,
            notes + " Original country-specific metadata and source flags are retained.")
        for year, value, status, flag, comparability in points:
            sid = legacy if year < 2021 else current
            observations.append(row(sid, year, value, source_id, status,
                                    notes + " Source flags: " + (flag or "none"), comparability))
    return definitions, observations, sources


def polish_affiliates(client):
    sid, observations, sources = "eurostat_poland:china", [], {}
    notes = ("Polish-resident ultimate control of enterprises/branches in China except Hong Kong, "
             "not the broader national survey of all foreign equity interests. Enterprise-unit "
             "implementation breaks the series in 2018. Preserve the selected legacy NACE2010 "
             "transition; 2021 is source-dataset lineage, not a second Polish unit break. "
             "Joint 50/50 and documented de-facto minority control can qualify; no size cutoff. "
             "2007 is unreported, not zero. No pre2008 or unreported current-year counts inferred.")
    definitions = {sid: series(
        "Poland", "controlled_enterprises", "enterprises",
        "Polish-resident UCI; enterprises/branches in China except Hong Kong; selected legacy NACE C-O except L, then NACE B-S except O",
        "annual", "reference year", notes)}
    _, methodology = client.get(POLAND_EUROSTAT_METHODS)
    for dataset, indicator, activity, first, last in [
        ("fats_out2", "indic_bp=ENT", "nace_r1=C-O_X_L", 2007, 2009),
        ("fats_out2_r2", "indic_bp=ENT", "nace_r2=B-S_X_O", 2010, 2020),
        ("fats_out_part", "indic_sbs=ENT_NR", "nace_r2=B-S_X_O", 2021, None),
    ]:
        url = EUROSTAT + dataset + "?lang=EN&freq=A&geo=PL&partner=CN_X_HK&" + indicator + "&" + activity
        body, asset = client.get(url)
        data, source_id = json.loads(body), dataset + "_PL"
        points = parse_eurostat(data, dataset, "PL")
        years = [point[0] for point in points]
        if (not years or years != list(range(first, (last or max(years)) + 1))
                or (last is None and max(years) < 2023)):
            raise ValueError("Polish affiliate history is incomplete; preserve the reviewed snapshot")
        sources[source_id] = source(
            "Eurostat / Statistics Poland", dataset + "; Poland; source updated " + data["updated"],
            [asset, methodology], "geo=PL; partner=CN_X_HK; " + indicator + "; " + activity, notes)
        for year, value, status, flag, _ in points:
            if "b" in flag.partition("|")[0].lower() and year not in {2010, 2018}:
                raise ValueError("Polish methodology break changed; preserve the reviewed snapshot")
            comparability = ("poland:legacy-nace-r1" if year < 2010 else
                             "poland:nace-r2:legal-units" if year < 2018 else
                             "poland:nace-r2:enterprise2018")
            observations.append(row(sid, year, value, source_id, status,
                                    notes + " Source flags: " + (flag or "none"), comparability))
    return definitions, observations, sources


def parse_snb_counts(metadata, data, native):
    if (metadata.get("cubeId") != SNB_CUBE or len(metadata.get("dimensions", [])) != 1
            or metadata["dimensions"][0].get("id") != "d0"):
        raise ValueError("SNB count cube or country dimension changed")
    labels = {}

    def visit(items, path):
        for item in items:
            label = path + [item["name"]]
            if item.get("dimensionItems"):
                visit(item["dimensionItems"], label)
            else:
                if item["id"] in labels:
                    raise ValueError("Duplicate SNB country identity")
                labels[item["id"]] = " - ".join(label)

    visit(metadata["dimensions"][0]["dimensionItems"], [])
    if (len(labels) != 92 or labels.get("CN") != "Alle Länder - Asien - China"
            or labels.get("HK") != "Alle Länder - Asien - Hongkong"):
        raise ValueError("SNB reviewed China/Hong Kong country scope changed")
    rows = list(csv.reader(io.StringIO(native.decode("utf-8-sig")), delimiter=";"))
    if rows[:4] != [["CubeId", SNB_CUBE], ["PublishingDate", "2025-12-12 09:00"],
                    [], ["Date", "d0", "Value"]]:
        raise ValueError("SNB native cube, measure, unit or reviewed publication changed")
    expected_years = {str(year) for year in range(2014, 2025)}
    native_values, json_values = {}, {}

    def number(value):
        if value is None or value == "":
            return None
        if isinstance(value, bool):
            raise ValueError("Invalid SNB count")
        try:
            result = Decimal(str(value))
        except InvalidOperation as exc:
            raise ValueError("Invalid SNB count") from exc
        if not result.is_finite() or result < 0:
            raise ValueError("Invalid SNB count")
        return result

    for item in rows[4:]:
        if len(item) != 3 or item[0] not in expected_years or item[1] not in labels:
            raise ValueError("SNB native country or measurement period changed")
        key = item[1], item[0]
        if key in native_values:
            raise ValueError("Duplicate SNB native count")
        native_values[key] = number(item[2])
    if set(data) != {"timeseries"} or len(data["timeseries"]) != len(labels):
        raise ValueError("SNB complete count query changed")
    for item in data["timeseries"]:
        meta = item["metadata"]
        match = re.fullmatch(r"EPB@SNB\." + SNB_CUBE + r"\{([^{}]+)\}", meta["key"])
        if not match or meta["frequency"] != "P1Y" or meta["scale"] != "":
            raise ValueError("SNB cube, count unit, scale or annual frequency changed")
        code = match[1]
        if code not in labels or item["header"] != [
                {"dim": "Länder/Ländergruppe", "dimItem": labels[code]}]:
            raise ValueError("SNB country labels and identities disagree")
        for point in item["values"]:
            if set(point) != {"date", "value"} or point["date"] not in expected_years:
                raise ValueError("SNB measurement years or observation flags changed")
            key = code, point["date"]
            if key in json_values:
                raise ValueError("Duplicate SNB JSON count")
            json_values[key] = number(point["value"])
    expected = {(code, year) for code in labels for year in expected_years}
    if set(native_values) != expected or json_values != native_values:
        raise ValueError("SNB complete native/JSON values must reconcile without rounding or zero filling")
    china = []
    for year in sorted(expected_years):
        value = json_values["CN", year]
        if value is None or value != value.to_integral_value():
            raise ValueError("SNB China counts are missing or require a precision review")
        china.append((int(year), int(value)))
    return china


def snb_affiliates(client):
    bodies, assets = {}, []
    document_markers = {
        "methods": "3.3. Anzahl Unternehmen", "control": "4.3. Beherrschung",
        "countries": "Länder-, Branchen- und Sektorendefinitionen"}
    for name, url in [
        ("dimensions", SNB_API + "dimensions/de"), ("data", SNB_API + "data/json/de"),
        ("native", SNB_API + "data/csv/de"), ("release", SNB_API + "lastUpdate"),
        ("methods", SNB_METHODS), ("control", SNB_CONTROL_METHODS), ("countries", SNB_COUNTRY_METHODS),
    ]:
        if name in document_markers and hasattr(client, "get_public_document"):
            bodies[name], asset = client.get_public_document(url, document_markers[name])
        else:
            bodies[name], asset = client.get(url)
        if (asset.get("request_url") != url or asset.get("url") != url
                or asset.get("sha256") != hashlib.sha256(bodies[name]).hexdigest()):
            raise ValueError("SNB source asset identity or checksum changed")
        assets.append(asset)
    method = " ".join(BeautifulSoup(bodies["methods"], "html.parser").get_text(" ", strip=True).split())
    control = " ".join(BeautifulSoup(bodies["control"], "html.parser").get_text(" ", strip=True).split())
    countries = " ".join(BeautifulSoup(bodies["countries"], "html.parser").get_text(" ", strip=True).split())
    if not all(token in method for token in (
            "Die AMNE/FATS erfassen nur Mehrheitsbeteiligungen",
            "Beteiligungsverhältnisse grösser als 50%", "nur Daten schweizerisch beherrschter Unternehmen",
            "in der Schweiz und dem Fürstentum Liechtenstein", "gruppenweit 10 Mio. Franken übersteigt",
            "Die Anzahl Unternehmen umfasst alle direkten und indirekten Mehrheitsbeteiligungen",
            "inkl. rechtlich unselbständiger Filialen und Zweigniederlassungen")):
        raise ValueError("SNB reporting frame or branch/control definition changed; reviewed rendered methods required")
    if ("Unternehmen gelten dann als schweizerisch beherrscht" not in control
            or "letztlich Berechtigte im Inland" not in control
            or "Liechtenstein, das statistisch zum Inland gehört" not in countries
            or "Geonomenklatur von Eurostat" not in countries
            or any(not re.search(r"\b" + name + r"\s+" + code + r"\b", countries)
                   for name, code in (("China", "CN"), ("Hongkong", "HK"), ("Macau", "MO")))):
        raise ValueError("SNB combined reporting area or native country definitions changed")
    if json.loads(bodies["release"]) != {
            "editionDate": "20251212_0900", "publicSinceDate": "20251212_0900"}:
        raise ValueError("SNB source release changed; review the new vintage before replacing frozen observations")
    points = parse_snb_counts(json.loads(bodies["dimensions"]),
                              json.loads(bodies["data"], parse_float=Decimal), bodies["native"])
    notes = (
        "Switzerland and Liechtenstein reporting area, not a pure-Switzerland slice or all firms. "
        "SNB-defined Swiss-controlled groups; direct/indirect majority subsidiaries include unincorporated "
        "branches, not unique parents, employees or FDI. Current reporting-group cross-border capital "
        "exceeds CHF10m; not a per-affiliate threshold or verified historically constant. Native CN China, "
        "HK Hong Kong and MO Macao are separately defined. Year-end 2014-2024 counts; no pre-2014, "
        "2025 measurement or default 2019 growth inferred. Native revision/estimate flags are not exposed; "
        "reported values do not imply final/census status. Full native query preserves other-country "
        "fractional values without rounding; only the eleven integer CN counts enter this population.")
    definitions = {SNB_SERIES: series(
        "Switzerland/Liechtenstein reporting area", "controlled_affiliate_units", "affiliates and branches",
        "Switzerland and Liechtenstein reporting area; SNB-defined Swiss-controlled groups; "
        "direct/indirect majority subsidiaries including unincorporated branches; "
        "groupwide cross-border capital CHF10m frame; SNB CN China, HK/MO separately defined",
        "annual", "year-end", notes)}
    source_id = "snb_opanmuausuland_20251212"
    sources = {source_id: source(
        "Swiss National Bank", "Published 2025-12-12 09:00; reference years 2014-2024",
        assets, "opanmuausuland; native CN; all 92 series / 1012 full JSON/CSV cells reconciled; "
        "reviewed rendered AMNE/control/country methods", notes)}
    observations = [
        row(SNB_SERIES, year, value, source_id, "reported", notes,
            "snb:opanmuausuland:CN:source-levels-only")
        for year, value in points]
    return definitions, observations, sources


def country_counts(sheets, title):
    found = {}
    for sheet in sheets:
        if sheet.name.lower().startswith(("contents", "footnotes")):
            continue
        heading = " ".join(str(v) for i in range(min(sheet.nrows, 15)) for v in sheet.row_values(i))
        if title not in heading:
            continue
        parent_header = any(" ".join(str(sheet.cell_value(i, 1)).split()).lower() == "all industries"
                            for i in range(min(15, sheet.nrows)))
        affiliate_header = re.search(r"number\s+of\s+affiliates", heading, re.I)
        if not (parent_header if title == "Number of U.S. Parents" else affiliate_header):
            raise ValueError("Business count column definition changed")
        for i in range(sheet.nrows):
            label = str(sheet.cell_value(i, 0)).strip()
            if label in {"China", "Europe"}:
                if label in found:
                    raise ValueError("Ambiguous business country row")
                found[label] = count_cell(sheet.cell_value(i, 1))
    if set(found) != {"China", "Europe"}:
        raise ValueError("Missing China/Europe business-count rows")
    return found


def parse_norway(data):
    if (data.get("source") != "Statistics Norway" or data.get("class") != "dataset"
            or data.get("extension", {}).get("px", {}).get("tableid") != "09827"
            or data.get("id") != ["Land", "ContentsCode", "Tid"]
            or data.get("size", [])[:2] != [1, 1]):
        raise ValueError("Norwegian subsidiary dataset identity changed")
    dimensions = data["dimension"]
    if (dimensions["Land"]["category"]["index"] != {"CN": 0}
            or dimensions["ContentsCode"]["category"]["index"] != {"Foretak": 0}
            or dimensions["ContentsCode"]["category"]["unit"]["Foretak"] != {
                "base": "enterprises", "decimals": 0}):
        raise ValueError("Norwegian country, counting unit or measure changed")
    years = dimensions["Tid"]["category"]["index"]
    if (not years or data["size"] != [1, 1, len(years)]
            or not isinstance(data["value"], list) or len(data["value"]) != len(years)
            or set(years.values()) != set(range(len(years))) or data.get("status")):
        raise ValueError("Norwegian periods or source flags require review")
    return [(int(year), *count_cell(data["value"][position])) for year, position in sorted(years.items())]


def norwegian_affiliates(client):
    body, asset = client.post_json(SSB_COUNTRIES, SSB_CHINA_QUERY)
    data = json.loads(body)
    points = parse_norway(data)
    definitions = {"ssb_norway:china": series(
        "Norway", "controlled_enterprises", "enterprises",
        "Subsidiary enterprises in China, separately from Hong Kong and Macao, under direct or indirect Norwegian control; all industries",
        "annual", "reference year",
        "Statistics Norway's own complete country history. All-industry coverage is broader than the "
        "old Eurostat B-S except O series; those values are preserved separately, not spliced. "
        "Counts are subsidiary enterprises, not unique Norwegian parents or physical offices.")}
    sources = {"ssb_09827": source(
        "Statistics Norway", "Table 09827; updated " + data["updated"], [asset],
        "Public API POST: Land=CN; ContentsCode=Foretak; Tid=all; all-industry country table",
        " ".join(data.get("note", [])) + " Hong Kong and Macao are separately coded countries. "
        "The original Eurostat population is not silently extended with this broader population.")}
    observations = [row("ssb_norway:china", year, value, "ssb_09827", status,
                        "Direct Statistics Norway all-industry country count, not a subset allocated from Asia.",
                        "ssb:09827:CN:all_industries") for year, value, status in points]
    return definitions, observations, sources


def parse_germany_history(content):
    if not content.startswith(b"%PDF"):
        raise ValueError("Archived German count report is not the reviewed PDF")
    pages = pdf_text(content)
    text = " ".join(" ".join(pages).split())
    if (not all(token in text for token in (
            "November 2020", "ultimately controlled by German investors", "more than 50%",
            "exceed \u20ac3 million", "branches", "permanent establishments",
            "latest figures must in all cases be regarded as provisional"))
            or not re.search(r"NACE\s+Rev\s*\.?\s*2", text)
            or not re.search(r"do not in[- ]*clude Section A:", text)
            or "Section O:" not in text):
        raise ValueError("Archived German publication, control or reporting scope changed")
    tables = [page for page in pages if "1.1. Number of the affiliates abroad" in page
              and "China, People's Republic" in page]
    if (len(tables) != 1
            or not re.search(r"Countries of investment\s+2015\s+2016\s+2017\s+2018", tables[0])
            or not re.search(r"1\s+Excluding Hong Kong\.", tables[0])):
        raise ValueError("Archived German reference years or host-country definition changed")
    counts = re.findall(
        r"China, People's Republic of 1\)\s+"
        r"(\d{1,3} \d{3})\s+(\d{1,3} \d{3})\s+(\d{1,3} \d{3})\s+(\d{1,3} \d{3})\s+Hong Kong",
        tables[0])
    if len(counts) != 1:
        raise ValueError("Archived German China count row missing or ambiguous")
    return [(year, int(value.replace(" ", "")), "provisional" if year == 2018 else "observed")
            for year, value in zip(range(2015, 2019), counts[0])]


def parse_germany(content):
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    method = " ".join(str(v) for r in book["Methodology"].values for v in r if v is not None)
    if not all(s in method for s in ("ultimately controlled by German investors", "more than 50%", "exceed €3 million")):
        raise ValueError("German affiliate control or reporting threshold changed")
    activities = list(book["Economic activity"].values)
    if not any(r[0] == "BTSXO" and "without agriculture" in str(r[1]).lower() for r in activities):
        raise ValueError("German affiliate economic activity definition changed")
    records = list(book["DATA"].values)
    try:
        dates = [from_excel(v, book.epoch) for v in records[2][1:]]
    except (TypeError, ValueError, OverflowError) as exc:
        raise ValueError("German affiliate annual columns changed") from exc
    years = [d.year for d in dates]
    if not years or years != sorted(set(years)) or any((d.month, d.day) != (12, 31) for d in dates):
        raise ValueError("German affiliate annual columns changed")
    matches = [r for r in records if r[0] == "BBFOFATS.A.DE.ENT.BTSXO.CN"]
    if len(matches) != 1 or len(matches[0]) != len(years) + 1:
        raise ValueError("German China enterprise-count row missing or ambiguous")
    points = []
    for year, value in zip(years, matches[0][1:]):
        if value == "X":
            count, status = None, "missing"
        elif value == "/":
            count, status = None, "missing"
        else:
            count, status = count_cell(0 if value == "-" else value)
            if status == "observed" and year == max(years):
                status = "provisional"
        points.append((year, count, status))
    return points


def german_affiliates(client):
    body, landing = client.get(BUNDESBANK_FATS)
    links = [urljoin(BUNDESBANK_FATS, a["href"]) for a in BeautifulSoup(body, "html.parser").select("a[href]")
             if a["href"].endswith("/statistik-outward-fats-data.xlsx")]
    if len(set(links)) != 1:
        raise ValueError("Bundesbank outward-affiliate workbook link missing or ambiguous")
    body, asset = client.get(links[0])
    points = parse_germany(body)
    definitions = {"bundesbank_germany:china": series(
        "Germany", "controlled_enterprises", "enterprises",
        "Foreign affiliates in China, excluding separately coded Hong Kong and Macao, ultimately controlled "
        "by German investors; NACE Rev. 2 BTSXO",
        "annual", "reference year",
        "Direct or indirect control above 50%; affiliates with total assets above EUR 3 million, including "
        "branches and permanent establishments. Excludes agriculture, public administration and private households. "
        "Not distinct German parent companies, all European businesses or a reconstruction of confidential EU totals.")}
    sources = {"bundesbank_outward_fats": source(
        "Deutsche Bundesbank", f"Outward FATS workbook; reference years {points[0][0]}-{points[-1][0]}",
        [landing, asset], "DATA; BBFOFATS.A.DE.ENT.BTSXO.CN; Methodology and Economic activity",
        "Only the China enterprise-count series is used; employment and turnover columns are not company counts. "
        "Latest year provisional. Germany is an overlapping EU member, not an extra component to add to EU27.")}
    observations = [row("bundesbank_germany:china", year, value, "bundesbank_outward_fats", status,
                        "Above EUR 3 million in assets; German ultimate control. Latest year provisional.",
                        "bundesbank:BTSXO:CN:EUR3m") for year, value, status in points]
    body, historical = client.get(GERMANY_FATS_2020)
    sources["bundesbank_outward_fats_2020"] = source(
        "Deutsche Bundesbank", "Outward FATS, November 2020; reference years 2015-2018",
        [historical], "Page 7; II.1.1; China, People's Republic; Number of the affiliates abroad",
        "Archived original publication, not the archive-capture year. German ultimate control above 50%, "
        "assets above EUR 3m, NACE Rev.2; branches and permanent establishments included. "
        "Original latest-year note makes 2018 provisional. Later historical revisions and 2019/2020 "
        "counts are not inferred. Not spliced with the different Eurostat population.")
    observations.extend(row(
        "bundesbank_germany:china", year, value, "bundesbank_outward_fats_2020", status,
        "November 2020 publication vintage; 2018 provisional under the original latest-year note. "
        "German ultimate control; assets above EUR 3m. No 2019/2020 value is inferred.",
        "bundesbank:BTSXO:CN:EUR3m") for year, value, status in parse_germany_history(body))
    definitions["eurostat_germany:china"] = series(
        "Germany (Eurostat)", "controlled_enterprises", "enterprises",
        "German-controlled enterprises in China except Hong Kong; Eurostat NACE B-S except O",
        "annual", "reference year",
        "Eurostat's reported history is retained separately from Bundesbank's current workbook. "
        "Overlapping 2021-2023 counts differ; they are not spliced, averaged or treated as a matched revision. "
        "The 2021 EBS methodology boundary is disconnected. Germany overlaps the EU27 aggregate.")
    definitions["bundesbank_germany:china"]["label"] = "Germany (Bundesbank)"
    for dataset, indicator in [("fats_out2_r2", "indic_bp=ENT"), ("fats_out_part", "indic_sbs=ENT_NR")]:
        url = (EUROSTAT + dataset + "?lang=EN&freq=A&nace_r2=B-S_X_O&" + indicator
               + "&partner=CN_X_HK&geo=DE")
        body, asset = client.get(url)
        source_id = dataset + "_DE"
        sources[source_id] = source(
            "Eurostat", dataset + "; Germany reporting country", [asset],
            "geo=DE; partner=CN_X_HK; NACE B-S except O; enterprise counts",
            "Separate source population/vintage from the Bundesbank workbook; overlapping counts are not forced to agree.")
        observations.extend(row("eurostat_germany:china", year, value, source_id, status,
                                "Eurostat German reporting population; source flags: " + flag, comparability)
                            for year, value, status, flag, comparability in parse_eurostat(
                                json.loads(body), dataset, "DE"))
    return definitions, observations, sources


def collect_business(client):
    definitions, observations, sources = {}, [], {}
    for region, label in [("china", "United States")]:
        definitions["bea_affiliates:" + region] = series(
            label, "majority_owned_affiliates", "affiliates",
            f"Foreign affiliates in {'mainland China, excluding Hong Kong and Macau' if region == 'china' else 'BEA-defined Europe, not EU27'}; combined U.S.-parent ownership over 50%",
            "annual", "fiscal ending year",
            "Entities, not unique parent companies, offices or storefronts. U.S. parents are U.S.-resident, "
            "not necessarily ultimately U.S.-owned. Includes holding companies. Nonbank coverage before 2009; "
            "reporting cutoffs $7m through 2003, $10m in 2004-2008, $25m from 2009. "
            "Counts exclude benchmark-exempt affiliates. Lines break at benchmark regimes.")
        definitions["bea_parents:" + region] = series(
            label, "parent_companies", "U.S. parents",
            "Distinct U.S.-resident parent companies with affiliates in China; at least 10% foreign-affiliate ownership; BEA reporting exclusions apply",
            "benchmark", "fiscal ending year",
            "Benchmark observations only, not an annual census. A parent is counted once per host country/region, "
            "but can appear in both China and Europe: never add these counts. Country attribution and reporting "
            "exclusions follow each benchmark. Not restricted to ultimately American-owned companies. "
            "Do not substitute majority-owned affiliate counts for parent counts.")
    for year, url in BEA_ARCHIVES.items():
        body, asset = client.get(url)
        archive = zipfile.ZipFile(io.BytesIO(body))
        members = [n for n in archive.namelist()
                   if re.fullmatch(r"tab3a1[a-z]?(?:_new)?\.xls", PurePosixPath(n).name.lower())]
        if not members:
            raise ValueError(f"Missing historical majority-owned affiliate tables: {year}")
        sheets = [sheet for n in members for sheet in xlrd.open_workbook(file_contents=archive.read(n)).sheets()]
        counts = country_counts(sheets, "Selected Data for Foreign Affiliates")
        source_id = f"bea_affiliates_{year}"
        sources[source_id] = source(
            "U.S. Bureau of Economic Analysis", f"Revised {year} operations; historical benchmark regime",
            [asset], "Table III.A1, Number of affiliates; " + ", ".join(members),
            "Majority-owned nonbank affiliates; benchmark-exempt entities excluded. "
            "Not comparable across the 2004 reporting-regime or 2009 banking/threshold changes.")
        for region, label in [("china", "China")]:
            value, status = counts[label]
            observations.append(row("bea_affiliates:" + region, year, value, source_id, status,
                                    f"Nonbank affiliate count; {'1999' if year < 2004 else '2004'} benchmark reporting regime.",
                                    "bea_affiliates:" + ("1999" if year < 2004 else "2004")))
        if year == 2004:
            parent = [n for n in archive.namelist() if PurePosixPath(n).name.lower() == "tab1m2.xls"]
            if len(parent) != 1:
                raise ValueError("Missing 2004 parent-company country table")
            parent_sheets = xlrd.open_workbook(file_contents=archive.read(parent[0])).sheets()
            counts = country_counts(parent_sheets, "Number of U.S. Parents")
            sources["bea_parents_2004"] = source(
                "U.S. Bureau of Economic Analysis", "Revised 2004 benchmark", [asset],
                parent[0] + "; Table I.M2; All industries",
                "U.S. parents counted once in each country where they have affiliates. Reporting exclusions apply; "
                "the 2004 benchmark changed the affiliate reporting cutoff from $7m to $10m.")
            for region, label in [("china", "China")]:
                value, status = counts[label]
                observations.append(row("bea_parents:" + region, 2004, value, "bea_parents_2004", status,
                                        comparability="bea_parents:2004"))
    body, asset = client.get(BEA_CURRENT)
    book = xlrd.open_workbook(file_contents=body)
    for year in range(2009, 2024):
        sheet = book.sheet_by_name(str(year))
        header = " ".join(str(v) for i in range(5) for v in sheet.row_values(i))
        if str(year) not in header or "$25 million" not in header:
            raise ValueError("BEA annual year or affiliate reporting threshold changed")
        counts = country_counts([sheet], "Majority-Owned Foreign Affiliates")
        source_id = f"bea_affiliates_{year}"
        sources[source_id] = source(
            "U.S. Bureau of Economic Analysis", f"2009-2023 selected-data workbook; {year} fiscal year",
            [asset], f"Sheet {year}; China/Europe rows; Number of affiliates above $25m",
            "Majority-owned affiliates across all industries. Assets, sales or net income/loss must exceed $25m. "
            "Not all small businesses or physical offices; includes holding entities. "
            "2014 and 2019 benchmarks can change coverage: their boundaries are not connected.")
        for region, label in [("china", "China")]:
            value, status = counts[label]
            if year == 2023 and value is not None:
                status = "provisional"
            observations.append(row("bea_affiliates:" + region, year, value, source_id, status,
                                    "All-industry majority-owned affiliate count above $25m; "
                                    + ("preliminary source year." if year == 2023 else "revised source year."),
                                    "bea_affiliates:" + str(2019 if year >= 2019 else 2014 if year >= 2014 else 2009)))
    for year, url in PARENT_BOOKS.items():
        body, asset = client.get(url)
        book = xlrd.open_workbook(file_contents=body)
        counts = country_counts(book.sheets(), "Number of U.S. Parents")
        source_id = f"bea_parents_{year}"
        sources[source_id] = source(
            "U.S. Bureau of Economic Analysis", f"Revised {year} benchmark", [asset],
            "Table I.K2; Number of U.S. parents; country rows, All industries",
            "Unique within each host region, not across regions; includes parents with minority-owned affiliates. "
            "Benchmark reporting exclusions and coverage apply; these points are not an annual series.")
        for region, label in [("china", "China")]:
            value, status = counts[label]
            observations.append(row("bea_parents:" + region, year, value, source_id, status,
                                    comparability=f"bea_parents:{year}"))
    definitions["eurostat_affiliates:china"] = series(
        "EU27 (current members)", "controlled_enterprises", "enterprises",
        "China except Hong Kong; ultimately controlled from fixed EU27_2020 membership, excluding the UK; market producers NACE B-S except O",
        "annual", "reference year",
        "Eurostat enterprise units, not unique European parents or all Europe. Macao is a separate partner. "
        "Not matched to BEA's size cutoff, U.S.-resident parent basis or consolidation. "
        "EBS methodology changes from 2021; do not join across that boundary or treat confidential cells as zero.")
    definitions["eurostat_historical_eu:china"] = series(
        "EU27 (2007 members)", "controlled_enterprises", "enterprises",
        "China except Hong Kong; EU27_2007 control, including the UK and excluding Croatia",
        "annual", "reference year",
        "Historical EU membership, not the current EU27. NACE Rev. 1.1 through 2009 and Rev. 2 from 2010; "
        "both classification and membership boundaries remain disconnected. No splice into current-membership counts.")
    uk_definitions, uk_observations, uk_sources = uk_early_affiliates(client)
    definitions.update(uk_definitions)
    observations.extend(uk_observations)
    sources.update(uk_sources)
    definitions["eurostat_norway:china"] = series(
        "Norway", "controlled_enterprises", "enterprises",
        "Norwegian-controlled enterprises in China except Hong Kong; NACE B-S except O",
        "annual", "reference year",
        "Norway is outside the EU and is shown separately, not added to an all-Europe total. "
        "Eurostat's 2021 methodology boundary is disconnected.")
    queries = [
        ("fats_out2_r2", "EU27_2020", "eurostat_affiliates:china", None),
        ("fats_out_part", "EU27_2020", "eurostat_affiliates:china", None),
        ("fats_out2", "EU27_2007", "eurostat_historical_eu:china", 2009),
        ("fats_out2_r2", "EU27_2007", "eurostat_historical_eu:china", 2013),
        ("fats_out2_r2", "UK", "eurostat_uk:china", None),
        ("fats_out2_r2", "NO", "eurostat_norway:china", None),
        ("fats_out_part", "NO", "eurostat_norway:china", None),
    ]
    for dataset, geo, sid, end in queries:
        indicator = "indic_sbs=ENT_NR" if dataset == "fats_out_part" else "indic_bp=ENT"
        activity = "nace_r1=C-O_X_L" if dataset == "fats_out2" else "nace_r2=B-S_X_O"
        url = EUROSTAT + dataset + f"?lang=EN&geo={geo}&partner=CN_X_HK&" + indicator + "&" + activity
        body, asset = client.get(url)
        data = json.loads(body)
        points = parse_eurostat(data, dataset, geo)
        source_id = dataset if geo == "EU27_2020" else dataset + "_" + geo
        sources[source_id] = source(
            "Eurostat", data["label"] + "; updated " + data["updated"], [asset],
            f"A; {geo}; CN_X_HK; {activity}; " + indicator,
            "Keep each control geography, membership and activity classification explicit. "
            "Preserve estimated, break and confidential flags; methodology boundaries remain disconnected. "
            "No reconstruction of confidential totals from member-country values.")
        observations.extend(row(sid, year, value, source_id, status,
                                "Eurostat source flags: " + (flag or "none"), comparability)
                            for year, value, status, flag, comparability in points if end is None or year <= end)
    body, asset = client.get(ONS_DATA)
    ons = parse_ons(body)
    sources["ons_2019"] = source(
        "Office for National Statistics", "2018 revised and 2019 estimates; released May 20, 2022", [asset],
        "Table 1 (2018), Table 2 (2019); China row, TOTAL ACTIVITY / ENT",
        "Foreign enterprise counts, not employment or turnover. Total activity includes agriculture; "
        "survey-based estimates. Later ONS vintage supersedes overlapping Eurostat UK values.")
    for year, value, status in ons:
        observations = [r for r in observations if (r["series_id"], r["year"]) != ("eurostat_uk:china", year)]
        observations.append(row("eurostat_uk:china", year, value, "ons_2019", status,
                                "ONS total-activity enterprise estimate, including agriculture; 2018 revised.",
                                "ons:uk:total_activity:2018_2019"))
    _, asset = client.get(ONS_COVERAGE)
    sources["ons_coverage"] = source(
        "Office for National Statistics", "July 15, 2025 disclosure; checked with current dataset listing", [asset],
        "Response 2a: outward affiliate statistics unavailable from 2020",
        "Do not substitute new inward-affiliate statistics or FDI investment values for UK-controlled enterprises in China.")
    # The original methodology documents are retained as provenance, not redistributed.
    for name, url in [
        ("bea_current_notes", "https://www.bea.gov/international/usdia2023p-info"),
        ("bea_release", "https://www.bea.gov/news/2025/activities-us-multinational-enterprises-2023"),
        ("bea_geography", "https://apps.bea.gov/international/bp_web/geographic_area_definitions.pdf"),
        ("eurostat_methods", "https://ec.europa.eu/eurostat/cache/metadata/en/fats_out_sims.htm"),
    ]:
        _, asset = client.get(url)
        sources[name] = source("BEA" if name.startswith("bea") else "Eurostat",
                               "Reviewed source definitions", [asset], "Definitions and source coverage")
    norway_definitions, norway_observations, norway_sources = norwegian_affiliates(client)
    definitions.update(norway_definitions)
    observations.extend(norway_observations)
    sources.update(norway_sources)
    germany_definitions, germany_observations, germany_sources = german_affiliates(client)
    definitions.update(germany_definitions)
    observations.extend(germany_observations)
    sources.update(germany_sources)
    nl_definitions, nl_observations, nl_sources = netherlands_affiliates(client)
    definitions.update(nl_definitions)
    observations.extend(nl_observations)
    sources.update(nl_sources)
    nordic_definitions, nordic_observations, nordic_sources = nordic_affiliates(client)
    definitions.update(nordic_definitions)
    observations.extend(nordic_observations)
    sources.update(nordic_sources)
    france_definitions, france_observations, france_sources = french_affiliates(client)
    definitions.update(france_definitions)
    observations.extend(france_observations)
    sources.update(france_sources)
    finland_definitions, finland_observations, finland_sources = finnish_eurostat_affiliates(client)
    definitions.update(finland_definitions)
    observations.extend(finland_observations)
    sources.update(finland_sources)
    italy_definitions, italy_observations, italy_sources = italian_affiliates(client)
    definitions.update(italy_definitions)
    observations.extend(italy_observations)
    sources.update(italy_sources)
    sweden_definitions, sweden_observations, sweden_sources = swedish_eurostat_affiliates(client)
    definitions.update(sweden_definitions)
    observations.extend(sweden_observations)
    sources.update(sweden_sources)
    spain_definitions, spain_observations, spain_sources = spanish_affiliates(client)
    definitions.update(spain_definitions)
    observations.extend(spain_observations)
    sources.update(spain_sources)
    denmark_definitions, denmark_observations, denmark_sources = danish_affiliates(client)
    definitions.update(denmark_definitions)
    observations.extend(denmark_observations)
    sources.update(denmark_sources)
    austria_definitions, austria_observations, austria_sources = austrian_affiliates(client)
    definitions.update(austria_definitions)
    observations.extend(austria_observations)
    sources.update(austria_sources)
    poland_definitions, poland_observations, poland_sources = polish_affiliates(client)
    definitions.update(poland_definitions)
    observations.extend(poland_observations)
    sources.update(poland_sources)
    snb_definitions, snb_observations, snb_sources = snb_affiliates(client)
    definitions.update(snb_definitions)
    observations.extend(snb_observations)
    sources.update(snb_sources)
    coverage = []
    for sid in definitions:
        selected = [r for r in observations if r["series_id"] == sid]
        actual = [str(r["year"]) for r in selected if r["value"] is not None]
        confidential = [str(r["year"]) for r in selected if r["status"] == "suppressed"]
        coverage.append({"id": sid, "status": "partial_history" if sid == SNB_SERIES else
                         "benchmark_only" if sid.startswith("bea_parents") else "annual",
                         "description": "Verified values: " + ", ".join(actual) +
                         (". Confidential: " + ", ".join(confidential) if confidential else "") +
                         ". All series concern businesses in China; definitions and gaps are retained." +
                         (" " + definitions[sid]["notes"] if sid == SNB_SERIES else "")})
    return {
        "version": 1, "series": definitions,
        "observations": sorted(observations, key=lambda r: (r["series_id"], r["year"])),
        "sources": sources,
        "coverage": coverage,
        "context": {
            "question": "How has U.S. and European business presence in China changed over time?",
            "interpretation": "Company counts, affiliate entities and branch locations are distinct. U.S. parents and ultimately EU-controlled enterprises are not matched ownership populations; BEA and Eurostat levels must not be used to rank otherwise comparable business presence.",
            "geography": "China is the host. European coverage includes current/historical EU groups, Germany, the UK, Norway, Sweden, Finland, France, the Netherlands, Spain, Denmark, Italy, Austria and Poland. French, Finnish, Swedish, Dutch, Spanish, Italian, Austrian and Polish Eurostat host categories are China except Hong Kong. Denmark separately codes Hong Kong/Macao; its resident-parent population includes foreign ultimate owners, unlike Eurostat's Danish-UCI population. Finnish and Swedish Eurostat UCI histories remain separate from their national series. Spain's Eurostat financial/nonfinancial history is not spliced with INE's national nonfinancial counts. Italian legal-unit history reaches 2023; grouped Eurostat enterprises remain a separate population. Austrian legal units before 2021 are not spliced into current enterprise counts or broader national populations. Poland's UCI population is not the broader national foreign-equity survey. National series overlap EU aggregates and have different control/sector definitions; these are not summed into an all-Europe total. Other European origins remain unresolved." + SNB_CONTEXT_NOTE,
            "rights": "Publish aggregate statistics and original charts only; raw workbooks and reports remain local. Sources and retrieval checksums accompany every observation.",
            "release_timing": "As checked October 22, 2026, BEA's latest activity release covers 2023. The published schedule puts the 2024 activity release on November 20, 2026; no 2024 value is inferred.",
        },
    }


def parse_cbs_netherlands(body):
    soup = BeautifulSoup(body, "html.parser")
    selected = []
    for table in soup.select("table"):
        figure = table.find_parent("figure")
        title = figure.get_text(" ", strip=True) if figure else ""
        if ("5.4.1 Top 10 landen met Nederlandse dochterondernemingen" in title
                and "(dochterondernemingen)" in title):
            selected.append(table)
    if len(selected) != 1:
        raise ValueError("CBS outward affiliate count table missing or ambiguous")
    rows = [[c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"])]
            for tr in selected[0].select("tr")]
    if not rows or rows[0] != ["Land", "2024*", "2023"]:
        raise ValueError("CBS reference years or provisional marker changed")
    china = [cells for cells in rows[1:] if cells and cells[0] == "China"]
    if len(china) != 1 or len(china[0]) != 3:
        raise ValueError("CBS China row missing or ambiguous")
    if "* voorlopige cijfers" not in soup.get_text(" ", strip=True):
        raise ValueError("CBS provisional-status legend missing")
    values = []
    for cell in china[0][1:]:
        if not re.fullmatch(r"\d+(?: \d{3})*", cell):
            raise ValueError("CBS missing/confidential counts cannot become zero")
        values.append(int(cell.replace(" ", "")))
    return {2024: (values[0], "provisional"), 2023: (values[1], "observed")}


def netherlands_affiliates(client):
    sid = "eurostat_netherlands:china"
    definitions = {sid: series(
        "Netherlands", "controlled_enterprises", "enterprises",
        "Dutch-resident UCI; affiliates in China except Hong Kong; NACE Rev.1.1 C-O except L "
        "in 2007-2009 and Rev.2 B-S except O from 2010",
        "annual", "reference year",
        "Financial and nonfinancial affiliate enterprises, not unique parents or everyone trading with China. "
        "Official population estimates include nonresponse weighting. Classification changes in 2010; "
        "the 2021 CBS method is not directly comparable with earlier years. No 2019 comparison across "
        "that boundary. Netherlands overlaps EU27; do not sum them. No pre-2007 or 2025 count inferred.")}
    observations, sources = [], {}
    for dataset, query in (
            ("fats_out2", "indic_bp=ENT&nace_r1=C-O_X_L"),
            ("fats_out2_r2", "indic_bp=ENT&nace_r2=B-S_X_O"),
            ("fats_out_part", "indic_sbs=ENT_NR&nace_r2=B-S_X_O")):
        body, asset = client.get(EUROSTAT + dataset + "?lang=EN&geo=NL&partner=CN_X_HK&" + query)
        data = json.loads(body)
        source_id = dataset + "_NL"
        sources[source_id] = source(
            "Eurostat", data["label"] + "; updated " + data["updated"], [asset],
            "A; NL; CN_X_HK; " + query,
            "Dutch control and China except Hong Kong. Preserve NACE, source flags and the 2021 method boundary.")
        observations.extend(row(sid, year, value, source_id, status,
                                "Eurostat source flags: " + (flag or "none"), comparability)
                            for year, value, status, flag, comparability in parse_eurostat(data, dataset, "NL"))
    body, report_asset = client.get(CBS_NETHERLANDS_REPORT)
    counts = parse_cbs_netherlands(body)
    methods, methods_asset = client.get(CBS_NETHERLANDS_METHODS)
    text = BeautifulSoup(methods, "html.parser").get_text(" ", strip=True)
    if not all(term in text for term in (
            "UCI in Nederland", "SBI 2008 sectie B t/m S exclusief O",
            "financi\u00eble en niet financi\u00eble", "vanaf verslagperiode 2007",
            "Voor de cijfers van 2021 is een nieuwe methodiek gekozen")):
        raise ValueError("CBS Dutch control, sectors or methodology changed")
    by_year = {r["year"]: r for r in observations}
    if 2023 not in by_year or counts[2023][0] != by_year[2023]["value"]:
        raise ValueError("CBS and Eurostat Dutch overlap changed; review the source revision")
    sources["cbs_netherlands_2026"] = source(
        "Statistics Netherlands (CBS)", "Nederland Handelsland 2026; 2024 preliminary",
        [report_asset, methods_asset], "Original HTML figure 5.4.1, China row, 2024* and 2023; * means provisional.",
        "Dutch-resident UCI; financial/nonfinancial B-S except O. Same OFATS population as the verified "
        "2023 Eurostat overlap. This is a 2024 measurement, not 2026 data. The 2021 method is not "
        "one-to-one comparable with earlier years.")
    if 2024 in by_year:
        if by_year[2024]["value"] != counts[2024][0]:
            raise ValueError("CBS and Eurostat Dutch 2024 counts differ; review the source revision")
    else:
        observations.append(row(
            sid, 2024, counts[2024][0], "cbs_netherlands_2026", counts[2024][1],
            "CBS figure 5.4.1; provisional. Published 2023 control agrees with Eurostat.",
            by_year[2023]["comparability"]))
    return definitions, sorted(observations, key=lambda r: r["year"]), sources


def parse_eurostat(data, dataset, geo="EU27_2020"):
    indicator = "indic_sbs" if dataset == "fats_out_part" else "indic_bp"
    activity = {"nace_r1": "C-O_X_L"} if dataset == "fats_out2" else {"nace_r2": "B-S_X_O"}
    wanted = {"freq": "A", **activity, indicator: "ENT" if indicator == "indic_bp" else "ENT_NR",
              "partner": "CN_X_HK", "geo": geo}
    if data.get("extension", {}).get("id", "").lower() != dataset or set(data["id"]) != set(wanted) | {"time"}:
        raise ValueError("Eurostat business dataset identity changed")
    for name, code in wanted.items():
        if data["dimension"][name]["category"]["index"] != {code: 0} or data["size"][data["id"].index(name)] != 1:
            raise ValueError("Eurostat returned a different geography, population or measure")
    times = data["dimension"]["time"]["category"]["index"]
    if not times or len(times) != data["size"][data["id"].index("time")]:
        raise ValueError("Eurostat time axis does not reconcile")
    result, segment = [], 0
    for year, position in sorted(times.items()):
        flag = data.get("status", {}).get(str(position), "")
        flags = set(re.sub(r"[|\s]", "", flag))
        if flags - set("Cbdeipru"):
            raise ValueError("Unknown Eurostat business observation flag")
        value = data.get("value", {}).get(str(position))
        if "C" in flags:
            value, status = None, "suppressed"
        elif value is None:
            status = "missing"
        else:
            value, status = count_cell(value)
            status = "provisional" if "p" in flags else "estimated" if flags & {"e", "i"} else "low_reliability" if "u" in flags else status
        if "b" in flags or "d" in flags:
            segment += 1
        comparability = f"{dataset}:{geo}:{segment}"
        if "d" in flags:
            comparability += ":" + year
        result.append((int(year), value, status, flag, comparability))
    return result


def parse_ons(body):
    book = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    result = []
    for year, name in [(2018, "Table 1"), (2019, "Table 2")]:
        rows = list(book[name].iter_rows(min_col=1, max_col=4, values_only=True))
        if str(year) not in str(rows[0][0]) or rows[8][1] != "TOTAL ACTIVITY" or rows[9][1] != "ENT":
            raise ValueError("ONS year, sector total or enterprise-count column changed")
        matches = [r for r in rows if r[0] == "China"]
        if len(matches) != 1:
            raise ValueError("ONS China row missing or ambiguous")
        value, status = count_cell("(D)" if matches[0][1] == "C" else matches[0][1])
        result.append((year, value, "estimated" if value is not None else status))
    return result
