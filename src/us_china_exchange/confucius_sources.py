"""Reviewed national snapshots and public directory records, never inferred activity."""

import copy
import hashlib
import io
import json
import re
from datetime import date, datetime
from pathlib import Path
from urllib.parse import urlencode, urlsplit

from bs4 import BeautifulSoup
from openpyxl import load_workbook

from .language_sources import pdf_text
from .model import json_text


DIRECTORY = "https://ci.cn/open/site-tab/sitesByLabelIdAndSiteName?" + urlencode({
    "labelId": "C-Ame-United-States-of-America", "siteName": "", "language": "EN"})
URLS = {
    "mfa_2014": "https://np.china-embassy.gov.cn/eng/fyrth/201409/t20140929_1592935.htm",
    "nas_2017": "https://www.nas.org/wp-content/uploads/2025/10/NAS_confuciusInstitutes.pdf",
    "nas_2022": "https://www.nas.org/wp-content/uploads/2025/10/After_Confucius_Institutes_NAS.pdf",
    "crs_2022": "https://www.congress.gov/crs_external_products/IF/PDF/IF11180/IF11180.11.pdf",
    "crs_2023": "https://www.congress.gov/crs_external_products/IF/PDF/IF11180/IF11180.13.pdf",
    "senate_2019": "https://www.congress.gov/118/meeting/house/116334/documents/HHRG-118-ED14-20230919-SD002.pdf",
    "usc_2021": "https://china.usc.edu/confucius-institutes-united-states",
    "nas_classrooms_2024": "https://www.nas.org/wp-content/uploads/2025/10/China_and_Our_Children.pdf",
    "nas_roster_2023": "https://www.nas.org/wp-content/uploads/2023/06/Confucius20Institutes20in20the20US206-22-23.xlsx",
    "nas_closures_2023": "https://www.nas.org/wp-content/uploads/2023/06/Closed20Confucius20Institutes206-22-23.xlsx",
    "cief_directory": DIRECTORY,
}
SERIES = {
    "institutes_reported": {
        "category": "institutes", "measurement": "reported_stock", "unit": "institutes",
        "scope": "National reported counts; host coverage and announced closures vary by source. Disconnected snapshots, not a harmonized annual census."},
    "institutes_campus": {
        "category": "institutes", "measurement": "campus_stock", "unit": "institutes",
        "scope": "U.S. college and university campuses only; excludes other hosts."},
    "institutes_listed": {
        "category": "institutes", "measurement": "directory_listings", "unit": "listed institutes",
        "scope": "Institute-labelled entries in the operator's U.S. directory. Website publication switches are not operating status."},
    "classrooms_reported": {
        "category": "classrooms", "measurement": "reported_stock", "unit": "reported classrooms or school sites",
        "scope": "Historical national reports of K-12 Confucius Classrooms; programs and school sites are not necessarily identical units."},
    "classrooms_inventory": {
        "category": "classrooms", "measurement": "documented_inventory", "unit": "documented cases",
        "scope": "NAS inventory of schools and school systems, including current and former programs. Incomplete inventory, not operating stock."},
    "classrooms_listed": {
        "category": "classrooms", "measurement": "directory_listings", "unit": "listed classroom programs",
        "scope": "Classroom-labelled entries in the operator's U.S. directory, including a university-named entry. Not a K-12 school census or proof of operation."},
    "successors_documented": {
        "category": "successors", "measurement": "documented_replacements", "unit": "institutions",
        "scope": "U.S. institutions NAS documented as replacing a closed institute with a similar program. Cumulative documentation, not currently operating programs."},
    "successors_proposed": {
        "category": "successors", "measurement": "proposed_replacements", "unit": "institutions",
        "scope": "Institutions seeking a replacement in the 2022 NAS investigation; not completed replacements."},
    "partners_retained": {
        "category": "successors", "measurement": "retained_partnerships", "unit": "institutions",
        "scope": "Institutions maintaining a former Chinese partner relationship in the 2022 NAS investigation. May overlap replacement cases; not itself a rebrand."},
    "institutes_transferred": {
        "category": "institutes", "measurement": "documented_transfers", "unit": "institutions",
        "scope": "Institutions transferring an institute in the 2022 NAS investigation; not new programs or additional successors."},
}
ROOT_FIELDS = {"version", "series", "observations", "records", "sources", "context"}
ROW_FIELDS = {"id", "series_id", "reference", "reference_basis", "qualifier", "value",
              "bound", "source_id", "locator", "notes"}
REVIEW_URL = "https://raw.githubusercontent.com/hanshanley/china-us-relation-change/main/data/reviewed/confucius-snapshots.json"


def plain(content):
    if content.startswith(b"%PDF"):
        return " ".join(" ".join(pdf_text(content)).split())
    return " ".join(BeautifulSoup(content, "html.parser").get_text(" ", strip=True).split())


def require(text, patterns, label):
    if any(re.search(pattern, text, re.I) is None for pattern in patterns):
        raise ValueError("Reviewed Confucius source changed: " + label)


def source(name, published, url, locator, asset, notes=""):
    return {"name": name, "published_on": published, "publisher_url": url, "locator": locator,
            "assets": [asset], "evidence_kind": "publisher_response", "notes": notes}


def observation(identity, series, reference, qualifier, number, source_id, locator, notes="", basis="as_of"):
    bounded = qualifier in {"at_least", "less_than"}
    return {"id": identity, "series_id": series, "reference": reference, "reference_basis": basis,
            "qualifier": qualifier, "value": None if bounded else number,
            "bound": number if bounded else None, "source_id": source_id, "locator": locator, "notes": notes}


def reference_date(reference):
    if re.fullmatch(r"\d{4}", reference):
        return date(int(reference), 1, 1)
    if re.fullmatch(r"\d{4}-\d{2}", reference):
        return date.fromisoformat(reference + "-01")
    return date.fromisoformat(reference)


def reported(row):
    if row["qualifier"] == "unknown":
        return "Unknown"
    number = row["value"] if row["value"] is not None else row["bound"]
    return {"exact": "", "approximate": "About ", "at_least": "At least ", "less_than": "Fewer than "}[row["qualifier"]] + f"{number:,}"


def parse_directory(content):
    payload = json.loads(content)
    if payload.get("code") != 200 or not isinstance(payload.get("data"), list) or not payload["data"]:
        raise ValueError("Operator directory returned no verifiable U.S. records")
    records, seen = [], set()
    for row in payload["data"]:
        key, name = str(row["siteKey"]), row["name"].strip()
        if (row["labelId"] != "C-Ame-United-States-of-America" or row["language"] != "EN"
                or not re.fullmatch(r"\d+", key) or key in seen):
            raise ValueError("Duplicate, foreign or wrong-language directory record")
        seen.add(key)
        if re.search(r"\bConfucius Institute\b", name):
            category = "institutes"
        elif re.search(r"\bConfucius Classroom\b", name):
            category = "classrooms"
        else:
            raise ValueError("Unrecognized operator program label; manual review required")
        records.append({
            "id": "cief:" + key, "category": category, "institution": name,
            "source_id": "cief_directory", "source_record_id": key, "country": "USA",
            "url": "https://ci.cn/en/site/" + key + "/", "classification": "directory_listing",
            "operating_status": "not_verified", "website_enable": row.get("enable"),
            "directory_updated_at": row.get("updateTime"), "source_date": "",
            "notes": "A directory listing does not establish that the program remains open. Website enable is a publishing control, not operational status.",
        })
    return sorted(records, key=lambda r: r["id"])


def cell_date(value):
    return value.date().isoformat() if isinstance(value, datetime) else str(value or "").strip()


def parse_roster(content):
    sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
    rows = list(sheet.values)
    header = next((i for i, r in enumerate(rows) if r[0] == "American Institution"), None)
    if header is None or rows[header][1:3] != ("Chinese Partner Institution", "State"):
        raise ValueError("Institute roster columns changed")
    records = []
    for row in rows[header + 1:]:
        if not row[0]:
            continue
        name = row[0].strip()
        if not row[2] or len(row[2]) != 2:
            raise ValueError("Institute roster has a non-institution row")
        records.append({
            "id": "nas-listed:" + hashlib.sha256(name.encode()).hexdigest()[:14],
            "institution": name, "category": "institutes", "country": "USA",
            "source_id": "nas_roster_2023", "source_date": "2023-06-22",
            "state": row[2], "chinese_partner": row[1] or "", "opened_as_reported": cell_date(row[4]),
            "host_type_as_reported": row[6] or "", "url": row[5] or "",
            "classification": "historical_listing",
            "operating_status": "paused_as_reported" if "paused" in name.lower() else "listed_as_reported",
            "notes": "Historical source roster, not independently reverified current operation. Internal header says March 2023; linked June 22, 2023 edition.",
        })
    if len(records) != 10 or sum(r["operating_status"] == "paused_as_reported" for r in records) != 1:
        raise ValueError("Reviewed ten-entry roster and paused program changed")
    return records


def parse_closures(content):
    rows = list(load_workbook(io.BytesIO(content), read_only=True, data_only=True).active.values)
    header = next((i for i, r in enumerate(rows) if r[:5] == (
        "Institution", "State", "Chinese Partner", "Date Opened", "Date Closed")), None)
    if header is None or rows[header][6:9] != ("R", "M", "T"):
        raise ValueError("Post-closure ledger columns changed")
    records = []
    for row in rows[header + 1:]:
        if not row[0]:
            continue
        flags = [str(v or "").strip() for v in row[6:9]]
        for value, allowed in zip(flags, ({"", "R", "R?"}, {"", "M", "M?"}, {"", "T", "T?"})):
            if value not in allowed:
                raise ValueError("Unrecognized post-closure classification")
        name = row[0].strip()
        evidence = " ".join(str(v or "") for v in row[9:11])
        urls = sorted(set(u.rstrip(".,);") for u in re.findall(r"https?://[^\s<>\"\u201d]+", evidence)))
        records.append({
            "id": "nas-closure:" + hashlib.sha256(name.encode()).hexdigest()[:14],
            "institution": name, "category": "institute_post_closure", "country": "USA",
            "source_id": "nas_closures_2023", "source_date": "2023-06-22",
            "state": row[1] or "", "chinese_partner": row[2] or "",
            "opened_as_reported": cell_date(row[3]), "closed_as_reported": cell_date(row[4]),
            "replacement_flag": flags[0], "partnership_flag": flags[1], "transfer_flag": flags[2],
            "classification": "source_coded_post_closure_record", "operating_status": "not_verified",
            "evidence_urls": urls,
            "notes": "Source flags can overlap and include proposals or retained activities. Not an independently verified count of currently operating successors; announced and ambiguous dates remain verbatim.",
        })
    if len({r["id"] for r in records}) != len(records):
        raise ValueError("Duplicate institution in the post-closure ledger")
    return records


def collect(root, client):
    root = Path(root)
    result = {"version": 1, "series": copy.deepcopy(SERIES), "observations": [], "records": [],
              "sources": {}, "context": {
                  "definition": "Institutes, K-12 classroom reports and documented institute successors are separate categories. No combined center total is calculated.",
                  "dates": "Reference periods retain their original year/month/day precision. Reported-by dates and retrieval dates are not invented measurement dates.",
                  "history": "Dated observations, not annual estimates. Host coverage, source methods and counting units differ; no connected or interpolated trend is drawn.",
                  "current": "The operator directory is current as retrieved, but is not an active-program census. No current nationwide operating total is verified for these three categories.",
                  "successors": "Source-reported replacements are cumulative documentation, not a current census. Proposed replacements, retained partnerships and transfers are separate, potentially overlapping classifications.",
                  "rights": "Selected factual aggregates and factual record fields only; no report prose, artwork, raw reports, contact information or administrative directory fields are redistributed.",
              }}
    rows, sources = result["observations"], result["sources"]

    def download(key, name, published, locator, notes=""):
        content, asset = client.get(URLS[key])
        sources[key] = source(name, published, URLS[key], locator, asset, notes)
        return content

    text = plain(download("mfa_2014", "China Ministry of Foreign Affairs", "2014-09-29",
                          "Spokesperson press conference, U.S. Confucius Institute question"))
    require(text, [r"now 100 Confucius Institutes and 366 Confucius classrooms"], "MFA 2014")
    for sid, value in (("institutes_reported", 100), ("classrooms_reported", 366)):
        rows.append(observation("mfa_2014_" + sid, sid, "2014-09-29", "exact", value, "mfa_2014",
                                sources["mfa_2014"]["locator"], "Official statement; source counting definitions not independently harmonized."))
    text = plain(download("nas_2017", "National Association of Scholars", "2017-04",
                          "Outsourced to China, April 2017, amended June 2017; Summary",
                          "Original 2017 report. A 2025 upload path or 2026 webpage date is not its observation year."))
    require(text, [r"APRIL 2017", r"103 remain in operation", r"501 primary and secondary schools"], "NAS 2017")
    for sid, value in (("institutes_reported", 103), ("classrooms_reported", 501)):
        rows.append(observation("nas_2017_" + sid, sid, "2017", "exact", value, "nas_2017", "Summary",
                                "Report-period snapshot from the April 2017 report, amended June 2017; no exact measurement month or end-of-year census is inferred.",
                                basis="reported_by"))
    text = plain(download("nas_2022", "National Association of Scholars", "2022-06",
                          "After Confucius Institutes, Executive Summary and printed page 30",
                          "Original June 2022 investigation. Source classifications are attributed, not adopted as causal or political conclusions."))
    require(text, [r"Published June 2022", r"by 2007 to 31", r"2015, when there were 109",
                   r"28 institutions have replaced", r"12 have sought to replace", r"58 have maintained"], "NAS 2022")
    require(text, [r"One year later the number had jumped to eleven"], "NAS retrospective history")
    for year, value in ((2006, 11), (2007, 31), (2015, 109)):
        rows.append(observation(f"nas_history_{year}", "institutes_reported", str(year), "exact", value,
                                "nas_2022", "Printed page 30, historical discussion",
                                "Retrospective source count; no month/day or continuous annual series is inferred."))
    for sid, value in (("successors_documented", 28), ("successors_proposed", 12),
                       ("partners_retained", 58), ("institutes_transferred", 5)):
        rows.append(observation("nas_2022_" + sid, sid, "2022-06", "exact", value, "nas_2022",
                                "Executive Summary, printed page 7",
                                "Counts are source-coded institutions and may overlap; not an active-program census.",
                                basis="reported_by"))
    text = plain(download("crs_2022", "Congressional Research Service", "2022-05-20",
                          "IF11180, version 11, Confucius Institutes in the United States"))
    require(text, [r"Updated May 20, 2022", r"fell to 18 as of April 2022"], "CRS 2022")
    rows.append(observation("crs_apr2022", "institutes_reported", "2022-04", "exact", 18, "crs_2022",
                            "Page 2", "Includes four announced closures, one institute under review and four non-college hosts. CRS attributes the count to NAS."))
    text = plain(download("crs_2023", "Congressional Research Service", "2023-05-02",
                          "IF11180, version 13, count attributed to the National Academies"))
    require(text, [r"Updated May 2, 2023", r"approximately seven as of December 2022"], "CRS 2023")
    rows.append(observation("crs_dec2022", "institutes_campus", "2022-12", "approximate", 7, "crs_2023",
                            "Pages 1-2", "CRS reports the National Academies' college/university count, approximately seven."))
    text = plain(download("usc_2021", "USC US-China Institute", "2021-11-18",
                          "Confucius Institutes in the United States, dated article and accompanying map"))
    require(text, [r"November 18, 2021", r"only 30 remain open today"], "USC 2021")
    rows.append(observation("usc_nov2021", "institutes_reported", "2021-11-18", "exact", 30, "usc_2021",
                            "Article introduction", "Some included programs had announced future closures."))
    text = plain(download("senate_2019", "U.S. Senate Permanent Subcommittee on Investigations", "2019-02-27",
                          "China's Impact on the U.S. Education System, Executive Summary",
                          "Original February 2019 report, later deposited with a 2023 hearing; the hearing date is not the reference year."))
    require(text, [r"Confucius Classrooms are currently in 519", r"2019"], "Senate 2019")
    rows.append(observation("senate_classrooms_2019", "classrooms_reported", "2019-02", "exact", 519,
                            "senate_2019", "Executive Summary, printed page 1",
                            "Reported classroom school sites; not physical classroom rooms.", basis="reported_by"))
    text = plain(download("nas_classrooms_2024", "National Association of Scholars", "2024-07-16",
                          "China and Our Children, Executive Summary and classroom inventory discussion"))
    require(text, [r"documented 164 Confucius Classroom", r"exact number of CCs currently in operation is unknown"], "NAS classroom inventory")
    rows.append(observation("nas_classroom_inventory_2024", "classrooms_inventory", "2024-07-16", "exact",
                            164, "nas_classrooms_2024", "Executive Summary; inventory includes current/former cases",
                            "An incomplete inventory mixing schools and school systems, not 164 active classrooms.",
                            basis="reported_by"))
    roster = download("nas_roster_2023", "National Association of Scholars", "2023-06-22",
                      "Linked ten-entry institute roster, including one paused program",
                      "Internal sheet header says March 2023. The tracker explicitly linked this edition on June 22, 2023.")
    listed = parse_roster(roster)
    result["records"].extend(listed)
    rows.append(observation("nas_listed_jun2023", "institutes_reported", "2023-06-22", "exact", len(listed),
                            "nas_roster_2023", "Current sheet, ten institution rows",
                            "Includes non-university hosts and one paused program. Listed, not ten independently verified operating institutes.",
                            basis="reported_by"))
    ledger = download("nas_closures_2023", "National Association of Scholars", "2023-06-22",
                      "Linked closed-institute workbook; original R/M/T classifications",
                      "108 institution rows, not the webpage's 111 closures/closures-in-progress. Internal header says March 2022. Differences are retained, never filled or reconciled by invention.")
    result["records"].extend(parse_closures(ledger))

    review_path = root / "data/reviewed/confucius-snapshots.json"
    review = review_path.read_bytes()
    reviewed = json.loads(review)
    if reviewed["version"] != 1 or len(reviewed["facts"]) != 3:
        raise ValueError("Reviewed Confucius selected-facts transcription changed")
    reviewed_asset = {"url": REVIEW_URL, "request_url": REVIEW_URL,
                      "sha256": hashlib.sha256(review).hexdigest(),
                      "retrieved_at": reviewed["reviewed_on"] + "T00:00:00+00:00"}
    for fact in reviewed["facts"]:
        name = "U.S. Government Accountability Office" if fact["source_id"].startswith("gao_") else "National Association of Scholars"
        sources[fact["source_id"]] = source(name, fact["published_on"], fact["publisher_url"], fact["locator"],
                                          reviewed_asset, reviewed["method"])
        sources[fact["source_id"]]["evidence_kind"] = "reviewed_selected_facts"
        sid = "institutes_campus" if fact["source_id"].startswith("gao_") else "successors_documented"
        number = fact["value"] if fact["value"] is not None else fact["bound"]
        rows.append(observation(fact["id"], sid, fact["reference"], fact["qualifier"], number,
                                fact["source_id"], fact["locator"], fact["scope"], fact["reference_basis"]))
    directory = download("cief_directory", "Chinese International Education Foundation", "",
                         "Public global-network directory, U.S. filter, complete returned record set",
                         "Lists institute and classroom names. API enable controls website publishing, not program operation. No national operating count is inferred.")
    listings = parse_directory(directory)
    snapshot = sources["cief_directory"]["assets"][0]["retrieved_at"][:10]
    for record in listings:
        record["source_date"] = snapshot
    result["records"].extend(listings)
    for category in ("institutes", "classrooms"):
        rows.append(observation("cief_" + category, category + "_listed", snapshot, "exact",
                                sum(r["category"] == category for r in listings), "cief_directory",
                                "Complete U.S. directory response, classified by program name",
                                "Directory entries only, not a verified current operating count.", basis="retrieved"))
    result["observations"].sort(key=lambda r: (r["series_id"], r["reference"], r["id"]))
    result["records"].sort(key=lambda r: r["id"])
    return validate(result)


def validate(data):
    if set(data) != ROOT_FIELDS or data["version"] != 1 or data["series"] != SERIES:
        raise ValueError("Confucius categories, units and population definitions must remain separate")
    if not all(data["context"].get(k) for k in ("definition", "dates", "history", "current", "successors", "rights")):
        raise ValueError("Confucius definitions and currentness limitations are required")
    identities = set()
    for row in data["observations"]:
        if (set(row) != ROW_FIELDS or row["id"] in identities or row["series_id"] not in SERIES
                or row["source_id"] not in data["sources"] or not row["locator"]):
            raise ValueError("Duplicate or incomplete Confucius observation")
        identities.add(row["id"])
        reference_date(row["reference"])
        if row["reference_basis"] not in {"as_of", "reported_by", "retrieved"}:
            raise ValueError("Reference dates must distinguish measurement, publication and retrieval")
        qualifier, value, bound = row["qualifier"], row["value"], row["bound"]
        if qualifier not in {"exact", "approximate", "less_than", "at_least", "unknown"}:
            raise ValueError("Unknown Confucius count qualifier")
        if qualifier in {"exact", "approximate"}:
            number = value
            if value is None or bound is not None:
                raise ValueError("Exact/approximate counts cannot be bounds")
        elif qualifier in {"less_than", "at_least"}:
            number = bound
            if value is not None or bound is None:
                raise ValueError("A bounded count cannot become an exact observation")
        else:
            number = None
            if value is not None or bound is not None:
                raise ValueError("Unknown operating totals are not zero")
        if number is not None and (type(number) is not int or number < 0 or (qualifier == "less_than" and number == 0)):
            raise ValueError("Program counts must be nonnegative integers")
        sid = row["series_id"]
        allowed_sources = {
            "institutes_reported": {"mfa_2014", "nas_2017", "nas_2022", "crs_2022", "usc_2021", "nas_roster_2023"},
            "institutes_campus": {"gao_2019", "gao_2023", "crs_2023"},
            "classrooms_reported": {"mfa_2014", "nas_2017", "senate_2019"},
        }
        if sid in allowed_sources and row["source_id"] not in allowed_sources[sid]:
            raise ValueError("Source population cannot be relabelled as a different national or campus series")
        if sid.endswith("_listed") and (row["source_id"] != "cief_directory" or row["reference_basis"] != "retrieved"):
            raise ValueError("Directory listings are not active-program measurements")
        if sid == "classrooms_inventory" and row["source_id"] != "nas_classrooms_2024":
            raise ValueError("Mixed school/district inventories cannot be national active totals")
        if sid == "successors_documented" and row["source_id"] not in {"nas_2022", "nas_tracker_2023"}:
            raise ValueError("Replacement totals require reviewed evidence, not a count of arbitrary partnerships")
        if sid in {"successors_proposed", "partners_retained", "institutes_transferred"} and row["source_id"] != "nas_2022":
            raise ValueError("Post-closure categories require the original source definition")
    if not all(any(SERIES[r["series_id"]]["category"] == category for r in data["observations"])
               for category in ("institutes", "classrooms", "successors")):
        raise ValueError("All three requested program categories are required")
    seen = set()
    for record in data["records"]:
        if record["id"] in seen or record["source_id"] not in data["sources"] or record["country"] != "USA" or not record["institution"]:
            raise ValueError("Duplicate or unsupported program record")
        seen.add(record["id"])
        if record["operating_status"] not in {"not_verified", "paused_as_reported", "listed_as_reported"}:
            raise ValueError("A historical record or directory entry cannot assert current operation")
        if record["source_id"] == "cief_directory" and record["classification"] != "directory_listing":
            raise ValueError("Directory records cannot be independently verified successors")
        if record["source_id"] == "nas_closures_2023":
            if record["classification"] != "source_coded_post_closure_record" or record["category"] != "institute_post_closure":
                raise ValueError("Original post-closure flags must not become independently confirmed rebrands")
            for key, flag in (("replacement_flag", "R"), ("partnership_flag", "M"), ("transfer_flag", "T")):
                if record[key] not in {"", flag, flag + "?"}:
                    raise ValueError("Unknown post-closure source code")
    listing_rows = [r for r in data["records"] if r["source_id"] == "cief_directory"]
    for category in ("institutes", "classrooms"):
        matches = [r for r in data["observations"] if r["series_id"] == category + "_listed"]
        if len(matches) != 1 or matches[0]["value"] != sum(r["category"] == category for r in listing_rows):
            raise ValueError("Directory headline must reconcile with every selected query record")
        if any(r["source_date"] != matches[0]["reference"] for r in listing_rows):
            raise ValueError("Directory snapshot date and records disagree")
    for key, item in data["sources"].items():
        if not item["name"] or not item["publisher_url"] or not item["locator"] or not item["assets"]:
            raise ValueError("Program source provenance incomplete")
        if item["published_on"]:
            reference_date(item["published_on"])
        for asset in item["assets"]:
            if not re.fullmatch("[0-9a-f]{64}", asset["sha256"]) or datetime.fromisoformat(asset["retrieved_at"]).tzinfo is None:
                raise ValueError("Program source checksum or retrieval timestamp invalid")
            for field in ("url", "request_url"):
                parsed = urlsplit(asset[field])
                if parsed.scheme != "https" or not parsed.netloc or parsed.username:
                    raise ValueError("Program sources require public HTTPS URLs")
            if item["evidence_kind"] == "reviewed_selected_facts" and asset["request_url"] != REVIEW_URL:
                raise ValueError("A transcription checksum must not masquerade as a publisher-response checksum")
        if key in URLS and item["assets"][0]["request_url"] != URLS[key]:
            raise ValueError("Confucius source request identity changed")
    return data


def stable_directory(records):
    return [{k: v for k, v in r.items() if k != "source_date"} for r in records]
