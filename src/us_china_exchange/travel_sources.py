"""Manual review of the public NTTO resident-air-travel monitor."""

import io
import json
import re
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import load_workbook
from pypdf import PdfReader

from .model import digest, json_text, observation, validate
from .pipeline import write_atomic
from .sources import attach


MONITOR = "https://powerbi.trade.gov/#/workspace/IA-NTTO-PUBLIC/report/SIAT_Res_Final_Monitor"
COUNTRIES = {"CHN": "China, People's Republic of", "HKG": "Hong Kong"}
ANNUAL_YEARS = (*range(2012, 2021), 2022, 2023, 2024, 2025)
PARTIAL_YEAR, PARTIAL_QUARTER = 2026, 2
CHINA_2021_PROFILE = "https://trade.gov/sites/default/files/2022-10/U.S.%20Outbound%20China%20-%202021.xlsx"
CHINA_2021_SHA256 = "f637faa552c2370e6baadf9629b06ff99c409aba382451e09350ed4a034e2d73"
SIAT_DEFINITION = (
    "NTTO SIAT estimates of U.S.-resident air travelers visiting the destination; all purposes, "
    "including multi-destination trips; not all-mode border arrivals, unique people or U.S. citizenship counts")
HISTORICAL_REPORTS = {
    2009: ("20100701065448", "http://tinet.ita.doc.gov/outreachpages/download_data_table/2009_US_Travel_Abroad.pdf",
           "China (POC)", "OTTI U.S. Resident Travel Abroad 2009; revised June 2010"),
    2010: ("20120131080346", "http://tinet.ita.doc.gov:80/outreachpages/download_data_table/2010_US_Travel_Abroad.pdf",
           "China (PRC)", "OTTI U.S. Resident Travel Abroad 2010; revised November 2011"),
    2011: ("20120916084454", "http://tinet.ita.doc.gov:80/outreachpages/download_data_table/2011_US_Travel_Abroad.pdf",
           "People's Republic of China", "OTTI U.S. Resident Travel Abroad 2011; released July 2012"),
    2012: ("20161221124638", "http://tinet.ita.doc.gov/outreachpages/download_data_table/2012_US_Travel_Abroad.pdf",
           "China", "NTTO U.S. Resident Travel Abroad 2012; released December 2013"),
}


def parse_historical_siat(body, report_year):
    label = HISTORICAL_REPORTS[report_year][2]
    years = list(range(report_year - 9, report_year + 1))
    pages = [page.extract_text() for page in PdfReader(io.BytesIO(body)).pages]
    normalize = lambda text: " ".join(text.split())
    if ("United States Resident Travel Abroad" not in normalize(pages[0])
            or str(report_year) not in pages[0] or "(000)" not in "".join(pages)):
        raise ValueError("Historical SIAT report identity or thousands unit changed")
    header = "Region/Country " + " ".join(map(str, years))
    selected = [page for page in pages if header in normalize(page) and "Hong Kong" in page]
    if len(selected) != 1 or "Historical Outbound Visitation (one or more nights)" not in normalize(selected[0]):
        raise ValueError("Historical SIAT country/year table missing or ambiguous")
    result = []
    for geo, country in (("CHN", label), ("HKG", "Hong Kong")):
        matches = [re.fullmatch(re.escape(country) + r"\s+(.+)", normalize(line))
                   for line in selected[0].splitlines()]
        matches = [match for match in matches if match]
        if len(matches) != 1:
            raise ValueError("Historical SIAT destination row missing or ambiguous: " + country)
        cells = matches[0][1].split()
        if len(cells) != len(years) or any(
                not re.fullmatch(r"(?:\d{1,3}(?:,\d{3})*|\d+|\*\*|na)", cell) for cell in cells):
            raise ValueError("Historical SIAT response columns or suppression markers changed")
        for year, cell in zip(years, cells):
            result.append({"destination": geo, "year": year,
                           "value": None if cell in {"**", "na"} else int(cell.replace(",", "")) * 1000,
                           "status": "missing" if cell in {"**", "na"} else "estimated"})
    return result


def historical_siat(client, modern_rows):
    merged = {}
    for report_year, (timestamp, original, _, vintage) in HISTORICAL_REPORTS.items():
        url = "https://web.archive.org/web/" + timestamp + "id_/" + original
        body, asset = client.get(url)
        for row in parse_historical_siat(body, report_year):
            identity = row["destination"], row["year"]
            if identity in merged and merged[identity]["value"] != row["value"]:
                raise ValueError("Historical SIAT report overlaps disagree; review the source vintages")
            year = row["year"]
            threshold = ("Source publication threshold: at least 100 respondents. " if year <= 2005 else
                         "Source publication threshold: at least 400 respondents. " if year <= 2010 else "")
            observation_row = observation(
                "ntto_siat", "USA", row["destination"], "air_visitors", year, row["value"],
                SIAT_DEFINITION, status=row["status"],
                notes=f"Original {report_year} U.S. Resident Travel Abroad report, PDF page 7; "
                      "separately reported mainland China/Hong Kong country estimates, in thousands. "
                      "Historical Outbound Visitation (one or more nights); all travel purposes. "
                      + threshold + "Exact destination sample size is not supplied. "
                      "Every overlapping retained report agrees; the 2012 values also match the current monitor. "
                      "Not the BTS combined China/Hong Kong total, citizenship counts or nonstop airport passengers. "
                      "Original publisher URL: " + original)
            merged[identity] = attach([observation_row], asset, vintage)[0]
    current = {(row["destination"], row["year"]): row["value"] for row in modern_rows
               if row["frequency"] == "annual"}
    for geo in COUNTRIES:
        if (geo, 2012) not in current or current[geo, 2012] != merged[geo, 2012]["value"]:
            raise ValueError("Historical SIAT 2012 overlap differs from the current monitor")
    if set(merged) != {(geo, year) for geo in COUNTRIES for year in range(2000, 2013)}:
        raise ValueError("Historical SIAT 2000-2012 coverage is incomplete")
    return validate([row for (_, year), row in merged.items() if year < 2012])


def parse_china_2021_profile(body):
    book = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    if book.sheetnames != ["CHINA"]:
        raise ValueError("2021 SIAT profile country identity changed")
    sheet = book["CHINA"]
    if (sheet["C5"].value != "CHINA"
            or "U.S. resident travel to China" not in str(sheet["C14"].value)
            or sheet["C63"].value != "Volume Trends (in thousands)"
            or sheet["C66"].value != "Total Departures"
            or sheet["C94"].value != "Respondents"):
        raise ValueError("2021 SIAT profile population, volume unit or row identity changed")
    if [sheet.cell(64, column).value for column in range(6, 28)] != list(range(2000, 2022)):
        raise ValueError("2021 SIAT profile chronological year columns changed")
    values = {}
    for column, year in enumerate(range(2000, 2022), 6):
        cell = sheet.cell(66, column)
        raw = cell.value
        if (isinstance(raw, bool) or not isinstance(raw, (int, float))
                or cell.number_format != "#,##0"):
            raise ValueError("2021 SIAT profile estimate missing or published precision changed")
        value = Decimal(str(raw))
        if not value.is_finite() or value <= 0:
            raise ValueError("2021 SIAT profile estimate requires clarification, not an inferred zero")
        values[year] = int(value.quantize(Decimal("1"), rounding=ROUND_HALF_UP)) * 1000
    sample = sheet["AA94"].value
    if isinstance(sample, bool) or not isinstance(sample, int) or sample <= 0:
        raise ValueError("2021 SIAT profile respondent count missing")
    return values, sample


def china_2021_siat(client, retained_rows):
    body, asset = client.get(CHINA_2021_PROFILE)
    if digest(body) != CHINA_2021_SHA256:
        raise ValueError("2021 SIAT profile changed; review the original published vintage")
    values, sample = parse_china_2021_profile(body)
    retained = {r["year"]: r["value"] for r in retained_rows
                if r["destination"] == "CHN" and r["frequency"] == "annual"}
    if any(retained.get(year) != values[year] for year in range(2000, 2021)):
        raise ValueError("2021 SIAT profile 2000-2020 overlaps differ or are incomplete")
    result = observation(
        "ntto_siat", "USA", "CHN", "air_visitors", 2021, values[2021], SIAT_DEFINITION,
        status="low_reliability" if sample < 100 else "estimated",
        notes="Original NTTO 2021 China market profile; CHINA!AA64/AA66 and AA94. "
              f"Published 5 thousand trips; {sample} survey respondents; rounded to 1,000, not an exact census. "
              "Chronological year columns F:AA exclude the narrow AB duplicate-year template helper. "
              "All 21 reported 2000-2020 overlaps match the retained series at published precision. "
              "The profile describes U.S.-resident destination estimates calibrated to APIS/I-92, "
              "not a substituted direct-airport or all-mode arrival count. "
              "Its 1 August 2023 archived copy is byte-identical to the reviewed original. "
              "The current monitor's contradictory zero is not used; Hong Kong 2021 remains unreported.")
    return validate(attach([result], asset, "NTTO U.S. Outbound China 2021 market profile; archived 1 August 2023"))


def parse_monitor_card(text, country, year, frequency):
    match = re.search(r"Total Traveler Estimate\s+to Destination\s+([\d,]+)\s+(.+?)\s+"
                      r"SIAT Respondent Count\s+to Destination\s+([\d,]+)", text, re.S)
    if not match:
        raise ValueError("Public SIAT estimate/respondent cards missing")
    title = " ".join(match[2].split())
    if frequency not in {"annual", "ytd"} or (frequency == "annual" and year not in ANNUAL_YEARS):
        raise ValueError("Unreviewed SIAT annual period")
    if frequency == "ytd" and (year not in {2025, PARTIAL_YEAR} or not re.search(
            r"Last Quarter Imported:\s*2026\s*Q2(?:\s|$)", text)):
        raise ValueError("A historical half-year filter or imported-quarter boundary has not been verified")
    expected = f"{year} Q1 - {year} Q{4 if frequency == 'annual' else 2} United States Travelers to {country}"
    imported_only = frequency == "ytd" and year == PARTIAL_YEAR and title == (
        f"{year} Q1 - {year} Q4 United States Travelers to {country}")
    if title != expected and not imported_only:
        raise ValueError("SIAT country/year/quarter filters do not match the displayed result: " + title)
    update = re.search(r"Last Update:\s*([^\n]+)", text)
    if not update or "06-Oct-2026" not in update[1]:
        raise ValueError("SIAT monitor vintage changed; review quarter availability")
    value, sample = int(match[1].replace(",", "")), int(match[3].replace(",", ""))
    if not sample or value % 1000:
        raise ValueError("SIAT empty sample or published estimate precision changed")
    if value == 0:
        raise ValueError("SIAT zero volume with positive survey responses requires source clarification")
    return value, sample, update[1].strip()


def collect_siat(client):
    from playwright.sync_api import sync_playwright

    rows = []
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()

        try:
            page.goto(MONITOR, timeout=45000, wait_until="domcontentloaded")
            page.wait_for_timeout(18000)
            frame = next(f for f in page.frames if "reportEmbed" in f.url)
            imported = "\n".join(f.locator("body").inner_text() for f in page.frames)
            if not re.search(r"Last Quarter Imported:\s*2026\s*Q2", imported):
                raise ValueError("SIAT imported-quarter boundary changed; review before collection")
            for year in (2026, 2025):
                node = frame.locator(f'[role=treeitem][title="{year}"]')
                if node.get_attribute("aria-expanded") == "true":
                    node.locator(".expandButton").click()
            year_scroller = frame.locator('[role=treeitem][title="2025"]').evaluate_handle(
                "(e)=>{let p=e.parentElement;while(p&&p.scrollHeight<=p.clientHeight)p=p.parentElement;return p;}")

            def year_node(year):
                node = frame.locator(f'[role=treeitem][title="{year}"]')
                if not node.count():
                    year_scroller.evaluate("(e)=>e.scrollTop=0")
                    for _ in range(12):
                        page.wait_for_timeout(150)
                        if node.count():
                            break
                        year_scroller.evaluate("(e)=>e.scrollTop+=180")
                if not node.count():
                    raise ValueError(f"SIAT year {year} is not available in the public year selector")
                return node

            frame.locator('[role=treeitem][title="OS"] .expandButton').click()
            frame.locator('[role=treeitem][title="Asia"] .expandButton').click()
            scroller = frame.locator('[role=treeitem][title="Asia"]').evaluate_handle(
                "(e)=>{let p=e.parentElement;while(p&&p.scrollHeight<=p.clientHeight)p=p.parentElement;return p;}")
            expanded = None
            periods = [(year, "annual") for year in ANNUAL_YEARS] + [(2026, "ytd"), (2025, "ytd")]
            for geo, country in COUNTRIES.items():
                selector = '[role=treeitem][title^="China,"]' if geo == "CHN" else '[role=treeitem][title="Hong Kong"]'
                node = frame.locator(selector)
                for attempt in range(8):
                    if node.count():
                        break
                    scroller.evaluate("(e)=>e.scrollTop+=100")
                    page.wait_for_timeout(500)
                node.locator(".slicerText").click()
                for year, frequency in periods:
                    if expanded is not None:
                        previous = year_node(expanded)
                        if previous.get_attribute("aria-expanded") == "true":
                            previous.locator(".expandButton").click()
                    selected_year = year_node(year)
                    selected_year.locator(".slicerText").click()
                    expanded = year
                    if frequency == "ytd":
                        if selected_year.get_attribute("aria-expanded") != "true":
                            selected_year.locator(".expandButton").click()
                        # Child checkboxes toggle independently; modifier-click can open a macOS context menu.
                        for quarter in ("Q3", "Q4"):
                            frame.locator(f'[role=treeitem][aria-level="2"][title="{quarter}"] .slicerText').click()
                    page.wait_for_timeout(2000)
                    text = frame.locator("body").inner_text() + "\n" + re.search(
                        r"Last Quarter Imported:\s*2026\s*Q2", imported)[0]
                    value, sample, update = parse_monitor_card(text, country, year, frequency)
                    body = text.encode()
                    checksum = digest(body)
                    write_atomic(client.cache / checksum, body)
                    asset = {"url": MONITOR, "sha256": checksum,
                             "retrieved_at": datetime.now(timezone.utc).isoformat(timespec="seconds")}
                    row = observation(
                        "ntto_siat", "USA", geo, "air_visitors", year, value,
                        SIAT_DEFINITION,
                        frequency=frequency, month=6 if frequency == "ytd" else None,
                        status="low_reliability" if sample < 100 else "estimated",
                        notes=f"Public monitor, all U.S.A. source filters; {country}; "
                              + ("explicit Q1-Q2 selection, January–June, not a full-year total. " if frequency == "ytd"
                                 else "calendar-year Q1-Q4 selection. ")
                              + f"SIAT respondent count {sample:,}; published estimate rounded to 1,000. "
                              + ("Small respondent base (<100), interpret cautiously. " if sample < 100 else "")
                              + f"Monitor updated {update}; last imported quarter 2026 Q2. "
                              "Source checksum identifies the browser-captured visible report text with the selected filters, not an API response. "
                              "2021 is excluded: the monitor displays zero volume despite positive survey responses; "
                              "no literal zero is inferred. Survey and border-arrival populations are not spliced.")
                    rows.extend(attach([row], asset, "NTTO public SIAT outbound monitor; 6 October 2026"))
        finally:
            browser.close()
    return validate(rows)


def siat(client, refresh=False):
    cache = client.cache / "siat-reviewed.json"
    if cache.is_file() and not refresh:
        rows = validate(json.loads(cache.read_text()))
        for row in rows:
            body = (client.cache / row["raw_sha256"]).read_bytes()
            if digest(body) != row["raw_sha256"]:
                raise ValueError("Retained SIAT source response checksum mismatch")
            value, sample, _ = parse_monitor_card(
                body.decode(), COUNTRIES[row["destination"]], row["year"], row["frequency"])
            if value != row["value"] or row["status"] != ("low_reliability" if sample < 100 else "estimated"):
                raise ValueError("Retained SIAT observation differs from visible source evidence")
    else:
        rows = collect_siat(client)
        write_atomic(cache, json_text(rows))
    if {(r["destination"], r["year"], r["frequency"]) for r in rows} != {
            (geo, year, frequency) for geo in COUNTRIES for year, frequency in
            [(y, "annual") for y in ANNUAL_YEARS] + [(2026, "ytd"), (2025, "ytd")]}:
        raise ValueError("Reviewed SIAT destination/year coverage incomplete")
    rows = validate(historical_siat(client, rows) + rows)
    rows = validate(rows + china_2021_siat(client, rows))
    return rows, "NTTO SIAT public monitor, 6 October 2026; historical reports and original 2021 China profile; through 2026 Q2", (
        "U.S.-resident air-travel destination estimates, kept separate from all-mode visitor arrivals. "
        "Mainland annual 2000-2025; Hong Kong 2000-2020 and 2022-2025. "
        "2000-2011 recovered from original resident-travel reports, with "
        "all historical overlaps and both 2012 monitor values reconciled. Published thousands, country isolation "
        "and original sample/suppression thresholds are retained. "
        "Separately selected Q1-Q2 observations in 2025 and 2026 support matched January–June comparisons. "
        "Quarter checkboxes are selected directly, not inferred by subtracting rounded annual and half-year totals. "
        "Monitor total-traveler and respondent cards read from the public report with explicit year, quarter and destination filters; "
        "source checksums identify retained browser-visible report text, not API responses. "
        "Estimates are rounded to 1,000; respondent counts appear in each observation's notes. "
        "Under-100 respondent bases are conservatively flagged low reliability. "
        "The 2021 monitor cards show zero volume despite positive respondent counts (China 267; Hong Kong 44); "
        "these contradictory volume outputs are not interpreted as zero actual travel. "
        "Mainland 2021 is separately recovered as the original profile's published 5,000-trip estimate, "
        "with its checksum-pinned chronological columns and all 21 retained overlaps verified. "
        "Hong Kong 2021 remains unreported. No aggregation across destinations.")
