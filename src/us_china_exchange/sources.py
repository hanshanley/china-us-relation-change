"""Primary-source parsers. A changed document layout fails closed."""

import csv
import io
import re
import time
from datetime import datetime
from decimal import Decimal
from urllib.parse import urljoin

import xlrd
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from pypdf import PdfReader

from .model import START_YEAR, digest, observation


BEA = "https://www.bea.gov/data/intl-trade-investment/international-trade-goods-and-services"
BEA_HISTORY = "https://www.bea.gov/sites/default/files/2024-12/trad-geo-time-series-1024.xlsx"
NTTO = "https://www.trade.gov/i-94-arrivals-program"
USTR_MACAU = "https://ustr.gov/countries-regions/china-mongolia-taiwan/macau"
UIS = "https://api.uis.unesco.org/api/public/"
TAIWAN = "https://stat.taiwan.net.tw/data/opendata/6"
IIE_INDEX = "https://opendoorsdata.org/2025-country-and-area-fact-sheet-pdfs/"
IIE_FILES = {
    "CHN": "https://opendoorsdata.org/wp-content/uploads/2026/01/OD_China_Fact-Sheet-2025.pdf",
    "HKG": "https://opendoorsdata.org/wp-content/uploads/2026/01/OD_Hong-Kong_Fact-Sheet-2025.pdf",
    "MAC": "https://opendoorsdata.org/wp-content/uploads/2026/01/OD_Macau_Fact-Sheet-2025.pdf",
    "TWN": "https://opendoorsdata.org/wp-content/uploads/2026/01/OD_Taiwan_Fact-Sheet-2025.pdf",
}
SOURCE_POLICIES = {
    "iie": {
        "automatic_retrieval": False,
        "terms_url": "https://www.iie.org/terms-and-conditions/",
        "notice": (
            "IIE automatic retrieval is paused: its terms prohibit automated downloading and analysis. "
            "Previously reviewed observations are retained, not newly checked. Separate retrieval and "
            "reuse permission has not been verified; public access is not commercial-use clearance."),
        "figure_notice": "IIE: automatic refresh paused; reuse permission unverified.",
    },
    "macau_tourism": {
        "automatic_retrieval": True,
        "terms_url": "https://dataplus.macaotourism.gov.mo/Indicator/VisitorsSummary/SummaryBar?lang=E",
        "notice": (
            "MGTO site terms limit use to noncommercial purposes. Commercial reuse permission has not "
            "been verified; these exports do not grant publication rights."),
        "figure_notice": "MGTO: commercial reuse permission unverified.",
    },
}


class SourcePermissionRequired(RuntimeError):
    pass


def workbook(content):
    if not content.startswith(b"PK"):
        raise ValueError("Expected an XLSX workbook")
    return load_workbook(io.BytesIO(content), data_only=True, read_only=True)


def attach(rows, asset, vintage):
    for row in rows:
        row.update(source_url=asset["url"], raw_sha256=asset["sha256"],
                   retrieved_at=asset["retrieved_at"], source_vintage=vintage)
    return rows


def links(content, base):
    return [(urljoin(base, a["href"]), a.get_text(" ", strip=True))
            for a in BeautifulSoup(content, "html.parser").select("a[href]")]


class BEAEarlyHistoryError(ValueError):
    pass


def parse_bea(content, omit_ambiguous_early_history=False):
    book = workbook(content)
    rows = []
    vintage = str(book["Table 4"].cell(1, 1).value)
    if not vintage.startswith("Last updated "):
        raise ValueError("BEA release label missing")
    for number, measure, outbound in [(4, "goods", True), (5, "goods", False),
                                      (7, "services", True), (8, "services", False)]:
        table = list(book[f"Table {number}"].values)
        if "Millions of dollars" not in str(table[3][0]):
            raise ValueError("Unexpected BEA units")
        header = next(r for r in table if r[0] == "Period" and "China" in r)
        annual_start = next(i for i, r in enumerate(table) if r[0] == "Annual") + 1
        quarterly_start = next(i for i, r in enumerate(table) if r[0] == "Quarterly")
        annual_labels = [str(r[0]) for r in table[annual_start:quarterly_start]]
        if annual_labels.count("2000") == 2 and "2001" not in annual_labels:
            if not omit_ambiguous_early_history:
                raise BEAEarlyHistoryError("BEA repeats 2000 and omits 2001; require independently labelled history")
        section = None
        quarters = {}
        for cells in table:
            label = str(cells[0]).strip()
            if label in {"Annual", "Quarterly"}:
                section = label
                continue
            match = re.fullmatch(r"(\d{4})(?:\s+([1-4]))?(?:\s*\(R\))?", label)
            if not match:
                continue
            year = int(match[1])
            if year < START_YEAR:
                continue
            if omit_ambiguous_early_history and section == "Annual" and year in {2000, 2001}:
                continue
            for name, geo in [("China", "CHN"), ("Hong Kong", "HKG"), ("Taiwan", "TWN")]:
                value = cells[header.index(name)]
                if not isinstance(value, (float, int)):
                    raise ValueError(f"Unexpected BEA value: {name} {label}")
                origin, destination = ("USA", geo) if outbound else (geo, "USA")
                if section == "Annual" and match[2] is None:
                    rows.append(observation("bea", origin, destination, measure, year,
                                            value * 1_000_000,
                                            "BEA balance-of-payments basis; nominal USD; annual, not seasonally adjusted",
                                            notes="Source units: millions of USD. " + label))
                elif section == "Quarterly" and match[2]:
                    quarters[(geo, year, int(match[2]))] = value * 1_000_000
        if not quarters:
            raise ValueError("BEA quarterly block missing")
        latest_year = max(k[1] for k in quarters)
        end_quarter = max(k[2] for k in quarters if k[1] == latest_year)
        for name, geo in [("China", "CHN"), ("Hong Kong", "HKG"), ("Taiwan", "TWN")]:
            for year in range(START_YEAR, latest_year + 1):
                values = [quarters.get((geo, year, q)) for q in range(1, end_quarter + 1)]
                if all(v is not None for v in values):
                    origin, destination = ("USA", geo) if outbound else (geo, "USA")
                    rows.append(observation(
                        "bea", origin, destination, measure, year, sum(values),
                        "BEA balance-of-payments basis; nominal USD; seasonally adjusted quarterly sums",
                        frequency="ytd", month=end_quarter * 3,
                        notes=f"Derived from complete Q1-Q{end_quarter}; source may revise history."))
    if len([r for r in rows if r["frequency"] == "annual"]) < 12:
        raise ValueError("Incomplete BEA annual tables")
    return rows, vintage


def bea(client):
    page, _ = client.get(BEA)
    urls = [u for u, text in links(page, BEA) if "geo-time-series" in u and u.endswith(".xlsx")]
    if len(set(urls)) != 1:
        raise ValueError("Could not uniquely discover BEA geographic history")
    content, asset = client.get(urls[0])
    try:
        rows, vintage = parse_bea(content)
    except BEAEarlyHistoryError:
        historical, history_asset = client.get(BEA_HISTORY)
        history_rows, history_vintage = parse_bea(historical)
        history_rows = [r for r in history_rows if r["frequency"] == "annual" and r["year"] in {2000, 2001}]
        if len(history_rows) != 24:
            raise ValueError("BEA independently labelled 2000/2001 history incomplete")
        for row in history_rows:
            row["notes"] += " Archived release used because the current workbook repeats 2000 and omits 2001; no year inferred."
        rows, vintage = parse_bea(content, omit_ambiguous_early_history=True)
        rows = attach(rows, asset, vintage) + attach(history_rows, history_asset, history_vintage)
        return rows, vintage, (
            "Annual 2000/2001 use the December 5, 2024 geographic workbook with its own provenance; "
            "the current workbook repeats the 2000 label and omits 2001. Later years use the current release.")
    return attach(rows, asset, vintage), vintage, ""


COUNTRIES = {"China": "CHN", "Hong Kong": "HKG", "Macao SAR": "MAC", "Taiwan": "TWN"}
VISIT_DEFINITION = (
    "NTTO I-94 overnight visitor trips by country of residence, qualified visa types; "
    "all purposes, not unique people or leisure-only visitors")


def parse_ntto(content, monthly=False):
    book = workbook(content)
    table = list(book["Monthly" if monthly else "Annual"].values)
    headers = table[0]
    country_col = 1 if monthly else 0
    periods = []
    for col, label in enumerate(headers):
        month = None
        if monthly:
            if isinstance(label, datetime):
                year, month = label.year, label.month
            else:
                match = re.match(r"^(\d{4})-(\d{1,2})(?:\D|$)", str(label))
                if not match:
                    continue
                year, month = map(int, match.groups())
        else:
            if not isinstance(label, int) or isinstance(label, bool):
                continue
            year = label
        if year >= START_YEAR:
            periods.append((col, year, month, str(label)))
    rows, found = [], set()
    for cells in table:
        geo = COUNTRIES.get(cells[country_col])
        if not geo:
            continue
        found.add(geo)
        for col, year, month, label in periods:
            value = cells[col]
            if value is not None and not isinstance(value, (int, float)):
                raise ValueError(f"Unexpected NTTO value: {geo} {label}")
            rows.append(observation(
                "ntto", geo, "USA", "visitors", year, value, VISIT_DEFINITION,
                frequency="monthly" if monthly else "annual", month=month,
                status="missing" if value is None else (
                    "provisional" if "preliminary" in label.lower() else "observed"),
                notes=label if monthly else "Official annual total; may differ from revised monthly sums."))
    if found != set(COUNTRIES.values()) or not periods:
        raise ValueError("NTTO geography or period headers missing")
    if monthly:
        available = [r for r in rows if r["value"] is not None]
        year = max(r["year"] for r in available)
        end_month = max(int(r["period"][5:7]) for r in available if r["year"] == year)
        index = {(r["origin"], r["year"], int(r["period"][5:7])): r for r in rows}
        for geo in found:
            for y in range(START_YEAR, year + 1):
                months = [index.get((geo, y, m)) for m in range(1, end_month + 1)]
                if all(r is not None and r["value"] is not None for r in months):
                    rows.append(observation(
                        "ntto", geo, "USA", "visitors", y, sum(r["value"] for r in months),
                        VISIT_DEFINITION, frequency="ytd", month=end_month,
                        status="provisional" if any(r["status"] == "provisional" for r in months) else "observed",
                        notes="Derived sum of complete January-to-cutoff monthly observations."))
    vintage = "Workbook modified " + str(book.properties.modified or "date not supplied")
    return rows, vintage


def ntto(client):
    page, _ = client.get(NTTO)
    found = links(page, NTTO)
    rows, vintages = [], []
    for monthly in [False, True]:
        urls = [u for u, text in found if ".xlsx" in u.lower()
                and ("monthly" if monthly else "annual") in (u + text).lower()
                and ("residence" in (u + text).lower())
                and ("present" in (u + text).lower())]
        if not urls:
            raise ValueError("NTTO current history link missing")
        url = urls[0].replace("www.trade.gov//", "www.trade.gov/")
        content, asset = client.get(url)
        candidate, vintage = parse_ntto(content, monthly)
        rows.extend(attach(candidate, asset, vintage))
        vintages.append(vintage)
    return rows, "; ".join(vintages), "Workbook modification dates are not official publication dates."


def uis(client):
    content, _ = client.get(UIS + "definitions/indicators")
    catalog = client.json(content)
    if isinstance(catalog, dict):
        catalog = catalog.get("indicators", catalog.get("records", []))
    catalog_by_id = {str(x["indicatorCode"]): x for x in catalog}
    rows, updates = [], []
    for geo, indicator in [("CHN", "26528"), ("HKG", "26532"), ("MAC", "26547")]:
        for origin, destination, indicator_id in [(geo, "USA", indicator), ("USA", geo, "26503")]:
            content, asset = client.get(
                UIS + "data/indicators",
                params={"indicator": indicator_id, "geoUnit": destination,
                        "start": START_YEAR, "footnotes": "true"})
            data = client.json(content)
            if not isinstance(data.get("records"), list):
                raise ValueError("UIS records missing")
            info = catalog_by_id.get(indicator_id, {})
            vintage = info.get("lastDataUpdateDescription", "UIS release date not supplied")
            updates.append(vintage)
            candidate = []
            for rec in data["records"]:
                if str(rec["indicatorId"]) != indicator_id or rec["geoUnit"] != destination:
                    raise ValueError("UIS indicator/destination mismatch")
                value = rec["value"]
                magnitude, qualifier = rec.get("magnitude"), rec.get("qualifier")
                if magnitude not in {None, "NIL", "NA", "SUPP", "LOWREL", "INCLUDED", "INCLUDES"}:
                    raise ValueError("Unknown UIS magnitude flag")
                if qualifier not in {None, "NAT_EST", "UIS_EST"}:
                    raise ValueError("Unknown UIS qualifier flag")
                if magnitude in {"NIL", "NA"} and (isinstance(value, bool) or value != 0):
                    raise ValueError("UIS NIL/NA flags require a source zero")
                if magnitude in {"SUPP", "INCLUDED"} and value is not None:
                    raise ValueError("UIS suppressed/included values must be blank")
                if magnitude in {"LOWREL", "INCLUDES"} and value is None:
                    raise ValueError("UIS reliability/inclusion flags require a numeric value")
                if magnitude in {"NA", "SUPP", "INCLUDED"}:
                    status = {"NA": "not_applicable", "SUPP": "suppressed",
                              "INCLUDED": "included_elsewhere"}[magnitude]
                    value = None
                else:
                    status = ("missing" if value is None else "low_reliability" if magnitude == "LOWREL"
                              else "estimated" if qualifier else "observed")
                flag_note = (f"Magnitude: {magnitude}; source value: {rec['value']}; "
                             if magnitude is not None else "")
                candidate.append(observation(
                    "uis", origin, destination, "degree_students", int(rec["year"]), value,
                    "UIS internationally mobile tertiary degree/diploma students; destination-reported; reporting year",
                    period_basis="reporting", status=status,
                    notes=flag_note + "Qualifier: " + str(qualifier) +
                    "; footnotes: " + str(rec.get("footnotes", [])) +
                    "; fractional source estimates retained. Not a physical-presence or departure count."))
            rows.extend(attach(candidate, asset, vintage))
    if not rows:
        raise ValueError("UIS returned no observations across all requested series")
    return rows, "; ".join(sorted(set(updates))), "Taiwan not separately listed; U.S.-to-mainland records may be absent."


def parse_taiwan(content):
    reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
    if not {"年別", "細分", "合計", "觀光"}.issubset(reader.fieldnames or []):
        raise ValueError("Taiwan CSV column layout changed")
    records = list(reader)
    rows = []
    for rec in records:
        if not rec["細分"].endswith("USA"):
            continue
        for measure, col in [("visitors", "合計"), ("leisure_visitors", "觀光")]:
            value = int(rec[col].replace(",", ""))
            rows.append(observation(
                "taiwan_tourism", "USA", "TWN", measure, int(rec["年別"]), value,
                "Taiwan inbound visitor arrivals by residence, " +
                ("all purposes" if measure == "visitors" else "declared sightseeing purpose") +
                "; trips, not unique people; not restricted to overnight stays"))
    if not rows:
        raise ValueError("Taiwan CSV has no U.S. residence records")
    return rows


def taiwan(client):
    content, asset = client.get(TAIWAN)
    rows = parse_taiwan(content)
    vintage = "Official annual CSV through " + str(max(r["year"] for r in rows))
    return attach(rows, asset, vintage), vintage, "Public open-data CSV; scope differs from NTTO overnight trips."


def parse_iie(content, geo):
    if not content.startswith(b"%PDF"):
        raise ValueError("Expected an IIE PDF")
    text = "\n".join(re.sub(r"[ \t]+", "", line) for page in
                     PdfReader(io.BytesIO(content)).pages for line in page.extract_text().splitlines())
    rows = []
    for match in re.finditer(r"(20\d{2})/(\d{2})\n([\d,]+)\n([\d,]+|Null|N/A)", text):
        year = int(match[1]) + 1
        if int(match[2]) != year % 100:
            raise ValueError("IIE academic year ending mismatch")
        for measure, value, origin, destination in [
            ("international_students", match[3], geo, "USA"),
            ("study_abroad", match[4], "USA", geo),
        ]:
            missing = value in {"Null", "N/A"}
            pandemic = 2021 <= year <= 2023
            student_definition = (
                "IIE Open Doors international students, including Optional Practical Training (OPT); " +
                ("expanded pandemic population includes coursework from outside the U.S. regardless of visa status"
                 if pandemic else
                 "standard temporary-visa population; excludes students enrolled online from outside the U.S."))
            population_note = (
                " Open Doors 2021-2023 (2020/21-2022/23) expanded the population to coursework from outside "
                "the U.S. regardless of visa status; Open Doors 2024 returned to the pre-pandemic definition. "
                "These populations are not directly comparable.")
            rows.append(observation(
                "iie", origin, destination, measure, year,
                None if missing else int(value.replace(",", "")),
                (student_definition
                 if measure == "international_students" else
                 "IIE Open Doors U.S. students receiving academic credit for study abroad; participation, not degree enrollment"),
                period_basis="academic", status="missing" if missing else "observed",
                comparability=("iie:international_students:" + ("pandemic-expanded" if pandemic else "standard")
                               if measure == "international_students" else ""),
                notes="Academic year ending shown. Counts are not departures." +
                (population_note if measure == "international_students" else "")))
    enrollment = {}
    for level in ["Undergraduate", "Graduate", "Non-Degree"]:
        match = re.search(level + r"\n([\d,]+)\n([\d,]+)\n", text)
        if not match:
            raise ValueError("IIE enrollment breakdown layout changed: " + level)
        enrollment[level] = [int(match[n].replace(",", "")) for n in [1, 2]]
    latest = max((r["year"] for r in rows), default=0)
    if latest != 2025 or len(rows) != 16:
        raise ValueError("IIE reviewed 2025 fact-sheet layout or vintage changed")
    for i, year in enumerate([2024, 2025]):
        rows.append(observation(
            "iie", geo, "USA", "enrolled_students", year,
            sum(values[i] for values in enrollment.values()),
            "IIE Open Doors undergraduate + graduate + non-degree enrollment; excludes OPT; "
            "not necessarily physically present", period_basis="academic",
            notes="Derived by adding three published academic-level counts, never OPT."))
    return rows


def iie(client):
    raise SourcePermissionRequired(SOURCE_POLICIES["iie"]["notice"])


MACAU_INDEX = "https://dataplus.macaotourism.gov.mo/Publication/StatisticalTables?lang=E"

NBS_VISITOR_BOOKS = [
    (2007, "https://www.stats.gov.cn/sj/ndsj/2007/html/S1904e.xls", 2000, 2006),
    (2010, "https://www.stats.gov.cn/sj/ndsj/2010/html/Q1823e.xls", 2007, 2009),
    (2013, "https://www.stats.gov.cn/sj/ndsj/2013/html/Z1813E.xls", 2010, 2012),
    (2014, "https://www.stats.gov.cn/sj/ndsj/2014/zk/html/Z1713E.xls", 2013, 2013),
]
NBS_VISITOR_IMAGES = [
    (2015, "https://www.stats.gov.cn/sj/ndsj/2015/html/EN1713.jpg",
     "b8746ca462e96a2bc9c6529fd64d7b590a0bab453be1b6b69f4c900279960c4b", {2014: "209.32"}),
    (2018, "https://www.stats.gov.cn/sj/ndsj/2018/html/EN1713.jpg",
     "0d62746af9611e8c9718d57777a7d6d05da1c570b29b36157d016bbe9af00642",
     {2015: "208.58", 2016: "224.96", 2017: "231.29"}),
    (2019, "https://www.stats.gov.cn/sj/ndsj/2019/html/C1713.jpg",
     "9158f19cee30f78d45d6b24fa5cd6c960a341a163734251a2cfb2096f17bca54", {2018: "248.46"}),
]


def parse_nbs_visitor_records(records, first, last):
    if not any("10 000 person-times" in str(v) for r in records[:6] for v in r):
        raise ValueError("NBS visitor units changed")
    header = next(r for r in records if str(r[0]).strip() == "Region")
    columns = {int(value): i for i, value in enumerate(header) if isinstance(value, (int, float)) and 2000 <= value <= 2100}
    if not set(range(first, last + 1)) <= columns.keys():
        raise ValueError("NBS visitor year coverage incomplete")
    matches = [r for r in records if str(r[0]).strip() == "United States"]
    if len(matches) != 1:
        raise ValueError("NBS United States visitor row missing or ambiguous")
    values = {}
    for year in range(first, last + 1):
        value = matches[0][columns[year]]
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            raise ValueError("NBS visitor missingness must not become zero")
        count = Decimal(str(value)) * 10000
        if count < 0 or count != int(count):
            raise ValueError("NBS person-time count does not reconcile")
        values[year] = int(count)
    return values


def nbs_tourism(client):
    result = []

    def add(values, asset, edition, method):
        for year, value in values.items():
            result.extend(attach([observation(
                "nbs_tourism", "USA", "CHN", "visitors", year, value,
                "NBS/MCT overseas visitor arrivals attributed to the United States; all-purpose trips, not unique people or leisure-only tourists",
                notes=f"{method}; source unit 10,000 person-times, converted to trips. Published precision retained. "
                      "Country attribution is not a count of U.S. residents departing from the U.S. "
                      "This historical all-mode series is not extended using air-only survey estimates."
            )], asset, f"China Statistical Yearbook {edition}; overseas visitor arrivals by country/region"))

    for edition, url, first, last in NBS_VISITOR_BOOKS:
        body, asset = client.get(url)
        sheet = xlrd.open_workbook(file_contents=body).sheet_by_index(0)
        values = parse_nbs_visitor_records([sheet.row_values(i) for i in range(sheet.nrows)], first, last)
        add(values, asset, edition, "Machine-read official Excel United States row")
    for edition, url, checksum, cells in NBS_VISITOR_IMAGES:
        body, asset = client.get(url)
        if digest(body) != checksum:
            raise ValueError("Reviewed NBS visitor image changed; recheck the country row before publication")
        add({year: int(Decimal(value) * 10000) for year, value in cells.items()}, asset, edition,
            "Visually reviewed official table, United States row; image checksum pinned; cells rounded to 0.01 source units")
    if {r["year"] for r in result} != set(range(2000, 2019)):
        raise ValueError("Reviewed mainland visitor history incomplete")
    return result, "NBS/MCT historical visitor tables, 2000–2018", (
        "Official China Statistical Yearbooks 2007, 2010, 2013, 2014, 2015, 2018 and 2019. "
        "Selected historical source vintages; source units converted from 10,000 person-times. "
        "Later bilateral all-mode counts are not verified here. Not leisure-only visitors, unique people, "
        "U.S. departures, or a series spliced to U.S. resident air-travel estimates.")


def parse_macau(content, year):
    if not content.startswith(b"%PDF"):
        raise ValueError("Expected a Macau tourism PDF")
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(content)).pages)
    if str(year) not in text or "Visitor Arrivals" not in text:
        raise ValueError("Unexpected Macau report vintage or content")
    matches = re.findall(r"(?:USA|U\.S\.A\.?|United States(?: of America)?)\s+(\d[\d,]*)(?=\s)", text)
    values = {int(v.replace(",", "")) for v in matches
              if re.fullmatch(r"\d{1,3}(?:,\d{3})+|\d+", v)}
    if len(values) > 1:
        raise ValueError("Ambiguous U.S. totals in Macau report; review required")
    return next(iter(values)) if values else None


def parse_macau_market(content, edition):
    if not content.startswith(b"%PDF"):
        raise ValueError("Expected a Macau market-performance PDF")
    pages = [p.extract_text() for p in PdfReader(io.BytesIO(content)).pages]
    profiles = [t for t in pages if "VISITOR ARRIVALS FROM U.S.A." in t]
    basis = "place of residence" if edition < 2015 else "place of issue of travel document"
    if len(profiles) != 1 or basis not in " ".join(profiles[0].split()).lower():
        raise ValueError("Macau U.S. residence/issuing-document profile missing or ambiguous")
    values = {}
    for match in re.finditer(
            r"^(20\d{2})[ \t]+([\d,]+)[ \t]+([\d,]+)[ \t]+([\d,]+)[ \t]+([\d,]+)[ \t]+100\.0+%",
            profiles[0], re.MULTILINE):
        year, sea, land, air, total = (int(v.replace(",", "")) for v in match.groups())
        if year in values or sea + land + air != total:
            raise ValueError("Macau annual profile totals do not reconcile")
        values[year] = total
    if set(values) != set(range(edition - 5, edition + 1)):
        raise ValueError("Macau annual profile must cover six complete years")
    return values


def macau_observation(year, value, notes):
    basis = ("place of residence; includes non-resident workers and students" if year < 2008 else
             "place of residence" if year < 2015 else
             "mixed residence/document-issue transition" if year == 2015 else
             "place of issue of travel document")
    return observation(
        "macau_tourism", "USA", "MAC", "visitors", year, value,
        f"MGTO/DSEC U.S.-attributed arrivals by {basis}; same-day and overnight trips, not unique people",
        comparability="macau:" + basis,
        status="observed" if value is not None else "missing",
        notes=notes + " From 2008, arrivals exclude non-resident workers and students. "
        "Classification changed in April 2015; 2015 mixes methods. "
        "Document-issue attribution is not U.S. residence or a departure from the U.S.")


def macau(client):
    page, _ = client.get(MACAU_INDEX)
    urls = [urljoin("https://dataplus.macaotourism.gov.mo/", o["value"])
            for o in BeautifulSoup(page, "html.parser").select("option[value]")
            if "VAHightlights" in o["value"] and "Dec" in o["value"]]
    urls = [(int(re.search(r"/(\d{4})/", u)[1]), u) for u in urls]
    urls = [(year, u) for year, u in urls if year >= START_YEAR]
    if not urls:
        raise ValueError("Macau annual report archive missing")
    rows = {}
    for year, url in sorted(urls):
        content, asset = client.get(url)
        value = parse_macau(content, year)
        rows[year] = attach([macau_observation(
            year, value,
            "December report cumulative annual table; country may be absent from selected-market tables. "
            "2020 report contains a stale Jan-Nov section heading; surrounding text identifies annual 2020 totals."
        )], asset, f"MGTO December {year}")[0]
    profiles = []
    for option in BeautifulSoup(page, "html.parser").select("option[value]"):
        match = re.search(r"/Market Performance\s+(\d{4})\.pdf$", option["value"])
        if not match:
            match = re.search(r"/M12_Market(?:ing)?Performance_(\d{4})_\d+\.pdf$", option["value"])
        if match:
            profiles.append((int(match[1]), urljoin(MACAU_INDEX, option["value"])))
    if not profiles:
        raise ValueError("Macau complete annual market-performance archive missing")
    historical = [next((p for p in profiles if p[0] == year), None) for year in [2008, 2010, 2015]]
    if any(p is None for p in historical):
        raise ValueError("Macau historical 2008/2010/2015 profile missing")
    latest = max(profiles)
    editions = sorted(set(historical + [latest]))
    for edition, url in editions:
        content, asset = client.get(url)
        for year, value in parse_macau_market(content, edition).items():
            rows[year] = attach([macau_observation(
                year, value, f"Machine-read U.S. annual profile in Market Performance {edition}; "
                "sea + land + air reconciles to total. Complete annual profile takes priority over older highlights."
            )], asset, f"MGTO Market Performance {edition}")[0]
    return list(rows.values()), (
        f"MGTO annual highlights through {max(y for y, u in urls)}; Market Performance " +
        ", ".join(str(year) for year, _ in editions)), (
        "Selected-market PDFs; unlisted U.S. values are missing, never zero. "
        "Latest complete annual market profile fills selected-market omissions. "
        "Arrivals exclude non-resident workers/students from 2008. "
        "Residence classification changed to issuing place of travel document in April 2015; transition is not comparable. "
        "This adapter covers complete annual reports, not newer monthly highlights. " +
        SOURCE_POLICIES["macau_tourism"]["notice"])


def parse_ustr_macau_services(content):
    plain = " ".join(BeautifulSoup(content, "html.parser").get_text(" ", strip=True).split())
    if "Macau Trade Summary" not in plain:
        raise ValueError("USTR Macau country identity changed")
    patterns = {
        "exports": r"U\.S\. services exports to Macau in (20\d{2}) were \$([\d,.]+) (million|billion)",
        "imports": r"U\.S\. services imports from Macau in (20\d{2}) were \$([\d,.]+) (million|billion)",
        "total": r"U\.S\. services trade \(exports plus imports\) with Macau totaled an estimated \$([\d,.]+) (million|billion) in (20\d{2})",
        "balance": r"U\.S\. services trade (surplus|deficit) with Macau was \$([\d,.]+) (million|billion) in (20\d{2})",
    }
    figures, years, quotes, precisions = {}, set(), {}, {}
    for name, pattern in patterns.items():
        matches = list(re.finditer(pattern, plain))
        if len(matches) != 1:
            raise ValueError("USTR services amount/identity missing or ambiguous: " + name)
        match = matches[0]
        if name in {"exports", "imports"}:
            year, amount, unit = int(match[1]), match[2], match[3]
        elif name == "total":
            year, amount, unit = int(match[3]), match[1], match[2]
        else:
            year, amount, unit = int(match[4]), match[2], match[3]
        scale = Decimal(1_000_000 if unit == "million" else 1_000_000_000)
        number = Decimal(amount.replace(",", ""))
        figures[name] = int(number * scale)
        precisions[name] = scale * Decimal(10) ** number.as_tuple().exponent
        if name == "balance" and match[1] == "deficit":
            figures[name] *= -1
        quotes[name] = "$" + amount + " " + unit
        years.add(year)
    if len(years) != 1:
        raise ValueError("USTR services figures refer to different annual periods")
    year = years.pop()
    if year < START_YEAR or year >= datetime.now().year:
        raise ValueError("USTR services must be a reviewed complete year, not a forecast")
    for control, actual in [
        ("total", figures["exports"] + figures["imports"]),
        ("balance", figures["exports"] - figures["imports"]),
    ]:
        tolerance = (precisions["exports"] + precisions["imports"] + precisions[control]) / 2
        if abs(Decimal(actual) - figures[control]) > tolerance:
            raise ValueError("USTR service flows do not reconcile within quoted rounding precision")
    rows = []
    for flow, origin, destination in [("exports", "USA", "MAC"), ("imports", "MAC", "USA")]:
        rows.append(observation(
            "ustr", origin, destination, "services", year, figures[flow],
            "USTR U.S.-reported estimated cross-border services trade with Macau; nominal USD; "
            "annual exports/imports, not foreign-affiliate sales or investment",
            status="estimated", comparability="ustr:macau:cross_border_services",
            notes="Direct country services figure: " + quotes[flow] + ". Source estimated/rounded "
            "amount converted to USD, not single-dollar precision. Services exports + imports "
            "and the trade balance reconcile with the independently quoted controls within source "
            "precision. No goods total or regional remainder is allocated to services. "
            "Prior-year amounts are not inferred from percentages or rounded changes; older history "
            "remains unverified until another directly published annual source is reviewed."))
    return rows, year


def ustr_macau_services(client):
    content, asset = client.get(USTR_MACAU)
    rows, year = parse_ustr_macau_services(content)
    return attach(rows, asset, f"USTR Macau Trade Summary, {year}"), (
        f"USTR Macau Trade Summary; directly reported {year} estimated services"), (
        "Separate direct USTR country estimates fill Macau services missing from the selected BEA "
        "workbook. Exports, imports, total services and surplus/deficit are checked together. "
        "Rounded nominal USD; not regional allocations or affiliate sales. Longer history not imputed.")


def comtrade(client):
    rows = []
    latest = datetime.now().year - 1
    for year in range(START_YEAR, latest + 1):
        content, asset = client.get(
            "https://comtradeapi.un.org/public/v1/preview/C/A/HS",
            params={"period": year, "reporterCode": 842, "partnerCode": 446,
                    "cmdCode": "TOTAL", "flowCode": "X,M"})
        data = client.json(content)
        if not isinstance(data.get("data"), list):
            raise ValueError("Comtrade records missing")
        seen = set()
        for rec in data["data"]:
            flow = rec["flowCode"]
            if (str(rec.get("period")) != str(year) or rec.get("refYear") != year
                    or rec.get("freqCode") != "A" or rec.get("typeCode") != "C"):
                raise ValueError("Comtrade returned the wrong period, frequency or trade type")
            if (rec["reporterCode"], rec["partnerCode"], rec["cmdCode"]) != (842, 446, "TOTAL"):
                raise ValueError("Comtrade partner or aggregate mismatch")
            if flow not in {"X", "M"} or flow in seen:
                raise ValueError("Unexpected Comtrade flow")
            seen.add(flow)
            valuation = "FOB" if flow == "X" else "CIF"
            value = rec.get("fobvalue" if flow == "X" else "cifvalue")
            if not isinstance(value, (int, float)):
                raise ValueError("Comtrade documented " + valuation + " valuation unavailable")
            origin, destination = ("USA", "MAC") if flow == "X" else ("MAC", "USA")
            rows.extend(attach([observation(
                "comtrade", origin, destination, "goods", year, value,
                "UN Comtrade, U.S.-reported aggregate goods trade with Macao; " +
                valuation + " value, nominal USD; customs basis, not BEA balance-of-payments basis",
                notes="Reported aggregate; source isAggregate=" + str(rec.get("isAggregate")) +
                ", isReported=" + str(rec.get("isReported")) +
                ". Not comparable to BEA accounting basis across geographies."
            )], asset, "Comtrade annual " + str(year)))
        if data["data"] and seen != {"X", "M"}:
            raise ValueError("Comtrade incomplete bilateral flow pair")
        time.sleep(1)
    if not rows:
        raise ValueError("Comtrade returned no Macao goods records")
    return rows, "Comtrade annual records through " + str(max(r["year"] for r in rows)), (
        "U.S.-reported customs trade: exports FOB, imports CIF; intentionally not spliced with BEA balance-of-payments series.")


ADAPTERS = {"bea": bea, "ntto": ntto, "uis": uis,
            "taiwan_tourism": taiwan, "iie": iie,
            "macau_tourism": macau, "comtrade": comtrade,
            "ustr": ustr_macau_services}
