"""Manual review of official green-card tables; never infer migration from visits."""

import io
import json
import re
import zipfile
from pathlib import Path, PurePosixPath
from urllib.parse import urljoin

import xlrd
from openpyxl import load_workbook

from .language_sources import row, series, source


ASSETS = {
    2004: "https://ohss.dhs.gov/sites/default/files/2023-12/LPR2004_2.zip",
    2014: "https://ohss.dhs.gov/sites/default/files/2023-12/YRBK%25202014%2520LPR%2520Excel%2520Final_0.zip",
    2024: "https://ohss.dhs.gov/system/files/2026-06/2026_0604_ohss_yearbook_lawful_permanent_residents_fy2024.xlsx",
    2025: "https://ohss.dhs.gov/system/files/2026-06/2026_0604_ohss_legal-immigration-adjustment-of-status-fy-2025q4.xlsx",
}
ADMISSION_ASSETS = {
    2016: "https://ohss.dhs.gov/sites/default/files/2023-12/fy2016_q1q2q3q4_tables_d.xlsx",
    2017: "https://ohss.dhs.gov/sites/default/files/2023-12/fy2017_q1q2q3q4_tables_d.xlsx",
    2018: "https://ohss.dhs.gov/sites/default/files/2023-12/fy2018_q1q2q3q4_tables_d.xls_.xlsx",
    2019: "https://ohss.dhs.gov/sites/default/files/2023-12/fy2019_q4_d_final.xlsx",
    2020: "https://ohss.dhs.gov/sites/default/files/2023-12/fy_2020_q4_d_final_v2.pdf.xlsx",
    2021: "https://ohss.dhs.gov/sites/default/files/2023-12/22_0323_plcy_legal_immigration_and_adjustment_of_status_report_fy2021_q4_final_d_version.xlsx",
    2022: "https://ohss.dhs.gov/sites/default/files/2023-12/2023_0308_plcy_legal_immigration_adjustment_of_status_report_fy_2022q4_final_d_0.xlsx",
    2023: "https://ohss.dhs.gov/sites/default/files/2024-06/2024_0507_ohss_legal-immigration-adjustment-of-status-fy-2023q4.xlsx",
    2024: "https://ohss.dhs.gov/sites/default/files/2025-06/2025_0624_ohss_legal-immigration-adjustment-of-status-fy-2024q4_0.xlsx",
    2025: ASSETS[2025],
}
ADMISSION_YEARS = tuple(sorted(ADMISSION_ASSETS))
ADMISSION_VINTAGES_NOTE = (
    "Nationality counts retain the reviewed Q4 report-table versions, not a harmonized final-yearbook series. "
    "The FY2016-FY2018 Table 1A assets explicitly record a September 2019 correction from birthplace "
    "to nationality; their observation fiscal years are unchanged.")
BIRTH_YEARS = tuple(range(1989, 2025))
ARRIVAL_YEARS = tuple(range(2000, 2025))
ARRIVAL_GAPS = {2003, 2004, 2005}
ARRIVAL_ASSETS = {
    year: f"https://ohss.dhs.gov/sites/default/files/2023-12/fy{year}_tables8-11newadj_d.xlsx"
    for year in range(2006, 2019)}
ARRIVAL_ASSETS.update({
    2014: "https://ohss.dhs.gov/sites/default/files/2023-12/fy2014_tables8-11newadj_d_0.xlsx",
    2015: "https://ohss.dhs.gov/sites/default/files/2023-12/fy2015_tables8-11newadj_d_0.xlsx",
    2019: "https://ohss.dhs.gov/sites/default/files/2023-12/fy2019_tables8-11_newadj_d.xlsx",
    2020: "https://ohss.dhs.gov/sites/default/files/2023-12/2021_0920_plcy_tables8-11newadj_fy2020_d.xlsx",
    2021: "https://ohss.dhs.gov/sites/default/files/2023-12/2202_0405_plcy_tables8-11newadj_fy2021_d.xlsx",
    2022: "https://ohss.dhs.gov/sites/default/files/2023-12/plcy_tables8-11newadj_fy2022_d.xlsx",
    2023: "https://ohss.dhs.gov/system/files/2026-07/2025_0828_ohss_tables8-11newadj_fy2023_v2.xlsx",
    2024: "https://ohss.dhs.gov/data/profiles/lpr/2024/fy2024cobbook51.xlsx",
})
ARRIVAL_SOURCE = "ohss_birth_arrivals_history"
ARRIVAL_SERIES = "ohss_birth_arrivals:CHN"
ARRIVAL_NOTE = (
    "Country of birth, not nationality or previous residence; excludes in-country adjustments. "
    "FY2003-2005 have no verified arrivals breakdown in the reviewed tables and remain missing, not zero. "
    "Original annual source vintages are retained; no imputation or visa-issuance proxy is used.")
ROUNDING_NOTE = (
    "Birthplace all-grant counts for FY2015-2024 and nationality counts for FY2023-2025 are rounded to ten. "
    "Birthplace new-arrival counts retain exact published values through FY2022; FY2023-2024 use ten-person precision. "
    "Earlier source vintages retain their published exact counts.")
COUNTRIES = {"China, People's Republic": "CHN", "Hong Kong": "HKG", "Macau": "MAC"}
UN_STOCK = "https://www.un.org/development/desa/pd/sites/www.un.org.development.desa.pd/files/undesa_pd_2024_ims_stock_by_sex_destination_and_origin.xlsx"
STOCK_YEARS = [2000, 2005, 2010, 2015, 2020, 2024]
STOCK_PAIRS = {("CHN", "USA"): "B", ("HKG", "USA"): "B",
               ("USA", "CHN"): "C R", ("USA", "HKG"): "B R"}
MACAU_ABODE_YEARS = tuple(range(2007, 2026))
DSEC_YEARBOOKS = "https://www.dsec.gov.mo/StatisticsJSON/AllReleases_en-US_General.json"


def parse_macau_abode_2008(records):
    title = re.sub(r"\s+", " ", " ".join(str(v) for r in records[:4] for v in r if v))
    if "INDIVIDUALS AUTHORIZED TO RESIDE IN MACAO BY COUNTRY/TERRITORY OF ISSUANCE OF IDENTITY DOCUMENT" not in title:
        raise ValueError("Macau 2008 residence-approval population changed")
    notes = re.sub(r"\s+", " ", " ".join(str(v) for r in records for v in r if v))
    if "direct comparison with data prior to 2007 is not appropriate" not in notes:
        raise ValueError("Macau revised-compilation boundary missing")
    headers = [r for r in records[:8] if "Country / Territory" in str(r[0])]
    if len(headers) != 1:
        raise ValueError("Macau 2008 approval year header missing or ambiguous")
    columns = {int(str(v).removesuffix(".0")): i for i, v in enumerate(headers[0])
               if re.fullmatch(r"20\d{2}(?:\.0)?", str(v).strip())}
    if list(columns) != [2007, 2008]:
        raise ValueError("Macau 2008 reviewed year columns changed")
    matches = [r for r in records if str(r[8]).strip() == "United States"
               and str(r[6]).strip() == "Estados Unidos"]
    totals = [r for r in records if str(r[1]).strip().endswith("Total")]
    if len(matches) != 1 or len(totals) != 1:
        raise ValueError("Macau 2008 U.S. category or total missing or ambiguous")
    values = {year: count(matches[0][column])[0] for year, column in columns.items()}
    controls = {year: count(totals[0][column])[0] for year, column in columns.items()}
    if any(values[y] is None or controls[y] is None or values[y] > controls[y] for y in columns):
        raise ValueError("Macau 2008 U.S. approvals exceed the total or require suppression review")
    return values


def parse_macau_abode(records):
    title = " ".join(str(v).strip() for r in records[:4] for v in r if v)
    if not re.search(r"INDIVIDUALS (?:NEWLY )?GRANTED RIGHT OF ABODE BY COUNTRY OR TERRITORY", title):
        raise ValueError("Macau residence-approval population or country category changed")
    headers = [r for r in records[:8] if "Country or territory" in str(r[0])]
    if len(headers) != 1:
        raise ValueError("Macau residence-approval year header missing or ambiguous")
    columns = {int(str(v).removesuffix(".0")): i for i, v in enumerate(headers[0])
               if re.fullmatch(r"20\d{2}(?:\.0)?", str(v).strip())}
    if len(columns) != 5 or sorted(columns) != list(columns):
        raise ValueError("Macau residence-approval year columns changed")
    matches = [i for i, r in enumerate(records) if str(r[2]).strip() == "United States of America"]
    if len(matches) != 1 or matches[0] < 2 or str(records[matches[0] - 1][2]).strip() != "Estados Unidos da América":
        raise ValueError("Macau U.S. country row missing or ambiguous")
    us = records[matches[0] - 2]

    def counts(r):
        result = {}
        for year, column in columns.items():
            value = r[column]
            if isinstance(value, str) and value.isdigit():
                value = int(value)
            result[year] = count(value)[0]
        return result

    values = counts(us)
    total_rows = [i for i, r in enumerate(records) if str(r[1]).strip() == "Total"]
    if len(total_rows) != 1:
        raise ValueError("Macau total control row missing")
    totals = counts(records[total_rows[0] - 1])
    if any(values[y] is None or totals[y] is None or values[y] > totals[y] for y in columns):
        raise ValueError("Macau U.S. approvals exceed the published total or require suppression review")
    return values


def macau_abode(client):
    body, _ = client.get(DSEC_YEARBOOKS)
    records = json.loads(body)["YearbookOfStatistics"]["Containers"]["YearbookOfStatistics"]["Publication"]["Yearly"]
    publications = sorted((r for r in records if int(r["Name"]) in {2008, 2013, 2014} or int(r["Name"]) >= 2020),
                          key=lambda r: int(r["Name"]))
    observations, sources = {}, {}
    for publication in publications:
        edition = int(publication["Name"])
        files = [f for f in publication["File"]
                 if f["Type"] in {".xlsx", ".xls"} and f["Name"] == "Statistical Tables 1"]
        if len(files) != 1:
            raise ValueError("Macau population workbook missing or ambiguous")
        body, asset = client.get(urljoin(DSEC_YEARBOOKS, files[0]["Link"]))
        if files[0]["Type"] == ".xls":
            book = xlrd.open_workbook(file_contents=body)
            tables = [(s.name.strip(), [s.row_values(i) for i in range(s.nrows)]) for s in book.sheets()]
        else:
            book = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
            tables = [(s.title.strip(), list(s.values)) for s in book]
        table_name = "2.2.9" if edition == 2008 else "2.10"
        matches = [(name, rows) for name, rows in tables if name == table_name]
        if len(matches) != 1:
            raise ValueError("Macau residence-approval table missing or ambiguous")
        values = (parse_macau_abode_2008 if edition == 2008 else parse_macau_abode)(matches[0][1])
        source_id = f"dsec_abode_{edition}"
        sources[source_id] = source(
            "Macao Statistics and Census Service", f"Yearbook of Statistics {edition}",
            [asset], ("Table 2.2.9; United States identity-document issuing-country category; "
                      "2007-2008; revised-compilation note excludes comparison before 2007"
                      if edition == 2008 else
                      "Table 2.10; United States of America country/territory category; all genders"),
            "Individuals newly granted residence, not resident stocks, visa issuances or physical arrivals. "
            + ("The earliest English title says 'authorized to reside'; no " if edition == 2008 else
               "The English title uses 'right of abode'; no ") +
            "permanent-resident legal classification or U.S. "
            "citizenship/birthplace population is inferred. Overlapping annual tables are reconciled."
            + (" The 2008 original explicitly identifies document issuing country and a revised "
               "compilation from 2007; earlier-method figures are not joined." if edition == 2008 else ""))
        for year, value in values.items():
            if year < MACAU_ABODE_YEARS[0]:
                continue
            if year in observations and observations[year]["value"] != value:
                raise ValueError("Macau overlapping annual approval counts changed; review the revision")
            observations[year] = row(
                "dsec_abode:USA:MAC", year, value, source_id,
                notes="Calendar-year people approved; source U.S. country/territory category, "
                "not verified U.S. departures or a permanent-residence classification.")
    if set(observations) != set(MACAU_ABODE_YEARS):
        raise ValueError("Macau reviewed annual approval history incomplete or has unreviewed years")
    definitions = {"dsec_abode:USA:MAC": series(
        "U.S. country category", "residence_authorizations", "people approved",
        "Macau individuals newly granted residence in DSEC's United States of America country/territory category",
        "annual", "calendar",
        "Annual approvals, not the stock of U.S. residents, physical U.S. departures or visa issuances. "
        "Not labelled permanent-resident grants; the source country category does not establish citizenship or birthplace. "
        "The revised compilation starts in 2007; earlier-method figures are not joined. "
        "The earliest original explicitly identifies identity-document issuing country.")}
    return definitions, sorted(observations.values(), key=lambda r: r["year"]), sources


def parse_stock(records):
    header = records[10]
    if (header[4], header[6]) != ("Location code of destination", "Location code of origin"):
        raise ValueError("UN destination/origin identity columns changed")
    if [str(value) for value in header[9:15]] != [str(y) for y in STOCK_YEARS]:
        raise ValueError("UN both-sexes benchmark columns changed")
    codes = {156: "CHN", 344: "HKG", 446: "MAC", 840: "USA"}
    values = {}
    for record in records[11:]:
        pair = codes.get(record[6]), codes.get(record[4])
        if pair not in STOCK_PAIRS:
            continue
        if pair in values or record[3] != STOCK_PAIRS[pair]:
            raise ValueError("UN bilateral population is ambiguous or its data type changed")
        values[pair] = [count(value) for value in record[9:15]]
    if set(values) != set(STOCK_PAIRS):
        raise ValueError("UN reviewed bilateral stock coverage incomplete")
    return values


def stock(client):
    content, asset = client.get_public_download(UN_STOCK)
    records = list(load_workbook(io.BytesIO(content), read_only=True, data_only=True)["Table 1"].values)
    definitions, observations = {}, []
    labels = {("CHN", "USA"): "China-born in the U.S.", ("HKG", "USA"): "Hong Kong-born in the U.S.",
              ("USA", "CHN"): "U.S. citizens in mainland China", ("USA", "HKG"): "U.S.-born in Hong Kong"}
    for pair, values in parse_stock(records).items():
        origin, destination = pair
        sid = f"un_stock:{origin}:{destination}"
        definition = series(labels[pair], "migrant_stock", "estimated residents",
                            labels[pair] + "; UN destination-specific population definition",
                            "benchmark", basis="mid-year July 1",
                            notes="Estimated resident stock, not annual immigration or permanent-residency grants. "
                            "UN type C uses foreign citizenship; B uses foreign birth; R includes refugee adjustments. "
                            "The U.S. destination population includes the territories specified in UN coverage note 36.")
        definitions[sid] = dict(definition, origin=origin, destination=destination, data_type=STOCK_PAIRS[pair])
        for year, (value, status) in zip(STOCK_YEARS, values):
            observations.append(row(sid, year, value, "un_stock_2024",
                                    "estimated" if status == "observed" else status,
                                    "UN 2024 revision; 1 July benchmark. Not an annual flow.", sid))
    return definitions, observations, source(
        "United Nations DESA, Population Division", "International Migrant Stock 2024",
        [asset], "Table 1; both sexes combined; destination/origin codes 156, 344 and 840; mid-year 2000–2024",
        "CC BY 3.0 IGO. Adapted selected numerical aggregates; UN country codes and population definitions retained. "
        "Mainland China uses foreign-citizenship estimates (C R), Hong Kong foreign-born estimates (B R), "
        "and the U.S. foreign-born estimates (B). Benchmark stocks include UN estimates, not new yearly arrivals.")


def count(value, rounded=False):
    if isinstance(value, str):
        value = value.strip()
        if value == "D":
            return None, "suppressed"
        if value in {"", "N/A"}:
            return None, "missing"
        if value == "-":
            return 0, "rounded" if rounded else "observed"
    if isinstance(value, bool) or not isinstance(value, (int, float)) or value < 0 or int(value) != value:
        raise ValueError("Invalid reported immigration count")
    if rounded and value % 10:
        raise ValueError("OHSS rounded count is not a multiple of ten")
    return int(value), "rounded" if rounded else "observed"


def parse_birth(records, first, last, rounded=False):
    text = " ".join(str(v) for r in records[:5] for v in r if v).lower()
    if "country of birth" not in text or not any(s in text for s in ["immigrants admitted", "permanent resident"]):
        raise ValueError("Immigration table must count permanent-resident grants by birthplace")
    header = next(r for r in records if str(r[0]).strip().lower() == "region and country of birth")
    columns = {int(v): i for i, v in enumerate(header) if isinstance(v, (float, int)) and 1980 <= v <= 2100}
    if not set(range(first, last + 1)) <= columns.keys():
        raise ValueError("Immigration history year columns incomplete")
    if rounded and "nearest 10" not in " ".join(str(v) for r in records for v in r if v).lower():
        raise ValueError("OHSS rounding disclosure changed")
    result = []
    for name, geo in COUNTRIES.items():
        matches = [r for r in records if str(r[0]).strip() == name]
        if len(matches) != 1:
            raise ValueError("Immigration country row missing or ambiguous")
        result += [(geo, year, *count(matches[0][columns[year]], rounded)) for year in range(first, last + 1)]
    return result


def admission_rounding(records):
    text = " ".join(str(v) for r in records for v in r if v).lower()
    return bool(re.search(r"nearest\s+(?:10|ten)\b", text))


def parse_admissions(records, year, rounded=None):
    text = " ".join(str(v) for r in records[:8] for v in r if v).lower()
    if "country of nationality" not in text or f"fiscal year {year}" not in text:
        raise ValueError("Latest immigration nationality/year identity changed")
    headers = [i for i, r in enumerate(records) if len(r) >= 12
               and [r[j] for j in (1, 6, 11)] == ["Total", "Adjustments of status", "New arrivals"]]
    if len(headers) != 1:
        raise ValueError("Latest immigration admission-type columns changed")
    parent = headers[0]
    if not any(r[0] == "Region and country of nationality" for r in records[parent:parent + 2]):
        raise ValueError("Admission totals must retain the nationality population")
    if [records[parent + 1][i] for i in (1, 6, 11)] != ["Total"] * 3:
        raise ValueError("Use annual reported totals, not selected quarters")
    matches = [r for r in records if str(r[0]).strip() == "China, People's Republic"]
    if len(matches) != 1:
        raise ValueError("Latest China nationality row missing or ambiguous")
    rounded = admission_rounding(records) if rounded is None else rounded
    values = {kind: count(matches[0][i], rounded=rounded)[0]
              for kind, i in [("total", 1), ("adjustments", 6), ("new_arrivals", 11)]}
    if any(v is None for v in values.values()) or abs(values["total"] - values["adjustments"] - values["new_arrivals"]) > (10 if rounded else 0):
        raise ValueError("Reported admission types do not reconcile within rounding precision")
    return values


def parse_latest(records):
    return parse_admissions(records, 2025, rounded=True)


def parse_birth_arrivals(records, year, profile=False):
    records = [[re.sub(r"\s+", " ", v).strip() if isinstance(v, str) else v for v in r] for r in records]
    text = re.sub(r"\s+", " ", " ".join(str(v) for r in records[:8] for v in r if v)).lower()
    if ("country of birth" not in text.replace("region/country", "country") or f"fiscal year {year}" not in text
            or "country of last residence" in text or "country of nationality" in text):
        raise ValueError("Arrival birthplace/fiscal-year population changed")
    rounded = admission_rounding(records) or (profile and year == 2024)
    if profile:
        if "region/country: china, people's republic" not in text:
            raise ValueError("Arrival profile country changed")
        headers = [r for r in records if r[:2] == ["Characteristic", "Total"]]
        matches = [r for r in records if str(r[0]).strip() == "New arrivals"]
        if len(headers) != 1 or len(matches) != 1:
            raise ValueError("Arrival profile requires a separate new-arrivals total")
        return count(matches[0][1], rounded=rounded)
    matches = [r for r in records if str(r[0]).strip() == "China, People's Republic"]
    headers = [i for i, r in enumerate(records) if r[0] == "Region and country of birth"]
    if len(matches) != 1 or len(headers) != 1:
        raise ValueError("Arrival birthplace country/header missing or ambiguous")
    header = headers[0]
    if records[header][1] == "Adjustments of status" and records[header][8] == "New arrivals":
        column = 8
        if records[header + 1][column] != "Total":
            raise ValueError("Arrival table requires the all-class total")
    elif records[header][1] == "Total" and "new arrivals -" in text:
        column = 1
    else:
        raise ValueError("Arrival type must be new arrivals, not adjustments or all grants")
    value, status = count(matches[0][column], rounded=rounded)
    components = [count(v, rounded=rounded)[0] for v in matches[0][column + 1:column + 7]]
    if value is None or (None not in components and abs(value - sum(components)) > (30 if rounded else 0)):
        raise ValueError("Arrival classes do not reconcile to their total")
    return value, status


def birth_arrivals(client):
    seed = json.loads((Path(__file__).resolve().parents[2] / "data/reviewed/immigration-arrivals-early.json").read_text())
    assets, observations = [], []
    for table in seed["tables"]:
        record, year = table["row"], table["year"]
        if (year not in (2000, 2001, 2002) or record[0] != "China, People's Republic"
                or record[1] != sum(record[2:])):
            raise ValueError("Reviewed early arrival population or admission totals changed")
        assets.append(dict(
            url=seed["archive_url"], request_url=seed["archive_url"], sha256=table["sha256"],
            retrieved_at=table["retrieved_at"], member_path=table["member_path"],
            crc32=table["crc32"], checksum_scope="extracted XLS member, not the complete ZIP",
            dataset=seed["dataset"], licence=seed["licence"]))
        observations.append(row(ARRIVAL_SERIES, year, record[2], ARRIVAL_SOURCE,
                                notes="INS Table 11, country of birth; exact new-arrivals column. " + ARRIVAL_NOTE))
    for year, url in sorted(ARRIVAL_ASSETS.items()):
        body, asset = client.get(url)
        book = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
        names = [s for s in book if year == 2024 or s.title.replace(" ", "").lower().startswith("table10new")]
        if len(names) != 1:
            raise ValueError(f"FY{year} expanded country-of-birth arrivals worksheet missing or ambiguous")
        value, status = parse_birth_arrivals(list(names[0].iter_rows(max_col=15, values_only=True)), year, profile=year == 2024)
        assets.append(asset)
        observations.append(row(ARRIVAL_SERIES, year, value, ARRIVAL_SOURCE, status,
                                ("China birthplace profile, New arrivals row. " if year == 2024 else
                                 "Expanded Table 10, New arrivals total. ") + ARRIVAL_NOTE))
    for year, file in [(2003, "COBBook63.xls"), (2004, "COBBook63.xls"), (2005, "COBBook60.xls")]:
        body, asset = client.get(f"https://ohss.dhs.gov/data/profiles/lpr/{year}/{file}")
        sheet = xlrd.open_workbook(file_contents=body).sheet_by_index(0)
        records = [sheet.row_values(i) for i in range(sheet.nrows)]
        text = " ".join(str(v) for r in records for v in r if v).lower()
        if (f"fiscal year {year}" not in text or "region/country of birth" not in text
                or "region/country: china" not in text or "new arrivals" in text):
            raise ValueError("Early profile coverage changed; review the arrival gap")
        assets.append(asset)
        observations.append(row(ARRIVAL_SERIES, year, None, ARRIVAL_SOURCE, "missing", ARRIVAL_NOTE))
    definition = series(
        "Mainland China-born new arrivals", "permanent_resident_arrivals", "people",
        "U.S. new lawful permanent-resident arrivals; country of birth: China, People's Republic",
        "annual", basis="fiscal ending year", notes=ARRIVAL_NOTE)
    provenance = source(
        "U.S. INS / DHS / OHSS", "Reviewed fiscal-year arrivals tables, 2000-2024", assets,
        "FY2000-2002 INS Table 11 originals in Mendeley r2r444tpzs v2; FY2006-2023 expanded Table 10; "
        "FY2024 China birthplace profile, New arrivals row. FY2003-2005 profiles lack that breakdown.",
        ARRIVAL_NOTE + " Early XLS member SHA256/CRC32 are verified, not the complete research ZIP. "
        "The authors' grouped/imputed series is not used. FY2023-2024 counts are rounded to ten.")
    return {ARRIVAL_SERIES: definition}, sorted(observations, key=lambda r: r["year"]), {ARRIVAL_SOURCE: provenance}


def admissions(client):
    observations, sources = [], {}
    for year, url in ADMISSION_ASSETS.items():
        body, asset = client.get(url)
        book = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
        matches = [sheet for sheet in book if sheet.title.replace(" ", "").upper() == "TABLE1A"]
        if len(matches) != 1:
            raise ValueError(f"FY{year} nationality/admission-type table missing")
        records = list(matches[0].values)
        values = parse_admissions(records, year)
        rounded = admission_rounding(records)
        correction = ""
        if year in (2016, 2017, 2018):
            text = re.sub(r"\s+", " ", " ".join(str(v) for r in records for v in r if v)).lower()
            if ("previously published numbers in this table were computed using country of birth." not in text
                    or "on september 2019, numbers in this table were updated to account for country of nationality." not in text):
                raise ValueError("Early report nationality correction changed")
            correction = (
                "Table 1A was corrected from Country of Birth to Country of Nationality in September 2019. ")
        source_id = f"ohss_nationality_{year}"
        sources[source_id] = source(
            "U.S. Department of Homeland Security / OHSS",
            f"Legal Immigration and Adjustment of Status Report, FY{year} Q4"
            + ("; nationality corrected September 2019" if correction else ""), [asset],
            "Table 1A; China, People's Republic; full-fiscal-year totals for all grants, adjustments and new arrivals.",
            "Country of nationality, not birthplace or previous residence. "
            + ("Rounded to nearest ten. " if rounded else "Exact counts as published in this release. ") +
            correction
            + ("Each fiscal year retains its reviewed Q4 report-table version, not final yearbook revisions. "
               if correction else
               "Each fiscal year retains its Q4 preliminary administrative-report vintage, not final yearbook revisions. ")
            + "New arrivals enter as permanent residents; adjustments concern people already in the U.S. "
            "Total grants include both and are not all physical moves from mainland China.")
        for kind, value in values.items():
            sid = "ohss_nationality:" + kind
            observations.append(row(sid, year, value, source_id, "rounded" if rounded else "observed",
                                    "Fiscal year ending September 30. " +
                                    ("Rounded to nearest ten." if rounded else "Exact count from this source vintage."), sid))
    return observations, sources


def birthplace(client):
    definitions, observations, sources = {}, [], {}
    for year in (2004, 2014, 2024):
        body, asset = client.get(ASSETS[year])
        if year < 2024:
            archive = zipfile.ZipFile(io.BytesIO(body))
            names = [n for n in archive.namelist() if PurePosixPath(n).name.lower().replace(" ", "") == "table3d.xls"]
            if len(names) != 1:
                raise ValueError("Historical birthplace table missing or ambiguous")
            sheet = xlrd.open_workbook(file_contents=archive.read(names[0])).sheet_by_index(0)
            records = [sheet.row_values(i) for i in range(sheet.nrows)]
            first, last = (BIRTH_YEARS[0], 2004) if year == 2004 else (2005, 2014)
        else:
            book = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
            records = list(book["Table 3"].values)
            first, last = 2015, 2024
        sid = f"ohss_birth_{year}"
        sources[sid] = source("U.S. Department of Homeland Security / OHSS", f"{year} Yearbook of Immigration Statistics",
                              [asset], f"Table 3; country of birth; fiscal years {first}-{last}",
                              "Lawful permanent-resident status grants, including new arrivals and adjustments within the U.S. "
                              + ("Published cells rounded to nearest ten." if year == 2024 else "Historical reported counts; older source vintage."))
        for geo, reported_year, value, status in parse_birth(records, first, last, rounded=year == 2024):
            observations.append(row(f"ohss_birth:{geo}", reported_year, value, sid, status,
                                    "Fiscal year ending September 30. " + ("Rounded to nearest ten." if status == "rounded" else ""),
                                    f"ohss_birth:{geo}"))
    for name, geo in COUNTRIES.items():
        definitions[f"ohss_birth:{geo}"] = series(
            {"CHN": "Mainland China-born", "HKG": "Hong Kong-born", "MAC": "Macau-born"}[geo],
            "permanent_resident_grants", "status grants", f"U.S. lawful permanent-resident grants; country of birth: {name}",
            "annual", basis="fiscal ending year", notes="Includes arrivals and adjustments of status; not nationality, ethnicity or a count of new border crossings.")
    return definitions, observations, sources


def collect(client):
    definitions, observations, sources = birthplace(client)
    arrival_definitions, arrival_rows, arrival_sources = birth_arrivals(client)
    definitions.update(arrival_definitions)
    observations.extend(arrival_rows)
    sources.update(arrival_sources)
    admission_rows, admission_sources = admissions(client)
    sources.update(admission_sources)
    observations.extend(admission_rows)
    for kind, label in [("total", "Total grants"), ("new_arrivals", "New arrivals"), ("adjustments", "Adjustments within the U.S.")]:
        key = "ohss_nationality:" + kind
        definitions[key] = series(label, "permanent_resident_grants", "status grants",
                                  "DHS China, People's Republic nationality category; U.S. permanent-resident grants",
                                  "annual", basis="fiscal ending year",
                                  notes="Different population from the birthplace history. Not a visa issuance count.")
    stock_definitions, stock_rows, stock_source = stock(client)
    definitions.update(stock_definitions)
    observations.extend(stock_rows)
    sources["un_stock_2024"] = stock_source
    abode_definitions, abode_rows, abode_sources = macau_abode(client)
    definitions.update(abode_definitions)
    observations.extend(abode_rows)
    sources.update(abode_sources)
    return {
        "version": 1, "series": definitions, "observations": sorted(observations, key=lambda r: (r["series_id"], r["year"])),
        "sources": sources,
        "coverage": [
            {"id": "to_us_birthplace", "status": "available", "description":
             f"China, Hong Kong and Macau birthplace series, FY{BIRTH_YEARS[0]}-FY{BIRTH_YEARS[-1]}; "
             "all permanent-resident status grants."},
            {"id": "to_us_nationality", "status": "available", "description":
             f"China nationality category, FY{ADMISSION_YEARS[0]}-FY{ADMISSION_YEARS[-1]} annual totals "
             "with new arrivals and adjustments separate. Hong Kong and Macau are not separately identified in this table."},
            {"id": "from_us", "status": "not_verified", "description": "A comparable annual U.S.-to-mainland-China permanent-residency series has not been verified. All-foreigner visas, residence permits and border crossings are not substitutes."},
            {"id": "bilateral_resident_stocks", "status": "available", "description": "UN bilateral resident estimates in both directions for mainland China and Hong Kong: 2000, 2005, 2010, 2015, 2020 and 2024. Mainland destination uses U.S. citizenship, not U.S. birthplace. These stocks are not immigration flows."},
            {"id": "macau_resident_stock", "status": "not_verified", "description": "The reviewed UN bilateral table does not separately report the U.S.–Macau pair. No zero or allocation from mainland China is inferred."},
            {"id": "macau_residence_approvals", "status": "available",
             "description": f"DSEC U.S.-category residence approvals for every calendar year "
             f"{MACAU_ABODE_YEARS[0]}-{MACAU_ABODE_YEARS[-1]}. "
             "Separate from U.S. permanent-resident grants and UN resident stocks; not a permanent-residency classification."},
            {"id": "to_us_birthplace_arrivals", "status": "partial", "description":
             "Mainland China-born new permanent-resident arrivals, FY2000-2024. "
             "FY2003-2005 remain missing, not zero or imputed; nationality counts remain separate."},
        ],
        "context": {"definition": "The OHSS series count grants of U.S. permanent-resident status; the UN series count resident stocks. Neither is a measure of net migration. Grants include adjustments by people already in the United States, not only new physical arrivals. Birthplace is not a departure country.",
                    "periods": "U.S. fiscal years end September 30. Macau approvals use calendar years. UN stocks use 1 July benchmarks. Birthplace, nationality and source-country categories remain separate.",
                    "macau_approvals": "DSEC counts newly approved people in its U.S. country category, not verified U.S. citizens, U.S. departures or permanent-resident grants.",
                    "admissions_vintages": ADMISSION_VINTAGES_NOTE,
                    "rounding": ROUNDING_NOTE,
                    "resident_stocks": "UN stocks refer to 1 July in actual published benchmark years. China destination estimates use foreign citizenship; Hong Kong and U.S. destination estimates use foreign birth. Refugee adjustments are included where flagged R. These are not new immigration or net migration. No annual observations are invented between benchmarks.",
                    "rights": "Selected statistical aggregates and original figures only; no source workbooks are redistributed. UN estimates: United Nations DESA, Population Division, International Migrant Stock 2024, adapted under CC BY 3.0 IGO."},
    }
